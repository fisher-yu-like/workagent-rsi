"""Read structural facts for comparison;no allowed-edit policy or verdict."""

from copy import copy
import hashlib
from urllib.parse import quote
from typing import Any
from pydantic import TypeAdapter


def read_structure(path):
    facts, gaps = {}, []
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, data_only=False, keep_links=True)
        try:
            facts["excel/sheets"] = list(book.sheetnames)
            for sheet in book:
                base = "excel/sheet/" + quote(sheet.title, safe="")
                facts[base + "/state"] = sheet.sheet_state
                facts[base + "/merges"] = sorted(
                    str(v) for v in sheet.merged_cells.ranges
                )
                facts[base + "/row_dimensions"] = {
                    str(k): [v.height, v.hidden]
                    for k, v in sheet.row_dimensions.items()
                }
                facts[base + "/column_dimensions"] = {
                    str(k): [v.width, v.hidden]
                    for k, v in sheet.column_dimensions.items()
                }
                if len(sheet._cells) > 50000:
                    raise ValueError("preservation cell budget exceeded")
                for cell in sheet._cells.values():
                    if cell.value is not None or cell.has_style:
                        facts[base + "/cell/" + cell.coordinate] = {
                            "value": cell.value,
                            "format": cell.number_format,
                            "style": [
                                str(copy(getattr(cell, name)))
                                for name in [
                                    "font",
                                    "fill",
                                    "border",
                                    "alignment",
                                    "protection",
                                ]
                            ],
                        }
                if (
                    sheet._charts
                    or sheet._images
                    or sheet._pivots
                    or len(sheet.conditional_formatting)
                ):
                    gaps.append(
                        base
                        + ": charts/images/pivots/conditional formats require additional preservation support"
                    )
            if book.defined_names or book._external_links:
                gaps.append("Excel names/external links are outside structural-v1")
        finally:
            book.close()
    elif path.suffix.lower() == ".docx":
        from docx import Document
        from .artifact_locations import heading_level

        doc = Document(path)

        def para(p):
            runs = []
            for r in p.runs:
                style = {"bold": r.bold, "italic": r.italic, "underline": r.underline}
                if runs and {k: v for k, v in runs[-1].items() if k != "text"} == style:
                    runs[-1]["text"] += r.text
                elif r.text:
                    runs.append({"text": r.text, **style})
            return {
                "text": p.text,
                "style": p.style.style_id,
                "heading_level": heading_level(p),
                "runs": runs,
            }

        facts["word/paragraphs"] = [para(p) for p in doc.paragraphs]
        facts["word/tables"] = [
            [
                [
                    {
                        "text": cell.text,
                        "paragraphs": [para(p) for p in cell.paragraphs],
                    }
                    for cell in row.cells
                ]
                for row in table.rows
            ]
            for table in doc.tables
        ]
        facts["word/body_order"] = [
            "paragraph" if node.tag.endswith("}p") else "table"
            for node in doc.element.body
            if node.tag.endswith(("}p", "}tbl"))
        ]
        facts["word/sections"] = [
            [
                s.page_width,
                s.page_height,
                s.top_margin,
                s.bottom_margin,
                s.left_margin,
                s.right_margin,
            ]
            for s in doc.sections
        ]
        facts["word/headers_footers"] = [
            [p.text for part in [s.header, s.footer] for p in part.paragraphs]
            for s in doc.sections
        ]
        xml = doc.element.xml
        if any(
            token in xml
            for token in [
                "<w:drawing",
                "<w:fldChar",
                "<w:hyperlink",
                "<w:object",
                "<w:ins",
                "<w:del",
            ]
        ):
            gaps.append(
                "Word embedded drawings/fields/hyperlinks/tracked edits need additional preservation support"
            )
        if any(
            cell.tables
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        ):
            gaps.append("nested Word tables unsupported")
    elif path.suffix.lower() == ".pptx":
        from pptx import Presentation

        deck = Presentation(path)
        facts["powerpoint/slides"] = [slide.slide_id for slide in deck.slides]
        facts["powerpoint/canvas"] = [deck.slide_width, deck.slide_height]
        for slide in deck.slides:
            base = "powerpoint/slide/" + str(slide.slide_id)

            def shapes(items, parent):
                order = []
                for shape in items:
                    key = parent + "/object/" + str(shape.shape_id)
                    order.append(shape.shape_id)
                    value = {
                        "type": str(shape.shape_type),
                        "geometry": [
                            shape.left,
                            shape.top,
                            shape.width,
                            shape.height,
                            shape.rotation,
                        ],
                    }
                    if getattr(shape, "has_text_frame", False):
                        value["text"] = shape.text
                        value["runs"] = [
                            [
                                {
                                    "text": r.text,
                                    "bold": r.font.bold,
                                    "italic": r.font.italic,
                                    "size": r.font.size,
                                    "name": r.font.name,
                                }
                                for r in p.runs
                            ]
                            for p in shape.text_frame.paragraphs
                        ]
                    if getattr(shape, "has_table", False):
                        value["table"] = [
                            [c.text for c in row.cells] for row in shape.table.rows
                        ]
                    if hasattr(shape, "image"):
                        value["image_hash"] = hashlib.sha256(
                            shape.image.blob
                        ).hexdigest()
                    if getattr(shape, "has_chart", False):
                        gaps.append(key + ": chart preservation unsupported")
                    if hasattr(shape, "shapes"):
                        shapes(shape.shapes, key)
                    elif (
                        not getattr(shape, "has_text_frame", False)
                        and not getattr(shape, "has_table", False)
                        and not hasattr(shape, "image")
                    ):
                        gaps.append(key + ": unsupported object content")
                    if hasattr(shape, "fill"):
                        value["fill_type"] = str(shape.fill.type)
                        if str(shape.fill.type).startswith("SOLID"):
                            value["fill_color"] = (
                                str(shape.fill.fore_color.rgb)
                                if shape.fill.fore_color.type
                                and str(shape.fill.fore_color.type).startswith("RGB")
                                else str(shape.fill.fore_color.type)
                            )
                    facts[key] = value
                facts[parent + "/objects"] = order

            shapes(slide.shapes, base)
    else:
        raise ValueError("unsupported preservation format")
    if len(facts) > 50000 or any(
        isinstance(v, list) and len(v) > 5000 for v in facts.values()
    ):
        raise ValueError("preservation object budget exceeded")
    return TypeAdapter(Any).dump_python(facts, mode="json"), gaps
