from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field

from .contracts import TaskSpec
from .path_safety import validated_task_directory


class PilotSkillConfig(BaseModel):
    model_config = ConfigDict(extra="forbid")
    marker_source: Literal["task_id", "required_text"] = "task_id"
    domains: set[str] = Field(default_factory=lambda: {"excel", "word", "powerpoint"})


class SkillConfiguredOfficeAdapter:
    MEDIA_TYPES = {
        "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "word": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "powerpoint": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }

    def __init__(self, output_root: str | Path, config: PilotSkillConfig) -> None:
        self.output_root = Path(output_root)
        self.config = config

    def execute(self, task: TaskSpec, skill_id: str) -> Iterable[dict[str, Any]]:
        yield {"kind": "started", "skill_id": skill_id, "task_id": task.task_id}
        domain = task.domain.lower()
        if domain not in self.MEDIA_TYPES or domain not in self.config.domains:
            yield {"kind": "failure", "message": f"unsupported skill domain: {task.domain}", "retryable": False}
            return
        try:
            validated_task_directory(self.output_root, task.task_id)
        except ValueError as exc:
            yield {"kind": "failure", "message": str(exc), "retryable": False}
            return
        marker = task.task_id
        if self.config.marker_source == "required_text":
            marker = str(task.expected_constraints.get("required_text", task.task_id))
        self.output_root.mkdir(parents=True, exist_ok=True)
        if domain == "excel":
            from openpyxl import Workbook

            path = self.output_root / f"{task.task_id}.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Artifact"
            sheet["A1"] = "Generated Artifact"
            sheet["A2"] = marker
            workbook.save(path)
        elif domain == "word":
            from docx import Document

            path = self.output_root / f"{task.task_id}.docx"
            document = Document()
            document.add_heading("Generated Artifact", level=1)
            document.add_paragraph(marker)
            document.save(path)
        else:
            from pptx import Presentation

            path = self.output_root / f"{task.task_id}.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Generated Artifact"
            slide.placeholders[1].text = marker
            presentation.save(path)
        yield {"kind": "artifact", "artifact_path": str(path), "media_type": self.MEDIA_TYPES[domain]}
