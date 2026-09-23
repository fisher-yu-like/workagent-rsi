from __future__ import annotations

from pathlib import Path
import shutil
import tempfile
from typing import Sequence

from .contracts import ArtifactRef, EvaluationReport, TaskSpec


class BasicEvaluator:
    """Initial deterministic evaluator; Office-specific evaluators are Phase 1B follow-ups."""

    def evaluate(self, task: TaskSpec, artifacts: Sequence[ArtifactRef], trace_id: str) -> EvaluationReport:
        failures: list[str] = []
        evidence: list[str] = []
        if not artifacts:
            failures.append("no artifacts produced")
        content = b"".join(Path(ref.path).read_bytes() for ref in artifacts)
        required = task.expected_constraints.get("required_text")
        if required is not None and str(required).encode("utf-8") not in content:
            failures.append(f"required text missing: {required}")
        for ref in artifacts:
            evidence.append(f"artifact:{ref.artifact_id}")
        passed = not failures
        return EvaluationReport(
            passed=passed,
            score=1.0 if passed else 0.0,
            critical_failures=failures,
            dimensions={"structural_correctness": 1.0 if passed else 0.0},
            evidence=evidence + [f"trace:{trace_id}"],
        )


class OfficeArtifactEvaluator:
    """Reopen genuine Office packages and verify the required marker."""

    MEDIA_TYPES = {
        "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "word": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "powerpoint": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }

    def evaluate(self, task: TaskSpec, artifacts: Sequence[ArtifactRef], trace_id: str) -> EvaluationReport:
        failures: list[str] = []
        evidence: list[str] = []
        marker = str(task.expected_constraints.get("required_text", ""))
        expected_media_type = self.MEDIA_TYPES.get(task.domain.lower())
        if not artifacts:
            failures.append("no artifacts produced")
        for ref in artifacts:
            path = Path(ref.path)
            evidence.append(f"artifact:{ref.artifact_id}")
            if expected_media_type is None or not ref.media_type.startswith(expected_media_type):
                failures.append(f"unexpected format for {task.domain}: {ref.media_type}")
                continue
            try:
                suffix = {"excel": ".xlsx", "word": ".docx", "powerpoint": ".pptx"}[task.domain.lower()]
                with tempfile.TemporaryDirectory(prefix="workagent-office-eval-") as temp_dir:
                    reopen_path = Path(temp_dir) / f"artifact{suffix}"
                    shutil.copyfile(path, reopen_path)
                    if task.domain.lower() == "excel":
                        from openpyxl import load_workbook

                        workbook = load_workbook(reopen_path, read_only=True, data_only=False)
                        values = [cell.value for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row]
                        workbook.close()
                    elif task.domain.lower() == "word":
                        from docx import Document

                        document = Document(reopen_path)
                        values = [paragraph.text for paragraph in document.paragraphs]
                        values.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                    else:
                        from pptx import Presentation

                        presentation = Presentation(reopen_path)
                        values = [shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text")]
                text = "\n".join(str(value) for value in values if value is not None)
                if marker and marker not in text:
                    failures.append(f"required marker missing: {marker}")
            except Exception as exc:
                failures.append(f"format reopen failed: {exc}")
        passed = not failures
        return EvaluationReport(
            passed=passed,
            score=1.0 if passed else 0.0,
            critical_failures=failures,
            dimensions={"format_validity": 1.0 if passed else 0.0, "marker_correctness": 1.0 if passed else 0.0},
            evidence=evidence + [f"trace:{trace_id}"],
        )
