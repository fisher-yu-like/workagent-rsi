"""Actual production-check calibration with authored truth kept out of judging."""
from __future__ import annotations

from collections import Counter
import json
import math
from pathlib import Path

from .assessment_contracts import AcceptanceSpec, Requirement, Location
from .artifact_assessment import SharedAssessment
from .artifact_evaluator import ArtifactEvaluator
from .artifact_verifier import ArtifactVerifier
from .hashing import canonical_json_hash, sha256_file


def proportion(n: int, denominator: int) -> dict:
    if denominator == 0:
        return {'numerator': n, 'denominator': 0, 'value': None, 'interval_95': None}
    z = 1.959963984540054
    p, scale = n / denominator, 1 + z*z/denominator
    centre = (p+z*z/(2*denominator))/scale
    radius = z*math.sqrt(p*(1-p)/denominator+z*z/(4*denominator**2))/scale
    return {'numerator': n, 'denominator': denominator, 'value': p,
            'interval_95': [0 if n==0 else max(0, centre-radius), 1 if n==denominator else min(1, centre+radius)]}


def summarize_calibration(rows: list[dict]) -> dict:
    if len({r['case_id'] for r in rows}) != len(rows):
        raise ValueError('duplicate calibration case ID')
    for row in rows:
        if row['truth'] not in {'PASS','FAIL','INCOMPLETE'} or row['predicted'] not in {'PASS','FAIL','PARTIAL','INCOMPLETE'}:
            raise ValueError('unknown calibration status')
    artifacts = [r for r in rows if r['truth']!='INCOMPLETE']
    complete = [r for r in artifacts if r['predicted']!='INCOMPLETE']
    invalid = [r for r in complete if r['truth']=='FAIL']
    valid = [r for r in complete if r['truth']=='PASS']
    gaps = [r for r in rows if r['truth']=='INCOMPLETE']
    located = [r for r in rows if r['truth']=='FAIL' and r['predicted'] in {'FAIL','PARTIAL'} and r['expected_locations']]
    location_hits = sum(bool(set(r['expected_locations']) & set(r['observed_locations'])) for r in located)
    return {'case_count': len(rows), 'artifact_error_denominator': len(artifacts),
        'completed_artifact_denominator': len(complete), 'incomplete_artifact_cases': len(artifacts)-len(complete),
        'confusion_matrix': {t: dict(Counter(r['predicted'] for r in rows if r['truth']==t)) for t in ['PASS','FAIL','INCOMPLETE']},
        'false_accept_rate': proportion(sum(r['predicted']=='PASS' for r in invalid),len(invalid)),
        'false_reject_rate': proportion(sum(r['predicted']!='PASS' for r in valid),len(valid)),
        'gap_classification': proportion(sum(r['predicted']=='INCOMPLETE' and r.get('gap_origin')==r.get('expected_gap_origin')
                                             and r.get('expected_gap_origin') is not None for r in gaps),len(gaps)),
        'location_accuracy': proportion(location_hits,len(located)),
        'family_count': len({r['source_family'] for r in rows}), 'partial_policy':'not_accepted',
        'independent_accuracy_claim':False,
        'interval_scope':'Descriptive Wilson intervals over observed cases; cases share families and are not independent population samples.'}


def calibration_specs(case, corpus_root: Path) -> list[AcceptanceSpec]:
    """Compile public task/input facts only. No output inspection or label access.

Authorized alternatives each use the production checker. Their frozen decision
rule is PASS if any authorized form passes, FAIL if all fail, else INCOMPLETE.
This measures declared checks; human-reviewed whole-task admission is separate.
"""
    source = Path(corpus_root) / case.input_path
    reqs = []
    def req(name, check, expected, **location):
        return Requirement(requirement_id=name,description=name,check=check,
            location=Location(artifact='output',**location),expected=expected,critical=True,
            dimension='correctness',evidence_source='Task instruction and frozen input document')
    if case.domain=='excel':
        from openpyxl import load_workbook
        book=load_workbook(source,data_only=False)
        try:
            cells=[[c.value for c in row] for row in book['Ledger'].iter_rows(min_row=2,max_row=4,min_col=1,max_col=2)]
            protected=book['Notes']['A1'].value
        finally: book.close()
        total=sum(row[1] for row in cells)
        reqs=[req('rows','excel.range',cells,sheet='Ledger',cell='A2:B4'),req('source','excel.cell',protected,sheet='Notes',cell='A1')]
        alternatives=[req('result','excel.formula_exact','=SUM(B2:B4)',sheet='Ledger',cell='B6')]
        if 'A literal total in B6 is also allowed.' in case.instruction:
            alternatives.append(req('result','excel.cell',total,sheet='Ledger',cell='B6'))
    elif case.domain=='word':
        from docx import Document
        doc=Document(source)
        cells=[[cell.text for cell in row.cells] for row in doc.tables[0].rows[1:]]
        total=sum(int(row[1]) for row in cells)
        reqs=[req('source','word.paragraph',doc.paragraphs[0].text,paragraph=1)]
        for i,row in enumerate(cells,2):
            reqs.extend(req(f'row-{i}-{j}','word.table_cell',v,table=1,row=i,column=j) for j,v in enumerate(row,1))
        alternatives=[req('result','word.paragraph',f'Total {total}',paragraph=3)]
        if 'A final Total row in the body table is also allowed.' in case.instruction:
            alternatives.append(req('result','word.table_cell',str(total),table=1,row=5,column=2))
    else:
        from pptx import Presentation
        deck=Presentation(source)
        table=deck.slides[1].shapes[0].table
        cells=[[c.text for c in row.cells] for row in list(table.rows)[1:]]
        total=sum(int(row[1]) for row in cells)
        reqs=[req('source','text.contains',deck.slides[0].shapes[0].text,slide=1,object_id=2),
              req('rows','powerpoint.texts',[v for row in cells for v in row],slide=2,object_id=2)]
        alternatives=[req('result','text.contains',f'Total {total}',slide=2,object_id=3)]
        if 'A separate Total table on slide 2 is also allowed.' in case.instruction:
            alternatives.append(req('result','powerpoint.texts',['Total',str(total)],slide=2,object_id=3))
    # Ambiguity and absent capabilities are obligations in the public instruction.
    if 'Ambiguous reference:' in case.instruction or 'Ambiguous scope:' in case.instruction:
        reqs.append(req('unresolved-request','unsupported.ambiguous_requirement',True))
    elif 'profile has no renderer' in case.instruction:
        reqs.append(req('rendered-visibility','visual','Readable, unclipped result'))
    elif 'profile has no semantic judge' in case.instruction or 'verification service is unavailable' in case.instruction:
        reqs.append(req('unavailable-review','semantic','Required authorized verification evidence'))
    return [AcceptanceSpec(task_id=case.case_id,version='calibration-declared-checks-v1',
        artifacts={'output':case.domain},requirements=[*reqs,alt,*([req('total-label','word.table_cell','Total',table=1,row=5,column=1)]
            if case.domain=='word' and alt.check=='word.table_cell' else [])],dimension_weights={'correctness':1},
        input_hashes={str(source):sha256_file(source)},
        scope_note='Development declared-check calibration; not independently reviewed whole-task acceptance.') for alt in alternatives]


def _position(loc: Location) -> str:
    if loc.sheet: return f'{loc.sheet}!{loc.cell}'
    if loc.table: return f'body/table[{loc.table}]'
    if loc.paragraph: return f'body/paragraph[{loc.paragraph}]'
    if loc.slide:
        return f'slide[{loc.slide}]/' + ('source' if loc.slide==1 else 'table[1]' if loc.object_id==2 else 'result')
    return 'artifact'


def _truth_locations(case) -> list[str]:
    if case.domain=='excel':
        return [{'source_damage':'Notes!A1','omitted_last_row':'Ledger!A2:B4','wrong_value':'Ledger!A2:B4'}.get(case.defect_type,'Ledger!B6')]
    if case.domain=='word':
        return [{'source_damage':'body/paragraph[1]','omitted_last_row':'body/table[1]','wrong_order':'body/table[1]'}.get(case.defect_type,'body/paragraph[3]')]
    return [{'source_damage':'slide[1]/source','omitted_last_row':'slide[2]/table[1]'}.get(case.defect_type,'slide[2]/result')]


def calibrate(cases, evaluator_config: dict, output_root: str | Path) -> dict:
    root=Path(output_root)
    if root.exists(): raise FileExistsError(root)
    corpus=Path(evaluator_config['corpus_root'])
    root.mkdir(parents=True)
    shared=SharedAssessment()  # Missing semantic/visual is kept as a real gap.
    evaluator, verifier=ArtifactEvaluator(),ArtifactVerifier()
    identity=canonical_json_hash({'production':shared.identity(),'compiler':sha256_file(Path(__file__)),
        'alternatives':'any_PASS_else_all_FAIL_else_INCOMPLETE','partial':'not_accepted'})
    rows=[]
    for case in cases:
        output=corpus/case.artifact_path
        if sha256_file(output)!=case.artifact_hash: raise ValueError('calibration artifact hash changed')
        specs=calibration_specs(case,corpus)
        results=[]
        for spec in specs:
            assessment=shared.inspect(spec,{'output':output})
            score=evaluator.evaluate(assessment)
            issues=verifier.verify(assessment)
            results.append({'assessment':assessment.model_dump(mode='json'),'score':score.model_dump(mode='json'),
                            'issues':issues.model_dump(mode='json')})
        statuses=[r['score']['acceptance_status'] for r in results]
        predicted='PASS' if 'PASS' in statuses else 'FAIL' if all(s=='FAIL' for s in statuses) else 'INCOMPLETE'
        selected=next((r for r in results if r['score']['acceptance_status']==predicted),results[0])
        findings=[c for c in selected['assessment']['checks'] if c['status'] not in {'PASS','NOT_APPLICABLE'}]
        origins={c['failure_origin'] for c in findings}
        observed=[_position(Location.model_validate(c['location'])) for c in findings if c['completion'] is not None]
        row={'case_id':case.case_id,'domain':case.domain,'source_family':case.source_family,'category':case.category,
            'truth':case.oracle_status,'predicted':predicted,'artifact_hash':case.artifact_hash,
            'expected_locations':_truth_locations(case) if case.oracle_status=='FAIL' else [],
            'observed_locations':observed,'expected_gap_origin':case.gap_origin,
            'gap_origin':'spec_gap' if 'spec_gap' in origins else 'capability_gap' if predicted=='INCOMPLETE' else None,
            'evaluator_identity':identity,'measurement_scope':'declared-check calibration; task_acceptance=NOT_ASSESSED',
            'alternative_reports':results}
        rows.append(row)
    summary=summarize_calibration(rows)
    report={'status':'completed','evaluator_identity':identity,'summary':summary,
        'quality_qualified':False,'measurement_scope':'development declared checks; no human agreement or final accuracy claim',
        'known_critical_false_accepts':sum(r['truth']=='FAIL' and r['predicted']=='PASS' for r in rows),
        'permitted_equivalent_false_rejects':sum(r['category']=='equivalent' and r['predicted'] not in {'PASS','INCOMPLETE'} for r in rows)}
    with (root/'per-case-results.jsonl').open('x',encoding='utf-8') as stream:
        for row in rows: stream.write(json.dumps(row,ensure_ascii=False)+'\n')
    (root/'calibration-report.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    (root/'open-gaps.md').write_text('# Open calibration gaps\n\n'+ '\n'.join(
        f'- {r["case_id"]}: {r["gap_origin"]}; truth={r["truth"]}' for r in rows if r['predicted']=='INCOMPLETE')+
        '\n\nWhole-task human review, actual visual review, independent final families and natural human-label cohort remain separate.\n',encoding='utf-8')
    return report
