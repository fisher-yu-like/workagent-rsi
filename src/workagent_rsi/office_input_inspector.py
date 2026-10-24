"""Self-contained trusted inspection program, executed inside the tool sandbox."""
import hashlib
import json
from pathlib import Path


def inspect_office_file(path,domain):
    path=Path(path)
    remaining=[100000]; clipped=[False]
    def text(value):
        value=str(value);limit=min(2048,remaining[0]);result=value[:limit]
        if len(result)!=len(value):clipped[0]=True
        remaining[0]-=len(result)
        return result
    if domain=='excel':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False)
        sheets=[];complete=True;cells_left=4096
        for sheet in list(book)[:32]:
            complete=complete and sheet.max_row<=100 and sheet.max_column<=32
            column_count=min(sheet.max_column,32)
            row_count=min(sheet.max_row,100,cells_left//max(column_count,1))
            if row_count<sheet.max_row:complete=False
            rows=[[v if v is None or isinstance(v,(int,float,bool)) else text(v) for v in row]
                  for row in sheet.iter_rows(max_row=row_count,max_col=column_count,values_only=True)] if row_count else []
            cells_left-=row_count*column_count
            sheets.append({'name':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'values':rows})
        complete=complete and len(book.sheetnames)<=32;book.close()
        return {'domain':domain,'sheets':sheets,'complete':complete and not clipped[0]}
    if domain=='word':
        from docx import Document
        doc=Document(path)
        paragraphs=[{'index':i,'text':text(p.text),'style':p.style.name} for i,p in enumerate(doc.paragraphs[:100])]
        tables=[[[text(c.text) for c in row.cells[:32]] for row in table.rows[:100]] for table in doc.tables[:16]]
        complete=(len(doc.paragraphs)<=100 and len(doc.tables)<=16
                  and all(len(p.text)<=4000 for p in doc.paragraphs)
                  and all(len(t.rows)<=100 and len(t.columns)<=32 for t in doc.tables)
                  and all(len(c.text)<=4000 for t in doc.tables for row in t.rows for c in row.cells))
        return {'domain':domain,'paragraphs':paragraphs,'tables':tables,'complete':complete and not clipped[0]}
    if domain=='powerpoint':
        from pptx import Presentation
        deck=Presentation(path);slides=[];complete=len(deck.slides)<=100;shapes_left=1024
        for index,slide in enumerate(list(deck.slides)[:100]):
            selected=list(slide.shapes)[:min(100,shapes_left)]
            complete=complete and len(slide.shapes)==len(selected)
            shapes_left-=len(selected)
            shapes=[]
            for i,shape in enumerate(selected):
                raw_text=shape.text_frame.text if shape.has_text_frame else ''
                text_value=text(raw_text)
                complete=complete and len(raw_text)<=2048
                shapes.append({'index':i,'text':text_value,'left':shape.left,'top':shape.top,'width':shape.width,'height':shape.height})
            slides.append({'index':index,'shapes':shapes})
        return {'domain':domain,'slides':slides,'complete':complete and not clipped[0]}
    raise ValueError('unsupported inspection domain')


def run_inspection(config):
    path=Path(config['input'])
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    result={'input_path':config['input'],'input_sha256':digest,'structure':inspect_office_file(path,config['domain'])}
    if hashlib.sha256(path.read_bytes()).hexdigest()!=digest:raise ValueError('input changed during inspection')
    output=Path(config['output']);output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text(json.dumps(result,ensure_ascii=False,allow_nan=False),encoding='utf-8')
