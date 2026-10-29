"""Bounded, location-aware text evidence and semantic assessment without rendering."""

from __future__ import annotations

import json
from xml.etree.ElementTree import tostring

from pydantic import Field

from .assessment_contracts import AssessmentModel
from .judgement_provider import VisualJudgement


class SourceCitation(AssessmentModel):
    source_id: str = Field(min_length=1)
    quote: str = Field(min_length=1)


class SemanticObservation(AssessmentModel):
    evidence_id: str = Field(min_length=1)
    quote: str = Field(min_length=1)
    observation: str = Field(min_length=1)
    source_citations: list[SourceCitation] = Field(default_factory=list)


class SemanticJudgement(VisualJudgement):
    observations: list[SemanticObservation] = Field(min_length=1)


class EvidenceBudgetExceeded(ValueError):
    pass


def extract_evidence(obj, artifact, *, max_chars=60000, max_items=2000):
    """Never silently truncate. Unsupported embedded content is declared explicitly."""
    items, used = [], 0

    def add(text, **location):
        nonlocal used
        text = str(text)
        if not text.strip():
            return
        entry = {"evidence_id": f"e{len(items) + 1}", "location": {"artifact": artifact, **location}, "text": text}
        used += len(json.dumps(entry, ensure_ascii=False))
        if used > max_chars or len(items) >= max_items:
            raise EvidenceBudgetExceeded("semantic evidence exceeds configured character/item budget; split the requirement or increase the fixed budget")
        items.append(entry)

    limitations = ["Embedded images, drawings and page appearance require a visual check."]
    if hasattr(obj, "worksheets"):
        limitations.append("Formula strings are not recalculated values; numerical correctness belongs to deterministic checks.")
        for sheet in obj.worksheets:
            # Avoid iterating billions of empty cells in a formatting-expanded sheet.
            if sheet.max_row * sheet.max_column > max_items * 100:
                raise EvidenceBudgetExceeded("worksheet used range exceeds semantic extraction budget")
            add(f"Worksheet: {sheet.title}", sheet=sheet.title)
            for row in sheet:
                for cell in row:
                    if cell.value is not None:
                        add(cell.value, sheet=sheet.title, cell=cell.coordinate)
            for i, chart in enumerate(sheet._charts, 1):
                add(tostring(chart.to_tree(), encoding="unicode"), sheet=sheet.title, region=f"chart {i}")
    elif hasattr(obj, "paragraphs"):
        # Preserve body paragraph/table order rather than appending all tables at the end.
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        p_index = t_index = 0
        heading = None
        for child in obj.element.body:
            if child.tag.endswith("}p"):
                p_index += 1
                p = Paragraph(child, obj)
                style = p.style.name if p.style else ""
                if style.startswith("Heading"):
                    heading = p.text
                add(p.text, paragraph=p_index, heading=heading, region=f"style: {style}")
            elif child.tag.endswith("}tbl"):
                t_index += 1
                table = Table(child, obj)
                for r, row in enumerate(table.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        add(cell.text, table=t_index, row=r, column=c)
        limitations.append("Headers, footers, comments and embedded objects are outside this body-text extraction.")
    else:
        def shapes(items, parents=()):
            for shape in items:
                yield shape, parents
                if hasattr(shape, "shapes"):
                    yield from shapes(shape.shapes, parents+(shape.shape_id,))
        for i, slide in enumerate(obj.slides, 1):
            add(f"Slide {i}", slide=i)
            for shape, parents in shapes(slide.shapes):
                loc = {"slide": i, "object_id": shape.shape_id, **({"group_path":list(parents)} if parents else {})}
                if getattr(shape, "has_text_frame", False):
                    add(shape.text, **loc)
                if getattr(shape, "has_table", False):
                    for r, row in enumerate(shape.table.rows, 1):
                        for c, cell in enumerate(row.cells, 1):
                            add(cell.text, row=r, column=c, **loc)
                if getattr(shape, "has_chart", False):
                    chart = shape.chart
                    categories = [str(c.label) for c in chart.plots[0].categories]
                    for series in chart.series:
                        add(json.dumps({"series": series.name, "categories": categories, "values": list(series.values)}, ensure_ascii=False), **loc)
        limitations.append("Speaker notes are outside this slide-content extraction.")
    return {"items": items, "limitations": limitations, "complete_within_declared_scope": True}


def validate_semantic_evidence(decision, request):
    valid = {item["evidence_id"]: item for item in request["evidence"]["items"]}
    sources = {item["source_id"]: item for item in request.get("sources", [])}
    for observation in decision.observations:
        if observation.evidence_id not in valid or observation.quote not in valid[observation.evidence_id]["text"]:
            raise ValueError("semantic evidence ID or exact quote does not map to supplied content")
        for citation in observation.source_citations:
            if citation.source_id not in sources or citation.quote not in sources[citation.source_id]["text"]:
                raise ValueError("source citation does not map to the authorized source")


def source_coverage_complete(decision, request):
    required = request.get('requirement', {}).get('options', {}).get('verify_sources', False)
    return not required or bool(request.get('sources')) and all(o.source_citations for o in decision.observations)
