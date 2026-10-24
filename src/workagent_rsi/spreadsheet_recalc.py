"""Recalculate copies only through an explicit isolated LibreOffice worker."""
from pathlib import Path
import json,re,shutil
from typing import Literal,Any
from uuid import uuid4

from pydantic import Field

from .assessment_contracts import AssessmentModel
from .hashing import canonical_json_hash,sha256_file
from .view_safety import checked_path,ordinary_file,relative_reference


class RecalcResult(AssessmentModel):
    status: Literal['available','unavailable','error','timeout']
    source_sha256: str
    output_sha256: str | None = None
    engine_identity: str = 'not_configured'
    qualification: Literal['unverified','unit_fixture','isolated_run'] = 'unverified'
    values: dict[str,Any] = Field(default_factory=dict)
    environment: dict[str,Any] = Field(default_factory=dict)
    evidence_refs: list[str] = Field(default_factory=list)
    error: str | None = None


class RecalculationEngine:
    def __init__(self,*,sandbox=None,qualification='unverified',timeout_seconds=90,output_root=None):
        if qualification not in {'unverified','unit_fixture','isolated_run'}:raise ValueError('unknown recalculation qualification')
        if type(timeout_seconds) is not int or timeout_seconds<1:raise ValueError('positive recalculation timeout required')
        self.sandbox,self.qualification,self.timeout_seconds=sandbox,qualification,timeout_seconds
        self.output_root=Path(output_root) if output_root is not None else None

    def identity_config(self):
        return {'engine':'isolated-libreoffice-recalc-v1','qualification':self.qualification,
            'policy':self.sandbox.policy.model_dump(mode='json') if self.sandbox is not None else None,
            'worker_sha256':sha256_file(Path(__file__).with_name('libreoffice_recalc_worker.py')),'timeout':self.timeout_seconds}

    def recalculate(self,source:Path,output:Path)->RecalcResult:
        source=ordinary_file(source);output=checked_path(output)
        if output.exists() or output==source:raise FileExistsError('recalculation output must be new')
        digest=sha256_file(source);identity=canonical_json_hash(self.identity_config())
        def result(status,error=None,**kw):return RecalcResult(status=status,source_sha256=digest,engine_identity=identity,
            qualification=self.qualification,error=error,**kw)
        if source.suffix.lower()!='.xlsx' or output.suffix.lower()!='.xlsx':return result('unavailable','only xlsx recalculation is supported')
        if source.stat().st_size>104857600:return result('unavailable','fixed recalculation input budget exceeded')
        from openpyxl import load_workbook
        book=load_workbook(source,data_only=False,keep_links=True)
        try:
            if book._external_links:return result('unavailable','external workbook links are unavailable')
            for sheet in book:
                if len(sheet._cells)>50000:return result('unavailable','recalculation cell budget exceeded')
                for cell in sheet._cells.values():
                    if cell.data_type=='f' and (re.search(r'\[[^\]]+\]|https?://|\b(?:WEBSERVICE|RTD|NOW|TODAY|RAND|RANDBETWEEN|INDIRECT)\s*\(',str(cell.value),re.I)):
                        return result('unavailable','external or dynamic formula requires a separately authorized engine/data policy')
        finally:book.close()
        if self.sandbox is None:return result('unavailable','isolated recalculation worker is not configured')
        policy=self.sandbox.policy
        if policy.backend!='apptainer-ssh' or policy.network!='none':return result('unavailable','recalculation requires explicit no-network Apptainer worker')
        root=checked_path(self.output_root or output.parent)/('recalc-'+uuid4().hex)
        (root/'inputs').mkdir(parents=True,exist_ok=False)
        copied=root/'inputs/source.xlsx';shutil.copyfile(source,copied)
        if sha256_file(copied)!=digest:raise ValueError('recalculation source changed during copy')
        config={'source':'inputs/source.xlsx','source_sha256':digest,'worker_sha256':self.identity_config()['worker_sha256']}
        (root/'inputs/recalc-policy.json').write_text(json.dumps(config),encoding='utf-8')
        script=root/'recalculate.py';shutil.copyfile(Path(__file__).with_name('libreoffice_recalc_worker.py'),script)
        outcome=self.sandbox.run(script,root,self.timeout_seconds)
        if sha256_file(source)!=digest or sha256_file(copied)!=digest:raise ValueError('recalculation modified source or frozen copy')
        if outcome.status!='succeeded':return result('timeout' if outcome.status=='timeout' else 'unavailable','isolated worker '+outcome.status,evidence_refs=outcome.trace_refs)
        try:
            meta_path=ordinary_file(root/'outputs/recalc-metadata.json');meta=json.loads(meta_path.read_text(encoding='utf-8'))
            if meta.get('status')!='available':return result('unavailable',meta.get('error','worker capability unavailable'))
            if (meta.get('source_sha256')!=digest or meta.get('worker_sha256')!=config['worker_sha256'] or meta.get('calculate_all') is not True
                or meta.get('load_policy')!={'ReadOnly':True,'MacroExecutionMode':0,'UpdateDocMode':0} or meta.get('owned_cleanup') is not True):
                raise ValueError('worker did not establish source/calculate/load/cleanup policy')
            produced=ordinary_file(root/'outputs/recalculated.xlsx');out_hash=sha256_file(produced)
            if meta.get('output_sha256')!=out_hash or not meta.get('environment',{}).get('office_version'):
                raise ValueError('recalculated output/engine identity mismatch')
            if self.qualification=='isolated_run':
                traces=[json.loads(ordinary_file(root/Path(*relative_reference(ref).parts)).read_text()) for ref in outcome.trace_refs]
                if not any(t.get('cleanup',{}).get('confirmed_absent') is True and t.get('resource_observations',{}).get('passed') is True
                    and t.get('capability',{}).get('image_sha256')==policy.remote_image_sha256 for t in traces):
                    raise ValueError('no qualified runtime resource/image/cleanup proof')
                if meta['environment'].get('office_binary_sha256') is None:raise ValueError('missing engine binary identity')
            values={};book=load_workbook(produced,data_only=True,keep_links=False)
            try:
                for sheet in book:
                    if len(sheet._cells)>50000:raise ValueError('recalculated output cell budget exceeded')
                    for cell in sheet._cells.values():
                        if cell.value is not None:values[sheet.title+'!'+cell.coordinate]=cell.value
            finally:book.close()
            output.parent.mkdir(parents=True,exist_ok=True)
            with output.open('xb') as target,produced.open('rb') as stream:shutil.copyfileobj(stream,target)
            if sha256_file(output)!=out_hash or sha256_file(source)!=digest:raise ValueError('output copy or source identity changed')
            return result('available',output_sha256=out_hash,values=values,environment=meta['environment'],evidence_refs=[sha256_file(meta_path),*outcome.trace_refs])
        except (OSError,ValueError,KeyError,TypeError) as exc:return result('error',str(exc))
