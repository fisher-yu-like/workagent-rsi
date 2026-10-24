"""Frozen structural preservation for explicitly authorized Office edits."""
from __future__ import annotations

from copy import copy
from difflib import SequenceMatcher
from pathlib import Path
from typing import Literal
import hashlib
from urllib.parse import quote

from pydantic import Field,field_validator,TypeAdapter
from typing import Any

from .assessment_contracts import AssessmentModel,CheckRecord,Location
from .hashing import canonical_json_hash,sha256_file
from .view_safety import ordinary_file


class EditPermission(AssessmentModel):
    scope: str = Field(min_length=1)
    operation: Literal['replace','add','remove']

    @field_validator('scope')
    @classmethod
    def explicit_scope(cls,value):
        if value.split('/')[0] not in {'excel','word','powerpoint'} or value.startswith('/') or '*' in value or any(p in {'','..','.'} for p in value.split('/')):
            raise ValueError('preservation scope must be an explicit structural path')
        return value


class PreservationSpec(AssessmentModel):
    source_hash: str = Field(pattern=r'^[0-9a-f]{64}$')
    allowed_changes: list[EditPermission] = Field(default_factory=list)
    preserved_regions: list[str] = Field(default_factory=list)
    normalization_policy: Literal['office-structure-v1'] = 'office-structure-v1'

    @field_validator('preserved_regions')
    @classmethod
    def protected_scopes(cls,values):
        for value in values:EditPermission(scope=value,operation='replace')
        return values


def _under(path,scope):return path==scope or path.startswith(scope+'/')


def _snapshot(path):
    facts,gaps={},[]
    if path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,data_only=False,keep_links=True)
        try:
            facts['excel/sheets']=list(book.sheetnames)
            for sheet in book:
                base='excel/sheet/'+quote(sheet.title,safe='')
                facts[base+'/state']=sheet.sheet_state
                facts[base+'/merges']=sorted(str(v) for v in sheet.merged_cells.ranges)
                facts[base+'/row_dimensions']={str(k):[v.height,v.hidden] for k,v in sheet.row_dimensions.items()}
                facts[base+'/column_dimensions']={str(k):[v.width,v.hidden] for k,v in sheet.column_dimensions.items()}
                if len(sheet._cells)>50000:raise ValueError('preservation cell budget exceeded')
                for cell in sheet._cells.values():
                    if cell.value is not None or cell.has_style:
                        facts[base+'/cell/'+cell.coordinate]={'value':cell.value,'format':cell.number_format,
                            'style':[str(copy(getattr(cell,name))) for name in ['font','fill','border','alignment','protection']]}
                if sheet._charts or sheet._images or sheet._pivots or len(sheet.conditional_formatting):
                    gaps.append(base+': charts/images/pivots/conditional formats require additional preservation support')
            if book.defined_names or book._external_links:gaps.append('Excel names/external links are outside structural-v1')
        finally:book.close()
    elif path.suffix.lower()=='.docx':
        from docx import Document
        from .artifact_locations import heading_level
        doc=Document(path)
        def para(p):
            runs=[]
            for r in p.runs:
                style={'bold':r.bold,'italic':r.italic,'underline':r.underline}
                if runs and {k:v for k,v in runs[-1].items() if k!='text'}==style:runs[-1]['text']+=r.text
                elif r.text:runs.append({'text':r.text,**style})
            return {'text':p.text,'style':p.style.style_id,'heading_level':heading_level(p),'runs':runs}
        facts['word/paragraphs']=[para(p) for p in doc.paragraphs]
        facts['word/tables']=[[[{'text':cell.text,'paragraphs':[para(p) for p in cell.paragraphs]} for cell in row.cells] for row in table.rows] for table in doc.tables]
        facts['word/body_order']=['paragraph' if node.tag.endswith('}p') else 'table' for node in doc.element.body if node.tag.endswith(('}p','}tbl'))]
        facts['word/sections']=[[s.page_width,s.page_height,s.top_margin,s.bottom_margin,s.left_margin,s.right_margin] for s in doc.sections]
        facts['word/headers_footers']=[[p.text for part in [s.header,s.footer] for p in part.paragraphs] for s in doc.sections]
        xml=doc.element.xml
        if any(token in xml for token in ['<w:drawing','<w:fldChar','<w:hyperlink','<w:object','<w:ins','<w:del']):
            gaps.append('Word embedded drawings/fields/hyperlinks/tracked edits need additional preservation support')
        if any(cell.tables for table in doc.tables for row in table.rows for cell in row.cells):gaps.append('nested Word tables unsupported')
    elif path.suffix.lower()=='.pptx':
        from pptx import Presentation
        deck=Presentation(path)
        facts['powerpoint/slides']=[slide.slide_id for slide in deck.slides]
        facts['powerpoint/canvas']=[deck.slide_width,deck.slide_height]
        for slide in deck.slides:
            base='powerpoint/slide/'+str(slide.slide_id)
            def shapes(items,parent):
                order=[]
                for shape in items:
                    key=parent+'/object/'+str(shape.shape_id);order.append(shape.shape_id)
                    value={'type':str(shape.shape_type),'geometry':[shape.left,shape.top,shape.width,shape.height,shape.rotation]}
                    if getattr(shape,'has_text_frame',False):
                        value['text']=shape.text
                        value['runs']=[[{'text':r.text,'bold':r.font.bold,'italic':r.font.italic,'size':r.font.size,'name':r.font.name} for r in p.runs] for p in shape.text_frame.paragraphs]
                    if getattr(shape,'has_table',False):value['table']=[[c.text for c in row.cells] for row in shape.table.rows]
                    if hasattr(shape,'image'):value['image_hash']=hashlib.sha256(shape.image.blob).hexdigest()
                    if getattr(shape,'has_chart',False):gaps.append(key+': chart preservation unsupported')
                    if hasattr(shape,'shapes'):shapes(shape.shapes,key)
                    elif not getattr(shape,'has_text_frame',False) and not getattr(shape,'has_table',False) and not hasattr(shape,'image'):
                        gaps.append(key+': unsupported object content')
                    if hasattr(shape,'fill'):
                        value['fill_type']=str(shape.fill.type)
                        if str(shape.fill.type).startswith('SOLID'):
                            value['fill_color']=str(shape.fill.fore_color.rgb) if shape.fill.fore_color.type and str(shape.fill.fore_color.type).startswith('RGB') else str(shape.fill.fore_color.type)
                    facts[key]=value
                facts[parent+'/objects']=order
            shapes(slide.shapes,base)
    else:raise ValueError('unsupported preservation format')
    if len(facts)>50000 or any(isinstance(v,list) and len(v)>5000 for v in facts.values()):raise ValueError('preservation object budget exceeded')
    return TypeAdapter(Any).dump_python(facts,mode='json'),gaps


def compare_preservation(source,output,spec) -> list[CheckRecord]:
    spec=PreservationSpec.model_validate(spec.model_dump(mode='json') if isinstance(spec,PreservationSpec) else spec)
    source,output=ordinary_file(source),ordinary_file(output)
    if source.suffix.lower()!=output.suffix.lower():raise ValueError('source/output preservation formats differ')
    if max(source.stat().st_size,output.stat().st_size)>104857600:raise ValueError('preservation input budget exceeded')
    before,after=sha256_file(source),sha256_file(output)
    if before!=spec.source_hash:raise ValueError('frozen source hash mismatch')
    original,source_gaps=_snapshot(source);actual,output_gaps=_snapshot(output)
    for protected in spec.preserved_regions:
        if not any(_under(k,protected) for k in original):
            parent,_,index=protected.rpartition('/')
            if parent not in original or not isinstance(original[parent],list) or not index.isdigit() or not 1<=int(index)<=len(original[parent]):
                raise ValueError('preserved region does not resolve in source: '+protected)
    records=[]
    def allowed(path,operation):
        if any(_under(path,p) or _under(p,path) for p in spec.preserved_regions):return False
        return any(p.operation==operation and _under(path,p.scope) for p in spec.allowed_changes)
    def record(path,operation,left,right,permitted):
        observed={'scope':path,'operation':operation,'source':left,'output':right,'authorized':permitted}
        evidence=canonical_json_hash({'source':before,'output':after,'policy':spec.model_dump(mode='json'),'diff':observed})
        records.append(CheckRecord(requirement_id='preservation:'+str(len(records)+1),status='PASS' if permitted else 'FAIL',
            completion=1 if permitted else 0,location=Location(artifact='output'),expected='unchanged or explicitly authorized',
            observed=observed,evidence_refs=[evidence],failure_origin=None if permitted else 'artifact_defect'))
    def sequence(path,left,right):
        matcher=SequenceMatcher(a=[canonical_json_hash(v) for v in left],b=[canonical_json_hash(v) for v in right],autojunk=False)
        for tag,a,b,c,d in matcher.get_opcodes():
            if tag=='equal':continue
            operation={'insert':'add','delete':'remove','replace':'replace'}[tag]
            for i,value in enumerate(left[a:b] or [None],a+1):
                record(path+'/'+str(i),operation,value,right[c:d],allowed(path+'/'+str(i),operation))
    changed_containers={}
    for container in ['excel/sheets','powerpoint/slides']:
        if container in original:
            changed_containers[container]={'add':set(actual[container])-set(original[container]),'remove':set(original[container])-set(actual[container])}
    def subtree_permission(path,operation):
        for container,changes in changed_containers.items():
            for identity in changes.get(operation,set()):
                prefix=('excel/sheet/'+quote(identity,safe='')) if container=='excel/sheets' else 'powerpoint/slide/'+str(identity)
                if _under(path,prefix) and allowed(container,operation) and not any(_under(path,p) or _under(p,path) for p in spec.preserved_regions):return True
        return False
    for path in sorted(set(original)|set(actual)):
        left,right=original.get(path),actual.get(path)
        if path in original and path in actual and left==right:continue
        if path=='word/body_order' and any(p.scope in {'word/paragraphs','word/tables'} and p.operation in {'add','remove'} for p in spec.allowed_changes):
            permitted=not any(_under(path,p) or _under(p,path) for p in spec.preserved_regions)
            for tag,a,b,c,d in SequenceMatcher(a=left,b=right,autojunk=False).get_opcodes():
                if tag=='equal':continue
                permitted=permitted and all(allowed('word/'+('paragraphs' if v=='paragraph' else 'tables'),'remove') for v in left[a:b])
                permitted=permitted and all(allowed('word/'+('paragraphs' if v=='paragraph' else 'tables'),'add') for v in right[c:d])
            record(path,'replace',left,right,permitted);continue
        if path in original and path in actual and isinstance(left,list) and isinstance(right,list) and path in {'excel/sheets','powerpoint/slides','word/paragraphs','word/tables'}:
            sequence(path,left,right);continue
        operation='add' if path not in original else 'remove' if path not in actual else 'replace'
        record(path,operation,left,right,allowed(path,operation) or subtree_permission(path,operation))
    for message in sorted(set(source_gaps+output_gaps)):
        records.append(CheckRecord(requirement_id='preservation-gap:'+str(len(records)+1),status='NEEDS_REVIEW',
            location=Location(artifact='output'),expected='complete preservation support',observed=message,
            evidence_refs=[canonical_json_hash({'source':before,'output':after,'gap':message})],failure_origin='assessment_gap'))
    if not records:
        records.append(CheckRecord(requirement_id='preservation',status='PASS',completion=1,location=Location(artifact='output'),
            expected='unchanged or explicitly authorized',observed={'unchanged_structural_objects':len(original)},
            evidence_refs=[canonical_json_hash({'source':before,'output':after,'facts':original})]))
    if sha256_file(source)!=before or sha256_file(output)!=after:raise ValueError('preservation input changed during comparison')
    return records
