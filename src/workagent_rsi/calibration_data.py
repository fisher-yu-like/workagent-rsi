"""Project-generated development calibration materials, never final-test data.

Truth is authored in the checked-in manifest, not obtained by running an evaluator.
Only the recipes below construct files. Raw-XML integrity tests use the separate
literal claims in that manifest; this module never interprets those claims.
"""
from __future__ import annotations

from datetime import datetime
import json
import re
from pathlib import Path
from typing import Literal
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from pydantic import BaseModel, ConfigDict, Field, model_validator

from .hashing import sha256_file
from .view_safety import checked_path, ordinary_file, relative_reference


class CalibrationCase(BaseModel):
    model_config = ConfigDict(extra='forbid', frozen=True)
    case_id: str
    domain: Literal['excel', 'word', 'powerpoint']
    source_family: str
    defect_type: str | None
    oracle_status: Literal['PASS', 'FAIL', 'INCOMPLETE']
    expected_locations: list[str]
    expected_scope: dict
    source_refs: list[str]
    artifact_hash: str = Field(pattern=r'^[0-9a-f]{64}$')
    input_hash: str = Field(pattern=r'^[0-9a-f]{64}$')
    truth_source_hash: str = Field(pattern=r'^[0-9a-f]{64}$')
    artifact_path: str
    input_path: str
    instruction: str
    category: Literal['correct', 'single_defect', 'equivalent', 'gap']
    artifact_valid: bool
    gap_origin: Literal['spec_gap', 'capability_gap'] | None = None
    equivalence_authorization: str | None = None
    source_kind: Literal['project-generated'] = 'project-generated'
    seed: int

    @model_validator(mode='after')
    def label_consistency(self):
        expected = {'correct': 'PASS', 'single_defect': 'FAIL', 'equivalent': 'PASS', 'gap': 'INCOMPLETE'}
        if self.oracle_status != expected[self.category]:
            raise ValueError('category and oracle status disagree')
        if (self.category == 'single_defect') != (self.defect_type is not None):
            raise ValueError('only single-defect cases have a defect')
        if self.artifact_valid != (self.category != 'single_defect'):
            raise ValueError('artifact validity disagrees with the authored category')
        if (self.category == 'gap') != (self.gap_origin is not None):
            raise ValueError('gap origin required exactly for gaps')
        if self.category == 'equivalent' and (not self.equivalence_authorization or self.equivalence_authorization not in self.instruction):
            raise ValueError('equivalent solution needs explicit instruction authorization')
        if not all((self.expected_locations, self.expected_scope, self.source_refs)):
            raise ValueError('truth scope and provenance are required')
        return self


def _contained(root: Path, name: str) -> Path:
    relative_reference(name)
    path = checked_path(root / name)
    if not path.resolve().is_relative_to(root.resolve()):
        raise ValueError('calibration path escapes corpus')
    return path


def _freeze_zip(path: Path) -> None:
    """Remove library wall-clock metadata for byte-identical rebuilds."""
    with ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    with ZipFile(path, 'w', compression=ZIP_DEFLATED) as archive:
        for name, body in sorted(members.items()):
            if name == 'docProps/core.xml':
                body = re.sub(rb'(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)',
                              rb'\g<1>2026-10-16T00:00:00Z\g<2>', body)
            item = ZipInfo(name, (2026, 10, 16, 0, 0, 0))
            item.compress_type = ZIP_DEFLATED
            archive.writestr(item, body)


def _write_office(path: Path, domain: str, context: int, mutation: str) -> None:
    # Independent construction: no production evaluator, oracle or XML claim reader.
    values = ([12, 19, 7], [5, 17, 23], [8, 16, 21], [4, 11, 29], [9, 13, 22],
              [6, 18, 25], [10, 14, 27], [3, 20, 31], [15, 24, 32], [2, 26, 35])[context]
    total = sum(values)
    protected = f'Preserve source ledger {context + 1}.'
    damaged = mutation == 'source_damage'
    if domain == 'excel':
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = 'Ledger'
        sheet.append(['Item', 'Amount'])
        for label, value in zip(['North', 'South', 'West'], values):
            sheet.append([label, value])
        sheet['A6'] = 'Total'
        if mutation != 'input':
            sheet['B6'] = '=SUM(B2:B4)'
        if mutation == 'omitted_last_row':
            sheet['A4'] = sheet['B4'] = None
        elif mutation == 'wrong_range':
            sheet['B6'] = '=SUM(B2:B3)'
        elif mutation == 'misplaced_content':
            sheet['C6'], sheet['B6'] = sheet['B6'].value, None
        elif mutation == 'wrong_value':
            sheet['B2'] = values[0] + 1
        elif mutation == 'equivalent':
            sheet['B6'] = total
        book.create_sheet('Notes')['A1'] = 'DAMAGED' if damaged else protected
        book.properties.created = book.properties.modified = datetime(2026, 10, 16)
        book.save(path)
        book.close()
    elif domain == 'word':
        from docx import Document
        doc = Document()
        doc.add_paragraph('DAMAGED' if damaged else protected)
        doc.add_heading('Summary', level=1)
        result = f'Total {total + (mutation == "wrong_value")}'
        doc.add_paragraph(result if mutation not in {'input', 'misplaced_content', 'equivalent'} else '')
        table = doc.add_table(rows=1, cols=2)
        for cell, text in zip(table.rows[0].cells, ['Item', 'Amount']):
            cell.text = text
        pairs = list(zip(['North', 'South', 'West'], values))
        if mutation == 'wrong_order':
            pairs[0], pairs[1] = pairs[1], pairs[0]
        if mutation == 'omitted_last_row':
            pairs.pop()
        for label, value in pairs:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = label, str(value)
        if mutation == 'equivalent':
            cells = table.add_row().cells
            cells[0].text, cells[1].text = 'Total', str(total)
        if mutation == 'misplaced_content':
            doc.sections[0].footer.paragraphs[0].text = result
        doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 16)
        doc.save(path)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.dml.color import RGBColor
        deck = Presentation()
        source, summary = [deck.slides.add_slide(deck.slide_layouts[6]) for _ in range(2)]
        source.shapes.add_textbox(Inches(.5), Inches(.5), Inches(8), Inches(.5)).text = 'DAMAGED' if damaged else protected
        table = summary.shapes.add_table(4 if mutation != 'omitted_last_row' else 3, 2, Inches(.5), Inches(2), Inches(6), Inches(2)).table
        for col, text in enumerate(['Item', 'Amount']):
            table.cell(0, col).text = text
        for row, (label, value) in enumerate(zip(['North', 'South', 'West'], values), 1):
            if row < len(table.rows):
                table.cell(row, 0).text, table.cell(row, 1).text = label, str(value)
        if mutation == 'equivalent':
            # A separate result table is expressly permitted by these five tasks.
            out = summary.shapes.add_table(1, 2, Inches(.5), Inches(.5), Inches(6), Inches(.5)).table
            out.cell(0, 0).text, out.cell(0, 1).text = 'Total', str(total)
        elif mutation != 'input':
            target = source if mutation == 'misplaced_content' else summary
            target.shapes.add_textbox(Inches(.5), Inches(.5), Inches(6), Inches(.5)).text = f'Total {total + (mutation == "wrong_value")}'
        if mutation == 'obstruction':
            cover = summary.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.5), Inches(.5), Inches(6), Inches(.5))
            cover.fill.solid()
            cover.fill.fore_color.rgb = RGBColor(0, 0, 0)
            cover.line.fill.background()
        deck.core_properties.created = deck.core_properties.modified = datetime(2026, 10, 16)
        deck.save(path)
    _freeze_zip(path)


def build_calibration_cases(root: str | Path, truth_manifest: str | Path) -> list[CalibrationCase]:
    root, truth_manifest = checked_path(root), Path(truth_manifest)
    if root.exists() and any(root.iterdir()):
        raise FileExistsError(root)
    authored = json.loads(truth_manifest.read_text(encoding='utf-8'))
    root.mkdir(parents=True, exist_ok=True)
    (root / 'inputs').mkdir()
    (root / 'artifacts').mkdir()
    records = []
    for entry in authored['cases']:
        domain, context = entry['domain'], entry['recipe']['context']
        suffix = {'excel': 'xlsx', 'word': 'docx', 'powerpoint': 'pptx'}[domain]
        input_name = f'inputs/{domain}-{context}.{suffix}'
        source = _contained(root, input_name)
        if not source.exists():
            _write_office(source, domain, context, 'input')
        artifact_name = f'artifacts/{entry["case_id"]}.{suffix}'
        artifact = _contained(root, artifact_name)
        _write_office(artifact, domain, context, entry['recipe']['mutation'])
        metadata = {k: v for k, v in entry.items() if k not in {'recipe', 'xml_claims'}}
        record = CalibrationCase(**metadata, input_path=input_name, artifact_path=artifact_name,
            artifact_hash=sha256_file(artifact), input_hash=sha256_file(source),
            truth_source_hash=sha256_file(truth_manifest), seed=authored['seed'])
        records.append(record)
    payload = {'schema_version': 'calibration-corpus-v1', 'source_kind': 'project-generated',
               'purpose': 'development_calibration_only', 'human_agreement_measured': False,
               'seed': authored['seed'], 'cases': [c.model_dump(mode='json') for c in records]}
    (root / 'oracle-truth.json').write_bytes(truth_manifest.read_bytes())
    (root / 'case_manifest.json').write_text(json.dumps(payload, indent=2) + '\n', encoding='utf-8')
    (root / 'oracle-provenance.json').write_text(json.dumps({
        'truth_source': str(truth_manifest), 'truth_source_hash': sha256_file(truth_manifest),
        'construction_source_hash': sha256_file(Path(__file__)),
        'label_method': 'authored literal truth; independent raw-XML integrity checks',
        'source_families': sorted({c.source_family for c in records}),
        'limitation': 'Three project-generated source families; 90 cases are not 90 independent families.'
    }, indent=2) + '\n', encoding='utf-8')
    return records


def load_calibration_cases(manifest: str | Path) -> list[CalibrationCase]:
    manifest = Path(manifest)
    payload = json.loads(manifest.read_text(encoding='utf-8'))
    cases = [CalibrationCase.model_validate(c) for c in payload['cases']]
    truth_path = _contained(manifest.parent, 'oracle-truth.json')
    truth_hash = sha256_file(truth_path)
    truth = json.loads(truth_path.read_text(encoding='utf-8'))
    authored = {c['case_id']: c for c in truth['cases']}
    if len({c.case_id for c in cases}) != len(cases):
        raise ValueError('duplicate calibration case ID')
    if set(authored) != {c.case_id for c in cases}:
        raise ValueError('calibration truth case set mismatch')
    for case in cases:
        expected = {k: v for k, v in authored[case.case_id].items() if k not in {'recipe', 'xml_claims'}}
        if case.truth_source_hash != truth_hash or case.seed != truth['seed'] or any(case.model_dump()[k] != v for k, v in expected.items()):
            raise ValueError('calibration truth hash or label mismatch')
        for name, expected in [(case.artifact_path, case.artifact_hash), (case.input_path, case.input_hash)]:
            if sha256_file(ordinary_file(_contained(manifest.parent, name))) != expected:
                raise ValueError('calibration file hash mismatch')
    return cases
