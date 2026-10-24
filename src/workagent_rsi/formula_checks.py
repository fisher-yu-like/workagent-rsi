"""Exact syntax and independent value/variant checks have distinct evidence."""
from pathlib import Path
import json,math,shutil
from uuid import uuid4

from .assessment_contracts import Location
from .check_registry import validate_location
from .hashing import sha256_file
from .view_safety import checked_path


def validate_formula_options(requirement):
    if requirement.options.get('formula_policy','exact') not in {'exact','equivalent'}:raise ValueError('unknown formula policy')
    if requirement.options.get('formula_policy')!='equivalent':return
    if requirement.check!='excel.formula':raise ValueError('only the legacy formula check accepts explicit equivalent policy; exact/present/value semantics remain distinct')
    oracle=requirement.options.get('oracle_value');variants=requirement.options.get('input_variants')
    if type(oracle) not in {int,float} or not math.isfinite(oracle):raise ValueError('equivalent formula requires independent numeric oracle_value')
    if not isinstance(variants,list) or not 1<=len(variants)<=20:raise ValueError('equivalent formula requires 1 to 20 independent input variants')
    names=[]
    for variant in variants:
        if not isinstance(variant,dict) or set(variant)!={'name','inputs','expected_value'}:raise ValueError('invalid formula variant schema')
        if not isinstance(variant['name'],str) or not variant['name']:raise ValueError('variant name required')
        names.append(variant['name'])
        if type(variant['expected_value']) not in {int,float} or not math.isfinite(variant['expected_value']):raise ValueError('finite independent variant oracle required')
        if not isinstance(variant['inputs'],dict) or not 1<=len(variant['inputs'])<=1000:raise ValueError('variant needs 1 to 1000 explicit changed inputs')
        for key,value in variant['inputs'].items():
            if not isinstance(key,str) or '!' not in key:raise ValueError('variant input needs sheet!cell')
            sheet,cell=key.rsplit('!',1);validate_location(Location(artifact=requirement.location.artifact,sheet=sheet,cell=cell),'excel',fact=True)
            if sheet==requirement.location.sheet and cell.upper().replace('$','')==requirement.location.cell.upper().replace('$',''):raise ValueError('variant cannot overwrite the answer formula')
            if value is not None and (type(value) not in {int,float} or not math.isfinite(value)):raise ValueError('variant inputs must be finite numbers or null')
    if len(set(names))!=len(names):raise ValueError('duplicate variant names')


def inspect_formula(requirement,workbook,source,engine):
    loc=requirement.location;value=workbook[loc.sheet][loc.cell].value
    formula=isinstance(value,str) and value.startswith('=')
    if requirement.check=='excel.formula_present':
        passed=formula is requirement.expected
        return ('PASS' if passed else 'FAIL'),(1 if passed else 0),{'formula_present':formula,'formula_text':value},loc
    equivalent=requirement.options.get('formula_policy')=='equivalent'
    if not equivalent and requirement.check in {'excel.formula','excel.formula_exact'}:
        passed=value==requirement.expected
        return ('PASS' if passed else 'FAIL'),(1 if passed else 0),value,loc
    if equivalent and not formula:return 'FAIL',0,{'reason':'required formula replaced by literal','actual':value},loc
    if engine is None:return 'UNAVAILABLE',None,'qualified isolated recalculation engine is not configured',loc
    root=checked_path(engine.output_root or Path(source).parent)/('formula-'+uuid4().hex);root.mkdir(parents=True,exist_ok=False)
    cases=[('base',Path(source),requirement.options['oracle_value'] if equivalent else requirement.expected)]
    if equivalent:
        from openpyxl import load_workbook
        for index,variant in enumerate(requirement.options['input_variants'],1):
            copied=root/f'variant-{index}.xlsx';shutil.copyfile(source,copied);book=load_workbook(copied,data_only=False)
            try:
                for key,number in variant['inputs'].items():
                    sheet,cell=key.rsplit('!',1)
                    if sheet not in book.sheetnames:raise ValueError('variant worksheet absent')
                    book[sheet][cell]=number
                book.save(copied)
            finally:book.close()
            cases.append((variant['name'],copied,variant['expected_value']))
    observations=[];flags=[]
    for index,(name,path,expected) in enumerate(cases):
        result=engine.recalculate(path,root/f'result-{index}.xlsx')
        observations.append({'case':name,'source_sha256':sha256_file(path),'expected':expected,'recalculation':result.model_dump(mode='json')})
        if result.status!='available' or result.qualification!='isolated_run' or result.source_sha256!=sha256_file(path):
            return 'UNAVAILABLE',None,observations,loc
        key=loc.sheet+'!'+loc.cell
        if key not in result.values:return 'UNAVAILABLE',None,{'reason':'recalculated target value missing','cases':observations},loc
        actual=result.values[key]
        ok=(type(actual) in {int,float} and type(expected) in {int,float} and math.isfinite(actual) and abs(actual-expected)<=requirement.tolerance) if type(expected) in {int,float} else actual==expected
        flags.append(ok)
    return ('PASS' if all(flags) else 'FAIL'),(1 if all(flags) else 0),observations,loc
