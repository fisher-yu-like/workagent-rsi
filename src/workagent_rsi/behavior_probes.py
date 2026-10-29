"""Independent behavioral oracles for authored engineering probes, not a benchmark."""
from copy import deepcopy
import json
from pathlib import Path

from .contracts import TaskSpec
from .hashing import canonical_json_hash,sha256_file
from .repair_loop import write_record
from .task_observations import TaskObservation
from .view_safety import ordinary_file,checked_path,relative_reference


def _owned(root,value):
    path=ordinary_file(Path(value))
    if not path.resolve().is_relative_to(root.resolve()):raise ValueError('probe evidence escapes its owned run')
    return path


def _oracle(path,domain,expected):
    # The expected cells/text/style/counts are authored independently of the
    # generation algorithm. No aggregation algorithm is reused as the oracle.
    if domain=='excel':
        from openpyxl import load_workbook
        book=load_workbook(path,data_only=False)
        cells={key:book[sheet][cell].value for key in expected.get('cells',{}) for sheet,cell in [key.rsplit('!',1)]}
        signature={'cells':cells,'sheet_names':book.sheetnames}
        passed=cells==expected.get('cells',{}) and book.sheetnames==expected['sheet_names'];book.close()
        return passed,signature
    if domain=='word':
        from docx import Document
        doc=Document(path)
        signature={'paragraphs':[p.text for p in doc.paragraphs],
                   'styles':[p.style.name for p in doc.paragraphs],
                   'tables':[[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]}
        return all(signature[k]==v for k,v in expected.items()),signature
    from pptx import Presentation
    deck=Presentation(path)
    signature={'slide_count':len(deck.slides),'slide_texts':[
        [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text] for slide in deck.slides]}
    return all(signature[k]==v for k,v in expected.items()),signature


def run_behavior_probes(skill,probe_manifest,executor) -> dict:
    manifest=json.loads(Path(probe_manifest).read_text(encoding='utf-8')) if isinstance(probe_manifest,(str,Path)) else deepcopy(probe_manifest)
    cases=manifest.get('cases',[]);required=manifest.get('required_groups',[])
    if not cases or not required or len({c['case_id'] for c in cases})!=len(cases):raise ValueError('unique cases and required groups are mandatory')
    root=checked_path(Path(manifest['output_root']))
    if root.exists() and any(root.iterdir()):raise FileExistsError(root)
    root.mkdir(parents=True,exist_ok=True)
    report={'schema_version':'input-behavior-v1','qualification_scope':'authored_local_library',
        'model_calls':0 if executor is None else getattr(executor,'model_calls',None),
        'live_model_qualified':False,'sandbox_qualified':False,'rows':[],
        'skill_hash':canonical_json_hash(skill),'manifest_hash':canonical_json_hash(manifest)}
    if executor is None:
        report.update(status='unavailable',reason='no executor configured');write_record(root/'behavior-probe-report.json',report);return report
    report['qualification_scope']=getattr(executor,'evidence_kind','injected_executor')
    trace_index=[]
    for index,case in enumerate(cases):
        task=TaskSpec.model_validate(case['task']);attempt=root/f'case-{index:03d}';attempt.mkdir()
        inputs={str(p):sha256_file(ordinary_file(Path(p))) for p in task.input_files}
        row={'case_id':case['case_id'],'group':case['group'],'domain':task.domain,'input_hashes':inputs,
             'task_hash':canonical_json_hash(task.model_dump(mode='json')),'requested_changes':task.instruction}
        row['public_request_hash']=canonical_json_hash(case['public_request'])
        try:
            # Oracle fields and the manifest stay in this read-only verifier.
            public={'task':task.model_dump(mode='json',exclude={'expected_constraints','hidden_test'}),
                    'request':deepcopy(case['public_request'])}
            result=executor(deepcopy(skill),public,attempt)
            path=_owned(attempt,result['artifact']);observation_path=_owned(attempt,result['observation'])
            observation=TaskObservation.model_validate_json(observation_path.read_text(encoding='utf-8'))
            mapping={i['copied']:i['copy_sha256'] for i in result['input_manifest']['files']}
            if list(mapping.values())!=list(inputs.values()):raise ValueError('staged input hashes do not match original inputs')
            input_ok=(observation.status=='observed' and observation.input_hashes==mapping)
            records=observation_path.parent
            for item in observation.inspection_evidence:
                ref=relative_reference(item['ref']);evidence=_owned(attempt,records/Path(*ref.parts))
                if sha256_file(evidence)!=item['sha256']:raise ValueError('frozen read evidence changed')
                read=json.loads(evidence.read_text(encoding='utf-8'));name=read['input_path']
                if (mapping.get(name)!=read['input_sha256'] or read['execution']['status']!='succeeded'
                        or observation.inspected_structure.get(name)!=read['structure']
                        or item['script_sha256']!=read['script_sha256']):
                    raise ValueError('observation does not match the actual input read')
            if task.input_files and not observation.inspection_evidence:input_ok=False
            log=_owned(attempt,result['tool_log'])
            events=[json.loads(line) for line in log.read_text(encoding='utf-8').splitlines()]
            executed=[e['result'] for e in events if e.get('name')=='run_python' and e.get('result',{}).get('ok') is True]
            if not executed:raise ValueError('missing generation execution trace')
            generation=executed[-1];ref=relative_reference(generation['script_snapshot_ref'])
            script=_owned(attempt,records/Path(*ref.parts))
            if sha256_file(script)!=generation['script_sha256']:raise ValueError('generation source changed')
            traces=[]
            for reference in generation['execution']['trace_refs']:
                ref=relative_reference(reference);trace=_owned(attempt,Path(result['workspace'])/Path(*ref.parts));traces.append(str(trace))
            if not traces:raise ValueError('missing actual process trace')
            oracle,signature=_oracle(path,task.domain,case['expected'])
            if any(sha256_file(Path(p))!=digest for p,digest in inputs.items()):raise ValueError('source input changed')
            row.update(oracle_passed=oracle,input_evidence_passed=input_ok,observation_status=observation.status,
                output_signature=signature,artifact_sha256=sha256_file(path),generation_script_sha256=generation['script_sha256'],
                passed=oracle and input_ok,observation_sha256=sha256_file(observation_path))
            trace_index.append({'case_id':case['case_id'],'artifact':str(path),'observation':str(observation_path),
                'tool_log':str(log),'script':str(script),'execution_traces':traces})
        except Exception as exc:
            row.update(passed=False,oracle_passed=False,input_evidence_passed=False,observation_status='not_observed',error=str(exc))
        report['rows'].append(row)
    by_id={r['case_id']:r for r in report['rows']};groups={}
    for group in required:
        members=[r for r in report['rows'] if r['group']==group]
        changed=(len(members)>=2 and (list(members[0]['input_hashes'].values()),members[0]['public_request_hash'],members[0]['requested_changes'])
                 !=(list(members[1]['input_hashes'].values()),members[1]['public_request_hash'],members[1]['requested_changes']))
        variation=changed and members[0].get('output_signature')!=members[1].get('output_signature')
        groups[group]={'case_count':len(members),'input_or_request_change_produces_expected_change':variation,
                       'passed':variation and all(r['passed'] for r in members)}
    id_controls=[(by_id[c['case_id']],by_id[c['id_equivalent_to']]) for c in cases if c.get('id_equivalent_to')]
    id_ok=bool(id_controls) and all(a['passed'] and b['passed'] and a['output_signature']==b['output_signature'] for a,b in id_controls)
    code_ok=(len({r.get('generation_script_sha256') for r in report['rows']})==1 and all(r['passed'] for r in report['rows']))
    # Group-only unit probes can omit ID controls; full qualification must include them.
    complete_groups=set(required)=={'excel-sheet','excel-values-rows','excel-columns','word-target','word-paragraphs','word-table','ppt-count','ppt-target','ppt-append'}
    report.update(group_results=groups,id_invariance_passed=id_ok,same_code_dynamic_input_passed=code_ok,
        coverage_complete=complete_groups,
        status='passed' if all(g['passed'] for g in groups.values()) and (id_ok or not complete_groups) else 'failed')
    index=root/'generation-trace-index.json';write_record(index,{'rows':trace_index});report['trace_index']=str(index)
    write_record(root/'behavior-probe-report.json',report)
    return report
