from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path
import re
from typing import Any

from .contracts import TaskSpec


class MockWorkAgentAdapter:
    """Deterministic adapter used until a real WorkAgent provider is configured."""

    def execute(self, task: TaskSpec, skill_id: str) -> Iterable[dict[str, Any]]:
        yield {"kind": "started", "skill_id": skill_id, "task_id": task.task_id}
        if task.instruction.strip().lower() == "fail":
            yield {"kind": "failure", "message": "controlled failure", "retryable": False}
            return
        if task.instruction.strip().lower() == "timeout":
            yield {"kind": "failure", "message": "controlled timeout", "retryable": True}
            return
        yield {"kind": "artifact", "content": task.instruction, "media_type": "text/plain"}


class LocalOfficeAdapter:
    """Create genuine Office artifacts for local provider qualification."""

    MEDIA_TYPES = {
        "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "word": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "powerpoint": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }
    WINDOWS_RESERVED_NAMES = {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        "CLOCK$",
        *(f"COM{index}" for index in range(1, 10)),
        *(f"LPT{index}" for index in range(1, 10)),
    }

    def __init__(self, output_root: str | Path) -> None:
        self.output_root = Path(output_root)

    def execute(self, task: TaskSpec, skill_id: str) -> Iterable[dict[str, Any]]:
        yield {"kind": "started", "skill_id": skill_id, "task_id": task.task_id}
        domain = task.domain.lower()
        if domain not in self.MEDIA_TYPES:
            yield {"kind": "failure", "message": f"unsupported Office domain: {task.domain}", "retryable": False}
            return
        task_id_root = task.task_id.split(".", 1)[0].upper()
        if (
            not re.fullmatch(r"[A-Za-z0-9](?:[A-Za-z0-9._-]*[A-Za-z0-9])?", task.task_id)
            or ".." in task.task_id
            or task_id_root in self.WINDOWS_RESERVED_NAMES
        ):
            yield {
                "kind": "failure",
                "message": f"unsafe task_id for artifact path: {task.task_id}",
                "retryable": False,
            }
            return
        marker = str(task.expected_constraints.get("required_text", task.task_id))
        self.output_root.mkdir(parents=True, exist_ok=True)
        if domain == "excel":
            from openpyxl import Workbook

            path = self.output_root / f"{task.task_id}.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Artifact"
            sheet["A1"] = marker
            sheet["A2"] = task.instruction
            workbook.save(path)
        elif domain == "word":
            from docx import Document

            path = self.output_root / f"{task.task_id}.docx"
            document = Document()
            document.add_heading(task.instruction, level=1)
            document.add_paragraph(marker)
            document.save(path)
        else:
            from pptx import Presentation

            path = self.output_root / f"{task.task_id}.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = task.instruction
            slide.placeholders[1].text = marker
            presentation.save(path)
        yield {"kind": "artifact", "artifact_path": str(path), "media_type": self.MEDIA_TYPES[domain]}
