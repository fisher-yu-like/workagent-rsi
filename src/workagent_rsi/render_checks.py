"""Deterministic geometry checks for Office presentation artifacts.

The project does not assume that a GUI renderer is installed.  Geometry is
therefore checked from the OOXML package with ``python-pptx`` and the visual
renderer channel is reported separately as unavailable when it cannot be
probed.  An unavailable renderer never becomes a passing score.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

from pydantic import BaseModel, ConfigDict, Field


class GeometryReport(BaseModel):
    """Structured, read-only result of a PowerPoint geometry inspection."""

    model_config = ConfigDict(extra="forbid")

    path: str
    status: str
    passed: bool
    slide_count: int = 0
    overlap_count: int = 0
    overflow_count: int = 0
    empty_slide_count: int = 0
    warnings: list[str] = Field(default_factory=list)
    failures: list[str] = Field(default_factory=list)
    evidence: list[str] = Field(default_factory=list)
    channel_status: dict[str, str] = Field(default_factory=dict)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _intersects(left_a: int, top_a: int, right_a: int, bottom_a: int, left_b: int, top_b: int, right_b: int, bottom_b: int) -> bool:
    """Return true only for a positive-area intersection."""

    return max(left_a, left_b) < min(right_a, right_b) and max(top_a, top_b) < min(bottom_a, bottom_b)


def check_powerpoint_geometry(path: str | Path) -> GeometryReport:
    """Check slide bounds, shape overlap and empty slides without rendering.

    ``status`` describes whether the package could be inspected.  The
    ``channel_status`` map makes the distinction between deterministic
    geometry and an unavailable pixel renderer explicit.
    """

    source = Path(path)
    base = {
        "path": str(source),
        "channel_status": {"geometry": "unavailable", "rendering": "unavailable"},
    }
    if not source.exists() or not source.is_file():
        return GeometryReport(
            **base,
            status="failed",
            passed=False,
            failures=[f"PowerPoint file does not exist: {source}"],
        )

    try:
        from pptx import Presentation

        presentation = Presentation(str(source))
    except Exception as exc:
        return GeometryReport(
            **base,
            status="failed",
            passed=False,
            failures=[f"PowerPoint geometry inspection failed: {exc}"],
            evidence=[f"artifact_sha256:{_sha256(source)}"],
        )

    slide_width = int(presentation.slide_width)
    slide_height = int(presentation.slide_height)
    overlap_count = 0
    overflow_count = 0
    empty_slide_count = 0
    failures: list[str] = []
    evidence = [f"artifact_sha256:{_sha256(source)}", f"slide_count:{len(presentation.slides)}"]

    for slide_number, slide in enumerate(presentation.slides, 1):
        shapes = [shape for shape in slide.shapes if int(getattr(shape, "width", 0)) > 0 and int(getattr(shape, "height", 0)) > 0]
        if not shapes:
            empty_slide_count += 1
            failures.append(f"slide {slide_number} is empty")
            continue
        rectangles: list[tuple[int, int, int, int]] = []
        for shape_number, shape in enumerate(shapes, 1):
            left = int(shape.left)
            top = int(shape.top)
            right = left + int(shape.width)
            bottom = top + int(shape.height)
            rectangles.append((left, top, right, bottom))
            if left < 0 or top < 0 or right > slide_width or bottom > slide_height:
                overflow_count += 1
                failures.append(f"slide {slide_number} shape {shape_number} exceeds slide bounds")
        for first_index, first in enumerate(rectangles):
            for second in rectangles[first_index + 1 :]:
                if _intersects(*first, *second):
                    overlap_count += 1
        if len(rectangles) > 1:
            evidence.append(f"slide_{slide_number}_shape_count:{len(rectangles)}")

    if overlap_count:
        failures.append(f"{overlap_count} shape overlap(s) detected")
    if overflow_count:
        failures.append(f"{overflow_count} shape overflow(s) detected")

    warnings = [
        "pixel renderer unavailable; geometry results are package-level checks only",
    ]
    return GeometryReport(
        **{**base, "channel_status": {"geometry": "available", "rendering": "unavailable"}},
        status="available",
        passed=not failures,
        slide_count=len(presentation.slides),
        overlap_count=overlap_count,
        overflow_count=overflow_count,
        empty_slide_count=empty_slide_count,
        warnings=warnings,
        failures=failures,
        evidence=evidence,
    )


__all__ = ["GeometryReport", "check_powerpoint_geometry"]
