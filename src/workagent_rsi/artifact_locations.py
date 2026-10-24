"""Read-only, bounded Office text scopes with explicit observed coordinates."""
from __future__ import annotations

import re

from .assessment_contracts import Location
from .check_registry import validate_location


class AmbiguousScope(ValueError):
    pass


class ScopeBudgetExceeded(ValueError):
    pass


class ScopeNotFound(ValueError):
    pass


def resolve_shape(deck, location):
    if location.slide>len(deck.slides):
        return None
    shapes=deck.slides[location.slide-1].shapes
    for group_id in location.group_path or []:
        matches=[s for s in shapes if s.shape_id==group_id and hasattr(s,'shapes')]
        if len(matches)>1:
            raise AmbiguousScope('duplicate group identity')
        if not matches:
            return None
        shapes=matches[0].shapes
    return find_shape(shapes,location.object_id)


def heading_level(paragraph):
    """Stable built-in ID, paragraph/style outline level and inherited styles."""
    from docx.oxml.ns import qn
    properties = paragraph._p.pPr
    if properties is not None:
        outline = properties.find(qn('w:outlineLvl'))
        if outline is not None:
            value = int(outline.get(qn('w:val')))
            return value+1 if 0<=value<=8 else None
    style, visited = paragraph.style, set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        match = re.fullmatch(r'Heading([1-9])',style.style_id,re.I)
        if match:
            return int(match.group(1))
        properties = style.element.pPr
        if properties is not None:
            outline=properties.find(qn('w:outlineLvl'))
            if outline is not None:
                value=int(outline.get(qn('w:val')))
                return value+1 if 0<=value<=8 else None
        style=style.base_style
    return None


def find_shape(shapes, object_id):
    found=[]
    for shape in shapes:
        if shape.shape_id==object_id:
            found.append(shape)
        if hasattr(shape,'shapes'):
            child=find_shape(shape.shapes,object_id)
            if child is not None:
                found.append(child)
    if len(found)>1:
        raise AmbiguousScope('duplicate object identity inside slide')
    return found[0] if found else None


def resolve_text_scope(obj, location: Location, options: dict | None = None) -> list[dict]:
    """Return complete scoped text items or an explicit gap, never partial success.

    Missing valid positions yield an empty scope. Budgets are implementation
    bounds (50k items / 2m characters), not candidate-selected truncation knobs.
    """
    loc=Location.model_validate(location)
    domain='excel' if hasattr(obj,'worksheets') else 'word' if hasattr(obj,'paragraphs') else 'powerpoint'
    validate_location(loc,domain)
    if loc.region is not None:
        raise ValueError('region selector requires a supported visual channel')
    items, used = [], 0

    def add(value, **position):
        nonlocal used
        text='' if value is None else str(value)
        used+=len(text)
        if len(items)>=50000 or used>2000000:
            raise ScopeBudgetExceeded('text scope exceeds fixed item/character budget; no truncated result admitted')
        observed=Location(artifact=loc.artifact,**position)
        items.append({'text':text,'location':observed.model_dump(exclude_none=True)})

    if domain=='excel':
        from openpyxl.utils.cell import range_boundaries
        sheets=[obj[loc.sheet]] if loc.sheet in obj.sheetnames else [] if loc.sheet else list(obj.worksheets)
        for sheet in sheets:
            a,b,c,d=range_boundaries(loc.cell) if loc.cell else (1,1,sheet.max_column,sheet.max_row)
            if (c-a+1)*(d-b+1)>50000:
                raise ScopeBudgetExceeded('Excel text scope exceeds fixed 50000-cell budget')
            for row in sheet.iter_rows(min_col=a,min_row=b,max_col=c,max_row=d):
                for cell in row:
                    if cell.value is not None:
                        add(cell.value,sheet=sheet.title,cell=cell.coordinate)
    elif domain=='word':
        from docx.oxml.ns import qn
        paragraphs=list(obj.paragraphs)
        tables=list(obj.tables)
        paragraph_ids={p._p:(i,p) for i,p in enumerate(paragraphs,1)}
        table_ids={t._tbl:(i,t) for i,t in enumerate(tables,1)}
        body=[]
        for node in obj.element.body:
            if node in paragraph_ids:
                i,p=paragraph_ids[node];body.append(('paragraph',i,p))
            elif node in table_ids:
                i,t=table_ids[node];body.append(('table',i,t))
        if loc.heading is not None:
            headings=[i for i,(kind,_,p) in enumerate(body) if kind=='paragraph'
                and p.text==loc.heading and heading_level(p) is not None]
            if len(headings)>1:
                raise AmbiguousScope('heading selector matches multiple sections; explicit disambiguation is required')
            if not headings:
                return []
            start=headings[0];level=heading_level(body[start][2]);end=len(body)
            for i in range(start+1,len(body)):
                kind,_,p=body[i]
                if kind=='paragraph' and heading_level(p) is not None and heading_level(p)<=level:
                    end=i;break
            body=body[start:end]
        for kind,index,element in body:
            if kind=='paragraph':
                if loc.table is not None or loc.paragraph is not None and loc.paragraph!=index:
                    continue
                add(element.text,paragraph=index,heading=loc.heading)
            elif loc.paragraph is None and (loc.table is None or loc.table==index):
                for i,row in enumerate(element.rows,1):
                    for j,cell in enumerate(row.cells,1):
                        if loc.row is None or (i==loc.row and j==loc.column):
                            add(cell.text,table=index,row=i,column=j,heading=loc.heading)
    else:
        indices=[loc.slide] if loc.slide is not None else range(1,len(obj.slides)+1)
        for slide_index in indices:
            if slide_index>len(obj.slides):
                continue
            shapes=obj.slides[slide_index-1].shapes
            ancestry=[]
            if loc.group_path is not None:
                for group_id in loc.group_path:
                    matches=[s for s in shapes if s.shape_id==group_id and hasattr(s,'shapes')]
                    if len(matches)>1:
                        raise AmbiguousScope('duplicate group identity')
                    if not matches:
                        shapes=[];break
                    shapes=matches[0].shapes;ancestry.append(group_id)

            def walk(current, parents, selected=False):
                for shape in current:
                    here=selected or loc.object_id is None or shape.shape_id==loc.object_id
                    if hasattr(shape,'shapes'):
                        walk(shape.shapes,parents+[shape.shape_id],here)
                    elif here:
                        if getattr(shape,'has_table',False):
                            for i,row in enumerate(shape.table.rows,1):
                                for j,cell in enumerate(row.cells,1):
                                    if loc.row is None or (i==loc.row and j==loc.column):
                                        add(cell.text,slide=slide_index,object_id=shape.shape_id,
                                            group_path=parents or None,row=i,column=j)
                        elif loc.row is None and getattr(shape,'has_text_frame',False):
                            add(shape.text,slide=slide_index,object_id=shape.shape_id,group_path=parents or None)
            if loc.object_id is not None:
                # Verify uniqueness before allowing recursive selection.
                find_shape(shapes,loc.object_id)
            walk(shapes,ancestry)
    return items
