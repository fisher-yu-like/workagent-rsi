"""Read-only shared observations and standalone JSON entry point for both skills."""

from __future__ import annotations

import argparse
from datetime import date, datetime, time as datetime_time
import json
import math
from pathlib import Path
from typing import Any
import zipfile

from pydantic import TypeAdapter

from .assessment_contracts import AcceptanceSpec, Assessment, CheckRecord, Evidence, Location, Requirement
from .hashing import canonical_json_hash, sha256_file
from .office_checks import EXPECTED_SUFFIXES, OOXML_MEMBERS
from .visibility_checks import inspect_visibility
from .artifact_locations import (resolve_text_scope, heading_level, resolve_shape,
    AmbiguousScope, ScopeBudgetExceeded, ScopeNotFound)


KNOWN = {"PASS", "PARTIAL", "FAIL"}
JSON_VALUE = TypeAdapter(Any)


def assessment_identity(config: dict | None = None) -> str:
    root = Path(__file__).resolve().parents[2]
    names = ("artifact_feedback.py", "rsi_feedback.py", "orchestrator.py", "cli.py", "contracts.py", "run_contracts.py", "check.py", "harness_artifact_evaluator.py", "harness.py", "run.py", "evidence_partition.py", "visibility_checks.py", "assessment_contracts.py", "preservation_checks.py", "formula_checks.py", "spreadsheet_recalc.py", "libreoffice_recalc_worker.py", "requirement_coverage.py", "check_registry.py", "artifact_locations.py", "native_presentation_checks.py", "artifact_assessment.py", "artifact_evaluator.py", "artifact_verifier.py", "sales_assessment.py", "office_checks.py", "artifact_renderer.py", "libreoffice_worker.py", "view_safety.py", "artifact_com.py", "office_capabilities.py", "judgement_provider.py", "semantic_assessment.py", "model_judgement.py")
    names = (*names, 'word_package_checks.py')
    files = {name: sha256_file(Path(__file__).with_name(name)) for name in names if Path(__file__).with_name(name).exists()}
    for skill in ("office-artifact-evaluator", "office-artifact-verifier"):
        for path in sorted((root / "skills" / skill).rglob("*")):
            if path.is_file() and "__pycache__" not in path.parts:
                files[path.relative_to(root).as_posix()] = sha256_file(path)
    return canonical_json_hash({"files": files, "configuration": config or {}, "mode": "artifact-assessment-v1"})


def _equal(actual: Any, expected: Any, tolerance: float) -> bool:
    if isinstance(actual, (date, datetime, datetime_time)) and isinstance(expected, str):
        return actual.isoformat() == expected
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return isinstance(actual, (int, float)) and not isinstance(actual, bool) and math.isfinite(actual) and abs(actual - expected) <= tolerance
    return actual == expected


class SharedAssessment:
    def __init__(self, *, visual=None, semantic=None, recalculation=None):
        self.visual, self.semantic, self.recalculation = visual, semantic, recalculation

    def identity(self) -> str:
        return assessment_identity({"visual": self.visual.identity_config() if self.visual else "disabled",
                                    "semantic": self.semantic.identity_config() if self.semantic else "disabled",
                                    'recalculation':self.recalculation.identity_config() if self.recalculation else 'disabled'})

    def inspect(self, spec: AcceptanceSpec, artifacts: dict[str, str | Path]) -> Assessment:
        spec = AcceptanceSpec.model_validate(spec.model_dump(mode="json"))
        from .requirement_coverage import admission_for_spec
        admission = admission_for_spec(spec)
        from .check_registry import validate_acceptance_spec
        preflight = validate_acceptance_spec(spec, {
            'renderer': self.visual is not None and getattr(self.visual,'renderer',None) is not None,
            'visual_judge': self.visual is not None,
            'semantic_judge': self.semantic is not None,
            'spreadsheet_recalc': self.recalculation is not None})
        gaps = {f.requirement_id:f for f in preflight.findings}
        paths = {key: Path(value) for key, value in artifacts.items() if key in spec.artifacts}
        hashes = {key: sha256_file(paths[key]) if key in paths and paths[key].is_file() else None for key in spec.artifacts}
        for raw_path, digest in spec.input_hashes.items():
            if not Path(raw_path).is_file() or sha256_file(raw_path) != digest:
                raise ValueError("frozen assessment input changed: " + raw_path)
        parsed = {}
        parse_errors = {}
        needed = {r.location.artifact for r in spec.requirements if r.applicable and r.requirement_id not in gaps}
        for r in spec.requirements:
            if r.applicable and r.requirement_id not in gaps and r.check=='facts.consistent':
                needed.update(p['artifact'] for p in r.options['locations'])
        for key, domain in spec.artifacts.items():
            if key not in needed:
                continue
            path = paths.get(key)
            if path is None or not path.is_file():
                parse_errors[key] = "artifact missing"
                continue
            try:
                if path.suffix.lower() != EXPECTED_SUFFIXES[domain]:
                    raise ValueError("unexpected artifact extension")
                with zipfile.ZipFile(path) as package:
                    if not set(OOXML_MEMBERS[domain]).issubset(package.namelist()):
                        raise ValueError("required OOXML members missing")
                if domain == "excel":
                    from openpyxl import load_workbook
                    parsed[key] = load_workbook(path, data_only=False)
                elif domain == "word":
                    from docx import Document
                    parsed[key] = Document(path)
                else:
                    from pptx import Presentation
                    parsed[key] = Presentation(path)
            except Exception as exc:
                parse_errors[key] = f"artifact cannot be parsed: {type(exc).__name__}: {exc}"
        checks, evidence, telemetry = [], [], {"visual": [], "semantic": [], 'check_preflight':preflight.model_dump(mode='json')}
        try:
            for r in spec.requirements:
                location = r.location
                observed_locations = []
                message = ""
                origin = None
                if not r.applicable:
                    status, completion, observed = "NOT_APPLICABLE", None, r.applicability_reason
                elif r.requirement_id in gaps:
                    finding = gaps[r.requirement_id]
                    status, completion, observed = 'UNAVAILABLE', None, finding.message
                    origin = finding.origin
                elif r.location.artifact in parse_errors:
                    status, completion, observed = "FAIL", 0, parse_errors[r.location.artifact]
                else:
                    try:
                        if r.check in {"visual", "semantic"}:
                            channel = self.visual if r.check == "visual" else self.semantic
                            if channel is None:
                                status, completion, observed = "UNAVAILABLE", None, r.check + " channel disabled"
                            else:
                                source = paths[r.location.artifact] if r.check == "visual" else parsed[r.location.artifact]
                                result = channel.check(r, source, hashes[r.location.artifact])
                                status, completion, observed = result["status"], result.get("completion"), result.get("observed")
                                if result.get("location"):
                                    location = Location.model_validate(result["location"])
                                    observed_locations = [location]
                                telemetry[r.check].append(result.get("telemetry", {}))
                        else:
                            status, completion, observed, location = self._check(r, parsed, paths)
                            if isinstance(location,list):
                                observed_locations = [Location.model_validate(p) for p in location]
                            elif status in KNOWN:
                                observed_locations = [location]
                    except (AmbiguousScope,ScopeBudgetExceeded) as exc:
                        status, completion, observed = 'NEEDS_REVIEW',None,str(exc)
                    except ScopeNotFound as exc:
                        status, completion, observed = 'FAIL',0,str(exc)
                    except ImportError as exc:
                        status, completion, observed = "UNAVAILABLE", None, str(exc)
                    except Exception as exc:
                        status, completion, observed = "ERROR", None, f"{type(exc).__name__}: {exc}"
                if status not in KNOWN:
                    message = str(observed)
                    if status != 'NOT_APPLICABLE' and origin is None:
                        origin = 'assessment_gap'
                elif status != 'PASS':
                    origin = 'artifact_defect'
                json_observed = JSON_VALUE.dump_python(observed, mode="json")
                item = Evidence(evidence_id=canonical_json_hash({"spec": spec.fingerprint(), "requirement": r.requirement_id, "hashes": hashes, "status": status, "observed": json_observed}),
                                requirement_id=r.requirement_id, location=observed_locations[0] if observed_locations else r.location, source=r.evidence_source,
                                observed=observed, artifact_sha256=hashes[r.location.artifact], artifact_hashes=hashes,
                                requested_location=r.location,observed_locations=observed_locations)
                evidence.append(item)
                checks.append(CheckRecord(requirement_id=r.requirement_id, status=status, completion=completion,
                                          location=r.location, observed=observed, expected=r.expected,
                                          evidence_refs=[item.evidence_id], message=message,failure_origin=origin,
                                          requested_location=r.location,observed_locations=observed_locations))
        finally:
            for obj in parsed.values():
                if hasattr(obj, "close"):
                    obj.close()
        # A concurrently changed file invalidates the snapshot, not the artifact's score.
        if any(sha256_file(paths[k]) != h for k, h in hashes.items() if h is not None):
            raise ValueError("artifact changed during inspection")
        return Assessment(spec=spec, spec_hash=spec.fingerprint(), evaluator_identity=self.identity(),
                          artifact_hashes=hashes, checks=checks, evidence=evidence, telemetry=telemetry,
                          spec_admission=admission)

    def _check(self, r, parsed, paths):
        obj, loc = parsed[r.location.artifact], r.location
        if r.check=='preservation':
            from .preservation_checks import compare_preservation
            rows=compare_preservation(r.options['source_path'],paths[loc.artifact],r.options['preservation_spec'])
            status='FAIL' if any(row.status=='FAIL' for row in rows) else 'NEEDS_REVIEW' if any(row.completion is None for row in rows) else 'PASS'
            return status,(0 if status=='FAIL' else 1 if status=='PASS' else None),[row.model_dump(mode='json') for row in rows],loc
        if r.check == "file.readable":
            actual = True
        elif r.check == "file.format":
            reverse_suffixes = {suffix: domain for domain, suffix in EXPECTED_SUFFIXES.items()}
            actual = reverse_suffixes.get(paths[loc.artifact].suffix.lower(), "unknown")
        elif r.check == "file.name":
            actual = paths[loc.artifact].name
        elif r.check == 'word.package':
            from .word_package_checks import inspect_word_rules
            status,completion,observed=inspect_word_rules(paths[loc.artifact],r.expected,r.options.get('source_path'))
            return status,completion,observed,loc
        elif r.check == "powerpoint.native_objects":
            from .native_presentation_checks import inspect_native_objects
            return inspect_native_objects(obj, loc, r.expected)
        elif r.check == "excel.sheets":
            actual = list(obj.sheetnames)
            return self._fraction(r, [name in actual for name in r.expected], actual, loc)
        elif r.check.startswith("excel."):
            if loc.sheet not in obj.sheetnames:
                return "FAIL", 0, "worksheet missing", loc
            sheet = obj[loc.sheet]
            if r.check in {'excel.formula','excel.formula_exact','excel.formula_present','excel.recalculated_value'}:
                from .formula_checks import inspect_formula
                return inspect_formula(r,obj,paths[loc.artifact],self.recalculation)
            if r.check in {"excel.cell", "excel.number_format"}:
                cell = sheet[loc.cell]
                actual = cell.number_format if r.check == "excel.number_format" else cell.value
            elif r.check == "excel.range":
                from openpyxl.utils.cell import range_boundaries
                a,b,c,d=range_boundaries(loc.cell)
                actual = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=b,max_row=d,min_col=a,max_col=c)]
                flags = [i < len(actual) and j < len(actual[i]) and _equal(actual[i][j], value, r.tolerance) for i, row in enumerate(r.expected) for j, value in enumerate(row)]
                return self._fraction(r, flags, actual, loc)
            elif r.check == "excel.cached_value":
                from openpyxl import load_workbook
                cached = load_workbook(paths[loc.artifact], data_only=True)
                try:
                    actual = cached[loc.sheet][loc.cell].value
                finally:
                    cached.close()
                if actual is None:
                    return "UNAVAILABLE", None, "no cached value; openpyxl does not recalculate formulas", loc
            elif r.check == "excel.recalculated_value":
                return "UNAVAILABLE", None, "no recalculation engine configured", loc
            elif r.check == "excel.chart":
                actual = []
                for chart in sheet._charts:
                    for series in chart.series:
                        values = getattr(getattr(series, "val", None), "numRef", None)
                        categories = getattr(series, "cat", None)
                        categories = getattr(categories, "numRef", None) or getattr(categories, "strRef", None)
                        actual.append({"values": getattr(values, "f", None), "categories": getattr(categories, "f", None)})
                return self._fraction(r, [item in actual for item in r.expected], actual, loc)
            else:
                return "UNAVAILABLE", None, "unsupported check: " + r.check, loc
        elif r.check == "text.contains":
            texts = resolve_text_scope(obj,loc,r.options)
            matched = [row for row in texts if r.expected in row['text']]
            if hasattr(obj, 'slides'):
                return self._visible_texts(obj, r, texts, [r.expected])
            actual = matched[0]['text'] if matched else 'required content not found in requested scope'
            return ('PASS' if matched else 'FAIL'),(1 if matched else 0),actual,[row['location'] for row in matched]
        elif r.check == "word.headings":
            scoped = resolve_text_scope(obj,loc,r.options)
            positions = {row['location'].get('paragraph') for row in scoped}
            actual = [p.text for i,p in enumerate(obj.paragraphs,1) if i in positions and heading_level(p) is not None]
            return self._fraction(r, [item in actual for item in r.expected], actual, loc)
        elif r.check == "word.heading_style":
            scoped=resolve_text_scope(obj,loc,r.options)
            if not scoped:
                raise ScopeNotFound('requested heading is missing')
            index=scoped[0]['location']['paragraph']
            paragraph=obj.paragraphs[index-1]
            actual=paragraph.style.name
            passed=_equal(actual,r.expected,r.tolerance) or paragraph.style.style_id==r.expected.replace(' ','')
            return ('PASS' if passed else 'FAIL'),(1 if passed else 0),actual,[scoped[0]['location']]
        elif r.check == "word.references":
            actual = [rel.target_ref for rel in obj.part.rels.values() if rel.reltype.endswith("/hyperlink") and rel.is_external]
            return self._fraction(r, [ref in actual for ref in r.expected], actual, loc)
        elif r.check in {"word.paragraph", "word.table_cell", "powerpoint.object_text"}:
            actual = self._fact(obj, loc)
        elif r.check == "powerpoint.slide_count":
            actual = len(obj.slides)
        elif r.check == "powerpoint.texts":
            scoped=resolve_text_scope(obj,loc,r.options)
            return self._visible_texts(obj, r, scoped, r.expected)
        elif r.check == "powerpoint.chart_values":
            shape = resolve_shape(obj,loc)
            actual = [list(series.values) for series in shape.chart.series] if shape is not None and shape.has_chart else None
        elif r.check in {"powerpoint.bounds", "powerpoint.overlap"}:
            if loc.slide>len(obj.slides):
                raise ScopeNotFound('requested slide is missing')
            findings = inspect_visibility(obj, loc.slide)
            if r.check == "powerpoint.bounds":
                actual = sorted({f.object_id for f in findings if f.kind == 'off_canvas' and f.certainty == 'certain'})
                if not actual and any(f.kind in {'rotation', 'group_transform', 'off_canvas'} for f in findings):
                    return 'NEEDS_REVIEW', None, [f.model_dump() for f in findings], loc
            else:
                actual = []
                risks = [f.model_dump() for f in findings if f.kind in {'possible_occlusion', 'rotation', 'group_transform'}]
                if risks:
                    return "NEEDS_REVIEW", None, {"findings": risks, "note": "geometry alone cannot establish visual occlusion"}, loc
        elif r.check == "facts.consistent":
            actual, flags = [], []
            for payload in r.options["locations"]:
                position = Location.model_validate(payload)
                value = self._fact(parsed[position.artifact], position) if position.artifact in parsed else "artifact missing or unreadable"
                actual.append({"location": position.model_dump(exclude_none=True), "value": value})
                equal = _equal(value, r.expected, r.tolerance)
                if not equal:
                    loc = position
                flags.append(equal)
            return self._fraction(r, flags, actual, loc)
        else:
            return "UNAVAILABLE", None, "unsupported check: " + r.check, loc
        passed = _equal(actual, r.expected, r.tolerance)
        return ("PASS" if passed else "FAIL"), (1 if passed else 0), actual, loc

    def _visible_texts(self, obj, r, rows, wanted):
        findings = {slide: inspect_visibility(obj, slide) for slide in {row['location']['slide'] for row in rows}}
        flags, matched, risks = [], [], []
        for phrase in wanted:
            candidates = [row for row in rows if phrase in row['text']]
            visible, uncertain = [], []
            for row in candidates:
                location = row['location']
                selected = [f for f in findings[location['slide']] if f.object_id == location['object_id']
                            or f.kind == 'group_transform']
                if any(f.kind == 'off_canvas' and f.certainty == 'certain' for f in selected):
                    continue
                if selected:
                    uncertain.append(row)
                    risks.extend(f.model_dump() for f in selected)
                else:
                    visible.append(row)
            flags.append(bool(visible or uncertain))
            matched.extend(row['location'] for row in visible+uncertain)
            if visible:
                # A clear occurrence satisfies this phrase even if a duplicate is risky.
                continue
            if uncertain:
                risks.append({'phrase': phrase, 'requires_review': True})
        if all(flags) and any(f.get('requires_review') for f in risks):
            return 'NEEDS_REVIEW', None, {'texts': [row['text'] for row in rows], 'visibility': risks}, matched
        return self._fraction(r, flags, [row['text'] for row in rows], matched)

    @staticmethod
    def _fact(obj, loc):
        if loc.sheet:
            if loc.sheet not in obj.sheetnames:
                raise ScopeNotFound('requested worksheet is missing')
            return obj[loc.sheet][loc.cell].value
        scoped=resolve_text_scope(obj,loc)
        if not scoped:
            raise ScopeNotFound('requested paragraph/table cell/slide object is missing')
        return '\n'.join(row['text'] for row in scoped)

    @staticmethod
    def _fraction(r, flags, actual, loc):
        if not flags:
            raise ValueError("empty expected collection; mark the requirement not applicable instead")
        fraction = sum(flags) / len(flags)
        return ("PASS" if fraction == 1 else "FAIL" if fraction == 0 else "PARTIAL"), fraction, actual, loc

    @staticmethod
    def _texts(obj):
        return [(row['text'],{k:v for k,v in row['location'].items() if k!='artifact'})
                for row in resolve_text_scope(obj,Location(artifact='output'))]


def spec_from_task(task, *, assessment_profile=None) -> AcceptanceSpec:
    """Default public task admission; legacy declared checks require opt-in."""
    if assessment_profile not in {None,'legacy','structural_only','task_acceptance','research'}:
        raise ValueError('unknown assessment profile')
    spec=_spec_from_task(task)
    profile=None if assessment_profile=='legacy' else assessment_profile or spec.evaluation_profile or 'task_acceptance'
    return AcceptanceSpec.model_validate({**spec.model_dump(mode='json'),'evaluation_profile':profile})


def _spec_from_task(task) -> AcceptanceSpec:
    """Adapt existing constraints explicitly; unsupported constraints stay visible."""
    constraints = task.expected_constraints
    if constraints.get("sales_summary"):
        from .sales_assessment import sales_spec
        return _bind_task_admission(task, sales_spec(task))
    if "acceptance_spec" in constraints:
        spec = AcceptanceSpec.model_validate(constraints["acceptance_spec"])
        if spec.task_id != task.task_id:
            raise ValueError("acceptance specification task identity mismatch")
        return _bind_task_admission(task, spec)
    requirements = [Requirement(requirement_id="file", description="Readable required artifact", check="file.readable", location=Location(artifact="output"), expected=True, critical=True, dimension="correctness", evidence_source="OOXML package and library parser")]
    from .check_registry import LEGACY_ADAPTERS
    adapters, records = LEGACY_ADAPTERS, []
    for key, expected in constraints.items():
        if key in {'evaluation_profile','requirement_inventory'}:
            records.append({'key':key,'status':'converted','check_ids':[], 'reason':'task admission configuration'})
            continue
        if key == "marker" and constraints.get("required_text") == expected:
            records.append({'key':key,'status':'alias','check_ids':['required_text:'], 'reason':'same obligation as required_text'})
            continue
        check = adapters.get(key, "unsupported." + key)
        mapping = key in {'required_cells','required_formulas','required_formula_values','required_heading_styles'}
        malformed = mapping and (not isinstance(expected,dict) or not expected or any(not isinstance(k,str) for k in expected))
        if malformed:
            check='unsupported.'+key
        entries = expected.items() if mapping and not malformed else [("", expected)]
        check_ids=[]
        for target, value in entries:
            position = {"artifact": "output"}
            if "!" in target:
                position["sheet"], position["cell"] = target.rsplit("!", 1)
                position["sheet"] = position["sheet"].strip("'")
            elif key == "required_heading_styles":
                position["heading"] = target
            elif target:
                position["cell"] = target
            if key == "output_format":
                value = str(value).lower()
            rid=f'{key}:{target}';check_ids.append(rid)
            requirements.append(Requirement(requirement_id=rid, description=key, check=check, location=Location(**position), expected=value, critical=True, dimension="correctness", evidence_source="TaskSpec.expected_constraints"))
        records.append({'key':key,'status':'unsupported' if check.startswith('unsupported.') else 'converted',
            'check_ids':check_ids,'reason':'malformed mapping or unregistered legacy key' if check.startswith('unsupported.') else 'registered adapter; check syntax is separately preflighted'})
    domain = task.domain.lower()
    from .requirement_coverage import task_context
    return AcceptanceSpec(task_id=task.task_id, version="legacy-constraints-v1", artifacts={"output": domain}, requirements=requirements, dimension_weights={"correctness": 1},
        evaluation_profile=constraints.get('evaluation_profile', 'task_acceptance' if not constraints or 'requirement_inventory' in constraints else None),
        task_context=task_context(task),requirement_inventory=constraints.get('requirement_inventory'),
        legacy_adapter_records=records,
        scope_note='Legacy declared checks only; task requirements are not assumed complete without reviewed admission.')


def _bind_task_admission(task, spec):
    from .requirement_coverage import task_context
    constraints = task.expected_constraints
    payload = spec.model_dump(mode='json')
    # The actual caller task owns context; a nested acceptance spec cannot replace it.
    payload['task_context'] = task_context(task)
    existing = {row['key'] for row in payload['legacy_adapter_records']}
    for key,value in constraints.items():
        if key in existing:
            continue
        if key in {'acceptance_spec','sales_summary','evaluation_profile','requirement_inventory'}:
            ids = [r.requirement_id for r in spec.requirements] if key in {'acceptance_spec','sales_summary'} else []
            payload['legacy_adapter_records'].append({'key':key,'status':'converted','check_ids':ids,
                'reason':'explicit specification/template or admission configuration'})
        else:
            rid='unsupported-wrapper:'+key
            payload['requirements'].append(Requirement(requirement_id=rid,description=key,
                check='unsupported.'+key,expected=value,location=Location(artifact=next(iter(spec.artifacts))),
                critical=True,dimension=next(iter(spec.dimension_weights)),evidence_source='TaskSpec.expected_constraints').model_dump(mode='json'))
            payload['legacy_adapter_records'].append({'key':key,'status':'unsupported','check_ids':[rid],
                'reason':'extra constraint is not mapped by this explicit specification/template'})
    if 'evaluation_profile' in constraints:
        payload['evaluation_profile'] = constraints['evaluation_profile']
    if 'requirement_inventory' in constraints:
        payload['requirement_inventory'] = constraints['requirement_inventory']
        if payload['evaluation_profile'] is None:
            payload['evaluation_profile'] = 'task_acceptance'
    return AcceptanceSpec.model_validate(payload)


def assess(spec, artifacts, *, visual=None, semantic=None, recalculation=None):
    from .artifact_evaluator import ArtifactEvaluator
    from .artifact_verifier import ArtifactVerifier
    shared = SharedAssessment(visual=visual, semantic=semantic,recalculation=recalculation).inspect(spec, artifacts)
    return shared, ArtifactEvaluator().evaluate(shared), ArtifactVerifier().verify(shared)


def save_reports(output, shared, score, issues):
    root = Path(output)
    root.mkdir(parents=True, exist_ok=True)
    for name, report in (("assessment", shared), ("score", score), ("issues", issues)):
        target = root / f"{name}.json"
        with target.open("x", encoding="utf-8") as stream:
            stream.write(report.model_dump_json(indent=2))
    text = f"# Artifact assessment: {score.task_id}\n\nDeclared-check acceptance: {score.acceptance_status}\n\nTask acceptance: {score.task_acceptance}\n\nTask quality: {score.task_score}\n\nDeclared-check quality: {score.check_score}\n\nCheck completion: {score.coverage.completed}/{score.coverage.applicable}\n\nRequirement admission: {score.spec_admission.status if score.spec_admission else 'unavailable'}\n\nRequirement coverage: {score.requirement_coverage}\n\n{score.scope_note}\n\n"
    for issue in issues.issues:
        text += f"- {issue.requirement_id}: {issue.kind}; {issue.location.model_dump(exclude_none=True)}; observed={issue.observed!r}; expected={issue.expected!r}. {issue.repair_hint}\n"
    (root / "report.md").write_text(text, encoding="utf-8")


def main(argv=None, *, entry="both"):
    parser = argparse.ArgumentParser(description="Assess real Office artifacts with a frozen specification")
    parser.add_argument("request", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--model-config", type=Path, help="Trusted JSON model configuration (codex-cli, api or disabled)")
    args = parser.parse_args(argv)
    request = json.loads(args.request.read_text(encoding="utf-8-sig"))
    spec = AcceptanceSpec.model_validate(request["spec"])
    artifacts = {key: str((args.request.resolve().parent / value).resolve()) for key, value in request["artifacts"].items()}
    if args.model_config:
        from .model_judgement import configured_assessment
        from .artifact_evaluator import ArtifactEvaluator
        from .artifact_verifier import ArtifactVerifier
        config = json.loads(args.model_config.read_text(encoding="utf-8-sig"))
        engine = configured_assessment(config, args.output / "model-evidence")
        shared = engine.inspect(spec, artifacts)
        score, issues = ArtifactEvaluator().evaluate(shared), ArtifactVerifier().verify(shared)
    else:
        shared, score, issues = assess(spec, artifacts)
    save_reports(args.output, shared, score, issues)
    print((issues if entry == "verifier" else score).model_dump_json(indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
