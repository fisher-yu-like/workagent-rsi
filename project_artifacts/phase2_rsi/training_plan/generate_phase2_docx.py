"""Generate the Phase 2 RSI architecture and training plan report."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[3]
OUTPUT = ROOT / "project_artifacts" / "phase2_rsi" / "reports" / "RSI整体框架与训练计划.docx"


def font(run, size=10.5, bold=False, color="000000"):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(text), 15 if level == 1 else 12, True)


def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.2
    font(p.add_run(text))


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(text), 10)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    tc_pr.append(node)


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for cell, value in zip(t.rows[0].cells, headers):
        shade(cell, "1E3A8A")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(value), 8.5, True, "FFFFFF")
    for values in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, values):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            font(cell.paragraphs[0].add_run(str(value)), 8.5)
    return t


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style = doc.styles["Title"]
    ppr = style.element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None: ppr.remove(border)
    section = doc.sections[0]
    section.top_margin = Inches(0.65); section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75); section.right_margin = Inches(0.75)
    title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    font(title.add_run("WorkAgent-RSI 整体框架与训练计划"), 24, True)
    body(doc, "本报告定义面向 Excel、Word 和 PowerPoint 智能体的递归技能改进架构与实验计划。Phase 2 只交付可执行设计，不报告尚未运行的训练分数。")

    heading(doc, "1 RSI 定义与研究边界")
    body(doc, "本项目中的 RSI 指 Recursive Skill Improvement。系统把任务经验和评价反馈转化为版本化技能、工具策略或 Office 脚本的持久更新，并让通过验证的更新影响后续任务。初始范围为 L1 和受控 L2。人类保留目标、保护评估、资源权限和高风险发布权。")
    bullet(doc, "不修改基础模型权重。")
    bullet(doc, "不允许候选修改 evaluator、hidden tests、promotion policy 或历史证据。")
    bullet(doc, "不把 develop 分数提高等同于可迁移能力提高。")

    heading(doc, "2 系统组件")
    rows = [
        ("Generator", "基于失败与编辑历史提出原子候选", "候选 patch 与新增测试"),
        ("Verifier", "执行 schema、安全、格式和回归检查", "VerificationReport"),
        ("Evaluator", "冻结规则下评估 artifact 与 trace", "EvaluationReport"),
        ("Promotion", "执行非补偿式晋级门槛", "PromotionDecision"),
        ("Registry", "保存不可变版本、证据和回滚点", "SkillVersion"),
        ("Trace memory", "保存运行、失败、成本和决策", "ExperienceRecord"),
    ]
    table(doc, ["组件", "职责", "输出"], rows)

    heading(doc, "3 RRSI 正则化")
    body(doc, "候选侧使用退火的原子编辑预算、跨轮负证据和停滞时结构化探索。选择侧先筛查泄漏，再依据 champion 的经验噪声带、protected regression、hidden/OOD 结果和成本增量做保守选择。长期无正贡献的组件进入剪枝候选。")

    heading(doc, "4 数据与证据治理")
    body(doc, "每次运行绑定 TaskSpec、SkillVersion、EvaluatorVersion、环境配置、随机种子和 Git commit。原始 trace 与 artifact 为不可变证据，派生诊断必须引用源 hash。分割在模板族和来源组层面完成，并用 exact hash、规范化文本和结构指纹去重。")

    heading(doc, "5 训练与评估阶段")
    stage_rows = [
        ("A Pilot", "30 tasks", "验证环境、schema、grader 和方差"),
        ("B Excel", "闭环优化", "测试晋级、回归和 RRSI 消融"),
        ("C Word/PPT", "结构与视觉", "加入 OOXML、渲染、溢出和布局评估"),
        ("D Transfer", "冻结技能", "测试新模板、跨域和第二 executor"),
    ]
    table(doc, ["阶段", "范围", "目标"], stage_rows)

    heading(doc, "6 Baselines 与指标")
    body(doc, "比较固定技能、Self-Refine、generator-only、generator+verifier、generator+verifier+evaluator 和完整 WorkAgent-RSI。所有方法共享任务访问、模型、候选数和总预算。")
    bullet(doc, "主要指标：critical-gate task success、artifact correctness、regression rate、hidden/OOD delta、unsafe action rate。")
    bullet(doc, "次要指标：视觉质量、迭代次数、tokens、tool calls、latency、cost per validated gain、activation 与 faithful use。")

    heading(doc, "7 资源与预算")
    body(doc, "当前主机为 Intel Core i7-14650HX、16 GB RAM 和 8 GB RTX 5060 Laptop GPU，适合编排、Office 处理、渲染和轻量模型，不适合大型基础模型全量训练。Pilot baseline 约需 540 run units；Excel candidate evaluation 计划上限约为 2160 run units。正式成本需在 Phase 3 填入 provider 单价和实测运行时间后审批。")

    heading(doc, "8 Go No-Go 与风险")
    bullet(doc, "任务与 evaluator 必须达到可重复、可审计状态。")
    bullet(doc, "保护边界必须阻止候选写入 evaluator、hidden tests 和 evidence store。")
    bullet(doc, "候选必须跨重复运行复现，不增加 unsafe actions 或 critical regressions。")
    bullet(doc, "若 hidden/OOD 持续退化、视觉评价不稳定或成本超过批准预算，则停止并重设计。")
    body(doc, "剩余风险包括 provider 与 Office engine 差异、有限任务上的自适应过拟合、字体与渲染对视觉评估的影响，以及技能库增长导致的检索漂移。")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
