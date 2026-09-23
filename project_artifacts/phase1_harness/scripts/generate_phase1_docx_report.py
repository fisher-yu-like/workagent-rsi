"""Generate the Phase 1B DOCX report from recorded run evidence."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[3]
RESULTS = ROOT / "project_artifacts" / "phase1_harness" / "results" / "run_summary.json"
REPORT = ROOT / "project_artifacts" / "phase1_harness" / "reports" / "Harness详细设计与运行报告.docx"


def set_font(run, name: str, size: float, bold: bool = False, color: str = "000000") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def set_table_borders(table, color: str = "D9D9D9") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(5)
    set_font(paragraph.add_run(text), "Microsoft YaHei", 15, True, "000000")


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.25
    set_font(paragraph.add_run(text), "Microsoft YaHei", 10.5)


def main() -> int:
    results = json.loads(RESULTS.read_text(encoding="utf-8"))
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    title_style = document.styles["Title"]
    style_p_pr = title_style.element.get_or_add_pPr()
    style_border = style_p_pr.find(qn("w:pBdr"))
    if style_border is not None:
        style_p_pr.remove(style_border)
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(14)
    remove_paragraph_borders(title)
    set_font(title.add_run("Harness 详细设计与运行报告"), "Microsoft YaHei", 25, True, "000000")
    add_body(document, "本报告记录 Phase 1B 本地 Mock Pipeline 的实现与运行证据。当前尚未接入外部 WorkAgent provider，Office 文件也不代表真实办公任务执行结果。")

    add_heading(document, "设计与部署")
    add_body(document, "Pipeline 由 TaskSpec、Orchestrator、MockWorkAgentAdapter、ArtifactStore、TraceStore 和 BasicEvaluator 组成。Orchestrator 管理运行状态、受控重试和失败记录。ArtifactStore 使用 SHA-256 管理输出，TraceStore 使用 SQLite 保存事件。")

    add_heading(document, "运行记录")
    table = document.add_table(rows=1, cols=4)
    table.autofit = True
    set_table_borders(table)
    headers = ["Case", "State", "Run ID", "Seconds"]
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, "1E3A8A")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(paragraph.add_run(value), "Microsoft YaHei", 9.5, True, "FFFFFF")
    for result in results:
        row = table.add_row().cells
        values = [result["case_id"], result["state"], result["run_id"], f"{result['duration_seconds']:.3f}"]
        for index, (cell, value) in enumerate(zip(row, values)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (1, 3) else WD_ALIGN_PARAGRAPH.LEFT
            set_font(paragraph.add_run(value), "Microsoft YaHei", 9.5)

    add_heading(document, "复现与限制")
    add_body(document, "复现命令为 py -3.12 project_artifacts/phase1_harness/scripts/run_phase1_cases.py。每个 case 的 JSON 和 SQLite trace 位于 project_artifacts/phase1_harness/results。当前结果只证明本地 mock Harness 闭环，真实 WorkAgent 接入和 Office 引擎验证仍是下一步任务。")

    document.save(REPORT)
    print(REPORT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
