"""Prepare opaque review files without labels, model identities or judge scores."""
import argparse
from datetime import datetime, timezone
import html
import json
from pathlib import Path
import secrets
import shutil
import subprocess
import sys

ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT/'src'))
from workagent_rsi.calibration_data import load_calibration_cases
from workagent_rsi.environment_manifest import capture_environment
from workagent_rsi.hashing import sha256_file,canonical_json_hash
from workagent_rsi.view_safety import ordinary_file,relative_reference

RUBRIC={'version':'human-office-calibration-v1','language':'zh-CN',
    'labels':{'PASS':'满足当前任务的全部明确要求，允许任务授权的等价方案。',
        'FAIL':'有确定的任务缺陷；请写明位置和观察到的问题。',
        'PARTIAL':'只完成部分要求；请注明已完成和未完成的部分。',
        'INCOMPLETE':'任务标准有歧义，或缺少必要证据／能力，无法下结论。'},
    'instructions':'先读任务和输入，再检查成品及实际渲染页。不要根据文件名猜标签。空白表示未标注，不等于PASS。',
    'changes':'修改本标准后必须生成新的rubric_hash和评估器身份，重新运行整个可比cohort。'}


def prepare(manifest: Path, output: Path, *, natural_manifest: Path | None=None,
            allow_partial_natural: bool=False, include_standards: bool=True) -> dict:
    if output.exists(): raise FileExistsError(output)
    cases=load_calibration_cases(manifest)
    selected=[]
    for domain in ('excel','word','powerpoint'):
        for category in ('correct','single_defect','equivalent','gap'):
            options=[c for c in cases if c.domain==domain and c.category==category]
            if category=='single_defect':
                wanted={'excel':'wrong_range','word':'wrong_value','powerpoint':'obstruction'}[domain]
                options=[c for c in options if c.defect_type==wanted]
            selected.append(sorted(options,key=lambda c:c.case_id)[0])
    sources=[dict(original_id=c.case_id,domain=c.domain,instruction=c.instruction,
        artifact=manifest.parent/c.artifact_path,inputs=[manifest.parent/c.input_path],
        source_family=c.source_family,cohort='standards-12',oracle=c.oracle_status) for c in selected]
    if not include_standards:
        if natural_manifest is None: raise ValueError('natural-only review requires an actual natural manifest')
        sources=[]
    natural=[]
    if natural_manifest:
        data=json.loads(natural_manifest.read_text(encoding='utf-8'))
        for row in data['cases']:
            if row['source_kind'] not in {'user_supplied','natural_model_output','external_licensed'}:
                raise ValueError('controlled fixtures cannot substitute for a natural cohort')
            if row['source_family'] in {c.source_family for c in cases} or not row.get('source_refs'):
                raise ValueError('natural cohort needs independent family and provenance')
            relative_reference(row['artifact_path'])
            for p in row.get('input_paths',[]): relative_reference(p)
            artifact=ordinary_file(natural_manifest.parent/row['artifact_path'])
            if sha256_file(artifact)!=row['artifact_hash']: raise ValueError('natural artifact hash mismatch')
            natural.append(dict(original_id=row['case_id'],domain=row['domain'],instruction=row['instruction'],
                artifact=artifact,inputs=[ordinary_file(natural_manifest.parent/p) for p in row.get('input_paths',[])],
                source_family=row['source_family'],cohort='natural-blind',oracle=None))
        if len({r['original_id'] for r in natural}) != len(natural): raise ValueError('duplicate natural case identity')
        if len(natural)<20 and not allow_partial_natural: raise ValueError('natural blind cohort needs at least20 real finished artifacts')
    sources.extend(natural)
    output.mkdir(parents=True)
    public=output/'blind-review-pack';public.mkdir()
    private=output/'evaluator-only';private.mkdir()
    rubric_hash=canonical_json_hash(RUBRIC)
    mapping=[];public_rows=[]
    secrets.SystemRandom().shuffle(sources)
    for entry in sources:
        opaque='review-'+secrets.token_hex(8)
        artifact=ordinary_file(entry['artifact'])
        suffix={'excel':'.xlsx','word':'.docx','powerpoint':'.pptx'}[entry['domain']]
        if artifact.suffix.lower()!=suffix: raise ValueError('natural artifact domain mismatch')
        # Real corresponding libraries must reopen every supplied artifact.
        if suffix=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(artifact);book.close()
        elif suffix=='.docx':
            from docx import Document
            Document(artifact)
        else:
            from pptx import Presentation
            Presentation(artifact)
        name=opaque+suffix;shutil.copyfile(artifact,public/name)
        input_names=[]; input_labels=[]
        for i,path in enumerate(entry['inputs'],1):
            path=ordinary_file(path);n=f'{opaque}-input-{i}{path.suffix}'
            shutil.copyfile(path,public/n);input_names.append(n)
            input_labels.append(path.name if entry['cohort']=='natural-blind' else f'输入{i}')
        public_rows.append({'case_id':opaque,'domain':entry['domain'],'instruction':entry['instruction'],
            'artifact':name,'artifact_hash':sha256_file(public/name),'inputs':input_names,
            'input_labels':input_labels,'cohort':entry['cohort']})
        mapping.append({'case_id':opaque,'original_id':entry['original_id'],'source_family':entry['source_family'],
            'oracle':entry['oracle'],'artifact_hash':sha256_file(artifact),'cohort':entry['cohort']})
    def write(path,obj): path.write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    write(public/'rubric.json',RUBRIC)
    write(public/'review-manifest.json',{'rubric_hash':rubric_hash,'case_ids':[r['case_id'] for r in public_rows],'cases':public_rows})
    with (public/'human-labels.template.jsonl').open('x',encoding='utf-8') as stream:
        for row in public_rows: stream.write(json.dumps({'case_id':row['case_id'],'reviewer_id':None,'label':None,
            'defect_locations':[],'comment':'','rubric_hash':rubric_hash},ensure_ascii=False)+'\n')
    write(private/'case-key.json',{'rubric_hash':rubric_hash,'cases':mapping,'judge_results_revealed':False})
    cards=[]
    for row in public_rows:
        links=' '.join(f'<a href="{html.escape(n)}">{html.escape(label)}</a>' for n,label in zip(row['inputs'],row['input_labels']))
        cards.append(f'<article id="{row["case_id"]}"><h2>{row["case_id"]} / {row["domain"]}</h2><h3>任务说明</h3><p>{html.escape(row["instruction"])}</p>'
            f'<p>{links} <a href="{row["artifact"]}">打开成品</a></p><div class="pages"></div></article>')
    page='<!doctype html><meta charset="utf-8"><title>Office 盲标材料</title><style>body{max-width:1000px;margin:30px auto;font:16px sans-serif;line-height:1.6}article{border-top:1px solid #bbb;padding:20px 0}img{max-width:100%;border:1px solid #ddd}</style><h1>Office 成品审阅</h1><p>'+html.escape(RUBRIC['instructions'])+'</p><p>填写 human-labels.template.jsonl。PASS / FAIL / PARTIAL / INCOMPLETE；写明缺陷位置。</p>'+''.join(cards)
    page=page.replace('<h1>Office 成品审阅</h1>',
        '<h1>Office 成品审阅</h1><p>部分案例共享任务说明，但输入或成品版本可能不同。请分别核对每份成品；仅外观看起来正常，不足以确认所有要求均已满足。</p>')
    (public/'index.html').write_text(page,encoding='utf-8')
    attach_review_form(output)
    write(output/'human-labels.jsonl.status.json',{'status':'blocked','actual_human_labels':0,'reason':'Awaiting actual user/reviewer submission; no labels generated by the agent.'})
    write(output/'natural-intake-template.json',{'minimum_count':20,'independent_of_families':sorted({c.source_family for c in cases}),
        'cases':[],'required_case_fields':['case_id','domain','instruction','artifact_path','artifact_hash','source_family','source_kind','source_refs','input_paths']})
    report={'status':'blocked','standards_cases':12 if include_standards else 0,'natural_cases':len(natural),'required_natural_cases':20,
        'actual_human_labels':0,'rubric_hash':rubric_hash,'judge_results_revealed':False,
        'blind_pack_hash':canonical_json_hash({p.name:sha256_file(p) for p in public.iterdir() if p.is_file()}),
        'blockers':['Actual human labeling required']+([] if len(natural)>=20 else [f'Independent natural finished-artifact cohort needs20; actual{len(natural)}'])}
    write(output/'agreement-report.json',report)
    return report


def attach_review_form(output: Path) -> None:
    public=output/'blind-review-pack';index=public/'index.html'
    page=index.read_text(encoding='utf-8')
    if 'id="review-controls"' in page: return
    manifest=json.loads((public/'review-manifest.json').read_text(encoding='utf-8'))
    for case_id in manifest['case_ids']:
        start=page.index(f'<article id="{case_id}">');end=page.index('</article>',start)
        form=f'''<p><label>你的判断：<select id="label-{case_id}"><option value="">请选择</option><option value="PASS">PASS：符合要求</option><option value="FAIL">FAIL：确定有缺陷</option><option value="PARTIAL">PARTIAL：部分完成</option><option value="INCOMPLETE">INCOMPLETE：证据不足或标准不清</option></select></label></p><label>问题位置和说明（有问题时填写）：<br><textarea id="location-{case_id}" rows="3" style="width:95%"></textarea></label>'''
        page=page[:end]+form+page[end:]
    data=json.dumps({'case_ids':manifest['case_ids'],'rubric_hash':manifest['rubric_hash']}).replace('<','\\u003c')
    script='''<section id="review-controls"><h2>导出你的标注</h2><label>审阅者姓名或昵称：<input id="reviewer" autocomplete="name"></label> <button id="export-review">下载标注 JSON</button><p>完成所有选择后下载文件，再在聊天中发送该文件或告诉我保存路径。页面不会联网，也不会自动提交。</p></section><script id="review-data" type="application/json">'''+data+'''</script><script>
const reviewData=JSON.parse(document.getElementById('review-data').textContent);
document.getElementById('export-review').onclick=()=>{
 const reviewer=document.getElementById('reviewer').value.trim();
 if(!reviewer){alert('请填写审阅者姓名或昵称');return;}
 const timestamp=new Date().toISOString();
 const labels=reviewData.case_ids.map(case_id=>({case_id,reviewer_id:reviewer,
   label:document.getElementById('label-'+case_id).value,
   defect_locations:document.getElementById('location-'+case_id).value.split('\\n').map(x=>x.trim()).filter(Boolean),timestamp}));
 if(labels.some(x=>!x.label)){alert('还有案例没有选择标签，请先完成全部选择');return;}
 const payload={rubric_hash:reviewData.rubric_hash,reviewer_id:reviewer,labels};
 const url=URL.createObjectURL(new Blob([JSON.stringify(payload,null,2)],{type:'application/json'}));
 const link=document.createElement('a');link.href=url;link.download='office-human-labels.json';link.click();
 setTimeout(()=>URL.revokeObjectURL(url),1000);
};</script>'''
    index.write_text(page+script,encoding='utf-8')


def render_review_pack(output: Path, policy_path: Path) -> dict:
    """Only actual isolated renders become visible review pages."""
    from workagent_rsi.apptainer_sandbox import ApptainerPolicy,ApptainerSSHExecutor
    from workagent_rsi.artifact_renderer import OfficeArtifactRenderer
    policy=ApptainerPolicy.model_validate_json(policy_path.read_text(encoding='utf-8'))
    sandbox=ApptainerSSHExecutor(policy)
    public=output/'blind-review-pack'
    manifest=json.loads((public/'review-manifest.json').read_text(encoding='utf-8'))
    renderer=OfficeArtifactRenderer(output/'evaluator-only/rendering',sandbox=sandbox,dpi=96,max_pages=8,timeout_seconds=90)
    results=[]
    page=(public/'index.html').read_text(encoding='utf-8')
    for row in manifest['cases']:
        result=renderer.render(public/row['artifact'],row['domain'],row['artifact_hash'])
        images=[]
        for rendered in result.pages if result.status=='available' else []:
            source=ordinary_file(Path(rendered.path))
            if sha256_file(source)!=rendered.sha256: raise ValueError('render page hash mismatch')
            name=f'{row["case_id"]}-page-{rendered.page}.png'
            shutil.copyfile(source,public/name)
            images.append({'file':name,'sha256':rendered.sha256,'page':rendered.page})
        prefix=f'<article id="{row["case_id"]}">'
        start=page.index(prefix);end=page.index('</article>',start)
        section=page[start:end]
        content=''.join(f'<img src="{item["file"]}" alt="实际渲染第{item["page"]}页">' for item in images)
        if not images: content='<p>实际渲染不可用，请直接打开成品；不得把缺少图片视为产物缺陷。</p>'
        page=page[:start]+section.replace('<div class="pages"></div>','<div class="pages">'+content+'</div>')+page[end:]
        results.append({'case_id':row['case_id'],'status':result.status,'error':result.error,'pages':images,
            'renderer_identity':result.renderer_identity,'environment':result.environment})
        (output/'evaluator-only/render-progress.json').write_text(json.dumps(results,indent=2)+'\n',encoding='utf-8')
    (public/'index.html').write_text(page,encoding='utf-8')
    (public/'rendered-pages.json').write_text(json.dumps(results,indent=2)+'\n',encoding='utf-8')
    report=json.loads((output/'agreement-report.json').read_text())
    report.update(rendered_cases=sum(r['status']=='available' for r in results),rendered_pages=sum(len(r['pages']) for r in results),
        blind_pack_hash=canonical_json_hash({p.name:sha256_file(p) for p in public.iterdir() if p.is_file()}))
    (output/'agreement-report.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    return report


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--manifest',type=Path,required=True);parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--natural-manifest',type=Path)
    parser.add_argument('--allow-partial-natural',action='store_true',help='Prepare available files but retain the explicit20-file shortage blocker')
    parser.add_argument('--natural-only',action='store_true',help='Keep the earlier12 standards review separate; do not ask the user to label them twice')
    parser.add_argument('--render-policy',type=Path)
    parser.add_argument('--render-existing',action='store_true')
    parser.add_argument('--form-existing',action='store_true')
    args=parser.parse_args();start=datetime.now(timezone.utc).isoformat()
    if args.form_existing:
        attach_review_form(args.output)
        report=json.loads((args.output/'agreement-report.json').read_text())
        public=args.output/'blind-review-pack'
        report['blind_pack_hash']=canonical_json_hash({p.name:sha256_file(p) for p in public.iterdir() if p.is_file()})
        (args.output/'agreement-report.json').write_text(json.dumps(report,indent=2)+'\n')
    elif args.render_existing:
        if args.render_policy is None: parser.error('--render-existing requires --render-policy')
        report=render_review_pack(args.output,args.render_policy)
    else:
        report=prepare(args.manifest,args.output,natural_manifest=args.natural_manifest,
            allow_partial_natural=args.allow_partial_natural,include_standards=not args.natural_only)
        capture_environment(ROOT,args.output/'environment.json')
        if args.render_policy: report=render_review_pack(args.output,args.render_policy)
    stdout=json.dumps(report,sort_keys=True)
    (args.output/'stdout.txt').write_text(stdout+'\n');(args.output/'stderr.txt').write_text('')
    (args.output/('form-run.json' if args.form_existing else 'render-run.json' if args.render_existing else 'run.json')).write_text(json.dumps({'command':[sys.executable,*sys.argv],'cwd':str(ROOT),
        'started_at':start,'ended_at':datetime.now(timezone.utc).isoformat(),'exit_code':1,
        'git_commit':subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip(),
        'input_manifest_hash':sha256_file(args.manifest),'source_hash':sha256_file(Path(__file__))},indent=2)+'\n')
    print(stdout);return 1


if __name__=='__main__': raise SystemExit(main())
