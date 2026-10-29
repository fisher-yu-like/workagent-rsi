"""Authored T24 qualification. Runs only fixed trusted programs locally, no model."""
from copy import deepcopy
from datetime import datetime,timezone
import hashlib
import json
from pathlib import Path
import subprocess
import sys
from uuid import uuid4

ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT/'src'))
from workagent_rsi.contracts import TaskSpec
from workagent_rsi.hashing import sha256_file
from workagent_rsi.run_contracts import ExecutionOutcome
from workagent_rsi.task_observations import inspection_program
from workagent_rsi.workagent_office import WorkAgentOfficeAdapter
from workagent_rsi.workagent_provider import AgentResponse,ProviderOutcome,WorkAgentConfig,_execute_native_tool


class TrustedLocalExecutor:
    """Only the frozen authored generation program and exact trusted inspectors.

    This is an engineering fixture executor, not a model-code host fallback and
    not T04 sandbox evidence. No candidate program or skill text is executable.
    """
    qualification_only=True
    def __init__(self):
        self.program=Path(__file__).with_name('t24_authored_program.py').read_text(encoding='utf-8')
        self.allowed={hashlib.sha256(self.program.encode()).hexdigest()}
    def register_inspection(self,script,config):
        if script.read_text(encoding='utf-8')!=inspection_program(config):raise ValueError('untrusted inspection program')
        self.allowed.add(sha256_file(script))
    def run(self,script,workspace,timeout):
        if sha256_file(script) not in self.allowed:raise ValueError('authored executor refuses arbitrary source')
        started=datetime.now(timezone.utc)
        completed=subprocess.run([sys.executable,str(script)],cwd=workspace,capture_output=True,text=True,
            encoding='utf-8',timeout=timeout,check=False)
        ref='provider_records/authored-local/'+uuid4().hex+'.json';path=workspace/ref;path.parent.mkdir(parents=True,exist_ok=True)
        path.write_text(json.dumps({'qualification_scope':'authored_local_library','command':[sys.executable,str(script)],
            'cwd':str(workspace),'started_at':started.isoformat(),'ended_at':datetime.now(timezone.utc).isoformat(),
            'returncode':completed.returncode,'stdout':completed.stdout,'stderr':completed.stderr,
            'script_sha256':sha256_file(script)},indent=2),encoding='utf-8')
        return ExecutionOutcome(status='succeeded' if completed.returncode==0 else 'failed',
            failure_origin=None if completed.returncode==0 else 'agent_execution',trace_refs=[ref])


class AuthoredProgramProvider:
    qualification_only=True
    def __init__(self,request):self.request=request;self.sandbox=TrustedLocalExecutor()
    def run(self,prompt,workspace,records):
        records.mkdir(parents=True,exist_ok=True);(records/'prompt.md').write_text(prompt,encoding='utf-8')
        manifest=json.loads((workspace/'input_manifest.json').read_text(encoding='utf-8'));names=[i['copied'] for i in manifest['files']]
        def tool(name,args):
            result=_execute_native_tool(name,args,workspace,timeout_seconds=30,sandbox=self.sandbox,records=records)
            with (records/'tool_events.jsonl').open('a',encoding='utf-8') as stream:stream.write(json.dumps({'name':name,'arguments':args,'result':result})+'\n')
            if result.get('ok') is not True:raise ValueError(str(result.get('error')))
            return result
        for name in names:
            tool('inspect_office',{'path':name,'chosen_operations':['Locate targets from current structure and user request'],
                                  'verification_plan':['Reopen output and compare requested changes and preserved source']})
        request={**self.request,'inputs':names}
        tool('write_file',{'path':'scratch/request.json','content':json.dumps(request)})
        tool('write_file',{'path':'agent.py','content':self.sandbox.program})
        tool('run_python',{'path':'agent.py'})
        suffix={'excel':'.xlsx','word':'.docx','powerpoint':'.pptx'}[request['domain']]
        response=AgentResponse(status='completed',deliverables=['outputs/result'+suffix],summary='authored local qualification',input_files_used=names)
        tool('write_file',{'path':'deliverables.json','content':json.dumps({'deliverables':response.deliverables})})
        (workspace/'agent_response.json').write_text(response.model_dump_json())
        return ProviderOutcome(status='completed',response=response)


class AuthoredExecutor:
    qualification_only=True
    evidence_kind='authored_local_library'
    model_calls=0
    def __call__(self,skill,request,root):
        task=TaskSpec.model_validate(request['task'])
        adapter=WorkAgentOfficeAdapter(root,WorkAgentConfig(verify_com=False),input_base=ROOT,agent_instructions=skill.get('instructions',''))
        adapter.provider=AuthoredProgramProvider(request['request'])
        events=list(adapter.execute(task,'authored.generic.program'))
        (root/'adapter-events.json').write_text(json.dumps(events,indent=2),encoding='utf-8')
        failures=[e for e in events if e['kind'] in {'failure','unavailable'}]
        if failures:raise ValueError(failures[0]['message'])
        artifact=next(e['artifact_path'] for e in events if e['kind']=='artifact')
        return {'artifact':artifact,'workspace':str(root/'agent_workspace'),
            'observation':str(root/'provider_records/task-observation.json'),
            'tool_log':str(root/'provider_records/tool_events.jsonl'),
            'input_manifest':json.loads((root/'agent_workspace/input_manifest.json').read_text(encoding='utf-8'))}


def build_cases(root):
    root=Path(root).resolve();root.mkdir(parents=True,exist_ok=True);cases=[]
    def add(case_id,group,domain,instruction,source,request,expected):
        task=TaskSpec(task_id='task-'+uuid4().hex,domain=domain,instruction=instruction,input_files=(str(source),) if source else ())
        cases.append({'case_id':case_id,'group':group,'task':task.model_dump(mode='json'),
                      'public_request':{'domain':domain,**request},'expected':expected})
    def book(name,sheet,headers,rows):
        from openpyxl import Workbook
        path=root/(name+'.xlsx');wb=Workbook();ws=wb.active;ws.title=sheet;ws.append(headers)
        for row in rows:ws.append(row)
        wb.save(path);wb.close();return path
    excel=[('excel-sheet','Records',['Name','Revenue'],[['A',5],['B',11]],16),
           ('excel-sheet','Renamed Ledger',['Name','Revenue'],[['A',5],['B',11]],16),
           ('excel-values-rows','Data',['Name','Revenue'],[['A',5],['B',11]],16),
           ('excel-values-rows','Data',['Name','Revenue'],[['A',5],['B',11],['C',8]],24),
           ('excel-values-rows','Data',['Name','Revenue'],[['A',5],['B',19]],24),
           ('excel-columns','Records',['Name','Revenue'],[['A',5],['B',11]],16),
           ('excel-columns','Records',['Revenue','Extra','Name'],[[7,'x','A'],[13,'y','B']],20)]
    for i,(group,sheet,headers,rows,total) in enumerate(excel):
        source=book('input-'+str(i),sheet,headers,rows)
        add(group+'-'+str(i),group,'excel','Read every current Revenue record; preserve the source and write its total to Result!A1.',source,
            {'operation':'sum_column','value_header':'Revenue'},{'cells':{'Result!A1':total},'sheet_names':[sheet,'Result']})
    def word(name,paragraphs,table=None):
        from docx import Document
        path=root/(name+'.docx');doc=Document()
        for text in paragraphs:doc.add_paragraph(text)
        if table:
            t=doc.add_table(rows=len(table),cols=len(table[0]))
            for i,row in enumerate(table):
                for j,value in enumerate(row):t.cell(i,j).text=value
        doc.save(path);return path
    for i,target in enumerate(['Alpha','Beta']):
        source=word('target-'+str(i),['Alpha','Beta','Keep'])
        add('word-target-'+str(i),'word-target','word',f'Replace only the paragraph {target} with Edited.',source,
            {'operation':'replace_paragraph','target_text':target,'replacement':'Edited'},
            {'paragraphs':['Edited','Beta','Keep'] if i==0 else ['Alpha','Edited','Keep'],'styles':['Normal']*3,'tables':[]})
    for i,(paragraphs,level) in enumerate([(['Start'],2),(['Start','Extra'],3)]):
        source=word('paragraphs-'+str(i),paragraphs)
        add('word-paragraphs-'+str(i),'word-paragraphs','word',f'Preserve all paragraphs and append Added as Heading {level}.',source,
            {'operation':'append_heading','heading':'Added','level':level},
            {'paragraphs':paragraphs+['Added'],'styles':['Normal']*len(paragraphs)+['Heading '+str(level)],'tables':[]})
    for i,table in enumerate([[['Name','Value'],['A','old']],[['Value','Flag','Name'],['old','yes','A'],['spare','no','B']]]):
        source=word('table-'+str(i),['Register'],table)
        expected=[[['Name','Value'],['A','new']]] if i==0 else [[['Value','Flag','Name'],['new','yes','A'],['spare','no','B']]]
        add('word-table-'+str(i),'word-table','word','In the current table, change only Value for Name A to new.',source,
            {'operation':'update_table','key_header':'Name','value_header':'Value','key':'A','replacement':'new'},
            {'paragraphs':['Register'],'styles':['Normal'],'tables':expected})
    def deck(name,titles):
        from pptx import Presentation
        from pptx.util import Inches
        path=root/(name+'.pptx');obj=Presentation()
        for title in titles:
            slide=obj.slides.add_slide(obj.slide_layouts[6]);slide.shapes.add_textbox(Inches(.5),Inches(.4),Inches(8),Inches(.7)).text=title
        obj.save(path);return path
    for i,count in enumerate([2,4]):
        add('ppt-count-'+str(i),'ppt-count','powerpoint',f'Create exactly {count} slides, titled Item followed by its one-based slide number.',None,
            {'operation':'create_slides','count':count,'prefix':'Item'},
            {'slide_count':count,'slide_texts':[['Item 1'],['Item 2']] if count==2 else [['Item 1'],['Item 2'],['Item 3'],['Item 4']]})
    for i,target in enumerate([1,3]):
        source=deck('slides-'+str(i),['One','Two','Three'])
        expected=[['One','Added'],['Two'],['Three']] if target==1 else [['One'],['Two'],['Three','Added']]
        add('ppt-target-'+str(i),'ppt-target','powerpoint',f'Add Added only on slide {target}; keep all original slide text.',source,
            {'operation':'edit_slide','slide':target,'text':'Added'},{'slide_count':3,'slide_texts':expected})
    for i,titles in enumerate([['One','Two'],['One','Two','Three']]):
        source=deck('append-'+str(i),titles)
        add('ppt-append-'+str(i),'ppt-append','powerpoint','Append one new slide titled New; keep all existing slides.',source,
            {'operation':'append_slides','titles':['New']},{'slide_count':len(titles)+1,'slide_texts':[[t] for t in titles]+[['New']]})
    for group in ['excel-sheet','word-target','ppt-target']:
        original=next(c for c in cases if c['group']==group);clone=deepcopy(original)
        clone.update(case_id=group+'-random-id',id_equivalent_to=original['case_id'])
        clone['task']['task_id']='unrelated-'+uuid4().hex;cases.append(clone)
    return {'schema_version':'authored-input-probes-v1','provenance':'project-generated; independent literal oracles; no external benchmark',
            'required_groups':list(dict.fromkeys(c['group'] for c in cases)),'cases':cases,'output_root':str(root/'runs')}


if __name__=='__main__':
    import argparse
    from workagent_rsi.behavior_probes import run_behavior_probes
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--output-root',required=True);args=parser.parse_args()
    root=Path(args.output_root);manifest=build_cases(root/'cases')
    (root/'case-manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
    result=run_behavior_probes({'instructions':'Use the current inputs and task.'},manifest,AuthoredExecutor())
    print(json.dumps({'status':result['status'],'cases':len(result['rows']),'groups':len(result['group_results']),
                      'model_calls':result['model_calls'],'trace_index':result['trace_index']}))
    sys.exit(0 if result['status']=='passed' else 1)
