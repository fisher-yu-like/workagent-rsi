"""Freeze twenty genuine-output jobs and their input-only execution views.

No model call, source program execution, reference-output copying or human label
is performed here. Selected connected families leave every adaptive/transfer
split before any output is produced or rubric is calibrated.
"""
from __future__ import annotations

import argparse
from collections import Counter
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / 'src'))
from workagent_rsi.benchmark_corpus import export_execution_inputs
from workagent_rsi.hashing import canonical_json_hash, sha256_file

READABLE = {'.xlsx', '.docx', '.pptx', '.txt', '.md', '.json', '.csv'}


def preview(path: Path):
    """Bounded observations only; the generator must inspect full inputs in code."""
    suffix = path.suffix.lower()
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            return {'sheets': [{'name': s.title, 'dimensions': s.calculate_dimension(),
                'first_rows': [list(r) for r in s.iter_rows(min_row=1, max_row=min(s.max_row or 1, 15),
                                                        max_col=min(s.max_column or 1, 15), values_only=True)]}
                              for s in book.worksheets[:12]], 'partial': True}
        finally:
            book.close()
    if suffix == '.docx':
        from docx import Document
        doc = Document(path)
        return {'paragraphs': [p.text for p in doc.paragraphs[:120]],
                'tables': [[[c.text for c in row.cells[:12]] for row in table.rows[:25]] for table in doc.tables[:8]], 'partial': True}
    if suffix == '.pptx':
        from pptx import Presentation
        deck = Presentation(path)
        return {'slides': [[s.text for s in slide.shapes if s.has_text_frame] for slide in list(deck.slides)[:25]], 'partial': True}
    body = path.read_bytes()
    return {'text': body[:24000].decode('utf-8-sig', errors='replace'), 'partial': len(body) > 24000}


def prepare(manifest: Path, output: Path) -> dict:
    if output.exists():
        raise FileExistsError(output)
    data = json.loads(manifest.read_text(encoding='utf-8'))
    priority = {'spreadsheetbench-2': 0, 'workspace-bench-lite': 1, 'spreadsheetbench-v1': 2}
    eligible = []
    for task in data['tasks']:
        if (not task['variants'] or task['split'] == 'develop' or task.get('category') == 'Visualization'
                or len(task.get('output_files', ['output.xlsx'])) != 1):
            continue
        refs = task['variants'][0]['inputs']
        if all(Path(r['path']).suffix.lower() in READABLE for r in refs) and sum(r['size_bytes'] for r in refs) <= 4_000_000:
            eligible.append(task)
    selected = []
    # Native Lite contains only three eligible single-output PPT tasks. These
    # quotas are a descriptive pilot; no three-domain population claim follows.
    quotas = {'excel': 10, 'word': 7, 'powerpoint': 3}
    for domain, count in quotas.items():
        options = sorted((t for t in eligible if t['domain'] == domain), key=lambda t: (
            priority[t['benchmark']], canonical_json_hash(['human-pilot-20261025', t['base_task_id']])))
        if len(options) < count:
            raise ValueError(f'Not enough eligible {domain} tasks: {len(options)} < {count}')
        selected.extend(options[:count])
    families = {t['family_id'] for t in selected}
    for task in data['tasks']:
        if task['family_id'] in families:
            task['split'] = 'human_calibration'
    output.mkdir(parents=True)
    jobs = []
    for task in selected:
        variant = task['variants'][0]
        case_id = 'natural-' + canonical_json_hash(task['base_task_id'])[:16]
        workspace = output / 'execution-views' / case_id
        spec = export_execution_inputs(task, variant['variant_id'], Path(data['roots'][task['benchmark']]), workspace)
        instruction = spec.instruction
        output_name = task.get('output_files', ['output.xlsx'])[0]
        observations = {name: preview(workspace / name) for name in spec.input_files}
        request = {'task': spec.model_dump(mode='json'), 'required_output': 'outputs/' + output_name,
            'input_observations': observations,
            'scope': 'Observations may be clipped. Generated Python can read all staged inputs. No reference answer or evaluation rubric is supplied.'}
        request_path = output / 'model-requests' / (case_id + '.json')
        request_path.parent.mkdir(exist_ok=True)
        request_path.write_text(json.dumps(request, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')
        jobs.append({'case_id': case_id, 'base_task_id': task['base_task_id'], 'variant_id': variant['variant_id'],
            'domain': task['domain'], 'benchmark': task['benchmark'], 'source_family': task['family_id'],
            'source_refs': task['source_refs'], 'instruction': instruction,
            'request': request_path.relative_to(output).as_posix(), 'request_sha256': sha256_file(request_path),
            'workspace': workspace.relative_to(output).as_posix(), 'expected_output': output_name,
            'inputs': {n: sha256_file(workspace / n) for n in spec.input_files},
            'actual_output': None, 'actual_human_labels': 0})
    (output / 'evaluator-only').mkdir()
    (output / 'evaluator-only/reserved-task-manifest.json').write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    plan = {'schema_version': 'benchmark-natural-output-plan-v1', 'status': 'awaiting_new_bounded_model_approval',
        'source_manifest_sha256': sha256_file(manifest), 'jobs': jobs, 'job_count': len(jobs),
        'per_domain': dict(Counter(j['domain'] for j in jobs)),
        'independent_source_groups': len(families),
        'reserved_base_tasks': sum(t['split'] == 'human_calibration' for t in data['tasks']),
        'reserved_splits_hash': canonical_json_hash(data['tasks']),
        'actual_natural_artifacts': 0, 'actual_human_labels': 0, 'model_calls': 0,
        'proposed_execution_budget': {'provider': 'codex-cli', 'model': 'gpt-6.1-sol',
            'max_model_calls': 20, 'max_tool_calls': 40, 'max_wall_seconds': 3600,
            'model_tools_enabled': False, 'automatic_retries': 0,
            'output_token_instruction': 4096, 'output_token_hard_cap_available': False},
        'limits': ['Single-shot adapted execution; failures stay recorded and do not trigger extra requests.',
                   'Human calibration families are removed from RSI and OOD before use.',
                   'Reference outputs never substitute for actual generated artifacts.',
                   'Judge calls require labels and a separate approved bounded policy.']}
    (output / 'run-plan.json').write_text(json.dumps(plan, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return {k: v for k, v in plan.items() if k != 'jobs'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--manifest', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(prepare(args.manifest, args.output), ensure_ascii=False))
