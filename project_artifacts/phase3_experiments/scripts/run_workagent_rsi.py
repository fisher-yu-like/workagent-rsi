"""Freeze a project-generated Office RSI invocation; run only after qualified baseline."""

from __future__ import annotations

import json
import argparse
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "src"))

from workagent_rsi.candidate_provider import WorkAgentCandidateProvider
from workagent_rsi.contracts import TaskSpec
from workagent_rsi.evaluator import OfficeArtifactEvaluator
from workagent_rsi.hashing import canonical_json_hash, sha256_file
from workagent_rsi.workagent_experiment import WorkAgentExperimentRunner
from workagent_rsi.workagent_provider import WorkAgentConfig, WorkAgentSkill
from workagent_rsi.office_capabilities import verify_artifact_with_com

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches


RESULT_ROOT = ROOT / "project_artifacts/results/qualification/general-office"
CONFIG = ROOT / "project_artifacts/phase3_experiments/configs/general_office_rsi.json"
BASELINE = RESULT_ROOT / "20260929T195704Z-a4674234"
PILOT_CONFIG = ROOT / "project_artifacts/phase3_experiments/configs/general_office_pilot.json"
SCHEMA = ROOT / "project_artifacts/phase3_experiments/provider/candidate_patch.schema.json"
SKILL = WorkAgentSkill(instructions=(
    "Complete the Office task by creating or editing the requested real file. "
    "Use the staged input copy only. Write the final file inside outputs/, reopen it "
    "with the corresponding Python Office library, and list it in deliverables.json. "
    "Do not report completion until the file and manifest exist."
))
PROVIDER = WorkAgentConfig()
THRESHOLDS = {"develop_gain": 0.05, "regression_tolerance": 0.01,
              "hidden_degradation": 0.01, "ood_degradation": 0.01, "max_cost_delta": 1.0}
IDS = {f"{domain}-{kind}" for domain in ("excel", "word", "powerpoint") for kind in ("create", "edit")}
SHA256 = re.compile(r"[0-9a-f]{64}\Z")
INVOCATION_ID = re.compile(r"[0-9]{8}T[0-9]{6}Z-[0-9a-f]{8}\Z")


def resolve_baseline_invocation(invocation_id: str) -> Path:
    """Choose one regular direct-child qualification directory, never a path."""
    if not INVOCATION_ID.fullmatch(invocation_id):
        raise ValueError("baseline invocation must be a direct-child ID")
    path = RESULT_ROOT / invocation_id
    if path.is_symlink() or not path.is_dir() or path.resolve().parent != RESULT_ROOT.resolve():
        raise ValueError("baseline invocation does not name a regular results directory")
    return path


def _source_inputs(root: Path) -> dict[str, str]:
    inputs = root / "inputs"
    inputs.mkdir()
    book = Workbook()
    sheet = book.active
    sheet.title = "Transactions"
    for row in (("Item", "Revenue"), ("Alpha", 40), ("Beta", 55), ("Gamma", 65)):
        sheet.append(row)
    book.save(inputs / "transactions.xlsx")
    doc = Document()
    doc.add_heading("Draft Status", level=1)
    doc.add_paragraph("The pilot completed on 12 September.")
    doc.add_paragraph("This draft needs a final structure and an action item.")
    doc.save(inputs / "status_draft.docx")
    deck = Presentation()
    for label in ("Baseline: two regions", "Rollout options"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = label
    deck.save(inputs / "briefing_draft.pptx")
    return {path.name: sha256_file(path) for path in inputs.iterdir()}


def _write_json(path: Path, data: object) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def _candidate_runner(command: list[str], cwd: Path, timeout: int, attempt_path: Path) -> subprocess.CompletedProcess[str]:
    """Record launch intent immediately before crossing the subprocess boundary."""
    _write_json(attempt_path, {"attempted_at": datetime.now(timezone.utc).isoformat(),
                               "command": command, "cwd": str(cwd), "timeout_seconds": timeout})
    return subprocess.run(command, cwd=cwd, timeout=timeout, capture_output=True,
                          text=True, encoding="utf-8", errors="replace", check=False)


def _source_office_validation(root: Path, expected_hashes: dict[str, str], com_verify) -> dict:
    """Validate generated inputs only; this is not WorkAgent output qualification."""
    files = []
    for name, domain, reopen in (
        ("transactions.xlsx", "excel", lambda path: load_workbook(path, read_only=True).close()),
        ("status_draft.docx", "word", lambda path: Document(path)),
        ("briefing_draft.pptx", "powerpoint", lambda path: Presentation(path)),
    ):
        path = root / "inputs" / name
        row = {"name": name, "domain": domain, "sha256": sha256_file(path), "library_reopen_ok": False,
               "com_reopen": {"ok": False, "status": "not_run"}}
        try:
            reopen(path)
            row["library_reopen_ok"] = True
        except Exception as exc:
            row["library_error"] = str(exc)
        if row["library_reopen_ok"] and row["sha256"] == expected_hashes[name]:
            try:
                row["com_reopen"] = com_verify(path, domain)
            except Exception as exc:
                row["com_reopen"] = {"ok": False, "status": "probe_exception", "error": str(exc)}
        row["hash_unchanged"] = sha256_file(path) == expected_hashes[name]
        files.append(row)
    result = {"scope": "generated_source_inputs_not_workagent_outputs", "files": files,
              "passed": all(row["library_reopen_ok"] and row["hash_unchanged"] and row["com_reopen"].get("ok") is True for row in files)}
    _write_json(root / "source_office_validation.json", result)
    return result


def _read_json(path: Path) -> dict:
    if path.is_symlink() or not path.is_file():
        raise ValueError(f"missing or unsafe evidence: {path.name}")
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"invalid evidence: {path.name}")
    return value


def validate_baseline(root: Path) -> tuple[dict, list[str]]:
    """Fail closed on missing or contradictory six-task qualification evidence."""
    summary = _read_json(root / "summary.json")
    if summary.get("pilot_id") != "general-office-six-task-engineering-pilot-v1":
        raise ValueError("baseline pilot identity mismatch")
    if Path(summary.get("invocation", "")).resolve() != root.resolve():
        raise ValueError("baseline invocation path mismatch")
    if summary.get("config_sha256") != sha256_file(root / "config.json") or summary["config_sha256"] != sha256_file(PILOT_CONFIG):
        raise ValueError("baseline task config hash mismatch")
    if summary.get("evaluator_source_sha256") != OfficeArtifactEvaluator.evaluator_hash():
        raise ValueError("baseline evaluator hash mismatch")
    rows = summary.get("rows")
    if not isinstance(rows, list) or len(rows) != 6 or {row.get("task_id") for row in rows if isinstance(row, dict)} != IDS:
        raise ValueError("baseline must have exactly six distinct task rows")
    if summary.get("task_count") != 6 or summary.get("success_count") != sum(r["state"] == "SUCCEEDED" for r in rows) or summary.get("failure_count") != sum(r["state"] == "FAILED" for r in rows):
        raise ValueError("baseline row counts contradict summary")
    source_hashes = summary.get("source_hashes")
    if not isinstance(source_hashes, dict) or set(source_hashes) != {"transactions.xlsx", "status_draft.docx", "briefing_draft.pptx"}:
        raise ValueError("baseline source input hash records incomplete")
    provenance = _read_json(root / "input_provenance.json")
    if provenance.get("sha256") != source_hashes or provenance.get("provenance") != "project-generated":
        raise ValueError("baseline input provenance mismatch")
    for name, digest in source_hashes.items():
        if not isinstance(digest, str) or not SHA256.fullmatch(digest) or sha256_file(root / "inputs" / name) != digest:
            raise ValueError(f"baseline source input hash mismatch: {name}")
    gaps = []
    for row in rows:
        task_id = row["task_id"]
        domain, kind = task_id.split("-")
        task_root = root / task_id
        if row.get("domain") != domain or row.get("kind") != kind or row.get("state") not in ("SUCCEEDED", "FAILED"):
            raise ValueError(f"baseline task identity or terminal state invalid: {task_id}")
        if _read_json(task_root / "qualification_result.json") != row:
            raise ValueError(f"baseline persisted row mismatch: {task_id}")
        result = _read_json(task_root / "result.json")
        com = _read_json(task_root / "com_reopen.json")
        if result.get("state") not in ("SUCCEEDED", "FAILED") or row.get("harness_state") != result["state"]:
            raise ValueError(f"baseline harness evidence mismatch: {task_id}")
        input_names = {"excel": ["transactions.xlsx"], "word": ["status_draft.docx"], "powerpoint": ["briefing_draft.pptx"]}[domain] if kind == "edit" else []
        expected = {name: source_hashes[name] for name in input_names}
        if row.get("source_hashes_before") != expected or row.get("source_hashes_after") != expected or row.get("input_hashes_ok") is not True:
            raise ValueError(f"baseline row input hashes invalid: {task_id}")
        if row["state"] == "SUCCEEDED":
            evaluation = _read_json(task_root / "evaluation.json")
            artifacts = result.get("artifacts", [])
            if row.get("evaluation_passed") is not True or evaluation.get("passed") is not True or row.get("com_ok") is not True or com.get("ok") is not True or not row.get("suffix_ok") or len(artifacts) != 1 or len(row.get("artifact_sha256", [])) != 1:
                raise ValueError(f"baseline success lacks evaluator/artifact/COM evidence: {task_id}")
            artifact = Path(artifacts[0]["path"])
            if not artifact.is_file() or sha256_file(artifact) != artifacts[0].get("sha256") or row["artifact_sha256"][0] != artifacts[0]["sha256"]:
                raise ValueError(f"baseline artifact hash invalid: {task_id}")
        else:
            gaps.append(task_id)
    return summary, gaps


def run_invocation(baseline_root: Path, output_root: Path, *, runner_factory=WorkAgentExperimentRunner,
                   source_com_verify=verify_artifact_with_com) -> int:
    """Persist immutable preflight evidence before any model or candidate process."""
    root = Path(output_root)
    root.mkdir(parents=True, exist_ok=False)
    started = datetime.now(timezone.utc)
    config_bytes = CONFIG.read_bytes()
    config = json.loads(config_bytes)
    splits = config["splits"]
    if set(splits) != {"develop", "regression", "hidden", "ood_transfer"} or sum(map(len, splits.values())) != 12:
        raise ValueError("RSI task matrix must contain four splits and twelve tasks")
    tasks = {split: [TaskSpec.model_validate(row) for row in rows] for split, rows in splits.items()}
    (root / "rsi_tasks.json").write_bytes(config_bytes)
    source_hashes = _source_inputs(root)
    split_hashes = {split: canonical_json_hash([task.model_dump(mode="json") for task in rows]) for split, rows in tasks.items()}
    baseline_summary = Path(baseline_root) / "summary.json"
    baseline_summary_error = None
    try:
        baseline_summary_hash = sha256_file(baseline_summary) if baseline_summary.is_file() and not baseline_summary.is_symlink() else None
    except OSError as exc:
        baseline_summary_hash = None
        baseline_summary_error = str(exc)
    contract = {
        "experiment_id": config["experiment_id"], "provenance": "project-generated", "seed": config["seed"],
        "task_config_sha256": sha256_file(root / "rsi_tasks.json"), "split_hashes": split_hashes,
        "source_input_hashes": source_hashes, "evaluator_sha256": OfficeArtifactEvaluator.evaluator_hash(),
        "baseline_skill_sha256": canonical_json_hash(SKILL.model_dump()),
        "provider": "codex-cli/ollama", "model_identity": f"ollama:{PROVIDER.model}",
        "workagent_executable": PROVIDER.executable, "workagent_timeout_seconds": PROVIDER.timeout_seconds,
        "candidate_timeout_seconds": 180, "promotion_thresholds": THRESHOLDS,
        "baseline_invocation": str(Path(baseline_root).resolve()),
        "baseline_summary_sha256": baseline_summary_hash,
    }
    _write_json(root / "contract.json", contract)
    _write_json(root / "baseline_skill.json", SKILL.model_dump())
    source_validation = _source_office_validation(root, source_hashes, source_com_verify)
    reason = f"baseline evidence invalid: {baseline_summary_error}" if baseline_summary_error else None
    baseline = None
    try:
        baseline, gaps = validate_baseline(Path(baseline_root))
        if gaps and reason is None:
            reason = f"六任务基线缺少各格式成功的创建与编辑；失败任务：{', '.join(gaps)}"
    except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        reason = f"baseline evidence invalid: {exc}"
    if not source_validation["passed"]:
        reason = f"{reason}; generated source Office validation failed" if reason else "generated source Office validation failed"
    summary = {
        "status": "blocked" if reason else "running", "block_reason": reason,
        "baseline": {"invocation": str(Path(baseline_root).resolve()),
                     "success_count": baseline.get("success_count") if baseline else None,
                     "failure_count": baseline.get("failure_count") if baseline else None,
                     "task_count": baseline.get("task_count") if baseline else None},
        "rsi_started": False, "candidate_started": False, "token_usage": "unavailable",
        "started_at": started.isoformat(), "ended_at": datetime.now(timezone.utc).isoformat(),
        "wall_time_seconds": (datetime.now(timezone.utc) - started).total_seconds(),
    }
    if reason is None:
        try:
            runner = runner_factory(workagent_config=PROVIDER, input_base=root / "inputs")
            provider = WorkAgentCandidateProvider(SCHEMA, "local-1", executable=PROVIDER.executable,
                timeout_seconds=180, model_identity=f"ollama:{PROVIDER.model}",
                extra_args=["--ignore-user-config", "--oss", "--local-provider", "ollama", "--model", PROVIDER.model],
                runner=lambda command, cwd, timeout: _candidate_runner(command, cwd, timeout, root / "candidate_launch_attempt.json"))
            summary["rsi_started"] = True
            summary["rsi_result"] = runner.run(tasks, SKILL, provider, root / "rsi", rounds=1)
            summary["status"] = summary["rsi_result"].get("status", "incomplete")
        except Exception as exc:
            summary["status"] = "failed"
            summary["failure"] = f"{type(exc).__name__}: {exc}"
            reason = f"RSI execution failed: {exc}"
    summary["candidate_started"] = (root / "candidate_launch_attempt.json").is_file() or (
        any((root / "rsi").rglob("provider_record.json")) if (root / "rsi").exists() else False)
    summary["ended_at"] = datetime.now(timezone.utc).isoformat()
    summary["wall_time_seconds"] = (datetime.now(timezone.utc) - started).total_seconds()
    _write_json(root / "summary.json", summary)
    analysis = {"status": summary["status"], "block_reason": reason, "baseline": summary["baseline"],
                "rsi_started": summary["rsi_started"], "candidate_started": summary["candidate_started"], "claims_allowed": False,
                "external_benchmark": False, "unavailable_is_not_score": True, "token_usage": "unavailable"}
    _write_json(root / "analysis.json", analysis)
    report = ["# 通用 Office WorkAgent RSI 资格报告", "", "本次仅使用项目生成的输入，不是外部 benchmark。", "",
              f"基线：`{baseline_root}`", f"RSI invocation：`{root}`", "",
              f"状态：{'阻断' if summary['status'] == 'blocked' else '失败' if summary['status'] == 'failed' else summary['status']}",
              f"基线：{summary['baseline']['success_count']}/6 成功；失败 {summary['baseline']['failure_count']}/6。",
              f"原因：{reason or '资格门禁通过，已运行一轮 RSI。'}", "",
              "source_office_validation.json 仅检查本次新生成的输入文件，并非 WorkAgent 交付物。", "",
              "未把失败或不可用任务转换为分数；未声明外部 benchmark 成绩。", ""]
    (root / "qualification_report.md").write_text("\n".join(report), encoding="utf-8")
    summary["disk_bytes"] = sum(p.stat().st_size for p in root.rglob("*") if p.is_file())
    _write_json(root / "summary.json", summary)
    return 0 if summary["status"] in {"completed", "no_candidate_needed"} else 1


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline-invocation", default=BASELINE.name, metavar="INVOCATION_ID")
    args = parser.parse_args(argv)
    try:
        baseline = resolve_baseline_invocation(args.baseline_invocation)
    except ValueError as exc:
        parser.error(str(exc))
    invocation = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid4().hex[:8]
    root = RESULT_ROOT / invocation
    result = run_invocation(baseline, root)
    status = _read_json(root / "summary.json")["status"]
    print(json.dumps({"invocation": str(root), "status": status}))
    return result


if __name__ == "__main__":
    raise SystemExit(main())
