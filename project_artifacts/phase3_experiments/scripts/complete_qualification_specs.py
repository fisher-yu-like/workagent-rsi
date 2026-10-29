"""Add source-only executable checks and acceptance plans to the 72 T31 drafts.

No candidate artifacts, reference answers, server jobs or model calls are read.
The author does not sign the independent-review gate. Unresolved scopes remain
explicit, including broad financial-model completion and native rendering.
"""
import argparse
from collections import Counter
import copy
import hashlib
import json
from pathlib import Path
import re
import sys

ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT/'src'))
sys.path.insert(0,str(Path(__file__).parent))
from prepare_qualification_specs import draft_task
from workagent_rsi.assessment_contracts import AcceptanceSpec,Requirement,Location
from workagent_rsi.check_registry import validate_acceptance_spec
from workagent_rsi.hashing import canonical_json_hash,sha256_file
from workagent_rsi.requirement_coverage import admission_for_spec
from workagent_rsi.study_protocol import read_tasks
from workagent_rsi.view_safety import ordinary_file
from workagent_rsi.word_package_checks import _package,_text,NS,W,validate_word_rules


# Literal edits come from public instructions, never from completed outputs.
PAIRS={
 '02a00a9f7bd170512656':[('Alpha Technologies Inc.','Omega Systems Corp.'),('Beta Consulting Group LLC','Sigma Advisory Partners LLC')],
 '55b776efbc1cf3ef361e':[('Entreprise Générale S.A.R.L.','Entreprise Moderne S.A.S.'),('€500,000','€750,000')],
 'cc2fccd701c3ac869dfc':[('First content paragraph','Updated first paragraph')],
 'f4fdde5982eb8ebbc3ad':[('Acme Corp','Nova Industries'),('Draft - Not for Distribution','Final Version')],
 '0ca2e7837a28b8655b68':[('David Park','Jennifer Walsh'),('24 weeks','28 weeks')],
 'c67214beb186618811d4':[('2025','2026'),('Margaret A. Thornton','David R. Kimball'),('$2.3B','$2.7B')],
 '3c94244b1a2888d506c5':[('TechVentures Inc.','Aurora Systems Inc.'),('Delaware','New York')],
 'bd7fe0fab4968dc32780':[('30 minutes','15 minutes'),('12 characters','16 characters'),('100MB','250MB')],
 '8f2ad366a8017d3b62f2':[('Sequoia Capital','Andreessen Horowitz'),('$45M','$60M'),('$180M','$240M'),('Amanda Liu','Marc Benioff')],
 '1e1b47f9e68b4214b87c':[('Pinnacle Strategy Group LLC','[Service Provider]'),('Meridian Healthcare Systems Inc.','[Client Name]'),
 ('350 Fifth Avenue, Suite 4820, New York, NY 10118','[Provider Address]'),('200 Clarendon Street, 52nd Floor, Boston, MA 02116','[Client Address]'),
 ('James Richardson','[Provider Contact]'),('Lisa Park','[Client Contact]'),('April 1, 2026','[Effective Date]'),('May 1, 2026','[Start Date]'),('October 31, 2026','[End Date]'),('$45,000','[Fee Amount]')],
 'e7448d57763245b44271':[('[Enter client name]','Riverside Manufacturing LLC'),('[Enter address]','450 Industrial Blvd, Portland OR 97201'),
 ('[Select matter type]','Contract Dispute'),('[Enter date]','March 15, 2026'),('[Enter opposing party]','Granite Supply Co.')],
}
FINANCIAL_TARGETS={
 '09_03':['EBITDA','Reported EPS (Basic)','Total Capital Spending','Implied P/E','Present Value of Equity','Value per Share'],
 '18_04':['Return on Capital Employed','Average ISCR','Net Cash from Operating Activities','Price Index Multiple','Cost to Concessionaire'],
 '04_05':['Equity Turnover','Tax Rate','NOPAT','Other Long-Term Assets','Beauty','Grooming','Health Care','Snacks','Fabric Care','Baby Care'],
 '10_01':['YoY Growth','EBITDA Margin','Depreciation','Terminal Value','Debt/Equity','Earnings Yield'],
 '15_04':['PBT','EPS Growth','Interest Expenses','Other Non-operating Income','Check'],
}


def word_rules(task):
    identity=task.task_id.split('/')[1];parts=_package(Path(task.input_files[0]));rules=[];limits=[]
    pairs=PAIRS.get(identity,[])
    for old,new in pairs:
        rule={'kind':'replace','scope':'all','before':old,'after':new}
        if identity=='02a00a9f7bd170512656':rule['preserve_bold']=True
        rules.append(rule)
    if pairs:
        # Exact text and paragraph inventory preserves unrelated boilerplate,
        # without claiming visual formatting or all OOXML properties.
        expected={name:[_text(p) for p in xml.xpath('.//w:p',namespaces=NS)] for name,xml in parts.items() if name!='word/comments.xml'}
        for values in expected.values():
            for i,text in enumerate(values):
                for old,new in pairs:text=text.replace(old,new)
                values[i]=text
        rules.append({'kind':'paragraphs_equal','parts':expected})
    if identity=='cc2fccd701c3ac869dfc':rules.append({'kind':'blank_paragraphs'})
    if identity in {'4fde28fde3ca3d1fce6e','a4ed292e8ca46156df5b'}:
        comments=parts.get('word/comments.xml');selected=[]
        for c in comments.findall('w:comment',NS) if comments is not None else []:
            if identity=='4fde28fde3ca3d1fce6e' and c.get('{'+W+'}id')=='0' or identity=='a4ed292e8ca46156df5b' and c.get('{'+W+'}author')=='Sarah Chen' and re.match(r'^(?:Fixed|No longer relevant)\b',_text(c)):selected.append(c.get('{'+W+'}id'))
        wanted=1 if identity=='4fde28fde3ca3d1fce6e' else 2
        if len(selected)==wanted:rules.append({'kind':'comments_preserved','remove_ids':selected})
        else:limits.append('Source comment selection is ambiguous;do not infer IDs from order.')
    if identity=='4fde28fde3ca3d1fce6e':
        rules.extend([{'kind':'replace','scope':'headers','before':'DRAFT','after':'FINAL'},
            {'kind':'footnote_suffix','note_id':'1','suffix':' as amended in 2026'}])
    if identity=='394d4298a0f47f54e059':
        rules.append({'kind':'comments','author':'Reviewer','min_count':3})
        limits.append('Anchored comment count does not prove specific actionable critique of the four named clauses;semantic comment-text review is required.')
    if identity=='788ee5bc9b61152ede71':
        rules.append({'kind':'tracked_clean'})
        # w:delText is omitted and w:ins/w:t retained in this accepted-text view.
        expected=[_text(p) for p in parts['word/document.xml'].xpath('.//w:p',namespaces=NS)]
        expected=[re.sub(r'\{\+(.*?)\+\}',r'\1',re.sub(r'\[-.*?-\]','',t)) for t in expected]
        rules.append({'kind':'paragraphs_equal','parts':{'word/document.xml':expected}})
    tables=[[[_text(c) for c in row.findall('w:tc',NS)] for row in t.findall('w:tr',NS)] for t in parts['word/document.xml'].xpath('.//w:tbl',namespaces=NS)]
    updated=copy.deepcopy(tables);changed=False
    for table in updated:
        for i,row in list(enumerate(table)):
            if identity=='4fde28fde3ca3d1fce6e' and 'Enterprise' in row:
                if '$500' in row:row[row.index('$500')]='$750';changed=True
            if identity=='5c3287777b802adfb021':
                if 'Revenue' in row and '$1.2M' in row:row[row.index('$1.2M')]='$1.5M';changed=True
                if 'Operating Expenses' in row:table.insert(i+1,['Cloud Infrastructure','$120K','$145K']);changed=True
            if identity=='77240a705ffdbfe894f1' and 'West' in row:
                table.insert(i+1,['International','$600','$750','$820','$2,170','$2,000','+8.5%']);changed=True
            if identity=='f70a53d16d35d89c086b' and 'Professional' in row and '$1,999' in row:
                row[row.index('$1,999')]='$2,499'
                rules.append({'kind':'table_row','cells':row,'table_contains':'Monthly Cost'})
    if changed:rules.append({'kind':'tables_equal','tables':updated})
    if identity=='b47b04d19bea10fbc814':
        paragraphs=[_text(p) for p in parts['word/document.xml'].xpath('.//w:p',namespaces=NS)]
        before=[p for p in paragraphs if re.match(r'^ARTICLE 8\b',p,re.I)]
        after=[p for p in paragraphs if re.match(r'^ARTICLE 9\b',p,re.I)]
        if len(before)==len(after)==1:
            rules.append({'kind':'ordered_text','scope':'body','texts':[before[0],'ARTICLE 8A: DATA PROTECTION','8A.1 Data Processing.','8A.2 Security Measures.','8A.3 Breach Notification.','8A.4 Data Deletion.',after[0]]})
        else:limits.append('Article 8/9 source anchors are ambiguous;ordering cannot be inferred.')
        limits.append('Order and labels are necessary;full clause wording and survival cross-reference still require grounded semantic review.')
    if identity=='dbcd47f3ef2e36413aec':
        paragraphs=[_text(p) for p in parts['word/document.xml'].xpath('.//w:p',namespaces=NS)]
        indices=[i for i,t in enumerate(paragraphs) if t.strip().lower().endswith('exclusivity')]
        if len(indices)==1 and indices[0]+1<len(paragraphs):
            i=indices[0];rules.append({'kind':'paragraphs_equal','parts':{'word/document.xml':paragraphs[:i]+paragraphs[i+2:]}})
        else:limits.append('Exclusivity section boundaries need explicit source review.')
    if identity=='f70a53d16d35d89c086b':
        rules.extend([{'kind':'table_row','cells':['PATCH','/accounts/{id}','Partial update','50/min'],'table_contains':'/accounts'},
            {'kind':'table_row','cells':['tags','string[]','Yes','User-defined classification tags'],'table_contains':'parent_id'},
            {'kind':'replace','scope':'body','before':'3.2.1','after':'3.3.0'},
            {'kind':'text','scope':'body','text':'March 25, 2026'},{'kind':'text','scope':'body','text':'Released'}])
        limits.append('Version/date/status need exact source-located metadata binding;presence of words elsewhere is only a necessary check.')
    if not rules:limits.append('This task requires a source-grounded semantic edit plan,not a global literal replacement.')
    if 'gray fields' in task.instruction:limits.append('Placeholder text is checked;retaining gray-field shading needs a formatting check.')
    validate_word_rules(rules) if rules else None
    return rules,limits


def excel_bindings(task):
    from openpyxl import load_workbook
    source=ordinary_file(task.input_files[0]);book=load_workbook(source,data_only=False)
    identity=task.task_id.split('/')[-2];targets=FINANCIAL_TARGETS.get(identity,[]) if '/Financial_Model/' in task.task_id else []
    normalize=lambda t:re.sub('[^a-z0-9]','',t.lower())
    bindings=[];sheets=[];warnings=[]
    try:
        for sheet in book:
            if len(sheet._cells)>200000:raise ValueError('source cell inspection budget exceeded')
            labels=[c for c in sheet._cells.values() if isinstance(c.value,str) and len(c.value)<=180 and not c.value.startswith('=')]
            years=[{'cell':c.coordinate,'text':c.value} for c in labels if re.fullmatch(r'(?:FY\s*)?(?:19|20)?\d{2}[AFE]?',c.value.strip(),re.I)]
            sheets.append({'sheet':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'year_labels':years[:200],
                'formula_count':sum(c.data_type=='f' for c in sheet._cells.values()),'chart_count':len(sheet._charts),
                'validation_count':len(sheet.data_validations.dataValidation)})
            for target in targets:
                matches=[c for c in labels if normalize(target) in normalize(c.value)]
                for c in matches:
                    # Row label coordinates are evidence, not automatic permission
                    # to edit every cell in that row or a resolved numerical oracle.
                    bindings.append({'requested_metric':target,'sheet':sheet.title,'label_cell':c.coordinate,'label_text':c.value,
                        'target_row':c.row,'target_columns':'bind from the explicit period instruction and source year headers',
                        'numerical_oracle_status':'not_derived_from_candidate_or_gold'})
        warnings=['A label match does not establish target columns,dependencies or an independent numerical oracle.',
            'Blank styled cells are not automatically answer cells. Broad audits and fill-all-empty tasks require explicit scope.',
            'Formula text/cached values alone do not prove recalculation;qualified engine and counterfactual checks remain required.']
    finally:book.close()
    return {'sheets':sheets,'metric_label_bindings':bindings,'limitations':warnings}


def complete_task(task):
    bindings=json.loads((Path(__file__).parent.parent/'configs/t31_instruction_bindings.json').read_text(encoding='utf-8'))
    if bindings['instruction_sha256'].get(task.task_id)!=hashlib.sha256(task.instruction.encode('utf-8')).hexdigest():
        raise ValueError('authored rules are not authorized for a different instruction: '+task.task_id)
    row=draft_task(task);spec=AcceptanceSpec.model_validate(row['spec']);checks=list(spec.requirements)
    inventory=spec.requirement_inventory.model_dump(mode='json');plans=[];extra=[];limits=[]
    def check(name,kind,expected,options=None):
        item=Requirement(requirement_id=name,description='Source-grounded necessary acceptance control: '+name,
            check=kind,location=Location(artifact='output'),expected=expected,options=options or {},critical=True,
            dimension='correctness',evidence_source='Public instruction and frozen authorized inputs;no completed outputs used')
        checks.append(item);extra.append(name);return name
    if task.domain=='word':
        rules,limits=word_rules(task)
        if rules:check('word-native-edits','word.package',rules,{'source_path':task.input_files[0]})
        row['native_word_rules']=rules
        exact_ids=set(PAIRS)-{'e7448d57763245b44271','1e1b47f9e68b4214b87c'}
        if rules and (task.task_id.split('/')[1] in exact_ids or task.task_id.split('/')[1] in {'5c3287777b802adfb021','77240a705ffdbfe894f1','a4ed292e8ca46156df5b','788ee5bc9b61152ede71','dbcd47f3ef2e36413aec','4fde28fde3ca3d1fce6e'}):
            for unit,inv in zip(row['obligations'],inventory['requirements']):
                if task.task_id.split('/')[1]=='788ee5bc9b61152ede71' and 'markup' not in unit['source_quote'].lower():continue
                unit.update(proposed_method='deterministic',check_ids=['word-native-edits'],gap_kind=None,gap_reason=None)
                inv.update(validation_method='deterministic',check_ids=['word-native-edits'])
    elif task.domain=='powerpoint':
        refs=[{'path':p,'sha256':spec.input_hashes[p]} for p in task.input_files if Path(p).suffix.lower()=='.png']
        if len(refs)!=1:raise ValueError('PPT source requires one unambiguous frozen PNG')
        cid=check('reference-reconstruction','visual',{'obligation':'Compare all source text,nodes,connectors,arrow direction,spatial relationships and process semantics against the reconstructed slide;do not accept decoy native objects.'},
            {'pages':[1],'reference_images':refs})
        row['reference_binding']=refs
        for unit,inv in zip(row['obligations'],inventory['requirements']):
            quote=unit['source_quote']
            caption_start=task.instruction.find('Source caption:')
            if quote.startswith('Preserve its text,') or caption_start>=0 and unit['start']>=caption_start:
                unit.update(proposed_method='visual',check_ids=[cid],gap_kind=None,gap_reason=None)
                inv.update(validation_method='visual',check_ids=[cid])
            elif quote=='No external raster resources are authorized for this task.':
                # The source task explicitly prohibits images. This control is
                # sufficient for that prohibition,not for diagram reconstruction.
                unit.update(proposed_method='deterministic',check_ids=['native-structure'],gap_kind=None,gap_reason=None)
                inv.update(validation_method='deterministic',check_ids=['native-structure'])
        limits.append('Source-bound visual checks require actual rendering and approved image-capable judge execution. Native object counts are necessary only;independent review must validate editable reconstruction coverage.')
    else:
        row['excel_source_bindings']=excel_bindings(task)
        names=[]
        for name in ('Graph Average Patrol Time','Forecasting Summary Data','Summary','Data'):
            if f'"{name}"' in task.instruction or f'“{name}”' in task.instruction:names.append(name)
        if names:check('declared-worksheet-names','excel.sheets',names)
        limits.extend(row['excel_source_bindings']['limitations'])
    for unit in row['obligations']:
        quote=unit['source_quote'];gaps=[]
        if task.domain=='word':
            procedure='Inspect body,table,header,footer,footnote and comments as applicable. Compare exact literal changes and original unrelated text;use anchored comments and native revision checks where requested.'
            if any(term in quote.lower() for term in ('perspective','problematic','emphasize','recommendation','regulatory','legal','survival')):gaps.append('semantic_content_and_scope_review')
        elif task.domain=='excel':
            procedure='Freeze exact sheets,cells,periods and authorized edit scope from source labels. Derive numerical expectations independently from source assumptions. Recalculate in isolation;test changed inputs for dynamic requirements. Inspect native chart/data-validation structure plus actual appearance.'
            gaps.append('explicit_target_cells_and_independent_oracle' if any(w in quote.lower() for w in ('calculat','model','empty','audit','dynamic','summary')) else 'native_chart_and_render_validation')
        else:procedure='Jointly compare the frozen source PNG and actual slide render;verify native objects separately. Caption is source context and does not substitute for the diagram.'
        plans.append({'requirement_id':unit['requirement_id'],'source_quote':quote,'source_span':[unit['start'],unit['end']],
            'mandatory':True,'acceptance_procedure':procedure,'mapped_checks':unit['check_ids'],
            'supporting_necessary_controls':extra,'unresolved':gaps or ([unit['gap_kind']] if unit['gap_kind'] else []),
            'negative_control':'Reject missing or incorrect required content;do not infer success from file readability or a self-reported completion.'})
    referenced={cid for item in inventory['requirements'] for cid in item['check_ids']}
    checks=[r for r in checks if not r.requirement_id.startswith('check-') or r.requirement_id in referenced]
    spec=spec.model_copy(update={'version':'t31-source-bound-rules-v3','requirements':checks,
        'requirement_inventory':spec.requirement_inventory.model_validate(inventory),
        'scope_note':'Source-grounded executable necessary controls and explicit acceptance procedures. Independent review is pending;all mandatory gaps remain blocking. No benchmark answer or candidate output used.'})
    row.update(spec=spec.model_dump(mode='json'),admission=admission_for_spec(spec).model_dump(mode='json'),
        capability_preflight=validate_acceptance_spec(spec,{}).model_dump(mode='json'),acceptance_plan=plans,
        necessary_controls=extra,remaining_assembly=limits,live_ready=False)
    return row


def prepare(tasks_file,output):
    batches,families=read_tasks(tasks_file);owners={v:f for f in families for v in f.variant_ids}
    tasks=[t for rows in batches.values() for t in rows]
    for t in tasks:
        if owners[t.task_id].variant_input_hashes.get(t.task_id)!=canonical_json_hash(sorted(sha256_file(ordinary_file(p)) for p in t.input_files)):
            raise ValueError('frozen family input changed: '+t.task_id)
    rows=[]
    for split,values in batches.items():
        for task in values:rows.append({**complete_task(task),'split':split,'family_id':owners[task.task_id].family_id})
    output=Path(output)
    if output.exists():raise FileExistsError(output)
    output.mkdir(parents=True)
    for row in rows:
        row['file']=canonical_json_hash(row['task_id'])[:20]+'.json'
        (output/row['file']).write_text(json.dumps(row,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    report={'schema_version':'t31-source-bound-acceptance-rules-v3','source_bundle':str(tasks_file),
        'source_bundle_sha256':sha256_file(tasks_file),'task_count':len(rows),
        'domain_counts':dict(Counter(r['domain'] for r in rows)),
        'source_units':sum(len(r['obligations']) for r in rows),
        'routes':dict(Counter(o['proposed_method'] for r in rows for o in r['obligations'])),
        'necessary_controls':sum(len(r['necessary_controls']) for r in rows),
        'word_rule_count':sum(len(r.get('native_word_rules',[])) for r in rows),
        'excel_metric_bindings':sum(len(r.get('excel_source_bindings',{}).get('metric_label_bindings',[])) for r in rows),
        'reference_bound_tasks':sum('reference_binding' in r for r in rows),
        'admission_counts':dict(Counter(r['admission']['status'] for r in rows)),
        'model_calls':0,'server_jobs':0,'reviewer_signatures_created':0,'candidate_outputs_read':False,'reference_answers_read':False,
        'tasks':[{'task_id':r['task_id'],'domain':r['domain'],'file':r['file'],'sha256':sha256_file(output/r['file']),
            'uncovered_units':sum(o['proposed_method']=='uncovered' for o in r['obligations'])} for r in rows]}
    (output/'report.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    return report


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--tasks',type=Path,required=True);p.add_argument('--output',type=Path,required=True)
    a=p.parse_args();result=prepare(a.tasks,a.output);print(json.dumps({k:v for k,v in result.items() if k!='tasks'}))
