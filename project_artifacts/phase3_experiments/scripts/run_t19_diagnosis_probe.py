"""Collect genuine failures of fixed trusted authored local programs.

These small constant fixtures never contain model-produced code. This is local
diagnosis evidence, not execution-sandbox or live-agent qualification.
"""
from datetime import datetime,timezone
import json
from pathlib import Path
import subprocess
import sys

REPO=next(p for p in Path(__file__).resolve().parents if (p/'pyproject.toml').is_file())
sys.path.insert(0,str(REPO/'src'))
from openpyxl import Workbook,load_workbook
from workagent_rsi.assessment_contracts import ArtifactIssue,Location
from workagent_rsi.environment_manifest import capture_environment
from workagent_rsi.hashing import sha256_file
from workagent_rsi.run_contracts import ExecutionOutcome
from workagent_rsi.trace_diagnosis import diagnose_failure


def evidence(kind,text):
    import hashlib
    return {'kind':kind,'name':'agent.py' if kind=='script' else kind,'text':text,
        'sha256':hashlib.sha256(text.encode()).hexdigest(),'authorized':True,'split':'develop','feedback_visibility':'develop'}


def main():
    root=REPO/'project_artifacts/results/hardening-v2/T19-probe'/datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    root.mkdir(parents=True);capture_environment(REPO,root/'environment.json')
    started=datetime.now(timezone.utc).isoformat()
    sources={
        'syntax':'for value in [1,2]\n    print(value)\n',
        'api':'from openpyxl import Workbook\nws=Workbook().active\nws.append_rows([[1]])\n',
        'sheet':'from openpyxl import Workbook\nwb=Workbook()\nws=wb["Missing"]\n',
        'range':'from openpyxl import load_workbook,Workbook\nwb=load_workbook("inputs/data.xlsx")\nws=wb["Data"]\ntotal=0\nfor row in range(2, ws.max_row):\n    total += ws.cell(row, 2).value\nout=Workbook()\nout.active.title="Data"\nout.active["D2"]=total\nout.save("outputs/answer.xlsx")\n',
        'overwrite':'from openpyxl import Workbook\nimport hashlib,json\ndef receipt():\n    print(json.dumps({"operation":"write","path":"outputs/answer.xlsx","sha256":hashlib.sha256(open("outputs/answer.xlsx","rb").read()).hexdigest()}))\nbook=Workbook()\nbook.active.title="Data"\nbook.active["D2"]=32\nbook.save("outputs/answer.xlsx")\nreceipt()\nother=Workbook()\nother.active.title="Data"\nother.active["D2"]=0\nother.save("outputs/answer.xlsx")\nreceipt()\n',
    }
    rows=[]
    for name,source in sources.items():
        work=root/name;(work/'inputs').mkdir(parents=True);(work/'outputs').mkdir()
        script=work/'agent.py';script.write_text(source,encoding='utf-8')
        book=Workbook();book.active.title='Data'
        for i,value in enumerate([10,15,7],2):book.active.cell(i,2,value)
        book.save(work/'inputs/data.xlsx');book.close()
        command=[sys.executable,'-X','utf8',str(script)]
        run=subprocess.run(command,cwd=work,capture_output=True,text=True,encoding='utf-8',timeout=20)
        (work/'stdout.txt').write_text(run.stdout,encoding='utf-8');(work/'stderr.txt').write_text(run.stderr,encoding='utf-8')
        items=[evidence('script',source),evidence('stderr',run.stderr)]
        observed=None;issue=None
        if name in {'range','overwrite'}:
            book=load_workbook(work/'outputs/answer.xlsx',data_only=True)
            observed=book['Data']['D2'].value;book.close()
            assert observed!=32
            issue=ArtifactIssue(issue_id=name,requirement_id='sum',kind='artifact_defect',location=Location(artifact='output',sheet='Data',cell='D2'),
                observed=observed,expected=32,evidence_refs=['independent-input-oracle'],severity='major',repair_hint='Verify sum.',recheck_ids=['sum'],feedback_visibility='develop')
            if name=='range':items.append(evidence('input_manifest',json.dumps({'sheet':'Data','first_data_row':2,'last_data_row':4,'row_count':3,'max_row':4})))
            else:
                events=[{**json.loads(line),'sequence':index} for index,line in enumerate(run.stdout.splitlines(),1)]
                items.append(evidence('file_events',json.dumps(events)))
        outcome=ExecutionOutcome(status='succeeded' if run.returncode==0 else 'failed',failure_origin=None if run.returncode==0 else 'agent_execution',
            artifact_refs=['outputs/answer.xlsx'] if observed is not None else [])
        trace={'run_id':'authored-'+name,'split':'develop','hidden_test':False,'evidence':items}
        diagnoses=diagnose_failure(issue,outcome,trace)
        assert diagnoses and 'unknown' not in diagnoses[0].causal_hypothesis.lower()
        rows.append({'case':name,'command':command,'cwd':str(work),'exit_code':run.returncode,
            'script_sha256':sha256_file(script),'observed_value':observed,'expected_value':32 if issue else None,
            'trace':trace,'execution':outcome.model_dump(mode='json'),'diagnoses':[d.model_dump(mode='json') for d in diagnoses]})
    record={'status':'passed','exit_code':0,'scope':__doc__,'started_at':started,'ended_at':datetime.now(timezone.utc).isoformat(),
        'command':[sys.executable,*sys.argv],'commit':subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip(),
        'script_sha256':sha256_file(Path(__file__)),'cases':rows}
    (root/'trace-linked-cases.json').write_text(json.dumps(record,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'evidence':str(root.relative_to(REPO)),'status':'passed','cases':len(rows)}))


if __name__=='__main__':main()
