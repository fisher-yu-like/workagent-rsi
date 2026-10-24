"""Run authored T05 Office fixtures through an explicitly frozen Apptainer worker.

This is infrastructure qualification, not an agent benchmark or scoring run.
No host Office is launched. Every invocation creates a new evidence directory.
"""
from datetime import datetime, timezone
import argparse
import json
from pathlib import Path
import subprocess
import sys
import shutil

REPO = next(p for p in Path(__file__).resolve().parents if (p / 'pyproject.toml').is_file())
sys.path.insert(0, str(REPO / 'src'))
from workagent_rsi.apptainer_sandbox import ApptainerPolicy, ApptainerSSHExecutor
from workagent_rsi.artifact_renderer import OfficeArtifactRenderer, RenderPolicy
from workagent_rsi.hashing import sha256_file
from workagent_rsi.candidate_workspace import CandidateWorkspaceBuilder
from workagent_rsi.contracts import TaskSpec
from workagent_rsi.workagent_office import prepare_execution_view


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2) + '\n', encoding='utf-8')


def fixtures(root):
    from openpyxl import Workbook
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    root.mkdir()
    book = Workbook()
    sheet = book.active
    sheet.title = 'T05'
    sheet.append(['T05 isolated rendering', 'Value'])
    sheet.append(['First', 7])
    sheet.append(['Second', 11])
    sheet.append(['Sum', '=SUM(B2:B3)'])
    sheet.column_dimensions['A'].width = 30
    sheet.column_dimensions['B'].width = 15
    book.save(root / 'fixture.xlsx')
    doc = Document()
    doc.add_heading('T05 isolated rendering', 0)
    doc.add_paragraph('Authored infrastructure fixture. Source bytes must remain unchanged.')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Input'
    table.cell(0, 1).text = 'Verified copy'
    doc.save(root / 'fixture.docx')
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = 'T05 isolated rendering'
    slide.shapes.add_textbox(Inches(1), Inches(3), Inches(8), Inches(1)).text = 'Private worker; owned cleanup; unchanged source.'
    deck.save(root / 'fixture.pptx')
    return {'excel': root / 'fixture.xlsx', 'word': root / 'fixture.docx', 'powerpoint': root / 'fixture.pptx'}


def boundary(root, source):
    # Sentinels are authored fixtures, never real hidden evaluator data.
    source.joinpath('hidden.jsonl').write_text('T05_AUTHORED_HIDDEN_SENTINEL', encoding='utf-8')
    source.joinpath('skill.json').write_text('{"instructions":"Use only declared inputs"}', encoding='utf-8')
    task = TaskSpec(task_id='t05-authored-boundary', domain='excel', instruction='Read the supplied workbook.',
                    input_files=('fixture.xlsx',))
    view = root / 'execution-view'
    manifest = prepare_execution_view(task, source, view)
    candidate = root / 'candidate-view'
    exported = CandidateWorkspaceBuilder().export(source, candidate, ['skill.json'], {})
    files = lambda directory: sorted(p.relative_to(directory).as_posix() for p in directory.rglob('*')
                                     if p.is_file() and '.git' not in p.relative_to(directory).parts)
    actual = files(view)
    expected = ['input_manifest.json', 'inputs/0001-fixture.xlsx']
    candidate_files = files(candidate)
    passed = actual == expected and candidate_files == ['AGENTS.md', 'skill.json']
    record = {'scope': 'authored local allowlist views', 'passed': passed, 'execution_files': actual,
              'expected_execution_files': expected, 'candidate_files': candidate_files,
              'input_manifest': manifest, 'candidate_manifest': exported.model_dump(mode='json')}
    write(root / 'sandbox/evaluation-boundary.json', record)
    return passed


PROCESS_PROBE = r'''
import json,runpy,subprocess,tempfile,time
from pathlib import Path
import uno
config=json.loads(Path('inputs/render-policy.json').read_text())
report={}
with tempfile.TemporaryDirectory(prefix='wrsi-sentinel-') as profile:
    pipe='wrsi_authored_sentinel'
    args=['/usr/lib/libreoffice/program/soffice.bin','-env:UserInstallation='+Path(profile).as_uri(),
          '--headless','--norestore','--nodefault','--nologo',
          '--accept=pipe,name='+pipe+';urp;StarOffice.ComponentContext']
    sentinel=subprocess.Popen(args,stdin=subprocess.DEVNULL,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    remote=None
    try:
        context=uno.getComponentContext()
        resolver=context.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver',context)
        deadline=time.monotonic()+25
        restarted=False
        while time.monotonic()<deadline:
            code=sentinel.poll()
            if code==81 and not restarted:
                sentinel.wait()
                sentinel=subprocess.Popen(args,stdin=subprocess.DEVNULL,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
                restarted=True
            elif code is not None: break
            try:
                remote=resolver.resolve('uno:pipe,name='+pipe+';urp;StarOffice.ComponentContext')
                break
            except Exception: time.sleep(.1)
        if remote is None: raise RuntimeError('authored sentinel did not start')
        before={'pid':sentinel.pid,'start_ticks':Path('/proc/'+str(sentinel.pid)+'/stat').read_text().split()[21]}
        result=runpy.run_path('inputs/render-worker.py')['render'](config)
        after={'pid':sentinel.pid,'start_ticks':Path('/proc/'+str(sentinel.pid)+'/stat').read_text().split()[21]}
        report.update(render_status=result['status'],render_owned_cleanup=result.get('owned_cleanup'),
                      unrelated_before=before,unrelated_after=after,
                      unrelated_survived=(sentinel.poll() is None and before==after))
    finally:
        if remote is not None:
            try: remote.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop',remote).terminate()
            except Exception: pass
        if sentinel.poll() is None: sentinel.terminate()
        try: sentinel.wait(timeout=5)
        except subprocess.TimeoutExpired:
            sentinel.kill();sentinel.wait(timeout=5)
        report['authored_sentinel_cleaned']=sentinel.poll() is not None
        Path('outputs/process-check.json').write_text(json.dumps(report,indent=2))
'''

TIMEOUT_PROBE = r'''
import subprocess,time,tempfile
from pathlib import Path
profile=tempfile.mkdtemp(prefix='wrsi-authored-timeout-')
child=subprocess.Popen(['/usr/lib/libreoffice/program/soffice','-env:UserInstallation='+Path(profile).as_uri(),
                        '--headless','--norestore','--nodefault','--nologo',
                        '--accept=pipe,name=wrsi_timeout;urp;StarOffice.ComponentContext'],start_new_session=True)
time.sleep(3)
print('owned detached Office launcher pid='+str(child.pid)+' poll='+str(child.poll()),flush=True)
time.sleep(120)
'''


def process_checks(root, policy, source):
    work = root / 'process-checks/normal'
    (work / 'inputs').mkdir(parents=True)
    shutil.copyfile(source, work / 'inputs/artifact.xlsx')
    shutil.copyfile(REPO/'src/workagent_rsi/libreoffice_worker.py', work/'inputs/render-worker.py')
    write(work/'inputs/render-policy.json', {'source':'inputs/artifact.xlsx','source_sha256':sha256_file(source),
          'domain':'excel','dpi':72,'max_pages':40,'policy':RenderPolicy().model_dump(mode='json')})
    script=work/'authored_cleanup.py'
    script.write_text(PROCESS_PROBE,encoding='utf-8')
    normal=ApptainerSSHExecutor(policy).run(script,work,90)
    report=json.loads((work/'outputs/process-check.json').read_text()) if (work/'outputs/process-check.json').exists() else {}
    timeout_work=root/'process-checks/timeout'
    timeout_work.mkdir()
    script=timeout_work/'authored_timeout.py'
    script.write_text(TIMEOUT_PROBE,encoding='utf-8')
    timeout=ApptainerSSHExecutor(policy).run(script,timeout_work,10)
    trace=json.loads((timeout_work/timeout.trace_refs[0]).read_text())
    passed=(normal.status=='succeeded' and report.get('render_status')=='available'
            and report.get('render_owned_cleanup') is True and report.get('unrelated_survived') is True
            and report.get('authored_sentinel_cleaned') is True and timeout.status=='timeout'
            and trace.get('cleanup',{}).get('confirmed_absent') is True
            and 'poll=None' in trace.get('stdout',''))
    record={'passed':passed,'normal':normal.model_dump(mode='json'),'normal_report':report,
            'timeout':timeout.model_dump(mode='json'),'timeout_cleanup':trace.get('cleanup'),
            'timeout_stdout':trace.get('stdout')}
    write(root/'process-cleanup-tests.json',record)
    (root/'process-cleanup-tests.txt').write_text(json.dumps(record,indent=2)+'\n',encoding='utf-8')
    return passed


def remote_boundary_checks(root, policy, source):
    work=root/'remote-boundary'
    (work/'inputs').mkdir(parents=True)
    shutil.copyfile(source,work/'inputs/declared.xlsx')
    (work/'hidden.jsonl').write_text('T05_AUTHORED_LOCAL_HIDDEN_SENTINEL',encoding='utf-8')
    script=work/'authored_boundary.py'
    script.write_text('''import json,os,sys,hashlib
from pathlib import Path
try:
    Path('inputs/forbidden-write').write_text('probe')
    readonly=False
except OSError:
    readonly=True
report={'host_sentinel_invisible':not Path(sys.argv[1]).exists(),
        'dummy_environment_absent':'WRSI_PROBE_DUMMY_SECRET' not in os.environ,
        'hidden_sibling_invisible':not Path('/task/hidden.jsonl').exists(),
        'task_files':sorted(p.relative_to('/task').as_posix() for p in Path('/task').rglob('*') if p.is_file()),
        'input_readonly':readonly,'input_sha256':hashlib.sha256(Path('inputs/declared.xlsx').read_bytes()).hexdigest()}
Path('outputs/boundary.json').write_text(json.dumps(report,indent=2))
''',encoding='utf-8')
    outcome=ApptainerSSHExecutor(policy,authored_boundary_probe=True).run(script,work,20)
    report=json.loads((work/'outputs/boundary.json').read_text()) if (work/'outputs/boundary.json').exists() else {}
    passed=(outcome.status=='succeeded' and all(report.get(k) is True for k in
            ['host_sentinel_invisible','dummy_environment_absent','hidden_sibling_invisible','input_readonly'])
            and report.get('input_sha256')==sha256_file(source)
            and report.get('task_files')==['authored_boundary.py','inputs/declared.xlsx'])
    path=root/'sandbox/evaluation-boundary.json'
    local=json.loads(path.read_text())
    local.update(remote={'passed':passed,'report':report,'outcome':outcome.model_dump(mode='json')})
    write(path,local)
    return passed


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--policy', type=Path, required=True)
    parser.add_argument('--domain', choices=['all', 'excel', 'word', 'powerpoint'], default='all')
    parser.add_argument('--timeout', type=int, default=90)
    parser.add_argument('--process-checks', action='store_true', help='Real unrelated Office survivor and detached timeout probes')
    args = parser.parse_args()
    policy = ApptainerPolicy.model_validate_json(args.policy.read_text(encoding='utf-8'))
    root = REPO / 'project_artifacts/results/hardening-v2/T05' / datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    root.mkdir(parents=True, exist_ok=False)
    paths = fixtures(root / 'fixtures')
    record = {'command': [sys.executable, *sys.argv], 'cwd': str(REPO),
              'started_at': datetime.now(timezone.utc).isoformat(),
              'git_commit': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=REPO, text=True).strip(),
              'source_hashes': {p.relative_to(REPO).as_posix(): sha256_file(p) for p in [
                  Path(__file__), REPO/'src/workagent_rsi/artifact_renderer.py',
                  REPO/'src/workagent_rsi/libreoffice_worker.py', REPO/'src/workagent_rsi/apptainer_sandbox.py',
                  REPO/'src/workagent_rsi/apptainer_worker.py']},
              'boundary_passed': boundary(root, root / 'fixtures'), 'renders': {}}
    write(root / 'sandbox-policy.json', policy.model_dump(mode='json'))
    worker = ApptainerSSHExecutor(policy)
    renderer = OfficeArtifactRenderer(root / 'renders', sandbox=worker, timeout_seconds=args.timeout)
    write(root / 'office-worker-policy.json', renderer.identity_config())
    write(root / 'probe.json', worker.probe())
    for domain, path in paths.items():
        if args.domain != 'all' and args.domain != domain:
            continue
        before = sha256_file(path)
        result = renderer.render(path, domain, before)
        row = {'result': result.model_dump(mode='json'), 'source_sha256_before': before,
               'source_sha256_after': sha256_file(path)}
        row['passed'] = result.status == 'available' and before == row['source_sha256_after']
        record['renders'][domain] = row
        write(root / 'run.json', record)
        print(json.dumps({'domain': domain, 'status': result.status, 'error': result.error,
                          'evidence': str(root.relative_to(REPO))}), flush=True)
        if not row['passed']:
            break
    record['ended_at'] = datetime.now(timezone.utc).isoformat()
    record['passed'] = record['boundary_passed'] and bool(record['renders']) and all(r['passed'] for r in record['renders'].values())
    if args.process_checks and record['passed']:
        record['remote_boundary_passed']=remote_boundary_checks(root,policy,paths['excel'])
        record['process_checks_passed']=process_checks(root,policy,paths['excel'])
        record['passed'] = record['passed'] and record['process_checks_passed'] and record['remote_boundary_passed']
        record['ended_at'] = datetime.now(timezone.utc).isoformat()
    write(root / 'run.json', record)
    return 0 if record['passed'] else 1


if __name__ == '__main__':
    raise SystemExit(main())
