"""Fixed trusted local probe program. It never evaluates skill text or task IDs."""
import json
from pathlib import Path

request=json.loads(Path('scratch/request.json').read_text(encoding='utf-8'))
domain=request['domain'];operation=request['operation'];inputs=request['inputs']
Path('outputs').mkdir(exist_ok=True)
if domain=='excel':
    from openpyxl import load_workbook
    book=load_workbook(inputs[0]);source=next(s for s in book if s.title!='Result')
    header=[c.value for c in source[1]];column=header.index(request['value_header'])+1
    values=[row[column-1] for row in source.iter_rows(min_row=2,values_only=True) if isinstance(row[column-1],(int,float))]
    summary=book.create_sheet('Result');summary['A1']=sum(values)
    path=Path('outputs/result.xlsx');book.save(path);book.close()
    check=load_workbook(path);assert check['Result']['A1'].value is not None;check.close()
elif domain=='word':
    from docx import Document
    doc=Document(inputs[0])
    if operation=='replace_paragraph':
        target=next(p for p in doc.paragraphs if p.text==request['target_text']);target.text=request['replacement']
    elif operation=='append_heading':
        doc.add_heading(request['heading'],level=request['level'])
    elif operation=='update_table':
        table=doc.tables[0];headers=[c.text for c in table.rows[0].cells]
        key=headers.index(request['key_header']);value=headers.index(request['value_header'])
        row=next(row for row in list(table.rows)[1:] if row.cells[key].text==request['key'])
        row.cells[value].text=request['replacement']
    else:raise ValueError('unknown Word operation')
    path=Path('outputs/result.docx');doc.save(path);assert Document(path).paragraphs
else:
    from pptx import Presentation
    from pptx.util import Inches
    deck=Presentation(inputs[0]) if inputs else Presentation()
    def text(slide,value,top=.4):
        slide.shapes.add_textbox(Inches(.5),Inches(top),Inches(8),Inches(.7)).text=value
    if operation=='create_slides':
        for index in range(request['count']):text(deck.slides.add_slide(deck.slide_layouts[6]),request['prefix']+' '+str(index+1))
    elif operation=='edit_slide':
        text(deck.slides[request['slide']-1],request['text'],2)
    elif operation=='append_slides':
        for title in request['titles']:text(deck.slides.add_slide(deck.slide_layouts[6]),title)
    else:raise ValueError('unknown PowerPoint operation')
    path=Path('outputs/result.pptx');deck.save(path);assert len(Presentation(path).slides)>0
print('Reopened actual output:',path.as_posix())
