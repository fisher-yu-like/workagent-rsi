"""Run the frozen six-case engineering pilot once through the public Harness."""

from __future__ import annotations

import argparse
import json
import platform
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "src"))

from workagent_rsi.contracts import TaskSpec
from workagent_rsi.evaluator import OfficeArtifactEvaluator
from workagent_rsi.harness import Harness
from workagent_rsi.hashing import sha256_file
from workagent_rsi.office_capabilities import probe_capabilities, verify_artifact_with_com


CONFIG = ROOT / "project_artifacts/phase3_experiments/configs/general_office_pilot.json"
RESULT_ROOT = ROOT / "project_artifacts/results/qualification/general-office"
SUFFIXES = {"excel": ".xlsx", "word": ".docx", "powerpoint": ".pptx"}


def _write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def _source_inputs(root: Path) -> dict[str, str]:
    """Create only project-owned input material within this invocation."""
    inputs = root / "inputs"
    inputs.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    for row in (("Item", "Revenue"), ("Alpha", 40), ("Beta", 55), ("Gamma", 65)):
        sheet.append(row)
    workbook.save(inputs / "transactions.xlsx")

    document = Document()
    document.add_heading("Draft Status", level=1)
    document.add_paragraph("The pilot completed on 12 September.")
    document.add_paragraph("This draft needs a final structure and an action item.")
    document.save(inputs / "status_draft.docx")

    deck = Presentation()
    for text in ("Baseline: two regions", "Rollout options"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        box.text = text
    deck.save(inputs / "briefing_draft.pptx")
    return {path.name: sha256_file(path) for path in inputs.iterdir()}


def _validate_config(config: dict) -> None:
    if config.get("evaluator_source_sha256") != OfficeArtifactEvaluator.evaluator_hash():
        raise ValueError("frozen evaluator source hash mismatch")
    tasks = config.get("tasks")
    if not isinstance(tasks, list) or len(tasks) != 6:
        raise ValueError("pilot must contain exactly six tasks")
    ids = [item["task_id"] for item in tasks]
    if len(set(ids)) != 6 or set(ids) != {f"{domain}-{kind}" for domain in SUFFIXES for kind in ("create", "edit")}:
        raise ValueError("pilot task IDs are not the frozen six-case matrix")
    for item in tasks:
        if item["domain"] not in SUFFIXES or item["kind"] not in ("create", "edit"):
            raise ValueError("unsupported pilot task")
        if not isinstance(item["expected_constraints"], dict):
            raise ValueError("expected constraints must be evaluator-only mapping")
        if item["kind"] == "create" and item["input_files"]:
            raise ValueError("create task cannot have inputs")
        if item["kind"] == "edit" and len(item["input_files"]) != 1:
            raise ValueError("edit task needs one generated input")


def _report(root: Path, summary: dict) -> None:
    lines = [
        "# General Office six-task qualification",
        "",
        "This is a project-generated engineering pilot, not an external benchmark.",
        f"Invocation: `{root}`",
        f"Evaluator source SHA-256: `{summary['evaluator_source_sha256']}`",
        f"Config SHA-256: `{summary['config_sha256']}`",
        f"COM versions: `{json.dumps(summary['com_versions'], sort_keys=True)}`",
        f"Outcome: {summary['success_count']} succeeded, {summary['failure_count']} failed, {summary['task_count']} attempted.",
        "",
        "| Task | State | Evaluation | COM reopen | Input hashes |",
        "| --- | --- | --- | --- | --- |",
    ]
    for row in summary["rows"]:
        lines.append(f"| {row['task_id']} | {row['state']} | {row['evaluation_passed']} | {row['com_ok']} | {row['input_hashes_ok']} |")
    lines += ["", "Failures are retained as observed; no task is rerun within this invocation.", ""]
    (root / "qualification_report.md").write_text("\n".join(lines), encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.parse_args(argv)
    invocation = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid4().hex[:8]
    root = RESULT_ROOT / invocation
    root.mkdir(parents=True, exist_ok=False)
    started = datetime.now(timezone.utc).isoformat()
    report = probe_capabilities()
    _write_json(root / "capability_report.json", report.model_dump(mode="json"))
    required = ("excel_com", "word_com", "powerpoint_com")
    versions = {name: report.capabilities[name].version for name in required if name in report.capabilities}
    summary = {
        "pilot_id": "general-office-six-task-engineering-pilot-v1",
        "invocation": str(root), "started_at": started, "ended_at": None,
        "config_sha256": sha256_file(CONFIG),
        "evaluator_source_sha256": OfficeArtifactEvaluator.evaluator_hash(),
        "com_versions": versions, "source_hashes": {}, "rows": [],
        "task_count": 0, "success_count": 0, "failure_count": 0,
        "python": sys.version, "platform": platform.platform(),
        "git_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
    }
    if not all(report.is_available(name) for name in required):
        summary["failure"] = "all three Office COM applications must be available before launching model"
        summary["ended_at"] = datetime.now(timezone.utc).isoformat()
        _write_json(root / "summary.json", summary)
        _report(root, summary)
        print(json.dumps({"status": "probe_failed", "invocation": str(root), "com_versions": versions}))
        return 1

    try:
        config = json.loads(CONFIG.read_text(encoding="utf-8"))
        _validate_config(config)
        (root / "config.json").write_bytes(CONFIG.read_bytes())
        source_hashes = _source_inputs(root)
        summary["source_hashes"] = source_hashes
        _write_json(root / "input_provenance.json", {"provenance": "project-generated", "sha256": source_hashes})
        for item in config["tasks"]:
            task_id = item["task_id"]
            before = {name: sha256_file(root / "inputs" / name) for name in item["input_files"]}
            if any(before[name] != source_hashes[name] for name in before):
                raise ValueError(f"source hash mismatch before {task_id}")
            task = TaskSpec(task_id=task_id, domain=item["domain"], instruction=item["instruction"],
                            input_files=tuple(item["input_files"]), expected_constraints=item["expected_constraints"])
            start = time.perf_counter()
            outcome = Harness(root, office=True, execution_provider="workagent", input_base=root / "inputs").run(task, run_id=task_id, max_attempts=1)
            run_root = Path(outcome["result_dir"])
            after = {name: sha256_file(root / "inputs" / name) for name in item["input_files"] if (root / "inputs" / name).is_file()}
            hashes_ok = before == after == {name: source_hashes[name] for name in before}
            artifacts = outcome.get("artifacts", [])
            suffix_ok = len(artifacts) == 1 and Path(artifacts[0]["path"]).suffix.lower() == SUFFIXES[item["domain"]]
            com = verify_artifact_with_com(artifacts[0]["path"], item["domain"]) if suffix_ok else {"ok": False, "status": "not_run", "error": "no single correct-suffix artifact"}
            _write_json(run_root / "com_reopen.json", com)
            if "evaluation" in outcome:
                _write_json(run_root / "evaluation.json", outcome["evaluation"])
            qualified = outcome["state"] == "SUCCEEDED" and outcome.get("evaluation", {}).get("passed") is True and suffix_ok and com.get("ok") is True and hashes_ok
            row = {"task_id": task_id, "domain": item["domain"], "kind": item["kind"],
                   "run_id": outcome["run_id"], "state": "SUCCEEDED" if qualified else "FAILED",
                   "harness_state": outcome["state"], "evaluation_passed": outcome.get("evaluation", {}).get("passed"),
                   "com_ok": com.get("ok"), "com_version": com.get("version"), "input_hashes_ok": hashes_ok,
                   "source_hashes_before": before, "source_hashes_after": after, "suffix_ok": suffix_ok,
                   "artifact_sha256": [artifact["sha256"] for artifact in artifacts],
                   "failure": outcome.get("failure"), "evaluation_failures": outcome.get("evaluation", {}).get("critical_failures", []),
                   "duration_seconds": round(time.perf_counter() - start, 3), "result_dir": str(run_root)}
            _write_json(run_root / "qualification_result.json", row)
            summary["rows"].append(row)
            summary["task_count"] += 1
            summary["success_count"] += int(qualified)
            summary["failure_count"] += int(not qualified)
            _write_json(root / "summary.json", summary)
    except Exception as exc:
        summary["failure"] = str(exc)
        summary["failure_count"] += 1
    summary["ended_at"] = datetime.now(timezone.utc).isoformat()
    _write_json(root / "summary.json", summary)
    _report(root, summary)
    print(json.dumps({"invocation": str(root), "task_count": summary["task_count"], "success_count": summary["success_count"], "failure_count": summary["failure_count"], "failure": summary.get("failure")}))
    return 0 if summary["task_count"] == 6 and summary["failure_count"] == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
