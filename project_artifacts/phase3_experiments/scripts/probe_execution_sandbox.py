"""Run authored T04 boundary probes only with --execute and an immutable image.

No model/API calls, image pulls, builds, installs or system configuration changes.
Default captures capability only. Missing prerequisites produce exit code 2.
"""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import time

REPO = next(p for p in Path(__file__).resolve().parents if (p / "pyproject.toml").is_file())
sys.path.insert(0, str(REPO / "src"))
from workagent_rsi.environment_manifest import capture_environment
from workagent_rsi.execution_sandbox import SandboxExecutor, SandboxPolicy
from workagent_rsi.apptainer_sandbox import ApptainerPolicy, ApptainerSSHExecutor
from workagent_rsi.hashing import sha256_file


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--image", help="Already-provisioned local sha256 image ID or repository@sha256 digest")
    parser.add_argument("--policy", type=Path, help="Explicit Docker or Apptainer SSH policy JSON; no credentials")
    parser.add_argument("--execute", action="store_true", help="Execute tiny authored probes; no model calls")
    parser.add_argument("--output", type=Path, help="New, exclusive evidence directory")
    args = parser.parse_args()
    if args.policy and args.image:
        parser.error("choose --policy or --image")
    config = json.loads(args.policy.read_text(encoding="utf-8-sig")) if args.policy else {"image": args.image}
    policy = ApptainerPolicy(**config) if config.get("backend") == "apptainer-ssh" else SandboxPolicy(**config)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    output = args.output or REPO / "project_artifacts/results/hardening-v2/T04" / stamp / "sandbox"
    output.mkdir(parents=True, exist_ok=False)
    capture_environment(REPO, output / "environment.json")
    executor = ApptainerSSHExecutor(policy, authored_boundary_probe=True) if isinstance(policy, ApptainerPolicy) else SandboxExecutor(policy)
    capability = executor.probe()
    record = {"command": [sys.executable, *sys.argv], "started_at": datetime.now(timezone.utc).isoformat(),
              "git_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=REPO, text=True).strip(),
              "backend": policy.backend, "docker_cli": shutil.which("docker"), "capability": capability, "policy": policy.model_dump(mode="json"),
              "real_boundary_status": "not_run", "t04_qualified": False}
    record["source_hashes"] = {str(p.relative_to(REPO)): sha256_file(p) for p in [
        Path(__file__), REPO / "src/workagent_rsi/execution_sandbox.py",
        REPO / "src/workagent_rsi/apptainer_sandbox.py", REPO / "src/workagent_rsi/apptainer_worker.py"]}
    (output / "policy.json").write_text(policy.model_dump_json(indent=2)+"\n", encoding="utf-8")
    (output / "image-digest.txt").write_text((getattr(policy, "remote_image_sha256", None) or policy.image or "unavailable")+"\n", encoding="utf-8")
    if args.execute and capability["status"] == "available":
        with tempfile.TemporaryDirectory(prefix="wrsi-probe-") as temporary:
            root = Path(temporary)
            workspace = root / "workspace"
            (workspace / "inputs").mkdir(parents=True)
            (root / "host-sentinel.txt").write_text("authored-parent-only", encoding="utf-8")
            source = workspace / "inputs/source.txt"
            source.write_text("authored-read-only-input", encoding="utf-8")
            before = sha256_file(source)
            script = workspace / "probe.py"
            script.write_text('''import json, os, socket, subprocess, sys
from pathlib import Path
def denied(path):
    try:
        Path(path).read_text()
        return False
    except OSError:
        return True
try:
    Path('inputs/source.txt').write_text('modified')
    readonly = False
except OSError:
    readonly = True
s = socket.socket()
s.settimeout(1)
network_denied = s.connect_ex(('192.0.2.1', 9)) != 0
s.close()
try:
    Path('/etc/wrsi-probe-write').write_text('authored')
    root_readonly = False
except OSError:
    root_readonly = True
process_status = dict(line.split(':',1) for line in Path('/proc/self/status').read_text().splitlines() if ':' in line)
facts = {'parent_denied': denied('../host-sentinel.txt'), 'host_env_absent': os.environ.get('WRSI_PROBE_DUMMY_SECRET') is None,
         'absolute_parent_denied': denied(sys.argv[1]) if len(sys.argv)>1 else True,
         'root_readonly': root_readonly, 'no_new_privileges': process_status['NoNewPrivs'].strip()=='1',
         'zero_effective_capabilities': int(process_status['CapEff'].strip(),16)==0,
         'network_denied': network_denied, 'input_readonly': readonly, 'nonroot': os.getuid() != 0}
facts['network_loopback_only'] = set(name for _,name in socket.if_nameindex()) <= {'lo'}
Path('outputs/probe.json').write_text(json.dumps(facts))
''', encoding="utf-8")
            os.environ["WRSI_PROBE_DUMMY_SECRET"] = "authored-dummy-only"
            try:
                outcome = executor.run(script, workspace, 15)
            finally:
                os.environ.pop("WRSI_PROBE_DUMMY_SECRET", None)
            facts = {}
            if outcome.status == "succeeded":
                facts = json.loads((workspace / "outputs/probe.json").read_text())
            facts["source_hash_unchanged"] = sha256_file(source) == before
            facts["execution_succeeded"] = outcome.status == "succeeded"
            boundary_trace = {}
            for ref in outcome.trace_refs:
                shutil.copyfile(workspace / ref, output / "boundary-execution.json")
                boundary_trace = json.loads((workspace / ref).read_text())
            facts['resource_policy_observed'] = (boundary_trace.get('resource_observations', {}).get('passed') is True
                                                  if isinstance(policy, ApptainerPolicy) else outcome.status == 'succeeded')
            # Prove the provisioned runtime creates and exports real Office files.
            script.write_text('''from pathlib import Path
import json,sys,openpyxl,docx,pptx,yaml
book=openpyxl.Workbook()
book.active['A1']='authored-runtime-probe'
book.save('outputs/runtime.xlsx')
document=docx.Document()
document.add_paragraph('authored-runtime-probe')
document.save('outputs/runtime.docx')
deck=pptx.Presentation()
deck.slides.add_slide(deck.slide_layouts[6])
deck.save('outputs/runtime.pptx')
Path('outputs/runtime-versions.json').write_text(json.dumps({'python':sys.version,'openpyxl':openpyxl.__version__,
 'python-docx':docx.__version__,'python-pptx':pptx.__version__,'PyYAML':yaml.__version__}))
''', encoding="utf-8")
            runtime = executor.run(script, workspace, 15)
            runtime_checks = {"execution_succeeded": runtime.status == "succeeded"}
            if runtime.status == 'succeeded':
                import openpyxl, docx, pptx
                book = openpyxl.load_workbook(workspace / 'outputs/runtime.xlsx')
                runtime_checks['xlsx_readable_and_correct'] = book.active['A1'].value == 'authored-runtime-probe'
                book.close()
                runtime_checks['docx_readable_and_correct'] = docx.Document(workspace / 'outputs/runtime.docx').paragraphs[0].text == 'authored-runtime-probe'
                runtime_checks['pptx_readable_and_correct'] = len(pptx.Presentation(workspace / 'outputs/runtime.pptx').slides) == 1
                runtime_checks['versions'] = json.loads((workspace / 'outputs/runtime-versions.json').read_text())
                artifact_dir = output / 'artifacts'
                artifact_dir.mkdir()
                for name in ('runtime.xlsx','runtime.docx','runtime.pptx','runtime-versions.json'):
                    shutil.copyfile(workspace / 'outputs' / name, artifact_dir / name)
            for ref in runtime.trace_refs:
                shutil.copyfile(workspace / ref, output / 'runtime-execution.json')
            (output / 'runtime-tests.json').write_text(json.dumps(runtime_checks, indent=2)+'\n', encoding='utf-8')
            script.write_text("import subprocess,time,sys\nsubprocess.Popen([sys.executable,'-c','import os,time;os.setsid();time.sleep(60)'])\ntime.sleep(60)\n", encoding="utf-8")
            timeout = executor.run(script, workspace, 3)
            cleanup = {}
            for ref in timeout.trace_refs:
                shutil.copyfile(workspace / ref, output / "timeout-execution.json")
                cleanup = json.loads((workspace / ref).read_text())["cleanup"]
            timeout_facts = {"status": timeout.status, "cleanup": cleanup,
                             "descendants_gone": cleanup.get("confirmed_absent") is True}
            (output / "boundary-tests.json").write_text(json.dumps(facts, indent=2)+"\n", encoding="utf-8")
            (output / "timeout-cleanup.json").write_text(json.dumps(timeout_facts, indent=2)+"\n", encoding="utf-8")
            runtime_passed = (len(runtime_checks) == 5 and all(v for k,v in runtime_checks.items() if k != 'versions'))
            record["t04_qualified"] = (len(facts) == 13 and all(facts.values()) and runtime_passed and timeout.status == "timeout"
                                        and timeout_facts["descendants_gone"])
            record["real_boundary_status"] = "passed" if record["t04_qualified"] else "failed"
    else:
        reason = capability.get("reason", "--execute not requested")
        for name in ("boundary-tests.json", "timeout-cleanup.json"):
            (output / name).write_text(json.dumps({"status": "not_run", "reason": reason}, indent=2)+"\n", encoding="utf-8")
    record["ended_at"] = datetime.now(timezone.utc).isoformat()
    record["file_hashes"] = {p.relative_to(output).as_posix(): sha256_file(p) for p in output.rglob('*') if p.is_file()}
    (output / "probe.json").write_text(json.dumps(record, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps({"evidence": str(output), "status": record["real_boundary_status"], "t04_qualified": record["t04_qualified"]}))
    return 0 if record["t04_qualified"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
