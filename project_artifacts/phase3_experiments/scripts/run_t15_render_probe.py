"""Render authored visibility controls and every page of a long deck remotely.

Annotations are authored from the fixture design before rendering. They are
not human study labels or a measured judge accuracy claim.
"""
import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
import subprocess
import sys

REPO=next(p for p in Path(__file__).resolve().parents if (p/'pyproject.toml').is_file())
sys.path.insert(0,str(REPO/'src'))
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from workagent_rsi.apptainer_sandbox import ApptainerPolicy,ApptainerSSHExecutor
from workagent_rsi.artifact_renderer import OfficeArtifactRenderer
from workagent_rsi.environment_manifest import capture_environment
from workagent_rsi.hashing import sha256_file
from workagent_rsi.visibility_checks import inspect_visibility


def write(path,value):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(value,indent=2)+'\n',encoding='utf-8')


def text(slide,value,x=1,y=2,color='000000'):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(7),Inches(1))
    run=box.text_frame.paragraphs[0].add_run()
    run.text=value
    run.font.size=Pt(28)
    run.font.color.rgb=RGBColor.from_string(color)
    return box


def rectangle(slide,color,alpha=None):
    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.8),Inches(1.8),Inches(7.5),Inches(1.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb=RGBColor.from_string(color)
    shape.line.fill.background()
    if alpha is not None:
        element=OxmlElement('a:alpha');element.set('val',str(alpha))
        shape.fill._xPr.solidFill.srgbClr.append(element)
    return shape


def fixtures(root):
    root.mkdir()
    deck=Presentation()
    annotations=[]
    controls=[('off_canvas',False),('white_on_white',False),('opaque_cover',False),
              ('rotated_readable',True),('transparent_cover',True),('legal_background',True),('plain_control',True)]
    for name,visible in controls:
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
        text(slide,'Visibility page '+str(len(deck.slides)),y=.3)
        if name=='legal_background':rectangle(slide,'003366')
        box=text(slide,'Required message: Revenue 25',x=11 if name=='off_canvas' else 1,
                 color='FFFFFF' if name in {'white_on_white','legal_background'} else '000000')
        if name=='opaque_cover':rectangle(slide,'FFFFFF')
        if name=='transparent_cover':rectangle(slide,'003366',alpha=0)
        if name=='rotated_readable':box.rotation=8
        annotations.append({'page':len(deck.slides),'case':name,'required_text':'Required message: Revenue 25',
            'region':'message box near x=1..8, y=2..3 inches; off_canvas control x=11..18',
            'expected_readable':visible,'object_id':box.shape_id})
    deck.save(root/'controls.pptx')
    long=Presentation()
    for i in range(1,18):
        slide=long.slides.add_slide(long.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
        text(slide,f'Long coverage page {i}',y=.3)
        text(slide,'Required message: Revenue 25',color='FFFFFF' if i==17 else '000000')
    long.save(root/'long-17.pptx')
    return {'controls':annotations,'long-17':[{'page':i,'expected_readable':i!=17,
        'required_text':'Required message: Revenue 25','region':'message box near x=1..8, y=2..3 inches'} for i in range(1,18)]}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--policy',type=Path,required=True)
    options=parser.parse_args()
    root=REPO/'project_artifacts/results/hardening-v2/T15-real'/datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    root.mkdir(parents=True)
    capture_environment(REPO,root/'environment.json')
    record={'started_at':datetime.now(timezone.utc).isoformat(),'command':[sys.executable,*sys.argv],
        'commit':subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip(),
        'scope':'actual Office rendering; independent authored labels, no human calibration',
        'script_sha256':sha256_file(Path(__file__))}
    policy=ApptainerPolicy.model_validate_json(options.policy.read_text(encoding='utf-8'))
    write(root/'sandbox-policy.json',policy.model_dump(mode='json'))
    annotations=fixtures(root/'sources')
    write(root/'visual/authored-annotations.json',annotations)
    renderer=OfficeArtifactRenderer(root/'renders',sandbox=ApptainerSSHExecutor(policy),dpi=96,max_pages=40)
    results={};geometry={};unchanged=True
    for name in annotations:
        source=root/'sources'/(name+'.pptx');before=sha256_file(source)
        deck=Presentation(source)
        geometry[name]=[[f.model_dump(mode='json') for f in inspect_visibility(deck,i)] for i in range(1,len(deck.slides)+1)]
        result=renderer.render(source,'powerpoint',before)
        results[name]=result.model_dump(mode='json')
        unchanged=unchanged and sha256_file(source)==before
        write(root/'visual/real-render-manifest.json',results)
    write(root/'visual/geometry-findings.json',geometry)
    traces=[json.loads(p.read_text()) for p in (root/'renders').rglob('provider_records/sandbox/*.json')]
    isolated=bool(traces) and all(t.get('cleanup',{}).get('confirmed_absent') is True and
        t.get('resource_observations',{}).get('passed') is True and
        t.get('capability',{}).get('image_sha256')==policy.remote_image_sha256 for t in traces)
    passed=unchanged and isolated and all(r['status']=='available' and len(r['pages'])==len(annotations[name]) for name,r in results.items())
    record.update(ended_at=datetime.now(timezone.utc).isoformat(),render_status='passed' if passed else 'failed',
        judge_status='not_run',original_sources_unchanged=unchanged,all_runtime_traces_qualified=isolated,
        rendered_pages={name:len(r['pages']) for name,r in results.items()},exit_code=0 if passed else 1)
    write(root/'run.json',record)
    print(json.dumps({'evidence':str(root.relative_to(REPO)),**record}))
    return record['exit_code']


if __name__=='__main__':sys.exit(main())
