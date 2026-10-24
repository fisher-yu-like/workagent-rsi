"""Qualify a frozen no-network Office worker with independent arithmetic oracles.

No model calls, host Office, installation, or synthetic formula caches.
Each invocation writes a fresh attempt and retains failed receipts.
"""
import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
import subprocess
import sys

REPO = next(p for p in Path(__file__).resolve().parents if (p / 'pyproject.toml').is_file())
sys.path.insert(0, str(REPO / 'src'))
from openpyxl import Workbook, load_workbook
from workagent_rsi.apptainer_sandbox import ApptainerPolicy, ApptainerSSHExecutor
from workagent_rsi.assessment_contracts import Requirement, Location
from workagent_rsi.environment_manifest import capture_environment
from workagent_rsi.formula_checks import inspect_formula, validate_formula_options
from workagent_rsi.hashing import sha256_file
from workagent_rsi.spreadsheet_recalc import RecalculationEngine


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--policy', type=Path, required=True)
    options = parser.parse_args()
    stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    root = REPO / 'project_artifacts/results/hardening-v2/T13-real' / stamp
    root.mkdir(parents=True)
    record = {'started_at': datetime.now(timezone.utc).isoformat(), 'command': [sys.executable, *sys.argv],
              'cwd': str(REPO), 'commit': subprocess.check_output(['git', 'rev-parse', 'HEAD'], text=True).strip(),
              'scope': 'authored engine qualification, not agent accuracy',
              'source_hashes': {p.name: sha256_file(p) for p in [Path(__file__), REPO/'src/workagent_rsi/spreadsheet_recalc.py',
                  REPO/'src/workagent_rsi/libreoffice_recalc_worker.py', REPO/'src/workagent_rsi/formula_checks.py']}}
    capture_environment(REPO, root/'environment.json')
    policy = ApptainerPolicy.model_validate_json(options.policy.read_text(encoding='utf-8'))
    write(root/'sandbox-policy.json', policy.model_dump(mode='json'))
    sandbox = ApptainerSSHExecutor(policy)
    probe = sandbox.probe()
    write(root/'recalc/runtime-probe.json', probe)
    if probe.get('status') != 'available':
        record.update(status='blocked', error='frozen runtime unavailable')
        write(root/'run.json', record)
        print(json.dumps({'evidence': str(root), **record}))
        return 1
    sources = root/'sources'
    sources.mkdir()
    for name, formula in [('equivalent', '=SUM(B2:B100)'), ('exact', '=B2+B3'), ('hardcoded', 25),
                          ('constant_formula', '=25'), ('stale_formula', '=SUM(B2:B3)')]:
        book = Workbook()
        book.active.title = 'Data'
        book.active['B2'] = 10
        book.active['B3'] = 15
        book.active['D2'] = formula
        book.save(sources/(name+'.xlsx'))
        book.close()
    original = {p.name: sha256_file(p) for p in sources.iterdir()}
    # This independent numeric probe precedes qualified use of the same frozen worker.
    engine = RecalculationEngine(sandbox=sandbox, output_root=root/'recalc', qualification='unverified')
    engine_probe = engine.recalculate(sources/'equivalent.xlsx', root/'engine-probe.xlsx')
    write(root/'recalc/engine-probe.json', engine_probe.model_dump(mode='json'))
    probe_ok = engine_probe.status == 'available' and engine_probe.values.get('Data!D2') == 25
    cases = []
    if probe_ok:
        engine = RecalculationEngine(sandbox=sandbox, output_root=root/'recalc', qualification='isolated_run')
        variants = [
            {'name':'last_row_added','inputs':{'Data!B4':7},'expected_value':32},
            {'name':'negative','inputs':{'Data!B2':-10},'expected_value':5},
            {'name':'null','inputs':{'Data!B2':None},'expected_value':15},
            {'name':'zero','inputs':{'Data!B2':0},'expected_value':15},
            {'name':'order','inputs':{'Data!B2':15,'Data!B3':10},'expected_value':25},
        ]
        write(root/'formula-variants.json', {'base_inputs':{'Data!B2':10,'Data!B3':15},
              'base_oracle':25,'variants':variants,'oracle_source':'authored arithmetic before engine output'})
        equivalent = Requirement(requirement_id='sum', description='Sum every numeric input row in B2:B100',
            check='excel.formula', location=Location(artifact='output',sheet='Data',cell='D2'),
            expected='=B2+B3', dimension='content', evidence_source='independent arithmetic',
            options={'formula_policy':'equivalent','oracle_value':25,'input_variants':variants})
        validate_formula_options(equivalent)
        for name, requirement, expected in [
            ('equivalent', equivalent, 'PASS'),
            ('exact', equivalent.model_copy(update={'check':'excel.formula_exact','options':{},'expected':'=SUM(B2:B3)'}), 'FAIL'),
            ('hardcoded', equivalent, 'FAIL'),
            ('constant_formula', equivalent, 'FAIL'),
        ]:
            source = sources/(name+'.xlsx')
            book = load_workbook(source, data_only=False)
            try:
                status, completion, observed, location = inspect_formula(requirement,book,source,engine)
            finally:
                book.close()
            cases.append({'name':name,'expected_status':expected,'status':status,'completion':completion,
                          'observed':observed,'location':location.model_dump(mode='json')})
            write(root/'recalc-report.json', cases)
        # A real recomputation must replace the old base result after the input changes.
        stale = load_workbook(root/'engine-probe.xlsx',data_only=False)
        stale['Data']['B2'] = 100
        stale.save(root/'changed-input.xlsx')
        stale.close()
        recomputed = engine.recalculate(root/'changed-input.xlsx',root/'changed-recomputed.xlsx')
        write(root/'recalc/stale-cache-control.json', recomputed.model_dump(mode='json'))
        cases.append({'name':'changed_input_recomputed','expected_status':'PASS',
            'status':'PASS' if recomputed.status=='available' and recomputed.values.get('Data!D2')==115 else 'FAIL'})
    hashes_unchanged = original == {p.name:sha256_file(p) for p in sources.iterdir()}
    traces = [json.loads(p.read_text()) for p in (root/'recalc').rglob('provider_records/sandbox/*.json')]
    isolated = bool(traces) and all(t.get('cleanup',{}).get('confirmed_absent') is True
        and t.get('resource_observations',{}).get('passed') is True
        and t.get('capability',{}).get('image_sha256') == policy.remote_image_sha256 for t in traces)
    passed = probe_ok and hashes_unchanged and isolated and len(cases)==5 and all(c['status']==c['expected_status'] for c in cases)
    record.update(ended_at=datetime.now(timezone.utc).isoformat(), status='passed' if passed else 'failed',
        engine_probe_passed=probe_ok, source_hashes_unchanged=hashes_unchanged, original_hashes=original,
        all_runtime_traces_qualified=isolated, actual_worker_runs=len(traces), cases=cases, exit_code=0 if passed else 1)
    write(root/'run.json',record)
    print(json.dumps({'evidence':str(root.relative_to(REPO)),'status':record['status'],'worker_runs':len(traces)}))
    return record['exit_code']


if __name__ == '__main__':
    sys.exit(main())
