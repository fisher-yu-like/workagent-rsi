"""Create minimal real Office files and prove the optional renderer receives them."""

import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
import platform
import shutil
import subprocess
import sys

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from workagent_rsi.artifact_renderer import OfficeArtifactRenderer


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    invocation = {
        "command": sys.argv,
        "cwd": str(Path.cwd()),
        "started_at": datetime.now(timezone.utc).isoformat(),
        "python": sys.version,
        "platform": platform.platform(),
        "git_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], text=True).strip(),
        "capabilities": {
            "powershell": shutil.which("powershell.exe"),
            "pdftoppm": shutil.which("pdftoppm"),
            "configured_visual_model_command": None,
            "ollama": shutil.which("ollama"),
        },
    }
    artifacts = args.output / "artifacts"
    artifacts.mkdir()

    book = Workbook()
    book.active.append(["Region", "Sales"])
    book.active.append(["East", 15])
    excel = artifacts / "sample.xlsx"
    book.save(excel)
    book.close()

    document = Document()
    document.add_heading("Visual assessment render", 1)
    document.add_paragraph("This file is rendered from Microsoft Word, not reconstructed from extracted text.")
    word = artifacts / "sample.docx"
    document.save(word)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Visual assessment render"
    slide.placeholders[1].text = "This image comes from the real PowerPoint file."
    powerpoint = artifacts / "sample.pptx"
    presentation.save(powerpoint)

    renderer = OfficeArtifactRenderer(args.output / "renders")
    results = {domain: renderer.render(path, domain).model_dump(mode="json") for domain, path in
               (("excel", excel), ("word", word), ("powerpoint", powerpoint))}
    (args.output / "render_results.json").write_text(json.dumps(results, indent=2), encoding="utf-8")
    print(json.dumps({key: {"status": value["status"], "pages": len(value["pages"]), "error": value["error"]} for key, value in results.items()}, indent=2))
    code = 0 if all(value["status"] == "available" and value["pages"] for value in results.values()) else 1
    invocation.update(ended_at=datetime.now(timezone.utc).isoformat(), exit_code=code)
    (args.output / "invocation.json").write_text(json.dumps(invocation, indent=2), encoding="utf-8")
    return code


if __name__ == "__main__":
    raise SystemExit(main())
