"""Fresh, opt-in real model calls over project-generated Office fixtures."""

import argparse
from datetime import datetime, timezone
import importlib.metadata
import json
from pathlib import Path
import platform
import subprocess
import sys

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation

from workagent_rsi.artifact_assessment import save_reports
from workagent_rsi.artifact_evaluator import ArtifactEvaluator
from workagent_rsi.artifact_verifier import ArtifactVerifier
from workagent_rsi.assessment_contracts import AcceptanceSpec, Location, Requirement
from workagent_rsi.hashing import sha256_file
from workagent_rsi.model_judgement import ModelReviewConfig, configured_assessment


def fixtures(root, *, flawed=False):
    root.mkdir(parents=True)
    document = Document()
    document.add_heading("Sales review", 1)
    document.add_paragraph("East sales were 15 and West sales were 10. Total sales were 25. East contributed 60% of total sales.")
    document.add_paragraph("This single-period comparison does not establish growth or causality. Obtain a prior-period baseline before changing the budget.")
    if flawed:
        document.add_paragraph("The same East 15 and West 10 values sum to 900. This proves 400% growth caused by the campaign.")
    document.save(root / "report.docx")
    book = Workbook()
    sheet = book.active
    sheet.title = "Summary"
    for row in [["Region", "Sales"], ["East", 15], ["West", 10], ["Total", 25]]:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    if flawed:
        sheet["B4"] = 900
        sheet["A6"] = "This single period proves 400% growth caused by the campaign."
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 16
    book.save(root / "report.xlsx")
    book.close()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Sales review"
    slide.placeholders[1].text = "East: 15 | West: 10 | Total: 25\nEast contributed 60% of sales.\nA single period cannot establish growth or causality."
    if flawed:
        from pptx.util import Inches
        slide.placeholders[1].text = "East: 15 | West: 10 | Total: 900\nThis proves 400% growth caused by the campaign."
        for shape in (slide.shapes.title, slide.placeholders[1]):
            shape.left, shape.width = Inches(9), Inches(5)
    deck.save(root / "report.pptx")
    return {"word": root / "report.docx", "excel": root / "report.xlsx", "powerpoint": root / "report.pptx"}


def make_spec(domain, visual):
    requirements = [
        Requirement(requirement_id="readable", description="Readable Office file", check="file.readable", location=Location(artifact="output"), expected=True, critical=True, dimension="correctness", evidence_source="OOXML parser"),
        Requirement(requirement_id="content", description="Clearly communicate regional sales and the total without unsupported conclusions", check="semantic", location=Location(artifact="output"), expected="East 15, West 10, total 25. If a share is stated, East is 60%. No unsupported growth or causal claims.", dimension="content", evidence_source="Extracted artifact text and table data", options={"rubric": "Full credit: values are clearly labelled, the summary is coherent, and claims stay within this single-period evidence. Penalize ambiguous labels, contradictions and unsupported inferences."}),
    ]
    weights = {"correctness": .2, "content": .8}
    if visual:
        requirements.append(Requirement(requirement_id="appearance", description="Readable, orderly page with clear hierarchy and no clipped or overlapping content", check="visual", location=Location(artifact="output"), expected="Legible text, distinguishable heading/header, aligned content, no clipped or obscured values.", dimension="presentation", evidence_source="Actual Office rendered page", options={"pages": [1]}))
        weights = {"correctness": .2, "content": .4, "presentation": .4}
    return AcceptanceSpec(task_id="model-smoke-" + domain, version="1", artifacts={"output": domain}, requirements=requirements, dimension_weights=weights,
                          scope_note="Project-generated engineering smoke test; not calibrated model-quality evidence.")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--model-config", type=Path, required=True)
    parser.add_argument("--domains", nargs="+", choices=["word", "excel", "powerpoint"], default=["word", "excel", "powerpoint"])
    parser.add_argument("--visual", action="store_true")
    parser.add_argument("--flawed", action="store_true", help="Deliberately incorrect values/inferences and clipped PPT objects")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    config = ModelReviewConfig.model_validate_json(args.model_config.read_text(encoding="utf-8-sig"))
    paths = fixtures(args.output / "artifacts", flawed=args.flawed)
    record = {"command": sys.argv, "cwd": str(Path.cwd()), "started_at": datetime.now(timezone.utc).isoformat(),
              "python": sys.version, "platform": platform.platform(), "config": config.model_dump(mode="json"),
              "dependencies": {d.metadata["Name"]: d.version for d in importlib.metadata.distributions()},
              "git_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], text=True).strip(),
              "source_hashes": {str(p): sha256_file(p) for p in Path("src/workagent_rsi").glob("*.py")},
              "claim_boundary": "project-generated real model smoke; no benchmark or human agreement claim"}
    (args.output / "invocation.json").write_text(json.dumps(record, indent=2), encoding="utf-8")
    engine = configured_assessment(config, args.output / "model-evidence")
    results = {}
    for domain in args.domains:
        spec = make_spec(domain, args.visual)
        (args.output / f"{domain}-request.json").write_text(json.dumps({"spec": spec.model_dump(mode="json"), "artifacts": {"output": str(paths[domain].resolve())}}, indent=2), encoding="utf-8")
        print(f"Assessing {domain}; visual={args.visual}", flush=True)
        shared = engine.inspect(spec, {"output": paths[domain]})
        score, issues = ArtifactEvaluator().evaluate(shared), ArtifactVerifier().verify(shared)
        save_reports(args.output / domain, shared, score, issues)
        results[domain] = {"total_score": score.total_score, "acceptance": score.acceptance_status,
                           "checks": {c.requirement_id: c.status for c in shared.checks}}
        print(json.dumps(results[domain]), flush=True)
        (args.output / "summary.json").write_text(json.dumps(results, indent=2), encoding="utf-8")
    code = 0 if all(all(s in {"PASS", "PARTIAL", "FAIL"} for s in row["checks"].values()) for row in results.values()) else 1
    record.update(ended_at=datetime.now(timezone.utc).isoformat(), exit_code=code)
    (args.output / "invocation.json").write_text(json.dumps(record, indent=2), encoding="utf-8")
    return code


if __name__ == "__main__":
    raise SystemExit(main())
