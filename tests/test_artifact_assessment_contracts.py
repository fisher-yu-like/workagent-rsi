import json
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from workagent_rsi.assessment_contracts import AcceptanceSpec, CheckRecord, Location, Requirement
from workagent_rsi.artifact_assessment import assess


def simple_spec(check="excel.cell", expected=42, **kwargs):
    return AcceptanceSpec(task_id="simple", version="1", artifacts={"book": "excel"}, dimension_weights={"correctness": 1}, requirements=[Requirement(requirement_id="answer", description="Answer from input", check=check, location=Location(artifact="book", sheet="Sheet", cell="A1"), expected=expected, critical=True, dimension="correctness", evidence_source="independent fixture", **kwargs)])


def test_spec_serialization_and_duplicate_scoring_rejected():
    spec = simple_spec()
    assert AcceptanceSpec.model_validate_json(spec.model_dump_json()) == spec
    payload = spec.model_dump()
    payload["requirements"].append(payload["requirements"][0].copy())
    with pytest.raises(ValidationError, match="duplicate"):
        AcceptanceSpec.model_validate(payload)
    with pytest.raises(ValidationError):
        simple_spec(weight=-1)


def test_unknown_check_keeps_null_score_and_gap(tmp_path):
    path = tmp_path / "book.xlsx"
    Workbook().save(path)
    shared, score, issues = assess(simple_spec("unsupported.feature"), {"book": path})
    assert score.total_score is None
    assert score.acceptance_status == "INCOMPLETE"
    assert score.coverage.ratio == 0
    assert issues.issues[0].kind == "assessment_gap"
    assert issues.issues[0].skill_improvement_hint is None


def test_missing_artifact_is_failure_not_unavailable(tmp_path):
    _, score, issues = assess(simple_spec(), {"book": tmp_path / "missing.xlsx"})
    assert score.acceptance_status == "FAIL"
    assert score.total_score == 0
    assert issues.issues[0].location.cell == "A1"


@pytest.mark.parametrize("entry", ["evaluator/scripts/evaluate.py", "verifier/scripts/verify.py"])
def test_both_skill_entrypoints_read_real_file(tmp_path, entry):
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = 42
    workbook.save(path)
    request = tmp_path / "request.json"
    request.write_text(json.dumps({"spec": simple_spec().model_dump(mode="json"), "artifacts": {"book": "book.xlsx"}}), encoding="utf-8")
    root = Path(__file__).resolve().parents[1]
    script = root / "skills" / ("office-artifact-" + entry)
    output = tmp_path / "report"
    result = subprocess.run([sys.executable, str(script), str(request), "--output", str(output)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    score = json.loads((output / "score.json").read_text())
    assert score["total_score"] == 100
    assert json.loads((output / "issues.json").read_text())["issues"] == []


def test_unknown_status_cannot_have_numeric_completion():
    with pytest.raises(ValidationError):
        CheckRecord(requirement_id="a", status="UNAVAILABLE", completion=1, location=Location(artifact="a"), evidence_refs=["e"])


def test_check_facade_calls_both_skills_over_real_artifact(tmp_path):
    from workagent_rsi.check import Check
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = 42
    workbook.save(path)
    score, issues = Check().assess_artifact(simple_spec(), {"book": path})
    assert score.total_score == 100
    assert issues.issues == []


def test_datetime_observation_is_json_evidence(tmp_path):
    from datetime import datetime
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    value = datetime(2026, 9, 30, 8, 30)
    workbook.active["A1"] = value
    workbook.save(path)
    shared, score, _ = assess(simple_spec(expected=value), {"book": path})
    assert score.total_score == 100
    assert shared.evidence[0].artifact_hashes["book"] == shared.artifact_hashes["book"]
