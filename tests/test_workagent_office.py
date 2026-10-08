import json
import subprocess
import urllib.request
from pathlib import Path

import pytest
import yaml
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pydantic import ValidationError

from workagent_rsi.contracts import TaskSpec
from workagent_rsi.cli import main as cli_main
from workagent_rsi.hashing import canonical_json_hash, sha256_file
from workagent_rsi.harness import Harness
from workagent_rsi.office_checks import inspect_office_file
from workagent_rsi.workagent_office import build_task_prompt, copy_task_inputs, validate_deliverables, verify_task_input_hashes
from workagent_rsi.workagent_provider import (
    AgentResponse,
    CodexOfficeProvider,
    ProviderOutcome,
    WorkAgentConfig,
    _execute_native_tool,
    build_agent_response_schema,
)
from workagent_rsi.workagent_provider import WorkAgentSkill
from workagent_rsi.workagent_evaluator import WorkAgentFrozenEvaluator
from workagent_rsi.rsi_contracts import EvaluationContract
from workagent_rsi.evaluator import OfficeArtifactEvaluator


def _rsi_contract(tasks_by_split: dict[str, list[TaskSpec]]) -> EvaluationContract:
    hashes = {split: canonical_json_hash([task.model_dump(mode="json") for task in tasks])
              for split, tasks in tasks_by_split.items()}
    return EvaluationContract(contract_id="test", dataset_hash=canonical_json_hash(hashes),
        split_hashes=hashes, evaluator_hash=OfficeArtifactEvaluator.evaluator_hash(),
        provider_policy="deterministic-test", seeds=[1], repeats=1, timeout_seconds=30,
        thresholds={"develop_gain": 0.05}, git_commit="test")


def _deterministic_office_provider(self, prompt, workspace, records):
    workspace.mkdir(parents=True, exist_ok=True)
    outputs = workspace / "outputs"
    outputs.mkdir()
    if "Excel" in prompt or "excel" in prompt:
        artifact = outputs / "result.xlsx"
        book = Workbook()
        book.active["A1"] = "Ready" if "good instructions" in prompt else "Wrong"
        book.save(artifact)
    elif "Word" in prompt or "word" in prompt:
        artifact = outputs / "result.docx"
        doc = Document()
        doc.add_paragraph("Status ready" if "good instructions" in prompt else "Wrong")
        doc.save(artifact)
    else:
        artifact = outputs / "result.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text = "Review complete" if "good instructions" in prompt else "Wrong"
        deck.save(artifact)
    response = AgentResponse(status="completed", deliverables=["outputs/" + artifact.name], summary="done", input_files_used=[])
    (workspace / "agent_response.json").write_text(response.model_dump_json(), encoding="utf-8")
    (workspace / "deliverables.json").write_text(json.dumps({"deliverables": response.deliverables}), encoding="utf-8")
    return ProviderOutcome(status="completed", response=response)


def test_workagent_frozen_evaluator_real_office_files_and_hashes(monkeypatch, tmp_path: Path):
    monkeypatch.setattr(CodexOfficeProvider, "run", _deterministic_office_provider)
    tasks = [
        TaskSpec(task_id="excel", domain="excel", instruction="Create Excel", expected_constraints={"required_cells": {"Sheet!A1": "Ready"}}),
        TaskSpec(task_id="word", domain="word", instruction="Create Word", expected_constraints={"required_text": "Status ready"}),
        TaskSpec(task_id="powerpoint", domain="powerpoint", instruction="Create PowerPoint", expected_constraints={"required_slide_count": 1, "required_shape_text": ["Review complete"]}),
    ]
    contract = _rsi_contract({"develop": tasks})
    evaluator = WorkAgentFrozenEvaluator(contract.evaluator_hash, input_base=tmp_path, workagent_config=WorkAgentConfig(verify_com=False))
    result = evaluator.evaluate_split(tasks, "develop", WorkAgentSkill(instructions="good instructions"), contract, tmp_path / "results", reveal_per_task=True)
    assert result["status"] == "completed" and result["score"] == 1.0
    assert len(result["rows"]) == 3
    for task in tasks:
        run = tmp_path / "results" / task.task_id / "run"
        assert (run / "result.json").is_file()
        assert (run / "trace.sqlite").exists() or list(run.glob("*.db"))
        assert list((run / "artifacts").rglob("*"))
    with pytest.raises(ValueError, match="split hash"):
        evaluator.evaluate_split(tasks[:-1], "develop", WorkAgentSkill(instructions="good instructions"), contract, tmp_path / "bad")
    with pytest.raises(ValueError, match="evaluator hash"):
        WorkAgentFrozenEvaluator("wrong").evaluate_split(tasks, "develop", WorkAgentSkill(instructions="good instructions"), contract, tmp_path / "bad")


@pytest.mark.parametrize("evaluation", [
    {}, {"score": None}, {"score": "invalid"}, {"score": float("nan")}, {"score": float("inf")},
    {"score": 10**400},
])
def test_workagent_frozen_evaluator_malformed_terminal_score_is_incomplete(monkeypatch, tmp_path: Path, evaluation: dict):
    task = TaskSpec(task_id="malformed", domain="excel", instruction="Create Excel")
    contract = _rsi_contract({"develop": [task]})

    def malformed_harness_result(self, task, **kwargs):
        run = self.results_dir / "run"
        run.mkdir(parents=True)
        result = {"state": "SUCCEEDED", "evaluation": evaluation, "result_dir": str(run)}
        (run / "result.json").write_text(json.dumps(result), encoding="utf-8")
        return result

    monkeypatch.setattr(Harness, "run", malformed_harness_result)
    root = tmp_path / "evaluation"
    public = WorkAgentFrozenEvaluator(contract.evaluator_hash).evaluate_split(
        [task], "develop", WorkAgentSkill(instructions="test instructions"), contract, root, reveal_per_task=True)
    full = json.loads((root / "evaluation_full.json").read_text(encoding="utf-8"))
    assert public["status"] == full["status"] == "incomplete"
    assert public["completed_count"] == 0
    assert "score" not in public and "completed_wall_time_seconds" not in public
    assert "score" not in full["rows"][0] and "wall_time_seconds" not in full["rows"][0]
    assert full["rows"][0]["status"] == "incomplete"
    assert (root / "malformed" / "run" / "result.json").is_file()


def test_workagent_frozen_evaluator_rejects_live_hash_drift_before_run(monkeypatch, tmp_path: Path):
    task = TaskSpec(task_id="hash-drift", domain="excel", instruction="Create Excel")
    contract = _rsi_contract({"develop": [task]})
    frozen = WorkAgentFrozenEvaluator(contract.evaluator_hash)
    monkeypatch.setattr(OfficeArtifactEvaluator, "evaluator_hash", classmethod(lambda cls: "f" * 64))
    monkeypatch.setattr(Harness, "run", lambda *args, **kwargs: pytest.fail("Harness must not run after evaluator hash drift"))
    root = tmp_path / "must-not-exist"
    with pytest.raises(ValueError, match="evaluator hash"):
        frozen.evaluate_split([task], "develop", WorkAgentSkill(instructions="test instructions"), contract, root)
    assert not root.exists()


def test_general_office_rsi_config_has_twelve_domain_balanced_tasks():
    config_path = Path(__file__).resolve().parents[1] / "project_artifacts/phase3_experiments/configs/general_office_rsi.json"
    splits = json.loads(config_path.read_text(encoding="utf-8"))["splits"]
    assert set(splits) == {"develop", "regression", "hidden", "ood_transfer"}
    assert {split: len(tasks) for split, tasks in splits.items()} == {
        "develop": 3, "regression": 3, "hidden": 3, "ood_transfer": 3,
    }
    assert len({task["task_id"] for tasks in splits.values() for task in tasks}) == 12
    for split, tasks in splits.items():
        assert {task["domain"] for task in tasks} == {"excel", "word", "powerpoint"}
        for task in tasks:
            assert TaskSpec.model_validate(task).task_id == task["task_id"]
            assert task["instruction"] and task["expected_constraints"]
            assert bool(task["input_files"]) == (split == "regression")
            assert not any(key in task["instruction"] for key in task["expected_constraints"])


def _rsi_tasks():
    return {
        "develop": [TaskSpec(task_id="dev-excel", domain="excel", instruction="Create Excel", expected_constraints={"required_cells": {"Sheet!A1": "Ready"}})],
        "regression": [TaskSpec(task_id="reg-word", domain="word", instruction="Create Word", expected_constraints={"required_text": "Status ready"})],
        "hidden": [TaskSpec(task_id="hidden-powerpoint", domain="powerpoint", instruction="Create PowerPoint", expected_constraints={"required_slide_count": 1, "required_shape_text": ["Review complete"]})],
        "ood_transfer": [TaskSpec(task_id="ood-excel", domain="excel", instruction="Create Excel", expected_constraints={"required_cells": {"Sheet!A1": "Ready"}})],
    }


def test_workagent_rsi_no_diagnosis_does_not_generate(monkeypatch, tmp_path: Path):
    from workagent_rsi.workagent_experiment import WorkAgentExperimentRunner

    monkeypatch.setattr(CodexOfficeProvider, "run", _deterministic_office_provider)

    class ForbiddenProvider:
        def generate(self, *args):
            pytest.fail("candidate provider must not run without observed develop failure")

    summary = WorkAgentExperimentRunner(input_base=tmp_path, workagent_config=WorkAgentConfig(verify_com=False)).run(
        _rsi_tasks(), WorkAgentSkill(instructions="good instructions"), ForbiddenProvider(), tmp_path / "experiment")
    assert summary["status"] == "no_candidate_needed"
    assert summary["rounds"][0]["champion_before"] == summary["rounds"][0]["champion_after"]
    assert "rows" not in summary["rounds"][0]["baseline"]["hidden"]
    assert "rows" not in summary["rounds"][0]["baseline"]["ood_transfer"]


def test_workagent_rsi_unavailable_baseline_stops_without_score(monkeypatch, tmp_path: Path):
    from workagent_rsi.workagent_experiment import WorkAgentExperimentRunner

    def unavailable(*args):
        return ProviderOutcome(status="unavailable", error="deterministic missing provider")

    monkeypatch.setattr(CodexOfficeProvider, "run", unavailable)

    class ForbiddenProvider:
        def generate(self, *args):
            pytest.fail("candidate provider must not run after incomplete baseline")

    summary = WorkAgentExperimentRunner(input_base=tmp_path, workagent_config=WorkAgentConfig(verify_com=False)).run(
        _rsi_tasks(), WorkAgentSkill(instructions="bad instructions"), ForbiddenProvider(), tmp_path / "experiment")
    assert summary["status"] == "incomplete"
    assert "initial_develop_score" not in summary["metrics"]
    assert "score" not in summary["incomplete_round"]["baseline"]["develop"]
    assert not (tmp_path / "experiment" / "rounds" / "round-01" / "candidate_workspace").exists()


def _instruction_candidate(workspace, diagnoses, parent_version, edit_budget, record_root):
    from datetime import datetime, timezone
    from workagent_rsi.rsi_contracts import AtomicEdit, CandidatePatch, ProviderRecord

    now = datetime.now(timezone.utc)
    patch = CandidatePatch(candidate_id="deterministic-instructions", parent_version=parent_version,
        provider="deterministic", provider_version="1", diagnosis_refs=[ref for diagnosis in diagnoses for ref in diagnosis.evidence_refs],
        atomic_edits=[AtomicEdit(component="prompt", target_path="skill.json", hypothesis="repair required content",
            expected_metric="task_success_rate", patch=json.dumps({"instructions": "good instructions"}))],
        edit_budget=edit_budget, created_at=now)
    record = ProviderRecord(provider="deterministic", provider_version="1", command=[],
        prompt_hash="a" * 64, workspace_hash="b" * 64, started_at=now, ended_at=now,
        exit_code=0, status="completed")
    return patch, record


def test_workagent_rsi_candidate_verification_and_rollback(monkeypatch, tmp_path: Path):
    from workagent_rsi.workagent_experiment import WorkAgentExperimentRunner
    from workagent_rsi.registry import SkillRegistry

    monkeypatch.setattr(CodexOfficeProvider, "run", _deterministic_office_provider)

    class PatchProvider:
        generate = staticmethod(_instruction_candidate)

    summary = WorkAgentExperimentRunner(input_base=tmp_path, workagent_config=WorkAgentConfig(verify_com=False)).run(
        _rsi_tasks(), WorkAgentSkill(instructions="bad instructions"), PatchProvider(), tmp_path / "experiment")
    assert summary["status"] == "completed"
    row = summary["rounds"][0]
    assert row["verification"]["passed"]
    assert row["metrics"]["candidate_develop"] > row["metrics"]["champion_develop"]
    assert row["metrics"]["cost_delta"] != 0.1
    assert "rows" not in row["candidate_reports"]["hidden"]
    assert "rows" not in row["candidate_reports"]["ood_transfer"]
    assert row["decision"]["decision"] in {"accept", "reject"}
    if row["decision"]["decision"] == "reject":
        registry = SkillRegistry(tmp_path / "experiment" / "registry")
        assert registry.champion("office.workagent").version == "1.0.0"


def test_workagent_rsi_unavailable_candidate_stops_before_promotion(monkeypatch, tmp_path: Path):
    from workagent_rsi.workagent_experiment import WorkAgentExperimentRunner

    real_provider = _deterministic_office_provider

    def conditional(self, prompt, workspace, records):
        if "good instructions" in prompt:
            return ProviderOutcome(status="unavailable", error="candidate provider unavailable")
        return real_provider(self, prompt, workspace, records)

    monkeypatch.setattr(CodexOfficeProvider, "run", conditional)

    class PatchProvider:
        generate = staticmethod(_instruction_candidate)

    summary = WorkAgentExperimentRunner(input_base=tmp_path, workagent_config=WorkAgentConfig(verify_com=False)).run(
        _rsi_tasks(), WorkAgentSkill(instructions="bad instructions"), PatchProvider(), tmp_path / "experiment")
    assert summary["status"] == "incomplete"
    row = summary["incomplete_round"]
    assert "score" not in row["candidate_reports"]["develop"]
    assert "metrics" not in row and "decision" not in row
    assert row["champion_before"] == row["champion_after"]


def test_cli_accepts_explicit_workagent_provider(monkeypatch, tmp_path: Path):
    task_file = tmp_path / "task.yaml"
    task_file.write_text(yaml.safe_dump({"task_id": "cli-explicit", "domain": "excel", "instruction": "Create a workbook"}), encoding="utf-8")

    def missing_provider(command, **kwargs):
        raise FileNotFoundError("codex not found")

    monkeypatch.setattr(subprocess, "run", missing_provider)
    exit_code = cli_main([str(task_file), "--results-dir", str(tmp_path / "results"), "--execution-provider", "workagent"])

    assert exit_code == 1
    result = json.loads(next((tmp_path / "results").glob("*/result.json")).read_text(encoding="utf-8"))
    assert result["state"] == "UNAVAILABLE"


def test_cli_defaults_office_tasks_to_workagent(monkeypatch, tmp_path: Path):
    task_file = tmp_path / "task.yaml"
    task_file.write_text(yaml.safe_dump({"task_id": "cli-default", "domain": "excel", "instruction": "Create a workbook"}), encoding="utf-8")

    def missing_provider(command, **kwargs):
        raise FileNotFoundError("codex not found")

    monkeypatch.setattr(subprocess, "run", missing_provider)
    exit_code = cli_main([str(task_file), "--results-dir", str(tmp_path / "results")])

    assert exit_code == 1
    result = json.loads(next((tmp_path / "results").glob("*/result.json")).read_text(encoding="utf-8"))
    assert result["state"] == "UNAVAILABLE"


def test_cli_resolves_relative_input_files_from_task_yaml_directory(monkeypatch, tmp_path: Path):
    task_dir = tmp_path / "task-directory"
    task_dir.mkdir()
    Workbook().save(task_dir / "source.xlsx")
    task_file = task_dir / "task.yaml"
    task_file.write_text(yaml.safe_dump({"task_id": "cli-input", "domain": "excel", "instruction": "Edit the workbook", "input_files": ["source.xlsx"]}), encoding="utf-8")
    monkeypatch.chdir(Path(__file__).resolve().parents[1])

    def missing_provider(command, **kwargs):
        raise FileNotFoundError("codex not found")

    monkeypatch.setattr(subprocess, "run", missing_provider)
    exit_code = cli_main([str(task_file), "--results-dir", str(tmp_path / "results")])

    assert exit_code == 1
    result_file = next((tmp_path / "results").glob("*/result.json"))
    assert json.loads(result_file.read_text(encoding="utf-8"))["state"] == "UNAVAILABLE"
    manifest = json.loads((result_file.parent / "agent_workspace/input_manifest.json").read_text(encoding="utf-8"))
    assert manifest["files"][0]["source"] == "source.xlsx"
    assert (result_file.parent / "agent_workspace/inputs/0001-source.xlsx").is_file()


def test_cli_accepts_output_inside_selected_results_run_directory(tmp_path: Path):
    task_file = tmp_path / "task.yaml"
    task_file.write_text(yaml.safe_dump({"task_id": "cli-output", "domain": "smoke", "instruction": "Run smoke task"}), encoding="utf-8")
    results_dir = tmp_path / "custom-results"
    output = results_dir / "run-one" / "result.json"

    exit_code = cli_main([str(task_file), "--results-dir", str(results_dir), "--output", str(output)])

    assert exit_code == 0
    assert output.is_file()
    assert json.loads(output.read_text(encoding="utf-8"))["result_dir"] == str(output.parent.resolve())


@pytest.mark.parametrize("kind", ["outside", "traversal"])
def test_cli_rejects_results_dir_outside_project_results(tmp_path: Path, kind: str):
    task_file = tmp_path / "task.yaml"
    task_file.write_text(yaml.safe_dump({"task_id": "cli-results-boundary", "domain": "smoke", "instruction": "Run smoke task"}), encoding="utf-8")
    results_root = Path(__file__).resolve().parents[1] / "project_artifacts/results"
    target = results_root.parent / f"escaped-results-{tmp_path.name}"
    results_dir = target if kind == "outside" else results_root / ".." / target.name

    with pytest.raises(SystemExit) as exc:
        cli_main([str(task_file), "--results-dir", str(results_dir)])

    assert exc.value.code == 2
    assert not target.exists()


@pytest.mark.parametrize("kind", ["outside", "traversal", "results_root"])
def test_cli_rejects_output_outside_per_run_directory(tmp_path: Path, kind: str):
    task_file = tmp_path / "task.yaml"
    task_file.write_text(yaml.safe_dump({"task_id": "cli-output-boundary", "domain": "smoke", "instruction": "Run smoke task"}), encoding="utf-8")
    results_dir = tmp_path / "custom-results"
    if kind == "outside":
        output = results_dir.parent / "escaped-result.json"
    elif kind == "traversal":
        output = results_dir / "run-one" / ".." / ".." / "escaped-result.json"
    else:
        output = results_dir / "result.json"

    with pytest.raises(SystemExit) as exc:
        cli_main([str(task_file), "--results-dir", str(results_dir), "--output", str(output)])

    assert exc.value.code == 2
    assert not results_dir.exists()


def test_workagent_config_has_bounded_local_defaults():
    config = WorkAgentConfig()
    assert (config.model, config.executable, config.timeout_seconds, config.max_output_files) == (
        "qwen2.5:7b", "codex", 600, 8
    )
    assert config.max_artifact_bytes == 100 * 1024 * 1024
    assert config.verify_com is True


@pytest.mark.parametrize("field", ["timeout_seconds", "max_output_files", "max_artifact_bytes"])
def test_workagent_config_rejects_nonpositive_limits(field):
    with pytest.raises(ValidationError):
        WorkAgentConfig.model_validate({field: 0})


def test_agent_response_rejects_unknown_status_and_absolute_paths():
    with pytest.raises(ValidationError):
        AgentResponse.model_validate({"status": "maybe", "deliverables": [], "summary": "x", "input_files_used": []})
    for path in ("/root/report.xlsx", "C:/report.xlsx", "//server/share/report.xlsx"):
        with pytest.raises(ValidationError):
            AgentResponse(status="completed", deliverables=[path], summary="x", input_files_used=[])


@pytest.mark.parametrize("path", ["../report.xlsx", "outputs/../report.xlsx", "outputs//report.xlsx", "outputs/", "", r"outputs\report.xlsx", "./report.xlsx"])
@pytest.mark.parametrize("field", ["deliverables", "input_files_used"])
def test_agent_response_rejects_unsafe_relative_paths(path, field):
    payload = {"status": "completed", "deliverables": [], "summary": "x", "input_files_used": []}
    payload[field] = [path]
    with pytest.raises(ValidationError):
        AgentResponse.model_validate(payload)


def test_agent_response_accepts_posix_relative_paths_and_forbids_extras():
    response = AgentResponse(
        status="completed",
        deliverables=["outputs/report.xlsx"],
        summary="Created report",
        input_files_used=["inputs/0001-source.xlsx"],
    )
    assert response.deliverables == ["outputs/report.xlsx"]
    with pytest.raises(ValidationError):
        AgentResponse.model_validate({**response.model_dump(), "unexpected": True})


def test_agent_response_schema_resource_matches_builder():
    resource = Path(__file__).resolve().parents[1] / "src/workagent_rsi/schemas/agent_response.schema.json"
    schema = json.loads(resource.read_text(encoding="utf-8"))
    assert schema == build_agent_response_schema()
    assert set(schema["required"]) == {"status", "deliverables", "summary", "input_files_used"}
    assert schema["additionalProperties"] is False


def test_provider_outcome_rejects_unknown_status():
    with pytest.raises(ValidationError):
        ProviderOutcome(status="maybe")


def test_provider_uses_ollama_ephemeral_workspace_write_and_schema(tmp_path: Path):
    provider = CodexOfficeProvider(WorkAgentConfig())
    response_path = tmp_path / "agent_response.json"
    command = provider.command(tmp_path, response_path, provider.schema_path)

    assert command[:2] == ["codex", "exec"]
    assert command[command.index("--sandbox") + 1] == "workspace-write"
    assert command[command.index("--local-provider") + 1] == "ollama"
    assert command[command.index("--model") + 1] == "qwen2.5:7b"
    assert command[command.index("--output-schema") + 1] == str(provider.schema_path)
    assert command[command.index("--output-last-message") + 1] == str(response_path)
    assert command[command.index("--cd") + 1] == str(tmp_path)
    assert "--ignore-user-config" in command
    assert "--oss" in command
    assert "--json" in command
    assert "--ephemeral" in command
    assert command[-1] == "-"
    assert "--dangerously-bypass-approvals-and-sandbox" not in command


def test_provider_runs_prompt_on_stdin_and_persists_success_evidence(monkeypatch, tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    (workspace / "input.txt").write_text("source", encoding="utf-8")
    records = tmp_path / "records"
    response = {"status": "completed", "deliverables": ["outputs/report.docx"], "summary": "Created report", "input_files_used": ["inputs/source.txt"]}
    observed = {}

    def fake_run(command, **kwargs):
        observed.update(command=command, **kwargs)
        Path(command[command.index("--output-last-message") + 1]).write_text(json.dumps(response), encoding="utf-8")
        return subprocess.CompletedProcess(command, 0, stdout='{"type":"done"}\n', stderr="notice\n")

    monkeypatch.setattr(subprocess, "run", fake_run)
    outcome = CodexOfficeProvider(WorkAgentConfig()).run("make a report", workspace, records)

    assert outcome.status == "completed"
    assert outcome.response == AgentResponse.model_validate(response)
    assert outcome.record is not None
    assert outcome.record.model_identity == "ollama:qwen2.5:7b"
    assert outcome.record.prompt_hash == canonical_json_hash({"prompt": "make a report"})
    assert outcome.record.workspace_hash == canonical_json_hash({"input.txt": sha256_file(workspace / "input.txt")})
    assert outcome.record.exit_code == 0
    assert outcome.record.started_at <= outcome.record.ended_at
    response_path = Path(observed["command"][observed["command"].index("--output-last-message") + 1])
    assert outcome.record.output_ref == sha256_file(response_path)
    assert observed["input"] == "make a report"
    assert observed["capture_output"] is True
    assert observed["text"] is True
    assert observed["timeout"] == 600
    assert observed["command"][-1] == "-"
    assert (records / "prompt.md").read_text(encoding="utf-8") == "make a report"
    assert (records / "provider.stdout.jsonl").read_text(encoding="utf-8") == '{"type":"done"}\n'
    assert (records / "provider.stderr.txt").read_text(encoding="utf-8") == "notice\n"
    assert response_path.parent == workspace
    assert json.loads(response_path.read_text(encoding="utf-8")) == response
    assert json.loads((workspace / "agent_response.json").read_text(encoding="utf-8")) == response
    assert json.loads((records / "agent_response.json").read_text(encoding="utf-8")) == response
    assert json.loads((records / "command.json").read_text(encoding="utf-8")) == observed["command"]
    assert json.loads((records / "provider_record.json").read_text(encoding="utf-8")) == outcome.record.model_dump(mode="json")


def test_native_ollama_provider_executes_tool_calls_in_workspace(monkeypatch, tmp_path: Path):
    """A native Ollama tool call must cause a real workspace mutation before completion."""
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    records = tmp_path / "records"
    responses = iter([
        {
            "message": {
                "role": "assistant",
                "content": "",
                "tool_calls": [{"function": {"name": "write_file", "arguments": {"path": "outputs/probe.txt", "content": "TOOL_OK"}}}],
            },
            "done": True,
        },
        {
            "message": {
                "role": "assistant",
                "content": json.dumps({"status": "completed", "deliverables": ["outputs/probe.txt"], "summary": "created", "input_files_used": []}),
            },
            "done": True,
        },
    ])

    class FakeResponse:
        def __init__(self, payload):
            self.payload = payload

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self):
            return json.dumps(self.payload).encode("utf-8")

    def fake_urlopen(request, timeout):
        payload = json.loads(request.data.decode("utf-8"))
        assert payload["model"] == "qwen2.5:7b"
        assert payload["tools"]
        return FakeResponse(next(responses))

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    provider = CodexOfficeProvider(WorkAgentConfig(backend="ollama-native"))
    outcome = provider.run("create the probe", workspace, records)

    assert outcome.status == "completed"
    assert (workspace / "outputs/probe.txt").read_text(encoding="utf-8") == "TOOL_OK"
    assert outcome.response is not None
    assert (records / "tool_events.jsonl").is_file()


def test_native_ollama_surfaces_python_tool_failure_as_repair_instruction(monkeypatch, tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    records = tmp_path / "records"
    broken_script = "from pptx import Presentation\ntext = 'first line\nsecond line'\n"
    responses = iter([
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "write_file", "arguments": {"path": "create.py", "content": broken_script}}}
        ]}},
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "run_python", "arguments": {"path": "create.py"}}}
        ]}},
        {"message": {"role": "assistant", "content": json.dumps({
            "status": "completed", "deliverables": ["outputs/briefing.pptx"],
            "summary": "created", "input_files_used": []
        })}},
        {"message": {"role": "assistant", "content": ""}},
    ])
    requests = []

    class FakeResponse:
        def __init__(self, payload):
            self.payload = payload

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self):
            return json.dumps(self.payload).encode("utf-8")

    def fake_urlopen(request, timeout):
        requests.append(json.loads(request.data.decode("utf-8")))
        return FakeResponse(next(responses))

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    outcome = CodexOfficeProvider(WorkAgentConfig(backend="ollama-native", max_tool_turns=4)).run(
        "You are creating an Office deliverable for a powerpoint task.\nCreate a deck.",
        workspace,
        records,
    )

    assert outcome.status == "failed"
    repair_messages = [
        item["content"] for item in requests[2]["messages"]
        if item.get("role") == "user" and "tool" in item.get("content", "").lower()
    ]
    assert repair_messages
    assert "run_python" in repair_messages[-1]
    assert "SyntaxError" in repair_messages[-1]
    assert "do not claim completion" in repair_messages[-1].lower()


def test_native_ollama_provider_executes_batched_tool_calls_in_order(monkeypatch, tmp_path: Path):
    """A model response containing multiple calls must execute them in order."""
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    records = tmp_path / "records"
    responses = iter([
        {
            "message": {
                "role": "assistant",
                "content": "",
                "tool_calls": [
                    {"function": {"name": "write_file", "arguments": {"path": "script.py", "content": "from pathlib import Path\nPath('outputs/batched.txt').write_text('OK')"}}},
                    {"function": {"name": "run_python", "arguments": {"path": "script.py"}}},
                ],
            },
            "done": True,
        },
        {
            "message": {
                "role": "assistant",
                "content": json.dumps({"status": "completed", "deliverables": ["outputs/batched.txt"], "summary": "created", "input_files_used": []}),
            },
            "done": True,
        },
    ])

    class FakeResponse:
        def __init__(self, payload):
            self.payload = payload

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self):
            return json.dumps(self.payload).encode("utf-8")

    monkeypatch.setattr(urllib.request, "urlopen", lambda request, timeout: FakeResponse(next(responses)))
    provider = CodexOfficeProvider(WorkAgentConfig(backend="ollama-native"))
    outcome = provider.run("create the batched probe", workspace, records)

    assert outcome.status == "completed"
    assert (workspace / "outputs/batched.txt").read_text(encoding="utf-8") == "OK"
    events = [json.loads(line) for line in (records / "tool_events.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [event["name"] for event in events[:2]] == ["write_file", "run_python"]


def test_native_ollama_office_prompt_rejects_empty_early_completion(monkeypatch, tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    records = tmp_path / "records"
    script = (
        "from openpyxl import Workbook, load_workbook\n"
        "from pathlib import Path\n"
        "import json\n"
        "Path('outputs').mkdir(exist_ok=True)\n"
        "book = Workbook()\n"
        "book.active.title = 'Summary'\n"
        "book.active['B2'] = 42\n"
        "book.save('outputs/report.xlsx')\n"
        "check = load_workbook('outputs/report.xlsx', data_only=False)\n"
        "assert check['Summary']['B2'].value == 42\n"
        "Path('deliverables.json').write_text(json.dumps({'deliverables': ['outputs/report.xlsx']}))\n"
    )
    premature = {
        "status": "completed", "deliverables": [], "summary": "done", "input_files_used": []
    }
    responses = iter([
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "write_file", "arguments": {"path": "create.py", "content": script}}}
        ]}},
        {"message": {"role": "assistant", "content": json.dumps(premature)}},
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "run_python", "arguments": {"path": "create.py"}}}
        ]}},
        {"message": {"role": "assistant", "content": json.dumps({
            "status": "completed", "deliverables": ["outputs/report.xlsx"],
            "summary": "created", "input_files_used": []
        })}},
    ])
    requests = []

    class FakeResponse:
        def __init__(self, payload):
            self.payload = payload

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self):
            return json.dumps(self.payload).encode("utf-8")

    def fake_urlopen(request, timeout):
        requests.append(json.loads(request.data.decode("utf-8")))
        return FakeResponse(next(responses))

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    prompt = "You are creating an Office deliverable for a excel task.\nCreate Summary!B2=42."
    outcome = CodexOfficeProvider(WorkAgentConfig(backend="ollama-native")).run(prompt, workspace, records)

    assert outcome.status == "completed"
    assert outcome.response is not None
    assert outcome.response.deliverables == ["outputs/report.xlsx"]
    assert len(requests) == 4
    system_prompt = requests[0]["messages"][0]["content"]
    assert "write_file path exactly deliverables.json" in system_prompt
    assert "do not assert a cached numeric result" in system_prompt
    assert "assert every requested sheet, cell, and formula" in system_prompt
    assert "preserve every existing sheet and cell" in system_prompt
    assert "search all shapes and paragraphs" in system_prompt
    assert "input deck may contain blank slides" in system_prompt
    assert "do not use slide.shapes.title" in system_prompt
    assert "do not use fixed placeholders" in system_prompt
    assert "slide.shapes.add_textbox" in system_prompt
    assert "deliverables" in requests[2]["messages"][-1]["content"]
    assert (workspace / "outputs/report.xlsx").is_file()
    assert json.loads((workspace / "deliverables.json").read_text(encoding="utf-8")) == {
        "deliverables": ["outputs/report.xlsx"]
    }


def test_native_ollama_office_requires_manifest_tool_call_after_early_completion(monkeypatch, tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    records = tmp_path / "records"
    script = (
        "from openpyxl import Workbook, load_workbook\n"
        "from pathlib import Path\n"
        "Path('outputs').mkdir(exist_ok=True)\n"
        "book = Workbook()\n"
        "book.active['A1'] = 'verified'\n"
        "book.save('outputs/report.xlsx')\n"
        "check = load_workbook('outputs/report.xlsx', data_only=False)\n"
        "assert check.active['A1'].value == 'verified'\n"
    )
    responses = iter([
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "write_file", "arguments": {"path": "create.py", "content": script}}}
        ]}},
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "run_python", "arguments": {"path": "create.py"}}}
        ]}},
        {"message": {"role": "assistant", "content": json.dumps({
            "status": "completed", "deliverables": [], "summary": "created", "input_files_used": []
        })}},
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "write_file", "arguments": {
                "path": "deliverables.json",
                "content": '{"deliverables":["outputs/report.xlsx"]}'
            }}}
        ]}},
        {"message": {"role": "assistant", "content": json.dumps({
            "status": "completed", "deliverables": ["outputs/report.xlsx"],
            "summary": "created and verified", "input_files_used": []
        })}},
    ])
    requests = []

    class FakeResponse:
        def __init__(self, payload):
            self.payload = payload

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self):
            return json.dumps(self.payload).encode("utf-8")

    def fake_urlopen(request, timeout):
        requests.append(json.loads(request.data.decode("utf-8")))
        return FakeResponse(next(responses))

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    prompt = "You are creating an Office deliverable for a excel task.\nCreate a workbook."
    outcome = CodexOfficeProvider(WorkAgentConfig(backend="ollama-native")).run(prompt, workspace, records)

    assert outcome.status == "completed"
    assert outcome.response is not None
    assert outcome.response.deliverables == ["outputs/report.xlsx"]
    assert len(requests) == 5
    assert "must call write_file" in requests[3]["messages"][-1]["content"].lower()
    assert "read_file" not in {tool["function"]["name"] for tool in requests[0]["tools"]}
    events = [json.loads(line) for line in (records / "tool_events.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [event["name"] for event in events] == ["write_file", "run_python", "write_file"]


def test_native_ollama_tool_rejects_reading_binary_office_file(tmp_path: Path):
    workspace = tmp_path / "workspace"
    (workspace / "outputs").mkdir(parents=True)
    (workspace / "outputs/report.xlsx").write_bytes(b"PK\x03\x04binary")

    result = _execute_native_tool("read_file", {"path": "outputs/report.xlsx"}, workspace)

    assert result["ok"] is False
    assert "binary Office" in result["error"]


def test_native_ollama_manifest_requires_existing_workspace_deliverables(tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    manifest = '{"deliverables":["outputs/report.xlsx"]}'

    premature = _execute_native_tool(
        "write_file", {"path": "deliverables.json", "content": manifest}, workspace
    )

    assert premature["ok"] is False
    assert "outputs/report.xlsx" in premature["error"]
    assert not (workspace / "deliverables.json").exists()

    output = workspace / "outputs/report.xlsx"
    output.parent.mkdir()
    output.write_bytes(b"office artifact")
    accepted = _execute_native_tool(
        "write_file", {"path": "deliverables.json", "content": manifest}, workspace
    )

    assert accepted["ok"] is True
    assert json.loads((workspace / "deliverables.json").read_text(encoding="utf-8")) == {
        "deliverables": ["outputs/report.xlsx"]
    }


def test_native_ollama_invalid_intermediate_status_requires_manifest_tool_call(monkeypatch, tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    records = tmp_path / "records"
    script = (
        "from openpyxl import Workbook\n"
        "from pathlib import Path\n"
        "Path('outputs').mkdir(exist_ok=True)\n"
        "book = Workbook()\n"
        "book.active['A1'] = 'verified'\n"
        "book.save('outputs/report.xlsx')\n"
    )
    responses = iter([
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "write_file", "arguments": {"path": "create.py", "content": script}}}
        ]}},
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "run_python", "arguments": {"path": "create.py"}}}
        ]}},
        {"message": {"role": "assistant", "content": json.dumps({
            "status": "writing_manifest", "input_files_used": []
        })}},
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "write_file", "arguments": {
                "path": "deliverables.json",
                "content": '{"deliverables":["outputs/report.xlsx"]}'
            }}}
        ]}},
        {"message": {"role": "assistant", "content": json.dumps({
            "status": "completed", "deliverables": ["outputs/report.xlsx"],
            "summary": "created and verified", "input_files_used": []
        })}},
    ])
    requests = []

    class FakeResponse:
        def __init__(self, payload):
            self.payload = payload

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self):
            return json.dumps(self.payload).encode("utf-8")

    def fake_urlopen(request, timeout):
        requests.append(json.loads(request.data.decode("utf-8")))
        return FakeResponse(next(responses))

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    outcome = CodexOfficeProvider(WorkAgentConfig(backend="ollama-native")).run(
        "You are creating an Office deliverable for a excel task.\nCreate a workbook.",
        workspace,
        records,
    )

    assert outcome.status == "completed"
    assert outcome.response is not None
    assert outcome.response.deliverables == ["outputs/report.xlsx"]
    assert len(requests) == 5
    correction = requests[3]["messages"][-1]["content"].lower()
    assert "must call write_file" in correction
    assert "content argument must be a string" in correction
    assert "do not call tools" not in correction
    events = [json.loads(line) for line in (records / "tool_events.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [event["name"] for event in events] == ["write_file", "run_python", "write_file"]


def test_native_ollama_retries_completion_with_exact_staged_input_paths(monkeypatch, tmp_path: Path):
    workspace = tmp_path / "workspace"
    (workspace / "outputs").mkdir(parents=True)
    (workspace / "outputs/probe.txt").write_text("ready", encoding="utf-8")
    (workspace / "inputs").mkdir()
    (workspace / "inputs/0001-source.txt").write_text("source", encoding="utf-8")
    (workspace / "input_manifest.json").write_text(
        json.dumps({"files": [{"copied": "inputs/0001-source.txt"}]}), encoding="utf-8"
    )
    (workspace / "deliverables.json").write_text(
        json.dumps({"deliverables": ["outputs/probe.txt"]}), encoding="utf-8"
    )
    records = tmp_path / "records"
    responses = iter([
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "list_files", "arguments": {}}}
        ]}},
        {"message": {"role": "assistant", "content": json.dumps({
            "status": "completed", "deliverables": ["outputs/probe.txt"],
            "summary": "created", "input_files_used": [""]
        })}},
        {"message": {"role": "assistant", "content": json.dumps({
            "status": "completed", "deliverables": ["outputs/probe.txt"],
            "summary": "created", "input_files_used": ["inputs/0001-source.txt"]
        })}},
    ])
    requests = []

    class FakeResponse:
        def __init__(self, payload):
            self.payload = payload

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self):
            return json.dumps(self.payload).encode("utf-8")

    def fake_urlopen(request, timeout):
        requests.append(json.loads(request.data.decode("utf-8")))
        return FakeResponse(next(responses))

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    outcome = CodexOfficeProvider(WorkAgentConfig(backend="ollama-native")).run(
        "Create a text probe", workspace, records
    )

    assert outcome.status == "completed"
    assert outcome.response is not None
    assert outcome.response.input_files_used == ["inputs/0001-source.txt"]
    assert len(requests) == 3
    correction = requests[2]["messages"][-1]["content"]
    assert "input_files_used" in correction
    assert "inputs/0001-source.txt" in correction


def test_native_ollama_retries_invalid_final_json_but_fails_closed(monkeypatch, tmp_path: Path):
    workspace = tmp_path / "workspace"
    (workspace / "outputs").mkdir(parents=True)
    artifact = workspace / "outputs/probe.txt"
    artifact.write_text("ready", encoding="utf-8")
    records = tmp_path / "records"
    responses = iter([
        {"message": {"role": "assistant", "content": "", "tool_calls": [
            {"function": {"name": "list_files", "arguments": {}}}
        ]}},
        {"message": {"role": "assistant", "content": "The task is complete."}},
        {"message": {"role": "assistant", "content": "Still complete."}},
        {"message": {"role": "assistant", "content": "Still complete."}},
    ])
    requests = []

    class FakeResponse:
        def __init__(self, payload):
            self.payload = payload

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self):
            return json.dumps(self.payload).encode("utf-8")

    def fake_urlopen(request, timeout):
        requests.append(json.loads(request.data.decode("utf-8")))
        return FakeResponse(next(responses))

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    outcome = CodexOfficeProvider(WorkAgentConfig(backend="ollama-native", max_tool_turns=4)).run(
        "Create a probe", workspace, records
    )

    assert outcome.status == "failed"
    assert outcome.response is None
    assert len(requests) == 4
    assert "invalid provider response" in (outcome.error or "")
    assert "valid JSON" in requests[2]["messages"][-1]["content"]


def test_provider_timeout_preserves_partial_output(monkeypatch, tmp_path: Path):
    def fake_run(command, **kwargs):
        raise subprocess.TimeoutExpired(command, kwargs["timeout"], output=b"partial stdout", stderr=b"partial stderr")

    monkeypatch.setattr(subprocess, "run", fake_run)
    records = tmp_path / "records"
    outcome = CodexOfficeProvider(WorkAgentConfig()).run("prompt", tmp_path, records)

    assert outcome.status == "timeout"
    assert outcome.response is None
    assert outcome.record is not None
    assert outcome.record.model_identity == "ollama:qwen2.5:7b"
    assert outcome.record.exit_code is None
    assert (records / "provider.stdout.jsonl").read_text(encoding="utf-8") == "partial stdout"
    assert (records / "provider.stderr.txt").read_text(encoding="utf-8") == "partial stderr"
    assert json.loads((records / "provider_record.json").read_text(encoding="utf-8"))["status"] == "timeout"


def test_provider_missing_executable_is_unavailable(monkeypatch, tmp_path: Path):
    def fake_run(command, **kwargs):
        raise FileNotFoundError("codex not found")

    monkeypatch.setattr(subprocess, "run", fake_run)
    records = tmp_path / "records"
    outcome = CodexOfficeProvider(WorkAgentConfig()).run("prompt", tmp_path, records)

    assert outcome.status == "unavailable"
    assert outcome.response is None
    assert outcome.record is not None
    assert outcome.record.model_identity == "ollama:qwen2.5:7b"
    assert outcome.record.exit_code is None
    assert "codex not found" in (outcome.error or "")
    assert (records / "provider.stdout.jsonl").read_text(encoding="utf-8") == ""
    assert (records / "provider.stderr.txt").read_text(encoding="utf-8") == ""
    assert json.loads((records / "provider_record.json").read_text(encoding="utf-8"))["status"] == "unavailable"


def test_provider_reused_record_root_does_not_accept_stale_response(monkeypatch, tmp_path: Path):
    records = tmp_path / "records"
    response = {"status": "completed", "deliverables": [], "summary": "first run", "input_files_used": []}
    calls = []

    def fake_run(command, **kwargs):
        output_path = Path(command[command.index("--output-last-message") + 1])
        calls.append(output_path)
        if len(calls) == 1:
            output_path.write_text(json.dumps(response), encoding="utf-8")
        return subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setattr(subprocess, "run", fake_run)
    provider = CodexOfficeProvider(WorkAgentConfig())
    first = provider.run("first", tmp_path, records)
    second = provider.run("second", tmp_path, records)

    assert first.status == "completed"
    assert second.status == "failed"
    assert second.response is None
    assert second.record is not None
    assert second.record.output_ref is None
    assert calls[0] != calls[1]
    assert calls[0].exists()
    assert json.loads(calls[0].read_text(encoding="utf-8")) == response
    assert not (tmp_path / "agent_response.json").exists()


def _office_task(domain="excel", input_files=()):
    return TaskSpec(task_id="pilot-1", domain=domain, instruction="Build a report", input_files=input_files)


def test_input_copy_preserves_original_and_records_private_hashes(tmp_path: Path):
    input_base = tmp_path / "sources"
    input_base.mkdir()
    source = input_base / "source.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "original"
    workbook.save(source)
    before = sha256_file(source)
    workspace = tmp_path / "workspace"

    manifest = copy_task_inputs(_office_task(input_files=("source.xlsx",)), workspace, input_base)

    copied = workspace / "inputs/0001-source.xlsx"
    assert copied.is_file()
    assert sha256_file(source) == before == sha256_file(copied)
    assert manifest["files"][0]["source_sha256"] == before
    assert manifest["files"][0]["copy_sha256"] == before
    assert manifest["files"][0]["copied"] == "inputs/0001-source.xlsx"
    assert json.loads((workspace / "input_manifest.json").read_text(encoding="utf-8")) == manifest
    assert str(input_base) not in json.dumps(manifest)


def test_input_copy_rejects_cross_domain_and_nonregular_files(tmp_path: Path):
    source = tmp_path / "source.docx"
    Document().save(source)
    with pytest.raises(ValueError):
        copy_task_inputs(_office_task(input_files=("source.docx",)), tmp_path / "workspace", tmp_path)
    (tmp_path / "folder.xlsx").mkdir()
    with pytest.raises(ValueError):
        copy_task_inputs(_office_task(input_files=("folder.xlsx",)), tmp_path / "workspace", tmp_path)


@pytest.mark.parametrize("name", ["../source.xlsx", "missing.xlsx", "source.xls", "//server/share/source.xlsx", r"\\server\share\source.xlsx"])
def test_input_copy_rejects_unsafe_or_unsupported_paths(tmp_path: Path, name: str):
    with pytest.raises(ValueError):
        copy_task_inputs(_office_task(input_files=(name,)), tmp_path / "workspace", tmp_path)


def test_input_copy_accepts_explicit_local_absolute_path_without_leaking_it(tmp_path: Path):
    input_base = tmp_path / "unused-base"
    input_base.mkdir()
    source = tmp_path / "provided" / "source.xlsx"
    source.parent.mkdir()
    Workbook().save(source)
    task = _office_task(input_files=(str(source),))
    workspace = tmp_path / "workspace"

    manifest = copy_task_inputs(task, workspace, input_base)
    verify_task_input_hashes(task, workspace, input_base, manifest)
    prompt = build_task_prompt(task, manifest, "Use clear sheet names")

    assert sha256_file(workspace / "inputs/0001-source.xlsx") == sha256_file(source)
    assert str(source) not in json.dumps(manifest)
    assert str(source) not in (workspace / "input_manifest.json").read_text(encoding="utf-8")
    assert str(source) not in prompt


def test_input_copy_rejects_symlinked_parent(tmp_path: Path):
    real = tmp_path / "real"
    real.mkdir()
    Workbook().save(real / "source.xlsx")
    link = tmp_path / "link"
    try:
        link.symlink_to(real, target_is_directory=True)
    except OSError as exc:
        pytest.skip(f"symlinks unavailable: {exc}")
    with pytest.raises(ValueError):
        copy_task_inputs(_office_task(input_files=("link/source.xlsx",)), tmp_path / "workspace", tmp_path)


def test_input_copy_rejects_preexisting_destination_symlink(tmp_path: Path):
    Workbook().save(tmp_path / "source.xlsx")
    workspace = tmp_path / "workspace"
    (workspace / "inputs").mkdir(parents=True)
    destination = workspace / "inputs/0001-source.xlsx"
    try:
        destination.symlink_to(tmp_path / "source.xlsx")
    except OSError as exc:
        pytest.skip(f"symlinks unavailable: {exc}")
    with pytest.raises(ValueError):
        copy_task_inputs(_office_task(input_files=("source.xlsx",)), workspace, tmp_path)


@pytest.mark.parametrize("which", ["source", "copy"])
def test_input_hash_verification_rejects_post_provider_tampering(tmp_path: Path, which: str):
    source = tmp_path / "source.xlsx"
    Workbook().save(source)
    workspace = tmp_path / "workspace"
    task = _office_task(input_files=("source.xlsx",))
    manifest = copy_task_inputs(task, workspace, tmp_path)
    verify_task_input_hashes(task, workspace, tmp_path, manifest)
    target = source if which == "source" else workspace / "inputs/0001-source.xlsx"
    target.write_bytes(b"tampered")
    with pytest.raises(ValueError):
        verify_task_input_hashes(task, workspace, tmp_path, manifest)


def test_task_prompt_does_not_expose_evaluator_constraints_or_source_paths(tmp_path: Path):
    secret_path = str(tmp_path / "secret-source.xlsx")
    task = TaskSpec(task_id="pilot-1", domain="excel", instruction="Build an expense workbook",
                    input_files=(secret_path,), expected_constraints={"required_cells": {"Summary!B2": 999}})
    prompt = build_task_prompt(task, {"files": [{"copied": "inputs/0001-source.xlsx", "source": "source.xlsx"}]},
                               "Use clear sheet names")
    assert "Build an expense workbook" in prompt
    assert "excel" in prompt and "inputs/0001-source.xlsx" in prompt
    assert "outputs/" in prompt and "deliverables.json" in prompt
    assert "Use clear sheet names" in prompt
    assert "999" not in prompt and "expected_constraints" not in prompt
    assert secret_path not in prompt


def _deliverable_fixture(tmp_path: Path, paths=("outputs/report.xlsx",)):
    workspace = tmp_path / "workspace"
    (workspace / "outputs").mkdir(parents=True)
    response = AgentResponse(status="completed", deliverables=list(paths), summary="done", input_files_used=[])
    (workspace / "agent_response.json").write_text(response.model_dump_json(), encoding="utf-8")
    (workspace / "deliverables.json").write_text(json.dumps({"deliverables": list(paths)}), encoding="utf-8")
    for path in paths:
        if path.startswith("outputs/") and ".." not in path:
            target = workspace / path
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(b"office package")
    return workspace, response


def test_deliverable_accepts_only_manifested_domain_file(tmp_path: Path):
    workspace, response = _deliverable_fixture(tmp_path)
    assert validate_deliverables(_office_task(), workspace, response, WorkAgentConfig()) == [workspace / "outputs/report.xlsx"]


@pytest.mark.parametrize("path", ["outputs/../outside.xlsx", "C:/outside.xlsx", "/outside.xlsx", "outputs/report.docx"])
def test_deliverable_rejects_unsafe_path_or_cross_format(tmp_path: Path, path: str):
    workspace, response = _deliverable_fixture(tmp_path)
    (workspace / "deliverables.json").write_text(json.dumps({"deliverables": [path]}), encoding="utf-8")
    with pytest.raises(ValueError):
        validate_deliverables(_office_task(), workspace, response, WorkAgentConfig())


def test_deliverable_rejects_unlisted_output_and_missing_manifest(tmp_path: Path):
    workspace, response = _deliverable_fixture(tmp_path)
    (workspace / "deliverables.json").write_text(json.dumps({"deliverables": []}), encoding="utf-8")
    with pytest.raises(ValueError):
        validate_deliverables(_office_task(), workspace, response, WorkAgentConfig())
    (workspace / "deliverables.json").unlink()
    with pytest.raises(ValueError):
        validate_deliverables(_office_task(), workspace, response, WorkAgentConfig())


def test_deliverable_rejects_symlink_empty_and_oversize(tmp_path: Path):
    workspace, response = _deliverable_fixture(tmp_path)
    output = workspace / "outputs/report.xlsx"
    output.write_bytes(b"")
    with pytest.raises(ValueError):
        validate_deliverables(_office_task(), workspace, response, WorkAgentConfig())
    output.write_bytes(b"too large")
    with pytest.raises(ValueError):
        validate_deliverables(_office_task(), workspace, response, WorkAgentConfig(max_artifact_bytes=3))
    output.unlink()
    target = tmp_path / "outside.xlsx"
    target.write_bytes(b"content")
    try:
        output.symlink_to(target)
    except OSError as exc:
        pytest.skip(f"symlinks unavailable: {exc}")
    with pytest.raises(ValueError):
        validate_deliverables(_office_task(), workspace, response, WorkAgentConfig())


def test_deliverable_rejects_file_count_and_malformed_response(tmp_path: Path):
    workspace, response = _deliverable_fixture(tmp_path, ("outputs/one.xlsx", "outputs/two.xlsx"))
    with pytest.raises(ValueError):
        validate_deliverables(_office_task(), workspace, response, WorkAgentConfig(max_output_files=1))
    (workspace / "agent_response.json").write_text("{}", encoding="utf-8")
    with pytest.raises(ValueError):
        validate_deliverables(_office_task(), workspace, response, WorkAgentConfig())


def _harness_excel_task(input_files=()):
    return TaskSpec(
        task_id="harness-excel",
        domain="excel",
        instruction="Create a Summary workbook with B2 set to 42",
        input_files=input_files,
        expected_constraints={"required_cells": {"Summary!B2": 42}},
    )


def test_harness_workagent_stores_validated_excel_and_provider_trace(monkeypatch, tmp_path: Path):
    source = tmp_path / "source.xlsx"
    Workbook().save(source)
    source_hash = sha256_file(source)

    def fake_run(command, **kwargs):
        workspace = Path(command[command.index("--cd") + 1])
        assert kwargs["cwd"] == workspace
        assert (workspace / "inputs/0001-source.xlsx").is_file()
        assert "Summary workbook" in kwargs["input"]
        output = workspace / "outputs/report.xlsx"
        output.parent.mkdir()
        workbook = Workbook()
        workbook.active.title = "Summary"
        workbook.active["B2"] = 42
        workbook.save(output)
        response = {"status": "completed", "deliverables": ["outputs/report.xlsx"],
                    "summary": "Created workbook", "input_files_used": ["inputs/0001-source.xlsx"]}
        (workspace / "deliverables.json").write_text(json.dumps({"deliverables": response["deliverables"]}), encoding="utf-8")
        Path(command[command.index("--output-last-message") + 1]).write_text(json.dumps(response), encoding="utf-8")
        return subprocess.CompletedProcess(command, 0, stdout='{"type":"done"}\n', stderr="")

    monkeypatch.setattr(subprocess, "run", fake_run)
    harness = Harness(tmp_path / "results", execution_provider="workagent", workagent_config=WorkAgentConfig(verify_com=False))
    result = harness.run(_harness_excel_task((str(source),)))

    assert result["state"] == "SUCCEEDED"
    assert result["evaluation"]["passed"] is True
    assert sha256_file(source) == source_hash
    assert len(result["artifacts"]) == 1
    artifact = Path(result["artifacts"][0]["path"])
    assert artifact.is_file() and artifact.parent == Path(result["result_dir"]) / "artifacts"
    assert artifact.read_bytes() == (Path(result["result_dir"]) / "agent_workspace/outputs/report.xlsx").read_bytes()
    assert (Path(result["result_dir"]) / "provider_records/provider_record.json").is_file()
    saved = json.loads(Path(result["result_path"]).read_text(encoding="utf-8"))
    assert saved["run_id"] == result["run_id"] and saved["artifacts"] == result["artifacts"]
    events = harness.resume(result["result_dir"])["events"]
    kinds = [event["kind"] for event in events]
    assert kinds == ["run_started", "started", "input_manifest", "provider_output", "artifact", "com_reopen"]
    provider = events[3]["payload"]
    assert provider["record"]["provider"] == "codex-cli"
    assert provider["record"]["model_identity"] == "ollama:qwen2.5:7b"


def test_harness_workagent_unavailable_does_not_fall_back_to_template(monkeypatch, tmp_path: Path):
    def missing_provider(command, **kwargs):
        raise FileNotFoundError("codex not found")

    monkeypatch.setattr(subprocess, "run", missing_provider)
    result = Harness(tmp_path / "results", execution_provider="workagent").run(_harness_excel_task())

    assert result["state"] == "UNAVAILABLE"
    assert result["artifacts"] == []
    assert "evaluation" not in result
    assert not list((Path(result["result_dir"]) / "artifacts").iterdir())
    kinds = [event["kind"] for event in Harness().resume(result["result_dir"])["events"]]
    assert kinds == ["run_started", "started", "input_manifest", "provider_output", "unavailable"]
    assert not (Path(result["result_dir"]) / "generated").exists()


def test_harness_workagent_timeout_never_evaluates_or_emits_template(monkeypatch, tmp_path: Path):
    def timed_out(command, **kwargs):
        raise subprocess.TimeoutExpired(command, kwargs["timeout"], output=b"partial")

    monkeypatch.setattr(subprocess, "run", timed_out)
    result = Harness(tmp_path / "results", execution_provider="workagent").run(_harness_excel_task())

    assert result["state"] == "FAILED"
    assert result["failure"]["status"] == "timeout"
    assert result["artifacts"] == [] and "evaluation" not in result
    kinds = [event["kind"] for event in Harness().resume(result["result_dir"])["events"]]
    assert kinds[-1] == "failure" and "artifact" not in kinds
    assert not (Path(result["result_dir"]) / "generated").exists()


@pytest.mark.parametrize("target", ["source", "copy"])
def test_harness_workagent_rejects_input_tampering_after_provider_exit(monkeypatch, tmp_path: Path, target: str):
    source = tmp_path / "source.xlsx"
    Workbook().save(source)

    def tampering_provider(command, **kwargs):
        workspace = Path(command[command.index("--cd") + 1])
        changed = source if target == "source" else workspace / "inputs/0001-source.xlsx"
        changed.write_bytes(b"changed")
        output = workspace / "outputs/report.xlsx"
        output.parent.mkdir()
        workbook = Workbook()
        workbook.active.title = "Summary"
        workbook.active["B2"] = 42
        workbook.save(output)
        response = {"status": "completed", "deliverables": ["outputs/report.xlsx"],
                    "summary": "Created workbook", "input_files_used": ["inputs/0001-source.xlsx"]}
        (workspace / "deliverables.json").write_text(json.dumps({"deliverables": response["deliverables"]}), encoding="utf-8")
        Path(command[command.index("--output-last-message") + 1]).write_text(json.dumps(response), encoding="utf-8")
        return subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setattr(subprocess, "run", tampering_provider)
    result = Harness(tmp_path / "results", execution_provider="workagent").run(_harness_excel_task((str(source),)))

    assert result["state"] == "FAILED"
    assert "input" in result["failure"]["message"]
    assert result["artifacts"] == [] and "evaluation" not in result
    kinds = [event["kind"] for event in Harness().resume(result["result_dir"])["events"]]
    assert kinds[-1] == "failure" and "artifact" not in kinds


def test_harness_default_office_uses_workagent_and_separates_invocations(monkeypatch, tmp_path: Path):
    def missing_provider(command, **kwargs):
        raise FileNotFoundError("codex not found")

    monkeypatch.setattr(subprocess, "run", missing_provider)
    harness = Harness(tmp_path / "results")
    first = harness.run(_harness_excel_task())
    second = harness.run(_harness_excel_task())

    assert first["state"] == second["state"] == "UNAVAILABLE"
    assert first["result_dir"] != second["result_dir"]
    assert Path(first["result_path"]).is_file() and Path(second["result_path"]).is_file()


def test_harness_persisted_task_excludes_evaluator_constraints_and_source_paths(monkeypatch, tmp_path: Path):
    source = tmp_path / "private-source.xlsx"
    Workbook().save(source)

    def missing_provider(command, **kwargs):
        raise FileNotFoundError("codex not found")

    monkeypatch.setattr(subprocess, "run", missing_provider)
    task = TaskSpec(
        task_id="private-task",
        domain="excel",
        instruction="Build a workbook",
        input_files=(str(source),),
        expected_constraints={"required_cells": {"Summary!B2": "PRIVATE-EVALUATOR-ANSWER"}},
    )
    result = Harness(tmp_path / "results", execution_provider="workagent").run(task)

    task_text = (Path(result["result_dir"]) / "task.json").read_text(encoding="utf-8")
    saved = json.loads(task_text)
    assert saved["task_id"] == "private-task" and saved["instruction"] == "Build a workbook"
    assert "expected_constraints" not in saved and "input_files" not in saved
    assert "PRIVATE-EVALUATOR-ANSWER" not in task_text
    assert str(source) not in task_text


def test_harness_provider_permission_error_persists_evidence_and_one_unavailable(monkeypatch, tmp_path: Path):
    def denied_provider(command, **kwargs):
        raise PermissionError("provider launch denied")

    monkeypatch.setattr(subprocess, "run", denied_provider)
    result = Harness(tmp_path / "results", execution_provider="workagent").run(_harness_excel_task())

    assert result["state"] == "UNAVAILABLE"
    assert result["artifacts"] == [] and "evaluation" not in result
    root = Path(result["result_dir"])
    record = json.loads((root / "provider_records/provider_record.json").read_text(encoding="utf-8"))
    assert record["status"] == "unavailable"
    assert "provider launch denied" in record["error"]
    events = Harness().resume(root)["events"]
    kinds = [event["kind"] for event in events]
    assert kinds == ["run_started", "started", "input_manifest", "provider_output", "unavailable"]
    assert events[3]["payload"]["record"] == record
    assert "provider launch denied" in events[4]["payload"]["message"]


def test_evaluator_excel_pilot_checks_sheets_cells_and_formula(tmp_path: Path):
    path = tmp_path / "quarterly.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Summary"
    sheet["A1"] = "Quarter"
    sheet["B2"] = 120
    sheet["B5"] = "=SUM(B2:B4)"
    book.save(path)
    passing = TaskSpec(task_id="excel-pass", domain="excel", instruction="Create workbook", expected_constraints={
        "required_sheets": ["Summary"], "required_cells": {"Summary!A1": "Quarter", "Summary!B2": 120},
        "required_formulas": {"Summary!B5": "=SUM(B2:B4)"},
    })
    failing = passing.model_copy(update={"expected_constraints": {"required_sheets": ["Missing"]}})
    assert inspect_office_file(path, passing).passed
    assert any("required sheet" in item for item in inspect_office_file(path, failing).failures)


def test_evaluator_word_pilot_checks_headings_styles_and_text(tmp_path: Path):
    path = tmp_path / "status.docx"
    doc = Document()
    doc.add_heading("Status", level=1)
    doc.add_paragraph("Owner: Maya. Action: confirm launch date.")
    doc.save(path)
    passing = TaskSpec(task_id="word-pass", domain="word", instruction="Create report", expected_constraints={
        "required_headings": ["Status"], "required_heading_styles": {"Status": "Heading 1"},
        "required_text": "Action: confirm launch date.",
    })
    failing = passing.model_copy(update={"expected_constraints": {"required_heading_styles": {"Status": "Heading 2"}}})
    assert inspect_office_file(path, passing).passed
    assert any("expected style" in item for item in inspect_office_file(path, failing).failures)


def test_evaluator_powerpoint_pilot_checks_count_text_and_geometry(tmp_path: Path):
    path = tmp_path / "briefing.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    shape.text = "Launch briefing"
    deck.save(path)
    passing = TaskSpec(task_id="powerpoint-pass", domain="powerpoint", instruction="Create briefing", expected_constraints={
        "required_slide_count": 1, "required_shape_text": ["Launch briefing"],
    })
    failing = passing.model_copy(update={"expected_constraints": {"required_slide_count": 2}})
    report = inspect_office_file(path, passing)
    assert report.passed and report.channel_status["geometry"] == "available"
    assert any("slide count" in item for item in inspect_office_file(path, failing).failures)


def test_general_office_runner_refuses_pilot_before_all_com_apps_available(monkeypatch, tmp_path: Path):
    from project_artifacts.phase3_experiments.scripts import run_general_office_pilot as pilot
    from workagent_rsi.office_capabilities import CapabilityReport

    monkeypatch.setattr(pilot, "RESULT_ROOT", tmp_path)
    monkeypatch.setattr(pilot, "probe_capabilities", lambda: CapabilityReport.for_testing(powerpoint_com=None))
    called = []
    monkeypatch.setattr(pilot, "Harness", lambda *args, **kwargs: called.append("harness"))
    assert pilot.main([]) == 1
    assert called == []
    invocation = next(tmp_path.iterdir())
    assert (invocation / "capability_report.json").is_file()
    summary = json.loads((invocation / "summary.json").read_text(encoding="utf-8"))
    assert summary["task_count"] == 0
    assert "retry_lineage" not in summary
    assert "Retry lineage" not in (invocation / "qualification_report.md").read_text(encoding="utf-8")


def test_general_office_runner_task_filter_runs_only_selected_case(monkeypatch, tmp_path: Path):
    from project_artifacts.phase3_experiments.scripts import run_general_office_pilot as pilot
    from workagent_rsi.office_capabilities import CapabilityReport

    monkeypatch.setattr(pilot, "RESULT_ROOT", tmp_path)
    monkeypatch.setattr(pilot, "probe_capabilities", CapabilityReport.for_testing)
    called = []
    provider_configs = []

    class FailedHarness:
        def __init__(self, root, **kwargs):
            self.root = root
            provider_configs.append(kwargs["workagent_config"])

        def run(self, task, **kwargs):
            called.append(task.task_id)
            run_root = self.root / task.task_id
            run_root.mkdir()
            return {
                "run_id": task.task_id,
                "result_dir": str(run_root),
                "state": "FAILED",
                "artifacts": [],
                "failure": {"message": "controlled provider boundary", "status": "failed"},
            }

    monkeypatch.setattr(pilot, "Harness", FailedHarness)
    assert pilot.main(["--task", "excel-create"]) == 1
    assert called == ["excel-create"]
    invocation = next(tmp_path.iterdir())
    summary = json.loads((invocation / "summary.json").read_text(encoding="utf-8"))
    assert summary["task_filter"] == "excel-create"
    assert summary["task_count"] == 1
    assert summary["failure_count"] == 1
    assert set(summary["rows"][0]) >= {"task_id", "state", "failure"}
    assert (invocation / "qualification_report.md").read_text(encoding="utf-8").find("excel-create") >= 0
    assert provider_configs[0].max_tool_turns == 48


def test_general_office_runner_returns_success_for_qualified_task_filter(monkeypatch, tmp_path: Path):
    from project_artifacts.phase3_experiments.scripts import run_general_office_pilot as pilot
    from workagent_rsi.office_capabilities import CapabilityReport

    monkeypatch.setattr(pilot, "RESULT_ROOT", tmp_path)
    monkeypatch.setattr(pilot, "probe_capabilities", CapabilityReport.for_testing)
    monkeypatch.setattr(pilot, "verify_stored_office_artifacts", lambda *args, **kwargs: {
        "ok": True, "status": "verified", "version": "16.0", "artifact_type_ok": True
    })

    class SuccessfulHarness:
        def __init__(self, root, **kwargs):
            self.root = root

        def run(self, task, **kwargs):
            run_root = self.root / task.task_id
            artifacts_root = run_root / "artifacts"
            artifacts_root.mkdir(parents=True)
            artifact = artifacts_root / "quarterly_revenue.xlsx"
            Workbook().save(artifact)
            return {
                "run_id": task.task_id,
                "result_dir": str(run_root),
                "state": "SUCCEEDED",
                "artifacts": [{"path": str(artifact), "sha256": sha256_file(artifact)}],
                "evaluation": {"passed": True, "critical_failures": []},
            }

    monkeypatch.setattr(pilot, "Harness", SuccessfulHarness)

    assert pilot.main(["--task", "excel-create"]) == 0
    invocation = next(tmp_path.iterdir())
    summary = json.loads((invocation / "summary.json").read_text(encoding="utf-8"))
    assert summary["task_count"] == summary["success_count"] == 1
    assert summary["failure_count"] == 0


def test_general_office_runner_retry_help_and_metadata(monkeypatch, tmp_path: Path, capsys):
    from project_artifacts.phase3_experiments.scripts import run_general_office_pilot as pilot
    from workagent_rsi.office_capabilities import CapabilityReport

    monkeypatch.setattr(pilot, "RESULT_ROOT", tmp_path)
    with pytest.raises(SystemExit) as help_exit:
        pilot.main(["--help"])
    assert help_exit.value.code == 0
    assert "--retry-of" in capsys.readouterr().out

    prior_id = "20260929T190856Z-ab3365c5"
    prior = tmp_path / prior_id
    prior.mkdir()
    prior_summary = _general_office_prior_summary(prior)
    prior_bytes = json.dumps(prior_summary).encode("utf-8")
    (prior / "summary.json").write_bytes(prior_bytes)
    monkeypatch.setattr(pilot, "probe_capabilities", lambda: CapabilityReport.for_testing(powerpoint_com=None))

    assert pilot.main(["--retry-of", prior_id]) == 1
    current = next(path for path in tmp_path.iterdir() if path.name != prior_id)
    summary = json.loads((current / "summary.json").read_text(encoding="utf-8"))
    lineage = summary["retry_lineage"]
    assert lineage == {
        "retry_of_invocation_id": prior_id,
        "retry_of_invocation_path": str(prior.resolve()),
        "prior_config_sha256": "a" * 64,
        "prior_task_count": 6,
        "prior_success_count": 0,
        "prior_failure_count": 6,
        "new_config_sha256": sha256_file(pilot.CONFIG),
        "reason": "revised-prompt-and-criteria",
    }
    report = (current / "qualification_report.md").read_text(encoding="utf-8")
    for value in (prior_id, str(prior.resolve()), "a" * 64, lineage["new_config_sha256"], "revised-prompt-and-criteria", "0 succeeded", "6 failed"):
        assert value in report
    assert (prior / "summary.json").read_bytes() == prior_bytes


def _general_office_prior_summary(prior: Path) -> dict:
    """Test-sized shape of the recorded 20260929T190856Z-ab3365c5 pilot summary."""
    task_ids = ("excel-create", "excel-edit", "word-create", "word-edit", "powerpoint-create", "powerpoint-edit")
    return {
        "pilot_id": "general-office-six-task-engineering-pilot-v1",
        "invocation": str(prior.resolve()),
        "config_sha256": "a" * 64,
        "task_count": 6,
        "success_count": 0,
        "failure_count": 6,
        "rows": [{"task_id": task_id, "state": "FAILED"} for task_id in task_ids],
    }


@pytest.mark.parametrize("corruption", ["short_hash", "uppercase_hash", "wrong_pilot", "wrong_path", "wrong_task_count", "wrong_success_count", "wrong_failure_count", "unexplained_extra_failure"])
def test_general_office_runner_retry_rejects_foreign_or_inconsistent_summary(monkeypatch, tmp_path: Path, corruption: str):
    from project_artifacts.phase3_experiments.scripts import run_general_office_pilot as pilot

    monkeypatch.setattr(pilot, "RESULT_ROOT", tmp_path)
    monkeypatch.setattr(pilot, "probe_capabilities", lambda: pytest.fail("invalid summary must be rejected before COM probe"))
    monkeypatch.setattr(pilot, "Harness", lambda *args, **kwargs: pytest.fail("invalid summary must not launch Harness"))
    prior_id = "20260929T190856Z-ab3365c5"
    prior = tmp_path / prior_id
    prior.mkdir()
    summary = _general_office_prior_summary(prior)
    if corruption == "short_hash":
        summary["config_sha256"] = "a" * 63
    elif corruption == "uppercase_hash":
        summary["config_sha256"] = "A" * 64
    elif corruption == "wrong_pilot":
        summary["pilot_id"] = "another-pilot"
    elif corruption == "wrong_path":
        summary["invocation"] = str((tmp_path / "another-invocation").resolve())
    elif corruption == "wrong_task_count":
        summary["task_count"] = 5
    elif corruption == "wrong_success_count":
        summary["success_count"] = 1
    elif corruption == "wrong_failure_count":
        summary["failure_count"] = 5
    else:
        summary["failure_count"] = 7
    (prior / "summary.json").write_text(json.dumps(summary), encoding="utf-8")

    with pytest.raises(SystemExit) as exc:
        pilot.main(["--retry-of", prior_id])
    assert exc.value.code == 2
    assert [path.name for path in tmp_path.iterdir()] == [prior_id]


@pytest.mark.parametrize("case", ["com_preflight", "top_level_exception"])
def test_general_office_runner_retry_accepts_legitimate_runner_failure_counts(monkeypatch, tmp_path: Path, case: str):
    from project_artifacts.phase3_experiments.scripts import run_general_office_pilot as pilot
    from workagent_rsi.office_capabilities import CapabilityReport

    monkeypatch.setattr(pilot, "RESULT_ROOT", tmp_path)
    monkeypatch.setattr(pilot, "probe_capabilities", lambda: CapabilityReport.for_testing(powerpoint_com=None))
    monkeypatch.setattr(pilot, "Harness", lambda *args, **kwargs: pytest.fail("mocked COM preflight must not launch Harness"))
    prior_id = "20260929T190856Z-ab3365c5"
    prior = tmp_path / prior_id
    prior.mkdir()
    summary = _general_office_prior_summary(prior)
    if case == "com_preflight":
        summary.update(rows=[], task_count=0, success_count=0, failure_count=0, failure="COM unavailable")
    else:
        summary.update(rows=summary["rows"][:2], task_count=2, success_count=0, failure_count=3, failure="top-level exception")
    (prior / "summary.json").write_text(json.dumps(summary), encoding="utf-8")

    assert pilot.main(["--retry-of", prior_id]) == 1
    current = next(path for path in tmp_path.iterdir() if path.name != prior_id)
    lineage = json.loads((current / "summary.json").read_text(encoding="utf-8"))["retry_lineage"]
    assert lineage["prior_failure_count"] == summary["failure_count"]


@pytest.mark.parametrize("prior_id,make_directory", [("", False), ("../outside", False), ("C:/outside", False), ("20260929T190856Z-ab3365c5", True)])
def test_general_office_runner_retry_rejects_path_escape_and_missing_summary(monkeypatch, tmp_path: Path, prior_id: str, make_directory: bool):
    from project_artifacts.phase3_experiments.scripts import run_general_office_pilot as pilot
    from workagent_rsi.office_capabilities import CapabilityReport

    monkeypatch.setattr(pilot, "RESULT_ROOT", tmp_path)
    monkeypatch.setattr(pilot, "probe_capabilities", lambda: CapabilityReport.for_testing(powerpoint_com=None))
    monkeypatch.setattr(pilot, "Harness", lambda *args, **kwargs: pytest.fail("invalid retry must not launch Harness"))
    if make_directory:
        (tmp_path / prior_id).mkdir()
    with pytest.raises(SystemExit) as exc:
        pilot.main(["--retry-of", prior_id])
    assert exc.value.code == 2
    assert sorted(path.name for path in tmp_path.iterdir()) == ([prior_id] if make_directory else [])


@pytest.mark.parametrize("task_id", ["excel-create", "excel-edit", "word-create", "word-edit", "powerpoint-create", "powerpoint-edit"])
def test_evaluator_accepts_and_rejects_exact_general_office_pilot_constraints(tmp_path: Path, task_id: str):
    config_path = Path(__file__).resolve().parents[1] / "project_artifacts/phase3_experiments/configs/general_office_pilot.json"
    item = next(row for row in json.loads(config_path.read_text(encoding="utf-8"))["tasks"] if row["task_id"] == task_id)
    task = TaskSpec(task_id=task_id, domain=item["domain"], instruction=item["instruction"], expected_constraints=item["expected_constraints"])
    if task_id == "excel-create":
        path = tmp_path / "satisfying.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Summary"
        for cell, value in {"A1": "Quarter", "B1": "Revenue", "A2": "Q1", "B2": 120, "A3": "Q2", "B3": 150, "A4": "Q3", "B4": 180, "A5": "Total", "B5": "=SUM(B2:B4)"}.items():
            sheet[cell] = value
        book.save(path)
        assert inspect_office_file(path, task).passed
        sheet["B5"] = 450
        book.save(path)
        assert any("formula" in failure for failure in inspect_office_file(path, task).failures)
    elif task_id == "excel-edit":
        path = tmp_path / "satisfying.xlsx"
        book = Workbook()
        source = book.active
        source.title = "Transactions"
        for row in (("Item", "Revenue"), ("Alpha", 40), ("Beta", 55), ("Gamma", 65)):
            source.append(row)
        summary = book.create_sheet("Summary")
        summary["A1"], summary["B1"], summary["A2"], summary["B2"] = "Metric", "Amount", "Total Revenue", "=SUM(Transactions!B2:B4)"
        book.save(path)
        assert inspect_office_file(path, task).passed
        source["B3"] = 99
        book.save(path)
        assert any("Transactions!B3" in failure for failure in inspect_office_file(path, task).failures)
    elif task_id in {"word-create", "word-edit"}:
        path = tmp_path / "satisfying.docx"
        doc = Document()
        if task_id == "word-create":
            for title in ("Project Status", "Progress", "Risks", "Next Actions"):
                doc.add_heading(title, level=1)
            body = doc.add_paragraph("Action: confirm launch date with the sponsor by Friday.")
        else:
            for title in ("Executive Summary", "Completed Work"):
                doc.add_heading(title, level=1)
            fact = doc.add_heading("The pilot completed on 12 September.", level=2)
            doc.add_heading("Next Steps", level=1)
            doc.add_paragraph("Action: send the final report to the steering group.")
        doc.save(path)
        assert inspect_office_file(path, task).passed
        if task_id == "word-create":
            body.text = "No action recorded."
            failure_part = "required marker"
        else:
            fact.style = "Heading 1"
            failure_part = "expected style"
        doc.save(path)
        assert any(failure_part in failure for failure in inspect_office_file(path, task).failures)
    else:
        path = tmp_path / "satisfying.pptx"
        deck = Presentation()
        slide_texts = (("Context",), ("Plan",), ("Decision", "Takeaway: approve the phased rollout.")) if task_id == "powerpoint-create" else (("Baseline: two regions",), ("Next step: expand to three regions",))
        boxes = []
        for texts in slide_texts:
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for index, content in enumerate(texts):
                box = slide.shapes.add_textbox(Inches(1), Inches(1 + 2 * index), Inches(6), Inches(1))
                box.text = content
                boxes.append(box)
        deck.save(path)
        assert inspect_office_file(path, task).passed
        if task_id == "powerpoint-create":
            boxes[0].left = deck.slide_width + Inches(1)
            failure_part = "exceeds slide bounds"
        else:
            boxes[-1].text = "A different next step"
            failure_part = "shape texts are missing"
        deck.save(path)
        assert any(failure_part in failure for failure in inspect_office_file(path, task).failures)


def test_general_office_edit_tasks_make_required_changes_and_safe_placement_explicit():
    config_path = Path(__file__).resolve().parents[1] / "project_artifacts/phase3_experiments/configs/general_office_pilot.json"
    config = json.loads(config_path.read_text(encoding="utf-8"))
    instructions = {task["task_id"]: task["instruction"] for task in config["tasks"]}

    assert "Do not merely append" in instructions["word-edit"]
    assert "Heading 1" in instructions["word-edit"]
    assert "top=2.25" in instructions["powerpoint-edit"]
    assert "leave all source shape text unchanged" in instructions["powerpoint-edit"]


def test_general_office_runner_does_not_com_reopen_failed_evaluation(monkeypatch, tmp_path: Path):
    from project_artifacts.phase3_experiments.scripts import run_general_office_pilot as pilot
    from workagent_rsi.office_capabilities import CapabilityReport

    monkeypatch.setattr(pilot, "RESULT_ROOT", tmp_path)
    monkeypatch.setattr(pilot, "probe_capabilities", CapabilityReport.for_testing)
    monkeypatch.setattr(pilot, "verify_artifact_with_com", lambda *args: pytest.fail("COM must not reopen failed evaluation"))

    class FailedEvaluationHarness:
        def __init__(self, root, **kwargs):
            self.root = root

        def run(self, task, **kwargs):
            run_root = self.root / task.task_id
            run_root.mkdir()
            artifact = run_root / ("invalid" + {"excel": ".xlsx", "word": ".docx", "powerpoint": ".pptx"}[task.domain])
            artifact.write_bytes(b"invalid")
            return {"run_id": task.task_id, "result_dir": str(run_root), "state": "FAILED",
                    "evaluation": {"passed": False, "critical_failures": ["invalid Office structure"]},
                    "artifacts": [{"path": str(artifact), "sha256": "sample"}]}

    monkeypatch.setattr(pilot, "Harness", FailedEvaluationHarness)
    assert pilot.main([]) == 1
    summary = json.loads((next(tmp_path.iterdir()) / "summary.json").read_text(encoding="utf-8"))
    assert summary["task_count"] == 6
    assert all(row["com_ok"] is False for row in summary["rows"])


def test_general_office_runner_records_exception_and_continues_all_six(monkeypatch, tmp_path: Path):
    from project_artifacts.phase3_experiments.scripts import run_general_office_pilot as pilot
    from workagent_rsi.office_capabilities import CapabilityReport

    monkeypatch.setattr(pilot, "RESULT_ROOT", tmp_path)
    monkeypatch.setattr(pilot, "probe_capabilities", CapabilityReport.for_testing)

    class FailingHarness:
        def __init__(self, root, **kwargs):
            self.root = root

        def run(self, task, **kwargs):
            if task.task_id == "excel-edit":
                raise RuntimeError("controlled task exception")
            run_root = self.root / task.task_id
            run_root.mkdir()
            return {"run_id": task.task_id, "result_dir": str(run_root), "state": "FAILED", "artifacts": [],
                    "failure": {"message": "no output"}}

    monkeypatch.setattr(pilot, "Harness", FailingHarness)
    assert pilot.main([]) == 1
    invocation = next(tmp_path.iterdir())
    summary = json.loads((invocation / "summary.json").read_text(encoding="utf-8"))
    assert summary["task_count"] == 6 and len(summary["rows"]) == 6
    assert "controlled task exception" in summary["rows"][1]["failure"]["message"]
    assert (invocation / "excel-edit/qualification_result.json").exists()
    assert summary["rows"][-1]["task_id"] == "powerpoint-edit"


def test_office_prompt_requires_tool_action_before_claiming_completion():
    task = TaskSpec(task_id="prompt", domain="excel", instruction="Create an example workbook")
    prompt = build_task_prompt(task, {"files": []}, "")
    prompt_lower = prompt.lower()
    for requirement in ("create a Python script", "run the script", "reopen", "deliverables.json", "status to \"failed\""):
        assert requirement.lower() in prompt_lower


def test_office_prompt_shows_manifest_as_a_json_string_tool_argument():
    task = TaskSpec(task_id="excel-create", domain="excel", instruction="Create outputs/report.xlsx")
    prompt = build_task_prompt(task, {"files": []}, "")
    assert "content` argument must be a string, not a JSON object" in prompt
    assert '`{"deliverables":["outputs/report.xlsx"]}`' in prompt


def test_office_prompt_keeps_manifest_and_response_fields_separate():
    task = TaskSpec(task_id="excel-edit", domain="excel", instruction="Edit the supplied workbook.", input_files=("source.xlsx",))
    prompt = build_task_prompt(task, {"files": [{"copied": "inputs/0001-source.xlsx"}]}, "")
    assert 'exactly one key, `deliverables`' in prompt
    assert 'do not put `input_files_used` in this manifest' in prompt
    assert 'final structured response has its own input_files_used field' in prompt


def test_word_prompt_uses_document_object_for_body_paragraphs():
    task = TaskSpec(task_id="word-create", domain="word", instruction="Create a report.")
    prompt = build_task_prompt(task, {"files": []}, "")
    assert "Use document.add_paragraph(...) for body text" in prompt
    assert "add_heading(...) returns a Paragraph, not a Document" in prompt


def test_powerpoint_edit_prompt_gives_blank_slide_safe_edit_pattern():
    task = TaskSpec(
        task_id="powerpoint-edit",
        domain="powerpoint",
        instruction="Edit the supplied deck and preserve its slides.",
        input_files=("briefing_draft.pptx",),
    )
    prompt = build_task_prompt(
        task,
        {"files": [{"copied": "inputs/0001-briefing_draft.pptx", "source": "briefing_draft.pptx"}]},
        "",
    )
    assert "blank slides" in prompt
    assert "slide.shapes.add_textbox" in prompt
    assert "never call add_slide" in prompt
    assert "slide.shapes.title" in prompt
    assert "placeholders" in prompt
    assert "assert len(presentation.slides)" in prompt
    assert "save the phrase there" in prompt
    assert "never replace source text with the new phrase" in prompt
    assert "independent textbox" in prompt
    assert "choose exactly one existing slide" in prompt
    assert "do not add a textbox to every slide" in prompt
    assert "check every new textbox against every existing shape" in prompt


def test_powerpoint_prompt_requires_exact_slide_count_and_nonoverlap_for_creation():
    task = TaskSpec(
        task_id="powerpoint-create",
        domain="powerpoint",
        instruction="Create exactly three slides with all requested phrases.",
    )
    prompt = build_task_prompt(task, {"files": []}, "")
    assert "keep the slide count exactly equal to the requested count" in prompt
    assert "put all required phrases within those slides" in prompt
    assert "check every shape pair for overlap" in prompt
    assert "when no staged input exists, start with Presentation()" in prompt
    assert "do not branch on len(presentation.slides) inside a slide loop" in prompt
    assert "assert every required phrase is present" in prompt
    assert "use the blank slide layout" in prompt
    assert "map each requested phrase to its specified slide" in prompt.lower()
    assert "call presentation.slides.add_slide(...) once for each requested slide" in prompt.lower()
    assert "store each returned slide in its own variable" in prompt.lower()
    assert "do not put later slide content on the first slide" in prompt.lower()


def test_excel_prompt_gives_exact_cell_and_binary_verification_pattern():
    task = TaskSpec(
        task_id="excel-create",
        domain="excel",
        instruction="Create the requested workbook with exact cell values.",
    )
    prompt = build_task_prompt(task, {"files": []}, "")
    assert "write only the cells named in this task" in prompt
    assert "for this edit, write only to the Summary sheet" in prompt
    assert "load_workbook('outputs/" in prompt
    assert "never use read_file on .xlsx" in prompt
    assert "do not instantiate Workbook() for an edit" in prompt
    assert "wb.create_sheet('Summary')" in prompt
    assert "for an edit, write every requested summary label and formula cell before saving" in prompt.lower()
    assert "reopen the edited workbook and assert each requested summary cell" in prompt.lower()
    assert "for a creation task, start with openpyxl.workbook()" in prompt.lower()
    assert "write formulas by assigning the formula string to cell.value" in prompt.lower()
    assert "never use cell.formula" in prompt.lower()
    assert "use load_workbook only for an existing input or to reopen the saved output" in prompt.lower()
    assert "assign each requested literal cell value directly from the task" in prompt.lower()
    assert "do not compute cell values from row numbers, indexes, multiplication, or loops" in prompt.lower()
    assert "summary_sheet['b2'] = 120" in prompt.lower()


def test_excel_create_prompt_covers_total_label_and_formula_cells():
    task = TaskSpec(
        task_id="excel-create",
        domain="excel",
        instruction="Create the quarterly workbook with the requested total.",
    )
    prompt = build_task_prompt(task, {"files": []}, "").lower()

    assert "summary_sheet['a5'] = 'total'" in prompt
    assert "summary_sheet['b5'] = '=sum(b2:b4)'" in prompt
    assert "write the total label before writing the total formula" in prompt
    assert "assert summary_sheet['a5'].value == 'total'" in prompt


def test_office_prompt_requires_sequential_tool_execution_and_recovery():
    task = TaskSpec(task_id="excel-create", domain="excel", instruction="Create a workbook.")
    prompt = build_task_prompt(task, {"files": []}, "").lower()

    assert "do not call run_python until write_file confirms the script was written" in prompt
    assert "if a tool call fails, fix the script and run it again" in prompt
    assert "do not write the manifest until the script runs successfully" in prompt
    assert "verify the script after its final edit" in prompt
    assert "the first tool call must be write_file for the complete python script" in prompt
    assert "never write deliverables.json before the script has run successfully" in prompt
    assert "if run_python fails, fix or rewrite the script and run it again" in prompt
    assert "do not move on to a manifest or a completed response" in prompt


def test_office_prompt_prioritizes_task_mode_over_cross_domain_examples():
    edit_task = TaskSpec(
        task_id="rsi-regression-powerpoint-edit",
        domain="powerpoint",
        instruction="Edit the supplied deck and preserve its slides.",
        input_files=("briefing_draft.pptx",),
    )
    edit_prompt = build_task_prompt(
        edit_task,
        {"files": [{"copied": "inputs/0001-briefing_draft.pptx"}]},
        "",
    ).lower()
    assert "task mode: edit" in edit_prompt
    assert "a staged source file is authoritative" in edit_prompt
    assert "do not use any creation example below" in edit_prompt
    assert "do not add slides/sheets/sections unless the task explicitly requests it" in edit_prompt

    create_task = TaskSpec(
        task_id="rsi-develop-powerpoint-create",
        domain="powerpoint",
        instruction="Create a deck.",
    )
    create_prompt = build_task_prompt(create_task, {"files": []}, "").lower()
    assert "task mode: create" in create_prompt
    assert "there is no staged office source to preserve" in create_prompt


def test_office_prompt_requires_dynamic_task_constraints_for_general_rsi_tasks():
    excel_task = TaskSpec(
        task_id="rsi-hidden-excel-create",
        domain="excel",
        instruction="Create an XLSX workbook tracking monthly units on a Metrics worksheet. Set Metrics!A1 to Month.",
    )
    excel_prompt = build_task_prompt(excel_task, {"files": []}, "").lower()
    assert "set the worksheet title exactly to the name requested by the current task" in excel_prompt
    assert "assert the requested worksheet exists after reopening" in excel_prompt
    assert "do not assume the frozen summary example applies to another task" in excel_prompt

    ppt_create = TaskSpec(
        task_id="rsi-ood-powerpoint-create",
        domain="powerpoint",
        instruction="Create a PPTX update deck with exactly 4 nonempty slides.",
    )
    ppt_prompt = build_task_prompt(ppt_create, {"files": []}, "").lower()
    assert "derive the exact required slide count from the current task instruction" in ppt_prompt
    assert "create exactly one slide for that count" in ppt_prompt
    assert "the frozen three-slide skeleton applies only when the current task requests that exact mapping" in ppt_prompt

    ppt_edit = TaskSpec(
        task_id="rsi-regression-powerpoint-edit",
        domain="powerpoint",
        instruction="Edit the supplied PPTX deck and preserve exactly 2 nonempty slides.",
        input_files=("briefing_draft.pptx",),
    )
    ppt_edit_prompt = build_task_prompt(
        ppt_edit, {"files": [{"copied": "inputs/0001-briefing_draft.pptx"}]}, ""
    ).lower()
    assert "in task mode: edit, do not call presentation.slides.add_slide anywhere" in ppt_edit_prompt
    assert "the creation examples are not applicable in task mode: edit" in ppt_edit_prompt


def test_office_prompt_makes_generic_heading_two_and_single_slide_edit_explicit():
    word_task = TaskSpec(
        task_id="rsi-ood-word-create",
        domain="word",
        instruction="Create a decision memo with Heading 1 Recommendation and Heading 2 Rationale.",
    )
    word_prompt = build_task_prompt(word_task, {"files": []}, "").lower()
    assert "any required heading 2 must be created with document.add_heading(title, level=2)" in word_prompt
    assert "never use document.add_paragraph(..., style='heading 2') for a required heading" in word_prompt

    ppt_task = TaskSpec(
        task_id="rsi-regression-powerpoint-edit",
        domain="powerpoint",
        instruction="Edit the supplied deck and preserve exactly two slides; include two phrases.",
        input_files=("briefing_draft.pptx",),
    )
    ppt_prompt = build_task_prompt(
        ppt_task, {"files": [{"copied": "inputs/0001-briefing_draft.pptx"}]}, ""
    ).lower()
    assert "choose exactly one existing slide (slide 1 unless the task names another)" in ppt_prompt
    assert "do not iterate over every slide to add the new phrase" in ppt_prompt
    assert "add exactly one new textbox containing all missing phrases" in ppt_prompt
    assert "scan every existing shape on that slide before placing the textbox" in ppt_prompt


def test_word_prompt_has_unambiguous_heading_level_examples():
    task = TaskSpec(task_id="word-create", domain="word", instruction="Create a report.")
    prompt = build_task_prompt(task, {"files": []}, "").lower()

    assert "document.add_heading('project status', level=1)" in prompt
    assert "never use level=0 for a required heading 1" in prompt
    assert "use separate add_heading and add_paragraph calls for each required section" in prompt
    assert "do not put a literal line break inside a quoted python string" in prompt
    assert "call document.add_heading(title, level=1) separately for every required heading 1" in prompt
    assert "for a required heading 2, call document.add_heading(title, level=2)" in prompt
    assert "never attach heading or body text with add_run to a heading paragraph" in prompt


def test_word_prompt_forbids_mutating_heading_style_levels():
    task = TaskSpec(task_id="word-edit", domain="word", instruction="Reorganize the supplied report.", input_files=("status_draft.docx",))
    prompt = build_task_prompt(task, {"files": [{"copied": "inputs/0001-status_draft.docx"}]}, "").lower()

    assert "never set document.styles['heading 1'].level" in prompt
    assert "never modify the built-in heading 1 style level" in prompt
    assert "document.add_heading('the pilot completed on 12 september.', level=2)" in prompt


def test_word_edit_prompt_uses_document_level_calls_for_restructure():
    task = TaskSpec(
        task_id="word-edit",
        domain="word",
        instruction="Reorganize the supplied report with the required headings.",
        input_files=("status_draft.docx",),
    )
    prompt = build_task_prompt(
        task, {"files": [{"copied": "inputs/0001-status_draft.docx"}]}, ""
    ).lower()

    assert "call document.add_heading(...) directly for each required heading" in prompt
    assert "never assign the return value of add_heading to a document variable" in prompt
    assert "document.add_heading('the pilot completed on 12 september.', level=2)" in prompt
    assert "document.add_paragraph('action: send the final report to the steering group.')" in prompt


def test_powerpoint_prompt_uses_shape_text_and_library_import_names():
    task = TaskSpec(
        task_id="powerpoint-edit", domain="powerpoint", instruction="Edit the supplied deck.",
        input_files=("briefing_draft.pptx",),
    )
    prompt = build_task_prompt(
        task, {"files": [{"copied": "inputs/0001-briefing_draft.pptx"}]}, ""
    ).lower()

    assert "import pptx" in prompt
    assert "never use slide.text" in prompt
    assert "inspect shape.text_frame.text" in prompt
    assert "using an independent textbox; never replace source text with the new phrase" in prompt
    assert "do not put a literal line break inside a quoted python string" in prompt
    assert "reopen and assert the new phrase before saving the manifest" in prompt
    assert "capture existing_shapes = list(slide.shapes) before calling add_textbox" in prompt
    assert "never include the new textbox in the existing_shapes overlap check" in prompt
    assert "pass slide.shapes.add_textbox coordinates and dimensions as inches(...) values" in prompt
    assert "do not call shape.overlap" in prompt
    assert "compare rectangle edges with left, top, width, and height" in prompt
    assert "shape has no right, bottom, or shapes attributes" in prompt
    assert "old.left + old.width <= new.left" in prompt
    assert "old.top + old.height <= new.top" in prompt
    create_task = TaskSpec(task_id="powerpoint-create", domain="powerpoint", instruction="Create three slides.")
    create_prompt = build_task_prompt(create_task, {"files": []}, "").lower()
    assert "put context and plan in separate text boxes; do not combine them in a newline string" in create_prompt
    assert "assign separate .text values 'context' and 'plan' to two different add_textbox calls" in create_prompt


def test_powerpoint_create_prompt_shows_three_slide_variables():
    task = TaskSpec(task_id="powerpoint-create", domain="powerpoint", instruction="Create three slides.")
    prompt = build_task_prompt(task, {"files": []}, "").lower()

    assert "slide1 = presentation.slides.add_slide(presentation.slide_layouts[6])" in prompt
    assert "slide2 = presentation.slides.add_slide(presentation.slide_layouts[6])" in prompt
    assert "slide3 = presentation.slides.add_slide(presentation.slide_layouts[6])" in prompt
    assert "put decision on slide2" in prompt
    assert "put takeaway: approve the phased rollout. on slide3" in prompt


def test_native_ollama_system_prompt_forbids_literal_newlines_in_python_strings(monkeypatch, tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    records = tmp_path / "records"
    responses = iter([{"message": {"role": "assistant", "content": json.dumps({
        "status": "failed", "deliverables": [], "summary": "stop", "input_files_used": []
    })}}])
    requests = []

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self):
            return json.dumps(next(responses)).encode("utf-8")

    def fake_urlopen(request, timeout):
        requests.append(json.loads(request.data.decode("utf-8")))
        return FakeResponse()

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    CodexOfficeProvider(WorkAgentConfig(backend="ollama-native", max_tool_turns=1)).run(
        "You are creating an Office deliverable for a word task.\nCreate a report.",
        workspace,
        records,
    )

    system_prompt = requests[0]["messages"][0]["content"].lower()
    assert "do not put a literal line break inside a quoted python string" in system_prompt
    assert "write_file content for a .py file must be raw python source code, not a json object with a script field" in system_prompt
