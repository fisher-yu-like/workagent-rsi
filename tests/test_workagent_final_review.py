"""Regression coverage for the seven final-review findings; no live providers."""

import json
import re
import subprocess
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

from workagent_rsi.contracts import TaskSpec
from workagent_rsi.harness import Harness
from workagent_rsi.hashing import sha256_file
from workagent_rsi.office_capabilities import CapabilityReport
from workagent_rsi.office_checks import inspect_office_file
from workagent_rsi.workagent_provider import AgentResponse, CodexOfficeProvider, ProviderOutcome, WorkAgentConfig, WorkAgentSkill
from workagent_rsi.workagent_experiment import WorkAgentExperimentRunner
from test_workagent_office import _deterministic_office_provider, _instruction_candidate, _rsi_tasks


ROOT = Path(__file__).resolve().parents[1]


def _fixture_provider(self, prompt, workspace, records):
    """Produce known valid fixtures to test the real Harness/store/COM boundary."""
    from project_artifacts.phase3_experiments.scripts import run_general_office_pilot as pilot

    task = next(row for row in json.loads(pilot.CONFIG.read_text(encoding="utf-8"))["tasks"] if row["task_id"] == workspace.parent.name)
    constraints = task["expected_constraints"]
    outputs = workspace / "outputs"
    outputs.mkdir()
    domain = task["domain"]
    if domain == "excel":
        path = outputs / "result.xlsx"
        book = Workbook()
        book.remove(book.active)
        for name in constraints["required_sheets"]:
            book.create_sheet(name)
        for key, value in {**constraints["required_cells"], **constraints["required_formulas"]}.items():
            sheet, cell = key.split("!")
            book[sheet][cell] = value
        book.save(path)
    elif domain == "word":
        path = outputs / "result.docx"
        doc = Document()
        for title, style in constraints["required_heading_styles"].items():
            doc.add_heading(title, level=int(style[-1]))
        doc.add_paragraph(constraints["required_text"])
        doc.save(path)
    else:
        path = outputs / "result.pptx"
        deck = Presentation()
        count = constraints["required_slide_count"]
        for index in range(count):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            text = "\n".join(constraints["required_shape_text"][index::count])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text = text
        deck.save(path)
    response = AgentResponse(status="completed", deliverables=["outputs/" + path.name], summary="fixture", input_files_used=[])
    (workspace / "agent_response.json").write_text(response.model_dump_json(), encoding="utf-8")
    (workspace / "deliverables.json").write_text(json.dumps({"deliverables": response.deliverables}), encoding="utf-8")
    return ProviderOutcome(status="completed", response=response)


def test_pilot_content_addressed_artifacts_reach_com_and_qualify(tmp_path, monkeypatch):
    from project_artifacts.phase3_experiments.scripts import run_general_office_pilot as pilot

    monkeypatch.setattr(pilot, "RESULT_ROOT", tmp_path)
    monkeypatch.setattr(pilot, "probe_capabilities", CapabilityReport.for_testing)
    monkeypatch.setattr(CodexOfficeProvider, "run", _fixture_provider)
    reopened = []

    def reopen(path, domain):
        path = Path(path)
        assert path.suffix == {"excel": ".xlsx", "word": ".docx", "powerpoint": ".pptx"}[domain]
        assert inspect_office_file(path, TaskSpec(task_id="com-input", domain=domain, instruction="Reopen")).passed
        reopened.append(sha256_file(path))
        return {"ok": True, "status": "available", "version": "16.0"}

    monkeypatch.setattr(pilot, "verify_artifact_with_com", reopen)
    assert pilot.main([]) == 0
    root = next(tmp_path.iterdir())
    summary = json.loads((root / "summary.json").read_text(encoding="utf-8"))
    assert summary["success_count"] == 6
    assert len(reopened) == 6  # Pilot owns this explicit gate; no second COM call.
    for row in summary["rows"]:
        result = json.loads((root / row["task_id"] / "result.json").read_text(encoding="utf-8"))
        ref = result["artifacts"][0]
        assert Path(ref["path"]).suffix == ""
        assert ref["sha256"] in reopened


@pytest.mark.parametrize("com_result,state", [
    ({"ok": True, "status": "available", "version": "16.0"}, "SUCCEEDED"),
    ({"ok": False, "status": "unavailable", "error": "Office absent"}, "UNAVAILABLE"),
    ({"ok": False, "status": "failed", "error": "Office refused package"}, "FAILED"),
    ({"ok": False, "status": "timeout", "error": "Office reopen timed out"}, "UNAVAILABLE"),
])
def test_normal_harness_enforces_and_persists_required_com(tmp_path, monkeypatch, com_result, state):
    import workagent_rsi.office_capabilities as capabilities

    monkeypatch.setattr(CodexOfficeProvider, "run", _deterministic_office_provider)
    seen = []

    def reopen(path, domain):
        assert Path(path).suffix == ".xlsx"
        seen.append(sha256_file(path))
        return com_result

    monkeypatch.setattr(capabilities, "verify_artifact_with_com", reopen)
    result = Harness(tmp_path, agent_instructions="good instructions").run(_rsi_tasks()["develop"][0])
    assert result["state"] == state
    assert seen == [result["artifacts"][0]["sha256"]]
    saved = json.loads((Path(result["result_dir"]) / "com_reopen.json").read_text(encoding="utf-8"))
    assert saved["ok"] is com_result["ok"] and saved["status"] == com_result["status"]
    assert result["com_reopen"] == saved
    assert any(event["kind"] == "com_reopen" for event in Harness().resume(result["result_dir"])["events"])
    if state != "SUCCEEDED":
        assert result["evaluation"]["passed"] is False
    if com_result["status"] in {"unavailable", "timeout"}:
        assert "score" not in result["evaluation"]


@pytest.mark.parametrize("corruption", ["hash", "media_type"])
def test_com_gate_rejects_untrusted_artifact_reference(tmp_path, corruption):
    from workagent_rsi.artifact_com import verify_stored_office_artifacts
    from workagent_rsi.office_checks import MEDIA_TYPES
    from workagent_rsi.storage import ArtifactStore

    source = tmp_path / "input.xlsx"
    Workbook().save(source)
    ref = ArtifactStore(tmp_path / "artifacts").put_file(source, MEDIA_TYPES["excel"]).model_dump()
    if corruption == "hash":
        ref["sha256"] = "0" * 64
    else:
        ref["media_type"] = MEDIA_TYPES["word"]
    result = verify_stored_office_artifacts([ref], "excel", tmp_path, verify=lambda *args: pytest.fail("untrusted artifact reached COM"))
    assert result["ok"] is False and result["status"] == "failed"
    assert result["artifact_type_ok"] is False
    assert source.read_bytes() == Path(ref["path"]).read_bytes()


def test_required_com_unavailability_stops_rsi_without_score_or_candidate(tmp_path, monkeypatch):
    import workagent_rsi.office_capabilities as capabilities

    monkeypatch.setattr(CodexOfficeProvider, "run", _deterministic_office_provider)
    monkeypatch.setattr(capabilities, "verify_artifact_with_com", lambda *args: {"ok": False, "status": "unavailable", "error": "Office absent"})

    class ForbiddenCandidate:
        def generate(self, *args):
            pytest.fail("candidate must not run without required COM evidence")

    summary = WorkAgentExperimentRunner(input_base=tmp_path).run(
        _rsi_tasks(), WorkAgentSkill(instructions="good instructions"), ForbiddenCandidate(), tmp_path / "experiment")
    assert summary["status"] == "incomplete"
    assert "initial_develop_score" not in summary["metrics"]
    for report in summary["incomplete_round"]["baseline"].values():
        assert "score" not in report


def test_explicit_com_disable_is_persisted_without_probe(tmp_path, monkeypatch):
    import workagent_rsi.office_capabilities as capabilities

    monkeypatch.setattr(CodexOfficeProvider, "run", _deterministic_office_provider)
    monkeypatch.setattr(capabilities, "verify_artifact_with_com", lambda *args: pytest.fail("COM disabled"))
    result = Harness(tmp_path, workagent_config=WorkAgentConfig(verify_com=False), agent_instructions="good instructions").run(_rsi_tasks()["develop"][0])
    assert result["state"] == "SUCCEEDED"
    assert result["com_reopen"]["status"] == "disabled"
    assert result["com_reopen"]["ok"] is False


def test_rsi_regression_loss_cannot_be_masked_by_another_task_gain(tmp_path, monkeypatch):
    tasks = _rsi_tasks()
    tasks["regression"] = [
        TaskSpec(task_id="reg-kept", domain="word", instruction="Create Word Keep", expected_constraints={"required_text": "Status ready"}),
        TaskSpec(task_id="reg-gain", domain="word", instruction="Create Word Gain", expected_constraints={"required_text": "Status ready"}),
    ]

    def provider(self, prompt, workspace, records):
        if "Create Word Keep" in prompt:
            prompt = prompt.replace("good instructions", "bad instructions") if "good instructions" in prompt else prompt.replace("bad instructions", "good instructions")
        return _deterministic_office_provider(self, prompt, workspace, records)

    class Candidate:
        @staticmethod
        def generate(workspace, *args):
            context = "\n".join(path.read_text(encoding="utf-8") for path in workspace.rglob("*.json"))
            for private_task in ("reg-kept", "reg-gain", "hidden-powerpoint", "ood-excel"):
                assert private_task not in context
            return _instruction_candidate(workspace, *args)

    monkeypatch.setattr(CodexOfficeProvider, "run", provider)
    summary = WorkAgentExperimentRunner(workagent_config=WorkAgentConfig(verify_com=False), input_base=tmp_path).run(
        tasks, WorkAgentSkill(instructions="bad instructions"), Candidate(), tmp_path / "experiment")
    result = summary["rounds"][0]
    assert result["baseline"]["regression"]["success_count"] == result["candidate_reports"]["regression"]["success_count"] == 1
    assert result["metrics"]["critical_regressions"] == 1
    assert result["decision"]["decision"] == "reject"
    for group in ("baseline", "candidate_reports"):
        for split in ("regression", "hidden", "ood_transfer"):
            assert "rows" not in result[group][split]


@pytest.mark.parametrize("stderr,expected", [
    ("error: unexpected argument '--ignore-user-config' found", "unavailable"),
    ("error: invalid value 'workspace-write' for '--sandbox <SANDBOX_MODE>'", "unavailable"),
    ("Error: Could not connect to Ollama at http://localhost:11434: connection refused", "unavailable"),
    ("Error: model 'qwen2.5:7b' not found", "unavailable"),
    ("Error: qwen2.5:7b does not support tools", "unavailable"),
    ("Error: model qwen2.5:7b does not support JSON schema", "unavailable"),
    ("Traceback: ValueError: invalid cell value", "failed"),
    ("Failed to create workbook: output directory permission denied", "failed"),
    ("Task script: model.json not found", "failed"),
])
def test_nonzero_cli_environment_failures_are_unavailable(tmp_path, monkeypatch, stderr, expected):
    stdout = '{"type":"error","message":"provider exited"}\n'
    calls = []

    def process(command, **kwargs):
        calls.append(command)
        return subprocess.CompletedProcess(command, 2, stdout, stderr)

    monkeypatch.setattr(subprocess, "run", process)
    records = tmp_path / "records"
    outcome = CodexOfficeProvider(WorkAgentConfig()).run("Create workbook", tmp_path, records)
    assert outcome.status == expected
    assert outcome.record.status == expected and outcome.record.exit_code == 2
    assert (records / "provider.stdout.jsonl").read_text(encoding="utf-8") == stdout
    assert (records / "provider.stderr.txt").read_text(encoding="utf-8") == stderr
    assert len(calls) == 1
    assert calls[0][calls[0].index("--sandbox") + 1] == "workspace-write"


def _artifact_from_instructions(task, target, input_root):
    """A small instruction-only executor: it never sees frozen expected constraints."""
    text = task.instruction
    inputs = [input_root / name for name in task.input_files]
    if task.domain == "excel":
        book = load_workbook(inputs[0]) if inputs else Workbook()
        for sheet, cell, raw in re.findall(r'Set ([A-Za-z]+)!([A-Z]+[0-9]+) to ("[^"\n]*"|[0-9]+)\.', text):
            if sheet not in book.sheetnames:
                book.create_sheet(sheet)
            book[sheet][cell] = json.loads(raw)
        book.save(target)
    elif task.domain == "word":
        doc = Document(inputs[0]) if inputs else Document()
        for level, title in re.findall(r'Add a Heading ([12]) paragraph named "([^"]+)"\.', text):
            doc.add_heading(title, level=int(level))
        for body in re.findall(r'Include the exact sentence "([^"]+)"\.', text):
            doc.add_paragraph(body)
        doc.save(target)
    else:
        deck = Presentation(inputs[0]) if inputs else Presentation()
        match = re.search(r'exactly ([0-9]+) nonempty slides', text)
        count = int(match[1]) if match else max(1, len(deck.slides))
        while len(deck.slides) < count:
            deck.slides.add_slide(deck.slide_layouts[6])
        labels = re.findall(r'Include visible text "([^"]+)"\.', text)
        for index, slide in enumerate(deck.slides):
            # Preserve an input slide's content; add one disjoint text box.
            content = "\n".join(labels[index::count]) or "Update"
            slide.shapes.add_textbox(Inches(1), Inches(3), Inches(7), Inches(2)).text = content
        deck.save(target)


@pytest.mark.parametrize("split", ["develop", "regression", "hidden", "ood_transfer"])
@pytest.mark.parametrize("domain", ["excel", "word", "powerpoint"])
def test_rsi_instructions_are_sufficient_for_frozen_checks(tmp_path, split, domain):
    from project_artifacts.phase3_experiments.scripts import run_workagent_rsi as entry

    entry._source_inputs(tmp_path)
    config = json.loads(entry.CONFIG.read_text(encoding="utf-8"))
    for item in (row for row in config["splits"][split] if row["domain"] == domain):
        task = TaskSpec.model_validate(item)
        target = tmp_path / (task.task_id + {"excel": ".xlsx", "word": ".docx", "powerpoint": ".pptx"}[task.domain])
        _artifact_from_instructions(task.model_copy(update={"expected_constraints": {}}), target, tmp_path / "inputs")
        report = inspect_office_file(target, task)
        assert report.passed, report.failures
