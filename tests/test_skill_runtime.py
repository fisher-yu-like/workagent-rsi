from pathlib import Path

from workagent_rsi.contracts import TaskSpec
from workagent_rsi.evaluator import OfficeArtifactEvaluator
from workagent_rsi.skill_runtime import PilotSkillConfig, SkillConfiguredOfficeAdapter
from workagent_rsi.storage import ArtifactStore


def evaluate(tmp_path: Path, marker_source: str):
    task = TaskSpec(
        task_id="word-1",
        domain="word",
        instruction="Create the monthly operating update.",
        expected_constraints={"required_text": "REQUIRED-001"},
    )
    events = list(SkillConfiguredOfficeAdapter(tmp_path / marker_source, PilotSkillConfig(marker_source=marker_source)).execute(task, "office.marker"))
    artifact = events[-1]
    store = ArtifactStore(tmp_path / "store" / marker_source)
    ref = store.put_file(artifact["artifact_path"], artifact["media_type"])
    return OfficeArtifactEvaluator().evaluate(task, [ref], "trace")


def test_skill_marker_policy_changes_real_office_result(tmp_path: Path):
    assert evaluate(tmp_path, "task_id").passed is False
    assert evaluate(tmp_path, "required_text").passed is True


def test_skill_runtime_uses_task_instruction_as_word_heading_and_powerpoint_title(tmp_path: Path):
    for domain in ("word", "powerpoint"):
        task = TaskSpec(
            task_id=f"{domain}-instruction",
            domain=domain,
            instruction="Executive review for verified sales",
            expected_constraints={"required_text": "MARKER-001"},
        )
        event = list(SkillConfiguredOfficeAdapter(tmp_path / domain, PilotSkillConfig(marker_source="required_text")).execute(task, "office.local"))[-1]
        if domain == "word":
            from docx import Document

            assert Document(event["artifact_path"]).paragraphs[0].text == task.instruction
        else:
            from pptx import Presentation

            assert Presentation(event["artifact_path"]).slides[0].shapes.title.text == task.instruction
