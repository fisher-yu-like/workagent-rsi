import json
from pathlib import Path
import subprocess
import sys
from urllib.error import HTTPError

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
import pytest

from workagent_rsi.artifact_assessment import SharedAssessment, main
from workagent_rsi.artifact_evaluator import ArtifactEvaluator
from workagent_rsi.artifact_verifier import ArtifactVerifier
from workagent_rsi.assessment_contracts import AcceptanceSpec, Location, Requirement
from workagent_rsi.model_judgement import ModelJudgementProvider, ModelReviewConfig, configured_assessment
from workagent_rsi.semantic_assessment import EvidenceBudgetExceeded, extract_evidence
from test_visual_assessment import FakeRenderer


def config(provider="api", **kwargs):
    return ModelReviewConfig(provider=provider, model="test-model", codex_executable=sys.executable, retries=0, **kwargs)


def decision(request, *, score=1, invalid=False):
    if request.get("modality") == "semantic":
        item = request["evidence"]["items"][0]
        observations = [{"evidence_id": "invented" if invalid else item["evidence_id"], "quote": item["text"], "observation": "Content supports the criterion."}]
    else:
        observations = [{"image_id": "invented" if invalid else i["image_id"], "artifact_location": i["artifact_location"], "region": "whole page", "observation": "Readable."} for i in request["images"]]
    return {"status": "PASS" if score == 1 else "FAIL" if score == 0 else "PARTIAL", "completion": score,
            "observed": "Evidence-based review.", "confidence": .9, "observations": observations}


def api_transport(calls, *, score=1, invalid=False):
    def transport(url, headers, payload, timeout):
        request = json.loads(payload["input"][0]["content"][0]["text"])
        calls.append((request, payload, headers, url))
        result = decision(request, score=score, invalid=invalid)
        return {"status": "completed", "id": "mock-response", "model": "test-model", "usage": {"input_tokens": 20, "output_tokens": 30},
                "output": [{"type": "message", "content": [{"type": "output_text", "text": json.dumps(result)}]}]}
    return transport


def cli_runner(calls, *, score=1, invalid=False):
    def run(command, cwd, timeout, prompt):
        request = json.loads(prompt.split("Evidence package:\n", 1)[1])
        calls.append((command, cwd, request))
        output = Path(command[command.index("--output-last-message") + 1])
        output.write_text(json.dumps(decision(request, score=score, invalid=invalid)), encoding="utf-8")
        return subprocess.CompletedProcess(command, 0, json.dumps({"type": "turn.completed", "usage": {"input_tokens": 4}}), "")
    return run


def spec(domain="word", *, visual=False, deterministic=False):
    requirements = [Requirement(requirement_id="meaning", description="Clear, grounded regional sales analysis", check="semantic", location=Location(artifact="out"),
                                expected="East 15, West 10", dimension="content", evidence_source="frozen task facts")]
    if visual:
        requirements.append(Requirement(requirement_id="layout", description="Readable page", check="visual", location=Location(artifact="out"), dimension="content", evidence_source="render", options={"pages": [1]}))
    if deterministic:
        requirements.append(Requirement(requirement_id="number", description="Frozen exact value", check="excel.cell", location=Location(artifact="out", sheet="Sheet", cell="B2"), expected=999, critical=True, dimension="content", evidence_source="independent source"))
    return AcceptanceSpec(task_id="semantic-test", version="1", artifacts={"out": domain}, dimension_weights={"content": 1}, requirements=requirements)


def artifact(root, domain):
    if domain == "word":
        obj, suffix = Document(), ".docx"
        obj.add_heading("Sales analysis", 1)
        obj.add_paragraph("East 15, West 10. Single-period data cannot establish growth.")
        table = obj.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "East", "15"
    elif domain == "excel":
        obj, suffix = Workbook(), ".xlsx"
        obj.active.append(["Region", "Sales"])
        obj.active.append(["East", 15])
    else:
        obj, suffix = Presentation(), ".pptx"
        slide = obj.slides.add_slide(obj.slide_layouts[1])
        slide.shapes.title.text = "Sales analysis"
        slide.placeholders[1].text = "East 15, West 10"
    path = root / ("report" + suffix)
    obj.save(path)
    return path


@pytest.mark.parametrize("backend", ["api", "codex-cli"])
@pytest.mark.parametrize("domain", ["word", "excel", "powerpoint"])
def test_backend_switch_shared_semantic_and_visual_reports(tmp_path, monkeypatch, backend, domain):
    monkeypatch.setenv("OPENAI_API_KEY", "private-test-key")
    calls = []
    engine = configured_assessment(config(backend), tmp_path / "evidence", runner=cli_runner(calls), transport=api_transport(calls), renderer=FakeRenderer(tmp_path / "render"))
    path = artifact(tmp_path, domain)
    shared = engine.inspect(spec(domain, visual=True), {"out": path})
    assert ArtifactEvaluator().evaluate(shared).total_score == 100
    assert ArtifactVerifier().verify(shared).issues == []
    assert len(calls) == 2
    semantic_request = calls[0][0] if backend == "api" else calls[0][2]
    assert semantic_request["images"] == []
    assert semantic_request["evidence"]["items"][0]["location"]["artifact"] == "out"
    if backend == "api":
        assert calls[0][1]["input"][0]["content"][0]["type"] == "input_text"
        assert len(calls[0][1]["input"][0]["content"]) == 1
        assert calls[1][1]["input"][0]["content"][1]["image_url"].startswith("data:image/png;base64,")
        assert calls[0][3].endswith("/responses")
        for file in (tmp_path / "evidence").rglob("*.json"):
            assert "private-test-key" not in file.read_text(encoding="utf-8")
    else:
        command, cwd, _ = calls[0]
        assert "--ignore-user-config" in command and "read-only" in command
        assert cwd != tmp_path and not cwd.exists()
        assert command[-1] == "-" and "--image" not in command
        assert "--image" in calls[1][0]
    cached = engine.inspect(spec(domain, visual=True), {"out": path})
    assert len(calls) == 2
    assert cached.telemetry["semantic"][0]["model_calls"] == 0
    assert cached.telemetry["visual"][0]["model_calls"] == 0


def test_semantic_never_renders_and_locates_table(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    calls = []
    class ForbiddenRenderer:
        def identity_config(self): return {"renderer": "must-not-run"}
        def render(self, *args): raise AssertionError("semantic review rendered a page")
    engine = configured_assessment(config(), tmp_path / "cache", transport=api_transport(calls), renderer=ForbiddenRenderer())
    shared = engine.inspect(spec(), {"out": artifact(tmp_path, "word")})
    assert shared.checks[0].status == "PASS"
    assert any(i["location"].get("table") == 1 for i in calls[0][0]["evidence"]["items"])


def test_disabled_missing_key_budget_and_empty_text_are_incomplete(tmp_path, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    path = artifact(tmp_path, "word")
    engines = [configured_assessment({"provider": "disabled"}, tmp_path / "a"), configured_assessment(config(), tmp_path / "b"),
               configured_assessment(config(max_text_chars=1), tmp_path / "c")]
    for engine in engines:
        shared = engine.inspect(spec(), {"out": path})
        assert shared.checks[0].status == "UNAVAILABLE"
        assert ArtifactEvaluator().evaluate(shared).total_score is None
    empty = Document()
    empty.save(path)
    shared = configured_assessment(config(), tmp_path / "empty").inspect(spec(), {"out": path})
    assert shared.checks[0].status == "NEEDS_REVIEW"


@pytest.mark.parametrize("backend", ["api", "codex-cli"])
def test_bad_evidence_fails_closed(tmp_path, monkeypatch, backend):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    engine = configured_assessment(config(backend), tmp_path / "cache", transport=api_transport([], invalid=True), runner=cli_runner([], invalid=True))
    shared = engine.inspect(spec(), {"out": artifact(tmp_path, "word")})
    assert shared.checks[0].status == "ERROR"
    assert ArtifactEvaluator().evaluate(shared).total_score is None
    assert not list((tmp_path / "cache").rglob("cache.json"))


def test_model_high_score_cannot_override_critical_numeric_failure(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    shared = configured_assessment(config(), tmp_path / "cache", transport=api_transport([])).inspect(spec("excel", deterministic=True), {"out": artifact(tmp_path, "excel")})
    score = ArtifactEvaluator().evaluate(shared)
    assert score.acceptance_status == "FAIL" and score.total_score == 50
    assert ArtifactVerifier().verify(shared).issues[0].requirement_id == "number"


def test_model_defects_have_primary_source_locations(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    engine = configured_assessment(config(), tmp_path / "cache", transport=api_transport([], score=.25), renderer=FakeRenderer(tmp_path / "renders"))
    shared = engine.inspect(spec("powerpoint", visual=True), {"out": artifact(tmp_path, "powerpoint")})
    issues = ArtifactVerifier().verify(shared).issues
    assert all(issue.location.slide == 1 for issue in issues)
    assert "whole page" in issues[1].location.region


def test_timeout_retry_and_independent_disagreement(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    calls = []
    def transport(*args):
        calls.append(1)
        if len(calls) == 1:
            raise TimeoutError()
        return api_transport([], score=1 if len(calls) == 2 else .25)(*args)
    settings = config(repeats=2).model_copy(update={"retries": 1})
    shared = configured_assessment(settings, tmp_path / "cache", transport=transport).inspect(spec(), {"out": artifact(tmp_path, "word")})
    assert shared.checks[0].status == "NEEDS_REVIEW"
    assert len(calls) == 3
    telemetry = shared.telemetry["semantic"][0]
    assert telemetry["model_calls"] == 3 and telemetry["independent_reviews"] == 2
    assert telemetry["agreement"] is False


@pytest.mark.parametrize("failure", ["timeout", "invalid_json", "refusal", "unauthorized", "cli_exit", "cli_tools"])
def test_provider_failure_modes(tmp_path, monkeypatch, failure):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    def transport(*args):
        if failure == "timeout": raise TimeoutError()
        if failure == "unauthorized": raise HTTPError("https://example.test", 401, "unauthorized", {}, None)
        return {"status": "completed", "output": [{"type": "message", "content": [{"type": "refusal" if failure == "refusal" else "output_text", "text": "bad-json"}]}]}
    def runner(command, cwd, timeout, prompt):
        if failure == "cli_exit": return subprocess.CompletedProcess(command, 1, "", "failed")
        result = cli_runner([])(command, cwd, timeout, prompt)
        result.stdout = json.dumps({"type": "item.completed", "item": {"type": "command_execution"}})
        return result
    engine = configured_assessment(config("codex-cli" if failure.startswith("cli") else "api"), tmp_path / "cache", transport=transport, runner=runner)
    shared = engine.inspect(spec(), {"out": artifact(tmp_path, "word")})
    assert shared.checks[0].status == ("UNAVAILABLE" if failure == "timeout" else "ERROR")
    assert ArtifactEvaluator().evaluate(shared).total_score is None


def test_api_visual_bad_hash_and_budget_stop_before_transport(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    path = tmp_path / "image.png"
    path.write_bytes(b"not-the-declared-image")
    calls = []
    provider = ModelJudgementProvider(config(), transport=api_transport(calls))
    result, meta = provider.judge({"images": [{"path": str(path), "sha256": "wrong"}]}, tmp_path / "run")
    assert result is None and meta["model_calls"] == 0 and calls == []


def test_cache_identity_includes_model_rubric_and_content(tmp_path, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    calls = []
    path = artifact(tmp_path, "word")
    engine = configured_assessment(config(), tmp_path / "cache", transport=api_transport(calls))
    engine.inspect(spec(), {"out": path})
    changed = spec().model_copy(update={"requirements": [spec().requirements[0].model_copy(update={"expected": "a different rubric"})]})
    engine.inspect(changed, {"out": path})
    document = Document(path)
    document.add_paragraph("New evidence")
    document.save(path)
    engine.inspect(spec(), {"out": path})
    other = configured_assessment(config().model_copy(update={"model": "other"}), tmp_path / "cache", transport=api_transport(calls))
    other.inspect(spec(), {"out": path})
    assert len(calls) == 4 and other.identity() != engine.identity()


def test_extraction_preserves_order_and_limits_expanded_excel():
    doc = Document()
    doc.add_paragraph("before")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "middle"
    doc.add_paragraph("after")
    assert [i["text"] for i in extract_evidence(doc, "out")["items"]] == ["before", "middle", "after"]
    book = Workbook()
    book.active["XFD1048576"] = "far away"
    with pytest.raises(EvidenceBudgetExceeded): extract_evidence(book, "out")


def test_cli_entry_config_disabled_is_explicit(tmp_path, capsys):
    path = artifact(tmp_path, "word")
    request = tmp_path / "request.json"
    request.write_text(json.dumps({"spec": spec().model_dump(mode="json"), "artifacts": {"out": path.name}}), encoding="utf-8")
    settings = tmp_path / "config.json"
    settings.write_text('{"provider":"disabled"}', encoding="utf-8")
    assert main([str(request), "--output", str(tmp_path / "report"), "--model-config", str(settings)]) == 0
    assert json.loads(capsys.readouterr().out)["total_score"] is None
    assert (tmp_path / "report/issues.json").exists()


def test_config_rejects_secret_url_and_implicit_model():
    with pytest.raises(ValueError): ModelReviewConfig(provider="api")
    with pytest.raises(ValueError): config(api_base_url="https://user:password@example.test/v1")
    with pytest.raises(ValueError): config(api_base_url="http://external.test/v1")


def test_real_http_serialization_with_local_stub(tmp_path, monkeypatch):
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    from threading import Thread
    monkeypatch.setenv("OPENAI_API_KEY", "local-stub-key")
    received = []
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            payload = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            received.append((self.path, self.headers.get("Authorization"), payload))
            reply = api_transport([])("", {}, payload, 1)
            data = json.dumps(reply).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
        def log_message(self, *args): pass
    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        settings = config(api_base_url=f"http://127.0.0.1:{server.server_port}/v1")
        engine = configured_assessment(settings, tmp_path / "cache")
        shared = engine.inspect(spec(), {"out": artifact(tmp_path, "word")})
        assert shared.checks[0].status == "PASS"
        assert received[0][0] == "/v1/responses"
        assert received[0][1] == "Bearer local-stub-key"
        schema = received[0][2]["text"]["format"]["schema"]
        assert set(schema["required"]) == set(schema["properties"])
        assert schema["additionalProperties"] is False
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


def test_cli_timeout_logs_partial_output(tmp_path):
    def timeout(command, cwd, seconds, prompt):
        raise subprocess.TimeoutExpired(command, seconds, output='{"type":"turn.started"}', stderr="network timeout")
    engine = configured_assessment(config("codex-cli"), tmp_path / "cache", runner=timeout)
    shared = engine.inspect(spec(), {"out": artifact(tmp_path, "word")})
    assert shared.checks[0].status == "UNAVAILABLE"
    assert any("network timeout" in p.read_text(encoding="utf-8") for p in (tmp_path / "cache").rglob("stderr.txt"))


@pytest.mark.parametrize("cache_enabled", [True, False])
def test_semantic_rsi_repeat_distinguishes_cache_from_real_calls(tmp_path, monkeypatch, cache_enabled):
    from workagent_rsi.contracts import TaskSpec
    from workagent_rsi.frozen_evaluator import FrozenEvaluator
    from workagent_rsi.hashing import canonical_json_hash
    from workagent_rsi.rsi_contracts import EvaluationContract
    from workagent_rsi.skill_runtime import PilotSkillConfig
    monkeypatch.setenv("OPENAI_API_KEY", "test")
    # Reuse the same genuine bytes on each execution so the cache key is identical.
    path = artifact(tmp_path, "word")
    class Adapter:
        def __init__(self, *args): pass
        def execute(self, *args): return iter([{"kind": "artifact", "artifact_path": str(path)}])
    monkeypatch.setattr("workagent_rsi.frozen_evaluator.SkillConfiguredOfficeAdapter", Adapter)
    acceptance = spec().model_copy(update={"artifacts": {"output": "word"}, "requirements": [spec().requirements[0].model_copy(update={"location": Location(artifact="output")})]})
    task = TaskSpec(task_id=acceptance.task_id, domain="word", instruction="Sales review", expected_constraints={"acceptance_spec": acceptance.model_dump(mode="json")})
    calls = []
    def transport(*args):
        response = api_transport(calls)(*args)
        # Independent rationales need not be word-for-word equal to agree on score/status.
        part = response["output"][0]["content"][0]
        judgement = json.loads(part["text"])
        judgement["observed"] = f"Review number {len(calls)}"
        part["text"] = json.dumps(judgement)
        return response
    engine = configured_assessment(config(cache_enabled=cache_enabled), tmp_path / "cache", transport=transport)
    identity = engine.identity()
    split_hash = canonical_json_hash([task.model_dump(mode="json")])
    contract = EvaluationContract(contract_id="semantic", dataset_hash=split_hash, split_hashes={"develop": split_hash}, evaluator_hash=identity,
                                  assessment_mode="artifact-v1", assessment_identity=identity, acceptance_hashes={task.task_id: acceptance.fingerprint()},
                                  provider_policy="test", seeds=[1], repeats=2, timeout_seconds=60, thresholds={}, git_commit="test")
    report = FrozenEvaluator(identity, assessment=engine).evaluate_split([task], "develop", PilotSkillConfig(), contract, tmp_path / "results")
    assert len(calls) == (1 if cache_enabled else 2)
    assert report["assessment_complete"] is True and report["repeat_agreement"] is (not cache_enabled)
