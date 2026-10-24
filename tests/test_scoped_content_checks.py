"""Real local OOXML fixtures with independent position and section expectations."""
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
import pytest

from workagent_rsi.assessment_contracts import AcceptanceSpec, Location, Requirement
from workagent_rsi.artifact_assessment import assess,save_reports


def inspect(path,expected='NEEDLE',check='text.contains',**position):
    domain={'.docx':'word','.xlsx':'excel','.pptx':'powerpoint'}[path.suffix]
    spec=AcceptanceSpec(task_id='scope',version='independent-v1',artifacts={'output':domain},
        dimension_weights={'content':1},requirements=[Requirement(requirement_id='content',
            description='Required content at declared position',check=check,expected=expected,
            location=Location(artifact='output',**position),dimension='content',critical=True,
            evidence_source='authored position oracle')])
    reports=assess(spec,{'output':path})
    folder=path.parent/('observation-'+str(len(list(path.parent.glob('observation-*')))))
    save_reports(folder,*reports)
    return reports


def word(tmp_path):
    doc=Document();doc.add_paragraph('NEEDLE');doc.add_paragraph('OTHER')
    path=tmp_path/'report.docx';doc.save(path);return doc,path


def excel(tmp_path):
    book=Workbook();book.active.title='First';book.active['A1']='NEEDLE';book.active['B2']='OTHER'
    book.create_sheet('Second')['A1']='OTHER'
    path=tmp_path/'book.xlsx';book.save(path);return book,path


def deck(tmp_path):
    prs=Presentation()
    for text in ['NEEDLE','OTHER']:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(0,0,Inches(2),Inches(1)).text=text
    path=tmp_path/'deck.pptx';prs.save(path);return prs,path


@pytest.mark.parametrize('paragraph',[2,99])
def test_word_hit_in_another_paragraph_does_not_pass(tmp_path,paragraph):
    _,path=word(tmp_path);shared,score,_=inspect(path,paragraph=paragraph)
    assert shared.checks[0].status=='FAIL'
    assert score.acceptance_status=='FAIL'


def test_requested_location_is_preserved_separately_from_observation(tmp_path):
    _,path=word(tmp_path);shared,_,_=inspect(path)
    record=shared.checks[0]
    assert record.location.paragraph is None
    assert record.requested_location==shared.spec.requirements[0].location
    assert record.observed_locations[0].paragraph==1
    assert shared.evidence[0].requested_location==record.requested_location


def test_word_heading_section_does_not_search_next_peer(tmp_path):
    doc=Document();doc.add_heading('Target',1);doc.add_paragraph('OTHER')
    doc.add_heading('Next',1);doc.add_paragraph('NEEDLE')
    path=tmp_path/'headings.docx';doc.save(path)
    assert inspect(path,heading='Target')[0].checks[0].status=='FAIL'


def test_heading_scope_includes_subsections_and_tables(tmp_path):
    doc=Document();doc.add_heading('Target',1);doc.add_heading('Child',2)
    table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='NEEDLE'
    doc.add_heading('Next',1);path=tmp_path/'section.docx';doc.save(path)
    shared,_,_=inspect(path,heading='Target')
    assert shared.checks[0].status=='PASS'
    assert shared.checks[0].observed_locations[0].table==1


@pytest.mark.parametrize('check,expected',[('text.contains','NEEDLE'),('word.heading_style','Heading 1')])
def test_duplicate_heading_is_ambiguous(tmp_path,check,expected):
    doc=Document()
    for _ in range(2):doc.add_heading('Repeated',1);doc.add_paragraph('NEEDLE')
    path=tmp_path/'duplicate.docx';doc.save(path)
    shared,score,_=inspect(path,expected,check,heading='Repeated')
    assert shared.checks[0].status=='NEEDS_REVIEW'
    assert score.total_score is None


def test_localized_builtin_heading_uses_stable_style_identity(tmp_path):
    doc=Document();style=doc.styles['Heading 1'];style.name='标题 1';doc.add_paragraph('Localized',style=style)
    doc.add_paragraph('NEEDLE');path=tmp_path/'localized.docx';doc.save(path)
    assert inspect(path,['Localized'],'word.headings')[0].checks[0].status=='PASS'
    assert inspect(path,heading='Localized')[0].checks[0].status=='PASS'
    assert inspect(path,'Heading 1','word.heading_style',heading='Localized')[0].checks[0].status=='PASS'


def test_custom_outline_heading_is_recognized(tmp_path):
    doc=Document();style=doc.styles.add_style('SectionTitleLocal',WD_STYLE_TYPE.PARAGRAPH)
    outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),'0');style.element.get_or_add_pPr().append(outline)
    doc.add_paragraph('Outline section',style=style);doc.add_paragraph('NEEDLE')
    path=tmp_path/'outline.docx';doc.save(path)
    assert inspect(path,['Outline section'],'word.headings')[0].checks[0].status=='PASS'
    assert inspect(path,heading='Outline section')[0].checks[0].status=='PASS'


def test_word_table_cell_scope_excludes_other_cells(tmp_path):
    doc=Document();table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='NEEDLE';table.cell(0,1).text='OTHER'
    path=tmp_path/'table.docx';doc.save(path)
    assert inspect(path,table=1,row=1,column=2)[0].checks[0].status=='FAIL'
    assert inspect(path,table=1,row=1,column=1)[0].checks[0].status=='PASS'
    assert inspect(path,table=1)[0].checks[0].status=='PASS'


@pytest.mark.parametrize('position',[{'slide':2},{'slide':99},{'slide':1,'object_id':999}])
def test_ppt_scope_cannot_borrow_other_slide_or_object(tmp_path,position):
    _,path=deck(tmp_path)
    assert inspect(path,**position)[0].checks[0].status=='FAIL'


def test_powerpoint_text_list_obeys_slide_scope(tmp_path):
    _,path=deck(tmp_path)
    assert inspect(path,['NEEDLE'],'powerpoint.texts',slide=2)[0].checks[0].status=='FAIL'


def test_ppt_group_text_is_found_with_precise_observation(tmp_path):
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    group=slide.shapes.add_group_shape();nested=group.shapes.add_group_shape()
    shape=nested.shapes.add_textbox(0,0,Inches(2),Inches(1));shape.text='NEEDLE'
    path=tmp_path/'group.pptx';prs.save(path)
    shared,_,_=inspect(path,slide=1,object_id=group.shape_id)
    assert shared.checks[0].status=='PASS'
    loc=shared.checks[0].observed_locations[0]
    assert loc.object_id==shape.shape_id and loc.group_path==[group.shape_id,nested.shape_id]
    assert inspect(path,slide=1,object_id=shape.shape_id,group_path=[group.shape_id,nested.shape_id])[0].checks[0].status=='PASS'


def test_wrong_group_path_does_not_fall_back_to_global_search(tmp_path):
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    group=slide.shapes.add_group_shape();shape=group.shapes.add_textbox(0,0,10,10);shape.text='NEEDLE'
    path=tmp_path/'wrong-group.pptx';prs.save(path)
    assert inspect(path,slide=1,object_id=shape.shape_id,group_path=[999])[0].checks[0].status=='FAIL'


def test_ppt_table_cell_scope(tmp_path):
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    shape=slide.shapes.add_table(1,2,0,0,Inches(4),Inches(1))
    shape.table.cell(0,0).text='NEEDLE';shape.table.cell(0,1).text='OTHER'
    path=tmp_path/'ppt-table.pptx';prs.save(path)
    assert inspect(path,slide=1,object_id=shape.shape_id,row=1,column=2)[0].checks[0].status=='FAIL'
    shared,_,_=inspect(path,slide=1,object_id=shape.shape_id,row=1,column=1)
    assert shared.checks[0].status=='PASS' and shared.checks[0].observed_locations[0].column==1


@pytest.mark.parametrize('position',[{'sheet':'Second'},{'sheet':'First','cell':'B2'},
    {'sheet':'First','cell':'B2:C3'},{'sheet':'Missing'}])
def test_excel_scope_excludes_other_cells_and_sheets(tmp_path,position):
    _,path=excel(tmp_path)
    assert inspect(path,**position)[0].checks[0].status=='FAIL'


def test_excel_range_reports_actual_cell_without_replacing_range(tmp_path):
    _,path=excel(tmp_path);shared,_,_=inspect(path,sheet='First',cell='A1:B2')
    record=shared.checks[0]
    assert record.status=='PASS' and record.location.cell=='A1:B2'
    assert record.requested_location.cell=='A1:B2'
    assert record.observed_locations[0].cell=='A1'


@pytest.mark.parametrize('factory',[word,excel,deck])
def test_unspecified_location_keeps_full_document_search(tmp_path,factory):
    _,path=factory(tmp_path)
    assert inspect(path)[0].checks[0].status=='PASS'


def test_unbounded_scope_does_not_iterate_billions_of_cells(tmp_path):
    _,path=excel(tmp_path)
    shared,score,_=inspect(path,sheet='First',cell='A1:XFD1048576')
    assert shared.checks[0].status=='NEEDS_REVIEW' and score.total_score is None


def test_unsupported_region_selector_is_spec_gap(tmp_path):
    _,path=word(tmp_path);shared,_,_=inspect(path,region='top-right')
    assert shared.checks[0].failure_origin=='spec_gap'


def test_source_bytes_unchanged_by_scope_checks(tmp_path):
    from workagent_rsi.hashing import sha256_file
    _,path=word(tmp_path);before=sha256_file(path)
    inspect(path,paragraph=2)
    assert sha256_file(path)==before


def test_slide_geometry_out_of_range_is_failure_not_checker_error(tmp_path):
    _,path=deck(tmp_path)
    assert inspect(path,[],'powerpoint.bounds',slide=99)[0].checks[0].status=='FAIL'


def test_exact_paragraph_out_of_range_is_failure(tmp_path):
    _,path=word(tmp_path)
    assert inspect(path,'NEEDLE','word.paragraph',paragraph=99)[0].checks[0].status=='FAIL'


def test_direct_grouped_object_text_uses_same_scope(tmp_path):
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    group=slide.shapes.add_group_shape();shape=group.shapes.add_textbox(0,0,10,10);shape.text='NEEDLE'
    path=tmp_path/'exact-group.pptx';prs.save(path)
    assert inspect(path,'NEEDLE','powerpoint.object_text',slide=1,object_id=shape.shape_id)[0].checks[0].status=='PASS'


def test_duplicate_object_identity_requires_review(tmp_path):
    prs,path=deck(tmp_path);slide=prs.slides[0]
    first=slide.shapes[0];second=slide.shapes.add_textbox(0,0,10,10);second.text='OTHER'
    second._element.xpath('.//p:cNvPr')[0].set('id',str(first.shape_id));prs.save(path)
    assert inspect(path,slide=1,object_id=first.shape_id)[0].checks[0].status=='NEEDS_REVIEW'


def test_heading_section_can_be_refined_to_one_paragraph(tmp_path):
    doc=Document();doc.add_heading('Target',1);doc.add_paragraph('NEEDLE');doc.add_paragraph('OTHER')
    path=tmp_path/'refined.docx';doc.save(path)
    assert inspect(path,heading='Target',paragraph=3)[0].checks[0].status=='FAIL'
    assert inspect(path,heading='Target',paragraph=2)[0].checks[0].status=='PASS'


def test_missing_table_cell_does_not_match_another_cell(tmp_path):
    doc=Document();table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='NEEDLE'
    path=tmp_path/'missing-cell.docx';doc.save(path)
    assert inspect(path,table=1,row=9,column=1)[0].checks[0].status=='FAIL'


def test_altered_requested_location_is_rejected_by_evidence_binding(tmp_path):
    from workagent_rsi.assessment_contracts import Assessment
    _,path=word(tmp_path);shared,_,_=inspect(path,paragraph=1)
    payload=shared.model_dump(mode='json');payload['checks'][0]['requested_location']['paragraph']=2
    with pytest.raises(ValueError,match='requested location'):
        Assessment.model_validate(payload)
