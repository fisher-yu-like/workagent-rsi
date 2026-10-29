"""Authored protocol fixtures. No real judge accuracy is asserted here."""

from openpyxl import Workbook, load_workbook
from workagent_rsi.contracts import TaskSpec, ArtifactRef
from workagent_rsi.assessment_contracts import AcceptanceSpec, Requirement, Location
from workagent_rsi.hashing import sha256_file
from workagent_rsi.artifact_assessment import SharedAssessment
from test_llm_task_review import Provider, judgement


def authored_spec(*, visual=False, critical_weight=1):
    return AcceptanceSpec(
        task_id="parity",
        version="unit-fixture",
        artifacts={"output": "excel"},
        dimension_weights={"content": 1},
        requirements=[
            Requirement(
                requirement_id="visual" if visual else "answer",
                description="Authored fixture",
                check="visual" if visual else "excel.cell",
                location=Location(artifact="output"),
                expected=42,
                dimension="content",
                evidence_source="unit fixture",
            )
        ],
    )


def reviewed_task(spec):
    return TaskSpec(
        task_id=spec.task_id,
        domain="excel",
        instruction="Write 42 in Sheet!A1."
        + (
            " Make the layout readable."
            if any(r.check == "visual" for r in spec.requirements)
            else ""
        )
        + (
            " Keep B1 original."
            if any(r.requirement_id == "keep" for r in spec.requirements)
            else ""
        ),
        expected_constraints={"acceptance_spec": spec.model_dump(mode="json")},
    )


def book(root, value=42):
    p = root / "book.xlsx"
    b = Workbook()
    b.active["A1"] = value
    b.save(p)
    return p


def ref(path):
    return ArtifactRef(
        artifact_id="unit",
        path=str(path),
        sha256=sha256_file(path),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        size_bytes=path.stat().st_size,
    )


class FixtureAssessment(SharedAssessment):
    """Offline scripted model double for control-flow tests,not an evaluator."""

    def __init__(self):
        super().__init__(provider=Provider())

    def inspect(self, spec, artifacts):
        if "readable" in spec.task_context["instruction"]:
            return self.unavailable(spec, "Scripted missing visual evidence")
        if artifacts:
            p = next(iter(artifacts.values()))
            b = load_workbook(p)
            value = b.active["A1"].value
            kept = b.active["B1"].value
            b.close()
            issue = {
                "issue_id": "answer",
                "description": "Scripted model defect",
                "location": {"artifact": "output", "sheet": "Sheet", "cell": "A1"},
                "kind": "artifact_defect",
                "severity": "major",
                "repair_suggestion": "Repeat the requested edit.",
                "evidence_refs": ["output"],
            }
            self.provider.decision = judgement(
                "PASS" if value == 42 else "FAIL",
                1 if value == 42 else 0.2,
                [] if value == 42 else [issue],
            )
            if "Keep B1" in spec.task_context["instruction"] and kept != "original":
                issue.update(
                    issue_id="keep", description="Scripted preservation defect"
                )
                self.provider.decision = judgement("FAIL", 0.3, [issue])
        return super().inspect(spec, artifacts)
