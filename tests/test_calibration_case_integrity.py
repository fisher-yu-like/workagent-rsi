"""Authored truth is checked through raw OOXML, outside the production evaluator."""
import hashlib
import json
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest

from workagent_rsi.calibration_data import build_calibration_cases, load_calibration_cases


MANIFEST = Path(__file__).parent / 'fixtures/calibration/case_manifest.json'


@pytest.fixture(scope='module')
def corpus(tmp_path_factory):
    root = tmp_path_factory.mktemp('calibration')
    return root, build_calibration_cases(root, MANIFEST)


def test_balanced_corpus_with_separate_truth(corpus):
    root, cases = corpus
    assert len(cases) == 90
    assert Counter(c.domain for c in cases) == {'excel': 30, 'word': 30, 'powerpoint': 30}
    for domain in ('excel', 'word', 'powerpoint'):
        assert Counter(c.category for c in cases if c.domain == domain) == {
            'correct': 10, 'single_defect': 10, 'equivalent': 5, 'gap': 5}
    assert Counter(c.oracle_status for c in cases) == {'PASS': 45, 'FAIL': 30, 'INCOMPLETE': 15}
    assert len({c.case_id for c in cases}) == 90
    for c in cases:
        assert c.source_refs and c.expected_locations and c.expected_scope
        assert c.artifact_hash == hashlib.sha256((root / c.artifact_path).read_bytes()).hexdigest()
        assert c.input_hash == hashlib.sha256((root / c.input_path).read_bytes()).hexdigest()
        assert c.truth_source_hash == hashlib.sha256(MANIFEST.read_bytes()).hexdigest()
        assert c.source_kind == 'project-generated'
    assert load_calibration_cases(root / 'case_manifest.json') == cases


def test_every_office_file_reopens_and_matches_authored_xml_claims(corpus):
    from openpyxl import load_workbook
    from docx import Document
    from pptx import Presentation
    root, cases = corpus
    truth = {c['case_id']: c for c in json.loads(MANIFEST.read_text(encoding='utf-8'))['cases']}
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
          'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
          'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
          'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
    for c in cases:
        path = root / c.artifact_path
        if c.domain == 'excel':
            book = load_workbook(path)
            assert len(book.worksheets) == 2
            book.close()
        elif c.domain == 'word':
            assert Document(path).tables
        else:
            assert len(Presentation(path).slides) == 2
        with ZipFile(path) as archive:
            for claim in truth[c.case_id]['xml_claims']:
                tree = ET.fromstring(archive.read(claim['part']))
                nodes = tree.findall(claim['xpath'], ns)
                if 'count' in claim:
                    assert len(nodes) == claim['count'], (c.case_id, claim)
                if 'texts' in claim:
                    assert [n.text for n in nodes] == claim['texts'], (c.case_id, claim)


def test_gaps_and_equivalents_are_not_bad_artifacts(corpus):
    _, cases = corpus
    for c in cases:
        if c.category == 'gap':
            assert c.oracle_status == 'INCOMPLETE' and c.gap_origin in {'spec_gap', 'capability_gap'}
            assert c.defect_type is None and c.artifact_valid is True
        elif c.category == 'equivalent':
            assert c.oracle_status == 'PASS' and c.equivalence_authorization in c.instruction
            assert c.defect_type is None
        elif c.category == 'single_defect':
            assert c.oracle_status == 'FAIL' and c.defect_type and c.artifact_valid is False
    assert {'omitted_last_row', 'wrong_range', 'misplaced_content', 'obstruction', 'source_damage'} <= {
        c.defect_type for c in cases if c.defect_type}


def test_rebuild_is_byte_reproducible_and_tamper_fails(corpus, tmp_path):
    root, cases = corpus
    rebuilt = build_calibration_cases(tmp_path / 'rebuilt', MANIFEST)
    assert [(c.artifact_hash, c.input_hash) for c in rebuilt] == [(c.artifact_hash, c.input_hash) for c in cases]
    damaged = tmp_path / 'rebuilt' / rebuilt[0].artifact_path
    damaged.write_bytes(b'not the frozen Office file')
    with pytest.raises(ValueError, match='hash'):
        load_calibration_cases(tmp_path / 'rebuilt/case_manifest.json')
    with pytest.raises(FileExistsError):
        build_calibration_cases(root, MANIFEST)
