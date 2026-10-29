from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document

from workagent_rsi.word_package_checks import inspect_word_rules, validate_word_rules


def save(tmp_path, name, text):
    d=Document();d.add_paragraph(text);p=tmp_path/name;d.save(p);return p


def add_part(path,name,xml):
    with ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    parts[name]=xml.encode()
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)


def test_replacement_checks_each_source_scope_and_split_runs(tmp_path):
    s=save(tmp_path,'source.docx','Acme');o=save(tmp_path,'output.docx','Nova')
    xml='<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>Ac</w:t></w:r><w:r><w:t>me</w:t></w:r></w:p></w:hdr>'
    add_part(s,'word/header1.xml',xml);add_part(o,'word/header1.xml',xml)
    rule={'kind':'replace','scope':'all','before':'Acme','after':'Nova'}
    assert inspect_word_rules(o,[rule],s)[0]=='FAIL'
    add_part(o,'word/header1.xml',xml.replace('Ac</w:t>','No</w:t>').replace('me</w:t>','va</w:t>'))
    assert inspect_word_rules(o,[rule],s)[0]=='PASS'


def test_replacement_cannot_pass_by_deleting_original_paragraph(tmp_path):
    s=save(tmp_path,'s.docx','Acme');o=save(tmp_path,'o.docx','Unrelated')
    assert inspect_word_rules(o,[{'kind':'replace','before':'Acme','after':'Nova'}],s)[0]=='FAIL'


def test_bold_name_must_remain_bold(tmp_path):
    d=Document();d.add_paragraph().add_run('Acme').bold=True;s=tmp_path/'s.docx';d.save(s)
    o=save(tmp_path,'o.docx','Nova');r={'kind':'replace','before':'Acme','after':'Nova','preserve_bold':True}
    assert inspect_word_rules(o,[r],s)[0]=='NEEDS_REVIEW'
    d=Document();d.add_paragraph().add_run('Nova').bold=True;d.save(o)
    assert inspect_word_rules(o,[r],s)[0]=='PASS'
    d.paragraphs[0].runs[0].bold=False;d.save(o)
    assert inspect_word_rules(o,[r],s)[0]=='FAIL'


def test_bold_occurrence_count_cannot_hide_wrong_location(tmp_path):
    d=Document();p=d.add_paragraph();p.add_run('Acme').bold=True;p.add_run(' and ');p.add_run('Acme').bold=False
    s=tmp_path/'s.docx';d.save(s)
    d=Document();p=d.add_paragraph();p.add_run('Nova').bold=False;p.add_run(' and ');p.add_run('Nova').bold=True
    o=tmp_path/'o.docx';d.save(o)
    r={'kind':'replace','before':'Acme','after':'Nova','preserve_bold':True}
    assert inspect_word_rules(o,[r],s)[0]=='NEEDS_REVIEW'


def test_comments_author_and_exact_count_are_not_body_decoys(tmp_path):
    o=save(tmp_path,'o.docx','Reviewer: unlimited liability')
    rule={'kind':'comments','author':'Reviewer','count':3}
    assert inspect_word_rules(o,[rule])[0]=='FAIL'
    xml='<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'+''.join(f'<w:comment w:id="{i}" w:author="Reviewer"><w:p><w:r><w:t>Review clause {i}</w:t></w:r></w:p></w:comment>' for i in range(3))+'</w:comments>'
    add_part(o,'word/comments.xml',xml)
    # Unanchored comments are not accepted as comments attached to clauses.
    assert inspect_word_rules(o,[rule])[0]=='FAIL'


def test_table_row_order_and_all_values_are_required(tmp_path):
    d=Document();t=d.add_table(rows=0,cols=3)
    for row in [('West','1','2'),('International','600','750')]:
        for c,v in zip(t.add_row().cells,row):c.text=v
    o=tmp_path/'o.docx';d.save(o)
    r={'kind':'table_row','cells':['International','600','750'],'after_row':'West'}
    assert inspect_word_rules(o,[r])[0]=='PASS'
    r['cells'][-1]='820';assert inspect_word_rules(o,[r])[0]=='FAIL'


def test_blank_lines_and_real_tracked_markup_are_checked(tmp_path):
    s=save(tmp_path,'s.docx','Hello');d=Document(s);d.add_paragraph('');d.save(s)
    o=save(tmp_path,'o.docx','Hello')
    assert inspect_word_rules(o,[{'kind':'blank_paragraphs'}],s)[0]=='FAIL'
    with ZipFile(o) as z:xml=z.read('word/document.xml').decode()
    add_part(o,'word/document.xml',xml.replace('<w:t>Hello</w:t>','<w:ins><w:t>Hello</w:t></w:ins>'))
    assert inspect_word_rules(o,[{'kind':'tracked_clean'}])[0]=='FAIL'


@pytest.mark.parametrize('rule',[{'kind':'unknown'},{'kind':'text','text':'x','count':True},{'kind':'replace','before':'x','after':''},{'kind':'text','scope':'word/../x','text':'x'}])
def test_malformed_rules_rejected(rule):
    with pytest.raises(ValueError):validate_word_rules([rule])


def test_comment_removal_preserves_other_comments_and_requires_removed_anchors_gone(tmp_path):
    s=save(tmp_path,'s.docx','Clause');o=save(tmp_path,'o.docx','Clause')
    ns='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    comments=f'<w:comments xmlns:w="{ns}">'+''.join(f'<w:comment w:id="{i}" w:author="Reviewer"><w:p><w:r><w:t>Note {i}</w:t></w:r></w:p></w:comment>' for i in range(2))+'</w:comments>'
    add_part(s,'word/comments.xml',comments)
    add_part(o,'word/comments.xml',comments.replace('<w:comment w:id="0" w:author="Reviewer"><w:p><w:r><w:t>Note 0</w:t></w:r></w:p></w:comment>',''))
    with ZipFile(o) as z:body=z.read('word/document.xml').decode()
    anchors='<w:commentRangeStart w:id="1"/><w:commentRangeEnd w:id="1"/><w:r><w:commentReference w:id="1"/></w:r>'
    add_part(o,'word/document.xml',body.replace('<w:t>Clause</w:t>','<w:t>Clause</w:t>'+anchors))
    assert inspect_word_rules(o,[{'kind':'comments_preserved','remove_ids':['0']}],s)[0]=='PASS'
    add_part(o,'word/comments.xml',comments)
    assert inspect_word_rules(o,[{'kind':'comments_preserved','remove_ids':['0']}],s)[0]=='FAIL'


def test_footnote_suffix_checks_exact_note_and_original_text(tmp_path):
    s=save(tmp_path,'s.docx','Body');o=save(tmp_path,'o.docx','Body')
    xml='<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>Original note</w:t></w:r></w:p></w:footnote></w:footnotes>'
    add_part(s,'word/footnotes.xml',xml);add_part(o,'word/footnotes.xml',xml.replace('Original note','Wrong note as amended'))
    rule={'kind':'footnote_suffix','note_id':'1','suffix':' as amended'}
    assert inspect_word_rules(o,[rule],s)[0]=='FAIL'
    add_part(o,'word/footnotes.xml',xml.replace('Original note','Original note as amended'))
    assert inspect_word_rules(o,[rule],s)[0]=='PASS'


def test_package_check_is_registered_source_bound_and_used_by_assessment(tmp_path):
    from workagent_rsi.artifact_assessment import SharedAssessment
    from workagent_rsi.assessment_contracts import AcceptanceSpec,Requirement,Location
    from workagent_rsi.check_registry import validate_acceptance_spec
    from workagent_rsi.hashing import sha256_file
    s=save(tmp_path,'s.docx','Acme');o=save(tmp_path,'o.docx','Nova')
    r=Requirement(requirement_id='rename',description='Exact authorized rename',check='word.package',
        location=Location(artifact='output'),expected=[{'kind':'replace','before':'Acme','after':'Nova'}],
        options={'source_path':str(s)},dimension='content',evidence_source='Authored source instruction')
    spec=AcceptanceSpec(task_id='t',version='v1',artifacts={'output':'word'},dimension_weights={'content':1},requirements=[r])
    assert validate_acceptance_spec(spec,{}).status=='SPEC_INCOMPLETE'
    spec=spec.model_copy(update={'input_hashes':{str(s):sha256_file(s)}})
    assert SharedAssessment().inspect(spec,{'output':o}).checks[0].status=='PASS'
    s.write_bytes(b'Changed source')
    with pytest.raises(ValueError,match='frozen assessment input changed'):
        SharedAssessment().inspect(spec,{'output':o})
