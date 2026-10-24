"""T21 qualification uses authored executors and real openpyxl files, no model."""
import inspect
import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from test_assessment_entrypoint_parity import authored_spec, reviewed_task
from workagent_rsi.harness import Harness
from workagent_rsi.workagent_provider import WorkAgentConfig


def policy(**kwargs):
    assert 'repair_policy' in inspect.signature(Harness.run).parameters, 'T21 repair entry point is absent'
    from workagent_rsi.repair_loop import RepairPolicy
    return RepairPolicy(**kwargs)


def executor(monkeypatch, tmp_path, values):
    calls = []
    class AuthoredExecutor:
        def __init__(self, root, *args, **kwargs):
            self.root = Path(root)
        def execute(self, task, skill_id):
            index = len(calls)
            calls.append((task.model_dump(mode='json'), skill_id))
            value = values[min(index, len(values)-1)]
            value, kept = value if isinstance(value,tuple) else (value,'preserved')
            if isinstance(value, str):
                yield {'kind': 'failure', 'message': value, 'failure_origin': 'agent_execution'}
                return
            self.root.mkdir(parents=True, exist_ok=True)
            path = self.root/'authored.xlsx'
            wb = Workbook(); wb.active['A1'] = value; wb.active['B1'] = kept; wb.save(path); wb.close()
            yield {'kind': 'artifact', 'artifact_path': str(path), 'artifact_name': 'output',
                   'media_type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}
    monkeypatch.setattr('workagent_rsi.run.WorkAgentOfficeAdapter', AuthoredExecutor)
    return calls


def harness(tmp_path):
    return Harness(tmp_path/'results', workagent_config=WorkAgentConfig(verify_com=False), agent_instructions='Frozen skill')


def test_wrong_artifact_gets_one_repair_and_fresh_full_assessment(tmp_path, monkeypatch):
    calls = executor(monkeypatch, tmp_path, [7, 42])
    task = reviewed_task(authored_spec())
    result = harness(tmp_path).run(task, repair_policy=policy(max_repairs=1))
    assert len(calls) == 2 and result['state'] == 'SUCCEEDED'
    repair = result['repair_loop']
    assert repair['repairs'] == 1
    assert [a['evaluation']['task_acceptance'] for a in repair['attempts']] == ['FAIL', 'PASS']
    assert repair['attempts'][0]['artifacts'][0]['sha256'] != repair['attempts'][1]['artifacts'][0]['sha256']
    assert len({a['evaluation']['report_paths']['score'] for a in repair['attempts']}) == 2
    wb = load_workbook(BytesIO(Path(result['artifacts'][0]['path']).read_bytes())); assert wb.active['A1'].value == 42; wb.close()
    request = json.loads((Path(result['result_dir'])/'repair-loop/attempt-001/request.json').read_text())
    assert request['task']['instruction'] == task.instruction
    assert request['feedback']['issues'] and request['task_hash'] == repair['task_hash']
    assert calls[0] == calls[1]


def test_disabled_repair_does_not_change_max_attempts(tmp_path, monkeypatch):
    calls = executor(monkeypatch, tmp_path, [7, 42])
    result = harness(tmp_path).run(reviewed_task(authored_spec()), max_attempts=3)
    assert result['state'] == 'FAILED' and len(calls) == 1 and 'repair_loop' not in result


def test_code_failure_can_enter_repair_without_fake_quality(tmp_path, monkeypatch):
    calls = executor(monkeypatch, tmp_path, ['SyntaxError: missing colon', 42])
    result = harness(tmp_path).run(reviewed_task(authored_spec()), repair_policy=policy(max_repairs=1))
    assert len(calls) == 2 and result['state'] == 'SUCCEEDED'
    assert 'evaluation' not in result['repair_loop']['attempts'][0]


def test_new_regression_is_detected_by_all_requirements(tmp_path, monkeypatch):
    from workagent_rsi.assessment_contracts import Requirement, Location
    spec = authored_spec()
    spec = spec.model_copy(update={'requirements': [*spec.requirements,
        Requirement(requirement_id='keep', description='Keep B1 unchanged', check='excel.cell',
                    location=Location(artifact='output', sheet='Sheet', cell='B1'), expected='original',
                    critical=True, dimension='content', evidence_source='independent authored oracle')]})
    calls = executor(monkeypatch, tmp_path, [(7,'original'), (42,'destroyed')])
    result = harness(tmp_path).run(reviewed_task(spec), repair_policy=policy(max_repairs=1))
    assert len(calls) == 2 and result['state'] == 'FAILED'
    last = result['repair_loop']['attempts'][-1]
    first_score = json.loads(Path(result['repair_loop']['attempts'][0]['evaluation']['report_paths']['score']).read_text())
    assert {c['requirement_id']: c['status'] for c in first_score['criterion_results']} == {'answer':'FAIL','keep':'PASS'}
    score = json.loads(Path(last['evaluation']['report_paths']['score']).read_text())
    assert {c['requirement_id']: c['status'] for c in score['criterion_results']} == {'answer':'PASS', 'keep':'FAIL'}


def test_environment_failure_is_not_repaired(tmp_path, monkeypatch):
    calls = []
    class MissingWorker:
        def __init__(self, *args, **kwargs): pass
        def execute(self, *args):
            calls.append(1)
            yield {'kind':'unavailable','message':'worker unavailable','failure_origin':'environment'}
    monkeypatch.setattr('workagent_rsi.run.WorkAgentOfficeAdapter', MissingWorker)
    result = harness(tmp_path).run(reviewed_task(authored_spec()), repair_policy=policy())
    assert len(calls) == 1 and result['repair_loop']['repairs'] == 0


def test_assessment_gap_does_not_ask_executor_to_change_artifact(tmp_path, monkeypatch):
    calls = executor(monkeypatch, tmp_path, [7, 42])
    result = harness(tmp_path).run(reviewed_task(authored_spec(visual=True)), repair_policy=policy())
    assert len(calls) == 1 and result['repair_loop']['repairs'] == 0


def test_inner_repair_never_changes_global_skill_registry(tmp_path, monkeypatch):
    registry = tmp_path/'registry.json'; registry.write_text('{"champion":"frozen"}')
    before = registry.read_bytes()
    executor(monkeypatch, tmp_path, [7, 42])
    result = harness(tmp_path).run(reviewed_task(authored_spec()), repair_policy=policy())
    assert result['state'] == 'SUCCEEDED' and registry.read_bytes() == before


def test_attempt_manifest_is_auditable_and_resume_uses_selected_state(tmp_path, monkeypatch):
    executor(monkeypatch, tmp_path, [7, 42])
    h = harness(tmp_path); result = h.run(reviewed_task(authored_spec()), repair_policy=policy())
    root = Path(result['result_dir'])/'repair-loop'
    manifest = json.loads((root/'attempt-manifest.json').read_text())
    assert manifest['repairs'] == 1 and len(manifest['attempts']) == 2
    assert (root/'before-after-reports.json').is_file()
    assert h.resume(result['result_dir'])['state'] == result['state']


def test_real_adapter_receives_repair_data_in_a_fresh_workspace(tmp_path, monkeypatch):
    from workagent_rsi.workagent_provider import AgentResponse, ProviderOutcome
    captured = []
    def fake_provider(self, prompt, workspace, records):
        captured.append((prompt, workspace))
        records.mkdir(parents=True, exist_ok=True)
        (workspace/'outputs').mkdir(exist_ok=True)
        script = workspace/'agent.py'; script.write_text('# authored provider substitute\n')
        wb = Workbook(); wb.active['A1'] = 7 if len(captured)==1 else 42
        wb.save(workspace/'outputs/book.xlsx'); wb.close()
        response = AgentResponse(status='completed', deliverables=['outputs/book.xlsx'], summary='authored', input_files_used=[])
        (workspace/'agent_response.json').write_text(response.model_dump_json())
        (workspace/'deliverables.json').write_text(json.dumps({'deliverables':response.deliverables}))
        return ProviderOutcome(status='completed', response=response)
    monkeypatch.setattr('workagent_rsi.workagent_office.CodexOfficeProvider.run', fake_provider)
    result = harness(tmp_path).run(reviewed_task(authored_spec()), repair_policy=policy())
    assert result['state']=='SUCCEEDED' and len(captured)==2
    assert captured[0][1]!=captured[1][1]
    prompt, workspace = captured[1]
    assert 'Task repair request' in prompt and 'authored provider substitute' in prompt
    assert 'Frozen skill' in prompt and 'expected_constraints' not in prompt
    assert (workspace/'previous_artifacts/artifact-0.xlsx').is_file()


@pytest.mark.parametrize('newline', ['\n','\r\n'])
def test_execution_repair_request_keeps_trace_grounded_diagnosis(tmp_path,monkeypatch,newline):
    import hashlib
    import traceback
    from workagent_rsi.run import Run
    calls=executor(monkeypatch,tmp_path,['SyntaxError: missing colon',42])
    original=Run._execute_once
    def with_trace(self,*args,**kwargs):
        result=original(self,*args,**kwargs)
        if result.get('failure'):
            source=('for value in [1,2]\n    print(value)\n').replace('\n',newline)
            try:compile(source,'agent.py','exec')
            except SyntaxError as error:stderr=''.join(traceback.format_exception_only(type(error),error))
            digest=hashlib.sha256(source.encode()).hexdigest()
            records=self.root/'provider_records';(records/'executed-scripts').mkdir(parents=True)
            (records/'executed-scripts'/f'{digest}.py').write_bytes(source.encode('utf-8'))
            event={'name':'run_python','arguments':{'path':'agent.py'},'result':{'script_sha256':digest,
                'script_snapshot_ref':f'executed-scripts/{digest}.py','stderr':stderr,
                'execution':{'status':'failed','failure_origin':'agent_execution'}}}
            (records/'tool_events.jsonl').write_text(json.dumps(event)+'\n')
        return result
    monkeypatch.setattr(Run,'_execute_once',with_trace)
    result=harness(tmp_path).run(reviewed_task(authored_spec()),repair_policy=policy(max_repairs=1))
    request=json.loads((Path(result['result_dir'])/'repair-loop/attempt-001/request.json').read_text())
    assert len(calls)==2 and result['state']=='SUCCEEDED'
    diagnoses=request['feedback']['diagnoses']
    assert diagnoses and diagnoses[0]['confidence']>0
    assert hashlib.sha256(request['current_code'][0]['text'].encode('utf-8')).hexdigest()==request['current_code'][0]['sha256']


@pytest.mark.parametrize('kwargs', [{'max_repairs':-1}, {'total_deadline_seconds':0}, {'visibility':'hidden'}])
def test_invalid_repair_policy_is_rejected(kwargs):
    with pytest.raises(ValueError): policy(**kwargs)


def test_current_code_uses_latest_execution_instead_of_all_old_versions(tmp_path):
    import hashlib
    from workagent_rsi.repair_loop import _sources
    records=tmp_path/'provider_records';(records/'executed-scripts').mkdir(parents=True)
    events=[]
    for text in ['# old program\n','# current executed program\n']:
        digest=hashlib.sha256(text.encode()).hexdigest()
        (records/'executed-scripts'/f'{digest}.py').write_bytes(text.encode())
        events.append({'name':'run_python','result':{'script_sha256':digest,'script_snapshot_ref':f'executed-scripts/{digest}.py'}})
    (records/'tool_events.jsonl').write_text(''.join(json.dumps(e)+'\n' for e in events))
    sources=_sources({'_attempt_root':str(tmp_path)})
    assert len(sources)==1 and sources[0]['text']=='# current executed program\n'
