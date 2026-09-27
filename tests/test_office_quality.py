from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from workagent_rsi.contracts import ArtifactRef, TaskSpec
from workagent_rsi.evaluator import OfficeArtifactEvaluator
from workagent_rsi.office_checks import inspect_office_file
from workagent_rsi.render_checks import check_powerpoint_geometry


def test_excel_quality_check_validates_sheet_cell_and_formula(tmp_path: Path):
    path = tmp_path / "quality.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "MARKER"
    sheet["B1"] = "=SUM(B2:B3)"
    workbook.save(path)
    task = TaskSpec(
        task_id="excel-quality",
        domain="excel",
        instruction="create a summary",
        expected_constraints={
            "required_text": "MARKER",
            "required_sheets": ["Summary"],
            "required_cells": {"Summary!A1": "MARKER"},
            "required_formulas": {"Summary!B1": "=SUM(B2:B3)"},
        },
    )

    report = inspect_office_file(path, task)

    assert report.passed is True
    assert report.dimensions["structure_correctness"] == 1.0
    assert report.dimensions["formula_correctness"] == 1.0


def test_word_quality_check_validates_heading_and_style(tmp_path: Path):
    path = tmp_path / "quality.docx"
    document = Document()
    heading = document.add_heading("Overview", level=1)
    document.add_paragraph("MARKER")
    document.save(path)
    task = TaskSpec(
        task_id="word-quality",
        domain="word",
        instruction="create a document",
        expected_constraints={
            "required_text": "MARKER",
            "required_headings": ["Overview"],
            "required_heading_styles": {"Overview": "Heading 1"},
        },
    )

    report = inspect_office_file(path, task)

    assert report.passed is True
    assert report.dimensions["structure_correctness"] == 1.0
    assert any("paragraph" in evidence for evidence in report.evidence)


def test_powerpoint_geometry_reports_overlap_and_overflow(tmp_path: Path):
    path = tmp_path / "quality.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(1.5), Inches(4), Inches(2))
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, presentation.slide_width - Inches(1), Inches(1), Inches(2), Inches(2))
    presentation.save(path)

    report = check_powerpoint_geometry(path)

    assert report.overlap_count >= 1
    assert report.overflow_count >= 1
    assert report.passed is False


def test_office_evaluator_exposes_file_quality_dimensions(tmp_path: Path):
    path = tmp_path / "quality.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "MARKER"
    workbook.save(path)
    task = TaskSpec(
        task_id="evaluator-quality",
        domain="excel",
        instruction="create a workbook",
        expected_constraints={"required_text": "MARKER"},
    )
    artifact = ArtifactRef(
        artifact_id="a" * 64,
        path=str(path),
        sha256="b" * 64,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        size_bytes=path.stat().st_size,
    )

    report = OfficeArtifactEvaluator().evaluate(task, [artifact], "trace-quality")

    assert report.passed is True
    assert report.dimensions["format_validity"] == 1.0
    assert report.dimensions["structure_correctness"] == 1.0
    assert report.dimensions["formula_correctness"] == 1.0
