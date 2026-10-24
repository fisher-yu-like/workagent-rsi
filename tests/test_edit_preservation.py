"""Independent before/after Office edits and explicitly authorized changes."""
import importlib
from pathlib import Path

from docx import Document
from openpyxl import Workbook,load_workbook
from pptx import Presentation
from pptx.util import Inches
import pytest

from workagent_rsi.hashing import sha256_file


def compare(source,output,allowed=(),preserved=()):
    module=importlib.import_module('workagent_rsi.preservation_checks')
    spec=module.PreservationSpec(source_hash=sha256_file(source),allowed_changes=list(allowed),
        preserved_regions=list(preserved),normalization_policy='office-structure-v1')
    return module.compare_preservation(source,output,spec)


def passed(rows):return all(row.status=='PASS' for row in rows)


def excel_pair(tmp_path):
    source=tmp_path/'source.xlsx';output=tmp_path/'output.xlsx'
    book=Workbook();book.active.title='Data';book.active.append(['Keep',10]);book.active['C1']='=B1*2';book.save(source)
    book.save(output);book.close();return source,output


def test_changed_unrelated_cell_is_defect(tmp_path):
    source,output=excel_pair(tmp_path);book=load_workbook(output);book.active['B1']=99;book.save(output);book.close()
    rows=compare(source,output,[{'scope':'excel/sheet/Data/cell/A1','operation':'replace'}])
    assert not passed(rows) and any('B1' in str(r.observed) for r in rows)


def test_authorized_cell_edit_preserves_other_cells(tmp_path):
    source,output=excel_pair(tmp_path);book=load_workbook(output);book.active['B1']=99;book.save(output);book.close()
    assert passed(compare(source,output,[{'scope':'excel/sheet/Data/cell/B1','operation':'replace'}]))


def test_formula_change_is_not_metadata(tmp_path):
    source,output=excel_pair(tmp_path);book=load_workbook(output);book.active['C1']='=B1*3';book.save(output);book.close()
    assert not passed(compare(source,output))


def test_requested_new_sheet_is_allowed(tmp_path):
    source,output=excel_pair(tmp_path);book=load_workbook(output);book.create_sheet('Added')['A1']='New';book.save(output);book.close()
    assert passed(compare(source,output,[{'scope':'excel/sheets','operation':'add'}]))
    assert not passed(compare(source,output))


def test_sheet_removal_is_detected(tmp_path):
    source,output=excel_pair(tmp_path);book=load_workbook(output);book.create_sheet('Other');book.remove(book['Data']);book.save(output);book.close()
    assert not passed(compare(source,output))


def test_dropped_source_paragraph_is_defect(tmp_path):
    source,output=tmp_path/'source.docx',tmp_path/'out.docx';doc=Document()
    doc.add_heading('Keep heading',1);doc.add_paragraph('Keep business fact');doc.save(source)
    Document().save(output)
    assert not passed(compare(source,output))


def test_requested_added_paragraph_does_not_renumber_into_false_defects(tmp_path):
    source,output=tmp_path/'source.docx',tmp_path/'out.docx';doc=Document()
    doc.add_paragraph('Original first');doc.add_paragraph('Original second');doc.save(source)
    doc.paragraphs[0].insert_paragraph_before('New requested introduction');doc.save(output)
    assert passed(compare(source,output,[{'scope':'word/paragraphs','operation':'add'}]))


def test_heading_level_is_preserved(tmp_path):
    source,output=tmp_path/'source.docx',tmp_path/'out.docx';doc=Document();doc.add_heading('Heading',1);doc.save(source)
    doc.paragraphs[0].style='Normal';doc.save(output)
    assert not passed(compare(source,output))


def test_word_table_business_value_is_preserved(tmp_path):
    source,output=tmp_path/'source.docx',tmp_path/'out.docx';doc=Document();table=doc.add_table(rows=1,cols=1)
    table.cell(0,0).text='Amount 10';doc.save(source);table.cell(0,0).text='Amount 99';doc.save(output)
    assert not passed(compare(source,output))


def test_requested_added_slide_is_allowed(tmp_path):
    source,output=tmp_path/'source.pptx',tmp_path/'out.pptx';deck=Presentation()
    slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(0,0,Inches(2),Inches(1)).text='Keep';deck.save(source)
    slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(0,0,Inches(2),Inches(1)).text='New';deck.save(output)
    assert passed(compare(source,output,[{'scope':'powerpoint/slides','operation':'add'}]))
    assert not passed(compare(source,output))


def test_unrelated_slide_object_change_is_defect(tmp_path):
    source,output=tmp_path/'source.pptx',tmp_path/'out.pptx';deck=Presentation()
    slide=deck.slides.add_slide(deck.slide_layouts[6]);shape=slide.shapes.add_textbox(0,0,10,10);shape.text='Keep';deck.save(source)
    shape.text='Changed';deck.save(output)
    assert not passed(compare(source,output))


def test_metadata_only_change_is_not_business_defect(tmp_path):
    from datetime import datetime
    source,output=excel_pair(tmp_path);book=load_workbook(output);book.properties.modified=datetime(2026,1,1);book.properties.creator='Other serialization';book.save(output);book.close()
    assert passed(compare(source,output))


def test_preserved_region_overrides_broad_permission(tmp_path):
    source,output=excel_pair(tmp_path);book=load_workbook(output);book.active['B1']=99;book.save(output);book.close()
    assert not passed(compare(source,output,[{'scope':'excel/sheet/Data','operation':'replace'}],['excel/sheet/Data/cell/B1']))


def test_wrong_source_hash_is_not_an_artifact_defect(tmp_path):
    module=importlib.import_module('workagent_rsi.preservation_checks');source,output=excel_pair(tmp_path)
    spec=module.PreservationSpec(source_hash='0'*64,allowed_changes=[],preserved_regions=[],normalization_policy='office-structure-v1')
    with pytest.raises(ValueError,match='source hash'):module.compare_preservation(source,output,spec)


def test_comparison_does_not_modify_inputs(tmp_path):
    source,output=excel_pair(tmp_path);before=[sha256_file(p) for p in (source,output)]
    compare(source,output)
    assert before==[sha256_file(p) for p in (source,output)]


def test_shared_assessment_detects_unauthorized_edit(tmp_path):
    from workagent_rsi.assessment_contracts import AcceptanceSpec,Requirement,Location
    from workagent_rsi.artifact_assessment import assess
    source,output=excel_pair(tmp_path);book=load_workbook(output);book.active['B1']=99;book.save(output);book.close()
    digest=sha256_file(source)
    spec=AcceptanceSpec(task_id='edit',version='independent-v1',artifacts={'output':'excel'},
        input_hashes={str(source):digest},dimension_weights={'content':1},requirements=[Requirement(
            requirement_id='preserve',description='Keep original data',check='preservation',expected=True,
            location=Location(artifact='output'),critical=True,dimension='content',evidence_source='frozen source',
            options={'source_path':str(source),'preservation_spec':{'source_hash':digest,'allowed_changes':[],
                'preserved_regions':[],'normalization_policy':'office-structure-v1'}})])
    shared,score,_=assess(spec,{'output':output})
    assert shared.checks[0].status=='FAIL' and score.acceptance_status=='FAIL'


def test_authorized_paragraph_removal(tmp_path):
    source,output=tmp_path/'source.docx',tmp_path/'out.docx';doc=Document();doc.add_paragraph('Remove');doc.add_paragraph('Keep');doc.save(source)
    p=doc.paragraphs[0]._p;p.getparent().remove(p);doc.save(output)
    assert passed(compare(source,output,[{'scope':'word/paragraphs','operation':'remove'}]))


def test_add_only_permission_cannot_replace_existing_text(tmp_path):
    source,output=tmp_path/'source.docx',tmp_path/'out.docx';doc=Document();doc.add_paragraph('Keep');doc.save(source)
    doc.paragraphs[0].text='Changed';doc.save(output)
    assert not passed(compare(source,output,[{'scope':'word/paragraphs','operation':'add'}]))


def test_complex_chart_is_explicit_preservation_gap(tmp_path):
    from openpyxl.chart import BarChart,Reference
    source,output=excel_pair(tmp_path);book=load_workbook(source);chart=BarChart();chart.add_data(Reference(book.active,min_col=2,min_row=1));book.active.add_chart(chart,'E1');book.save(source);book.save(output);book.close()
    assert any(r.status=='NEEDS_REVIEW' for r in compare(source,output))


def test_equivalent_word_run_splitting_is_serialization_only(tmp_path):
    source,output=tmp_path/'source.docx',tmp_path/'out.docx';doc=Document();p=doc.add_paragraph();p.add_run('Hello');p.add_run(' world');doc.save(source)
    p.text='Hello world';doc.save(output)
    assert passed(compare(source,output))


def test_unknown_preserved_scope_is_explicit_error(tmp_path):
    source,output=excel_pair(tmp_path)
    with pytest.raises(ValueError,match='does not resolve'):compare(source,output,preserved=['excel/sheet/Missing'])
