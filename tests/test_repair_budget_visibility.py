"""T22 offline budget/visibility controls, including mocked transport boundaries."""
import inspect
import json
from pathlib import Path

import pytest

from test_artifact_repair_loop import executor, harness, policy
from test_assessment_entrypoint_parity import authored_spec, reviewed_task
from workagent_rsi.harness import Harness


def budget(**kwargs):
    assert 'run_budget' in inspect.signature(Harness.run).parameters, 'T22 shared runtime budget is absent'
    from workagent_rsi.repair_loop import RunBudget
    return RunBudget(**kwargs)


def test_reservation_blocks_before_call_and_failed_calls_remain_charged():
    b = budget(max_model_calls=1)
    ticket = b.reserve('model')
    b.reconcile(ticket)
    with pytest.raises(RuntimeError): b.reserve('model')
    assert b.snapshot()['used']['model'] == 1


def test_cost_cap_requires_known_upper_bound_and_retains_unknown_actual_cost():
    b = budget(max_cost_usd=1.0)
    with pytest.raises(RuntimeError): b.reserve('model')
    ticket = b.reserve('model', cost_upper_bound=.75)
    b.reconcile(ticket)
    with pytest.raises(RuntimeError): b.reserve('model', cost_upper_bound=.3)
    with pytest.raises(RuntimeError): b.reserve('tool', cost_upper_bound=.3)
    assert b.snapshot()['cost_reserved_usd'] == .75


def test_repair_budget_keeps_initial_and_counts_every_attempt(tmp_path, monkeypatch):
    calls = executor(monkeypatch, tmp_path, [7, 8, 42])
    result = harness(tmp_path).run(reviewed_task(authored_spec()), repair_policy=policy(max_repairs=9),
                                   run_budget=budget(max_repairs=1))
    assert len(calls)==2 and result['repair_loop']['repairs']==1
    assert result['repair_loop']['stop_reason']=='budget_exhausted'
    ledger = json.loads((Path(result['result_dir'])/'repair-loop/budget-ledger.json').read_text())
    assert ledger['used']['repair']==1


def test_failed_repair_cannot_replace_best_assessed_artifact(tmp_path, monkeypatch):
    calls = executor(monkeypatch, tmp_path, [7, 'SyntaxError: broken'])
    result = harness(tmp_path).run(reviewed_task(authored_spec()), repair_policy=policy(max_repairs=1), run_budget=budget())
    assert len(calls)==2 and result['artifacts']
    assert result['artifacts']==result['repair_loop']['attempts'][0]['artifacts']
    assert result['repair_loop']['attempts'][1]['failure']['failure_origin']=='agent_execution'
    assert 'evaluation' in result and result['repair_loop']['selected_attempt']==0


def test_interrupt_preserves_best_result_and_history(tmp_path, monkeypatch):
    from workagent_rsi.repair_loop import RepairLoop
    calls = executor(monkeypatch, tmp_path, [7])
    # Interrupt precisely after the first assessed result has been persisted.
    original = RepairLoop._persist
    def interrupt(self, *args, **kwargs):
        value = original(self, *args, **kwargs)
        if args[-1]=='running': raise KeyboardInterrupt
        return value
    monkeypatch.setattr(RepairLoop, '_persist', interrupt)
    result = harness(tmp_path).run(reviewed_task(authored_spec()), repair_policy=policy(), run_budget=budget())
    assert len(calls)==1 and result['artifacts'] and result['repair_loop']['stop_reason']=='interrupted'
    assert (Path(result['result_dir'])/'repair-loop/attempt-000/result.json').is_file()


@pytest.mark.parametrize('split', ['regression','hidden','promotion_validation','protected','ood_transfer','final_test'])
def test_protected_splits_never_receive_adaptive_repair(tmp_path, monkeypatch, split):
    assert 'split_name' in inspect.signature(Harness.run).parameters, 'T22 split boundary absent'
    calls = executor(monkeypatch, tmp_path, [7,42])
    result = harness(tmp_path).run(reviewed_task(authored_spec()), repair_policy=policy(), run_budget=budget(), split_name=split)
    assert len(calls)==1 and result['repair_loop']['stop_reason']=='protected_split'
    assert not list((Path(result['result_dir'])/'repair-loop').rglob('request.json'))


def test_hidden_flag_overrides_develop_label(tmp_path, monkeypatch):
    calls = executor(monkeypatch, tmp_path, [7,42])
    task = reviewed_task(authored_spec()).model_copy(update={'hidden_test':True})
    result = harness(tmp_path).run(task, repair_policy=policy(), run_budget=budget())
    assert len(calls)==1 and result['repair_loop']['stop_reason']=='protected_split'


def test_repeated_state_stops_before_maximum(tmp_path, monkeypatch):
    # Reuse exact OOXML bytes so ZIP timestamps are irrelevant to this test.
    from openpyxl import Workbook
    calls = []
    path=tmp_path/'same.xlsx';wb=Workbook();wb.active['A1']=7;wb.save(path);wb.close()
    class Same:
        def __init__(self,*args,**kwargs): pass
        def execute(self,*args):
            calls.append(1)
            yield {'kind':'artifact','artifact_path':str(path),'artifact_name':'output'}
    monkeypatch.setattr('workagent_rsi.run.WorkAgentOfficeAdapter',Same)
    result=harness(tmp_path).run(reviewed_task(authored_spec()),repair_policy=policy(max_repairs=8),run_budget=budget(max_repairs=8))
    assert len(calls)<=3 and result['repair_loop']['stop_reason']=='no_progress'


def test_unknown_split_is_rejected_before_execution(tmp_path, monkeypatch):
    calls=executor(monkeypatch,tmp_path,[42])
    with pytest.raises(ValueError):
        harness(tmp_path).run(reviewed_task(authored_spec()),repair_policy=policy(),run_budget=budget(),split_name='develpo')
    assert not calls


def test_conservative_cli_rejection_precedes_subprocess(tmp_path,monkeypatch):
    from workagent_rsi.workagent_provider import CodexOfficeProvider, WorkAgentConfig
    b=budget()
    monkeypatch.setattr('workagent_rsi.workagent_provider.subprocess.run',lambda *a,**k:pytest.fail('unbudgeted CLI started'))
    workspace=tmp_path/'workspace';workspace.mkdir()
    with b.activate(tmp_path/'budget.json'):
        outcome=CodexOfficeProvider(WorkAgentConfig()).run('authored',workspace,tmp_path/'records')
    assert outcome.status=='unavailable' and 'budget' in outcome.error.lower()
    assert b.snapshot()['used']['model']==0


def test_native_http_retry_consumes_shared_model_budget(tmp_path,monkeypatch):
    from urllib.error import HTTPError
    from workagent_rsi.workagent_provider import CodexOfficeProvider, WorkAgentConfig
    monkeypatch.setenv('AUTHORED_OFFLINE_KEY','test-only')
    calls=[]
    def reject(*args,**kwargs):
        calls.append(1); raise HTTPError('http://authored.invalid',429,'authored',{},None)
    monkeypatch.setattr('workagent_rsi.workagent_provider.urllib.request.urlopen',reject)
    cfg=WorkAgentConfig(backend='openai-chat',api_key_env='AUTHORED_OFFLINE_KEY',api_retry_delay_seconds=0,api_max_retries=4)
    workspace=tmp_path/'workspace';workspace.mkdir();b=budget(max_model_calls=1)
    with b.activate(tmp_path/'budget.json'):
        outcome=CodexOfficeProvider(cfg).run('authored',workspace,tmp_path/'records')
    assert len(calls)==1 and outcome.status=='unavailable' and b.snapshot()['used']['model']==1


def test_tool_budget_blocks_before_tool_side_effect(tmp_path):
    from workagent_rsi.workagent_provider import _execute_native_tool
    b=budget(max_tool_calls=0)
    with b.activate(tmp_path/'budget.json'), pytest.raises(RuntimeError):
        _execute_native_tool('write_file',{'path':'unwanted.txt','content':'bad'},tmp_path,timeout_seconds=10)
    assert not (tmp_path/'unwanted.txt').exists()


def test_deadline_denies_reservation_without_sleep():
    now=[0.0]; b=budget(max_wall_seconds=1,clock=lambda:now[0]);now[0]=2
    with pytest.raises(RuntimeError): b.reserve('model')
    assert b.snapshot()['used']['model']==0


def test_review_retries_share_execution_budget_and_count_only_started_calls(tmp_path, monkeypatch):
    from test_model_judgement import config, api_transport, spec, artifact
    from workagent_rsi.model_judgement import configured_assessment
    monkeypatch.setenv('OPENAI_API_KEY','authored-offline-only')
    calls=[]; cfg=config().model_copy(update={'retries':3})
    assessment=configured_assessment(cfg,tmp_path/'judge',transport=api_transport(calls,invalid=True))
    b=budget(max_model_calls=1)
    with b.activate(tmp_path/'budget.json'):
        report=assessment.inspect(spec(),{'out':artifact(tmp_path,'word')})
    assert len(calls)==1 and b.snapshot()['used']['model']==1
    assert report.telemetry['semantic'][0]['model_calls']==1


def test_disallowed_artifact_change_rejects_otherwise_passing_repair(tmp_path,monkeypatch):
    from openpyxl import Workbook
    from workagent_rsi.assessment_contracts import Requirement,Location
    spec=authored_spec()
    spec=spec.model_copy(update={'artifacts':{'output':'excel','other':'excel'},'requirements':[
        *spec.requirements,Requirement(requirement_id='other-readable',description='Keep the other workbook',
            check='file.readable',location=Location(artifact='other'),expected=True,critical=True,
            dimension='content',evidence_source='independent OOXML oracle')]})
    calls=[]
    class TwoBooks:
        def __init__(self,root,*args,**kwargs):self.root=Path(root)
        def execute(self,*args):
            calls.append(1);self.root.mkdir(parents=True,exist_ok=True)
            for name,value in [('output',7 if len(calls)==1 else 42),('other',100+len(calls))]:
                path=self.root/(name+'.xlsx');wb=Workbook();wb.active['A1']=value;wb.save(path);wb.close()
                yield {'kind':'artifact','artifact_path':str(path),'artifact_name':name}
    monkeypatch.setattr('workagent_rsi.run.WorkAgentOfficeAdapter',TwoBooks)
    result=harness(tmp_path).run(reviewed_task(spec),repair_policy=policy(max_repairs=1,allowed_artifacts=('output',)),run_budget=budget())
    assert len(calls)==2 and result['repair_loop']['selected_attempt']==0
    rejected=result['repair_loop']['attempts'][1]
    assert rejected['evaluation']['passed'] and rejected['selection_rejected']
    assert result['state']=='FAILED'


def test_interrupt_inside_executor_retains_charged_repair(tmp_path,monkeypatch):
    calls=executor(monkeypatch,tmp_path,[7])
    from workagent_rsi.run import Run
    execute=Run._execute_once
    def interrupted(self,*args,**kwargs):
        if kwargs.get('repair_request') is not None:raise KeyboardInterrupt
        return execute(self,*args,**kwargs)
    monkeypatch.setattr(Run,'_execute_once',interrupted)
    b=budget();result=harness(tmp_path).run(reviewed_task(authored_spec()),repair_policy=policy(),run_budget=b)
    assert len(calls)==1 and result['repair_loop']['stop_reason']=='interrupted'
    assert result['repair_loop']['repairs']==1 and result['repair_loop']['completed_repairs']==0
    assert b.snapshot()['used']['repair']==1 and b.snapshot()['entries'][0]['status']=='charged'
    assert result['artifacts'] and (Path(result['result_dir'])/'repair-loop/best-result.json').is_file()


def test_parallel_reservations_cannot_overrun_one_slot():
    from concurrent.futures import ThreadPoolExecutor
    b=budget(max_model_calls=1)
    def reserve():
        try:b.reserve('model');return True
        except RuntimeError:return False
    with ThreadPoolExecutor(max_workers=4) as pool:
        assert sum(pool.map(lambda _:reserve(),range(8)))==1


def test_evaluator_passes_protected_split_and_shared_budget_to_harness(tmp_path,monkeypatch):
    from test_execution_failure_feedback import setup
    from workagent_rsi.workagent_provider import WorkAgentSkill
    from workagent_rsi.hashing import canonical_json_hash
    tasks,evaluator,contract=setup(tmp_path,('only',))
    b=budget(); evaluator.repair_policy=policy();evaluator.run_budget=b
    evaluator.evaluator_hash=evaluator._configured_identity()
    contract=contract.model_copy(update={'evaluator_hash':evaluator.evaluator_hash,'split_hashes':{
        'regression':canonical_json_hash([t.model_dump(mode='json') for t in tasks])}})
    calls=executor(monkeypatch,tmp_path,[7,42])
    report=evaluator.evaluate_split(tasks,'regression',WorkAgentSkill(instructions='Frozen skill'),contract,tmp_path/'eval',reveal_per_task=False)
    assert len(calls)==1 and report['task_success_rate']==0 and b.snapshot()['used']['repair']==0
    manifests=list((tmp_path/'eval').rglob('attempt-manifest.json'))
    assert len(manifests)==1 and json.loads(manifests[0].read_text())['stop_reason']=='protected_split'


def test_environment_gap_retains_best_file_but_does_not_claim_complete_outcome(tmp_path,monkeypatch):
    from workagent_rsi.run_contracts import execution_from_legacy
    calls=[]
    from openpyxl import Workbook
    class Worker:
        def __init__(self,root,*args,**kwargs):self.root=Path(root)
        def execute(self,*args):
            calls.append(1)
            if len(calls)>1:
                yield {'kind':'unavailable','message':'worker disconnected','failure_origin':'environment'}
                return
            self.root.mkdir(parents=True,exist_ok=True);path=self.root/'bad.xlsx'
            wb=Workbook();wb.active['A1']=7;wb.save(path);wb.close()
            yield {'kind':'artifact','artifact_path':str(path)}
    monkeypatch.setattr('workagent_rsi.run.WorkAgentOfficeAdapter',Worker)
    result=harness(tmp_path).run(reviewed_task(authored_spec()),repair_policy=policy(),run_budget=budget())
    assert result['artifacts'] and result['repair_loop']['selected_attempt']==0
    assert execution_from_legacy(result).failure_origin=='environment' and result['state']=='UNAVAILABLE'


def test_outer_candidate_api_uses_the_same_budget(tmp_path,monkeypatch):
    from workagent_rsi.assessment_contracts import RSIFeedback
    from workagent_rsi.candidate_provider import OpenAIChatCandidateProvider
    (tmp_path/'skill.json').write_text('{"marker_source":"task_id"}')
    monkeypatch.setenv('OPENAI_API_KEY','authored-offline-only')
    monkeypatch.setattr('workagent_rsi.candidate_provider.urllib.request.urlopen',lambda *a,**k:pytest.fail('budget denial should precede transport'))
    provider=OpenAIChatCandidateProvider(model='authored',api_base_url='http://authored.invalid')
    b=budget(max_model_calls=0)
    with b.activate(tmp_path/'budget.json'):
        patch,record=provider.generate_with_feedback(tmp_path,RSIFeedback(tasks=[]),'0.1.0',1,tmp_path/'records')
    assert patch is None and record.status=='unavailable' and b.snapshot()['used']['model']==0


def test_elapsed_call_cannot_be_accepted_as_within_budget():
    from workagent_rsi.run_budget import budgeted_call
    now=[0.0];b=budget(max_wall_seconds=2,clock=lambda:now[0])
    with b.activate(),pytest.raises(RuntimeError):
        with budgeted_call('model'):
            now[0]=3
    assert b.snapshot()['used']['model']==1 and b.snapshot()['entries'][0]['status']=='charged'


def test_protected_feedback_injection_invalidates_cohort_before_repair(tmp_path,monkeypatch):
    from workagent_rsi.repair_loop import RepairLoop
    executor(monkeypatch,tmp_path,[7])
    task=reviewed_task(authored_spec())
    initial=harness(tmp_path).run(task)
    loop=RepairLoop(tmp_path/'isolated-loop',policy())
    result=loop.run(task,initial,lambda _: {'issues':[{'kind':'artifact_defect','feedback_visibility':'hidden','expected':'PRIVATE'}]},
                    lambda *args:pytest.fail('private feedback reached execution'),lambda result:result)
    assert result['repair_loop']['cohort_valid'] is False
    assert result['repair_loop']['stop_reason']=='visibility_violation'
    assert not list((tmp_path/'isolated-loop').rglob('request.json'))
