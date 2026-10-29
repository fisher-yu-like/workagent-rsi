"""T24 authored local library probes. No live model or remote sandbox claim."""
import importlib.util
import json
from pathlib import Path

import pytest


def features():
    assert importlib.util.find_spec('workagent_rsi.task_observations'), 'T24 task observations absent'
    assert importlib.util.find_spec('workagent_rsi.behavior_probes'), 'T24 behavior qualification absent'
    from workagent_rsi.behavior_probes import run_behavior_probes
    path=Path(__file__).resolve().parents[1]/'project_artifacts/phase3_experiments/scripts/run_t24_behavior_probe.py'
    spec=importlib.util.spec_from_file_location('t24_probe',path);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    return run_behavior_probes,module


@pytest.fixture(scope='module')
def full_qualification(tmp_path_factory):
    run,module=features();manifest=module.build_cases(tmp_path_factory.mktemp('T24-qualified'))
    return run({'instructions':'Inspect the current inputs.'},manifest,module.AuthoredExecutor())


@pytest.mark.parametrize('group',['excel-sheet','excel-values-rows','excel-columns',
    'word-target','word-paragraphs','word-table','ppt-count','ppt-target','ppt-append'])
def test_each_domain_has_three_input_sensitive_groups(full_qualification,group):
    report=full_qualification
    assert report['status']=='passed' and report['group_results'][group]['passed']
    assert report['model_calls']==0 and report['qualification_scope']=='authored_local_library'
    assert all(r['observation_status']=='observed' for r in report['rows'])
    assert len({r['generation_script_sha256'] for r in report['rows']})==1
    assert report['group_results'][group]['input_or_request_change_produces_expected_change']


def test_randomized_ids_keep_semantic_outputs_and_same_program_is_valid(full_qualification):
    report=full_qualification
    assert report['status']=='passed' and len(report['group_results'])==9
    assert report['id_invariance_passed'] and report['same_code_dynamic_input_passed']
    assert report['live_model_qualified'] is False and report['sandbox_qualified'] is False
    assert Path(report['trace_index']).is_file()


def test_probe_executor_never_receives_oracle_or_hidden_checks(tmp_path):
    run,module=features();manifest=module.build_cases(tmp_path/'cases');manifest['cases']=manifest['cases'][:2];manifest['required_groups']=['excel-sheet']
    actual=module.AuthoredExecutor();seen=[]
    def capture(skill,request,root):
        seen.append(request);assert 'expected' not in request and 'oracle' not in request
        assert not request['task'].get('expected_constraints')
        return actual(skill,request,root)
    report=run({},manifest,capture)
    assert report['status']=='passed' and len(seen)==2


def test_wrong_total_is_rejected_by_independent_oracle(tmp_path):
    run,module=features();manifest=module.build_cases(tmp_path/'cases');manifest['cases']=manifest['cases'][:2];manifest['required_groups']=['excel-sheet']
    actual=module.AuthoredExecutor()
    def corrupt(skill,request,root):
        from openpyxl import load_workbook
        result=actual(skill,request,root);path=Path(result['artifact'])
        wb=load_workbook(path);wb['Result']['A1']=999;wb.save(path);wb.close()
        return result
    report=run({},manifest,corrupt)
    assert report['status']=='failed' and all(not r['oracle_passed'] for r in report['rows'])


def test_no_executor_is_unavailable_and_never_a_pass(tmp_path):
    run,module=features();report=run({},module.build_cases(tmp_path/'cases'),None)
    assert report['status']=='unavailable' and not report['rows']


def test_missing_read_evidence_cannot_claim_input_driven_behavior(tmp_path):
    run,module=features();manifest=module.build_cases(tmp_path/'cases');manifest['cases']=manifest['cases'][:2];manifest['required_groups']=['excel-sheet']
    actual=module.AuthoredExecutor()
    def erase(skill,request,root):
        result=actual(skill,request,root)
        value=json.loads(Path(result['observation']).read_text());value.update(status='not_observed',inspected_structure={},inspection_evidence=[])
        Path(result['observation']).write_text(json.dumps(value))
        return result
    report=run({},manifest,erase)
    assert report['status']=='failed' and all(not r['input_evidence_passed'] for r in report['rows'])


def test_inspection_tool_uses_sandbox_and_reports_unavailable_without_fallback(tmp_path):
    features()
    from workagent_rsi.workagent_office import prepare_execution_view
    from workagent_rsi.workagent_provider import _execute_native_tool
    from workagent_rsi.contracts import TaskSpec
    from workagent_rsi.run_contracts import ExecutionOutcome
    from openpyxl import Workbook
    source=tmp_path/'source.xlsx';wb=Workbook();wb.save(source);wb.close()
    workspace=tmp_path/'workspace';manifest=prepare_execution_view(TaskSpec(task_id='a',domain='excel',instruction='Inspect',input_files=(str(source),)),tmp_path,workspace)
    calls=[]
    class Unavailable:
        def run(self,script,root,timeout):
            calls.append(script)
            return ExecutionOutcome(status='unavailable',failure_origin='environment')
    result=_execute_native_tool('inspect_office',{'path':manifest['files'][0]['copied']},workspace,sandbox=Unavailable(),records=tmp_path/'records')
    assert calls and result['ok'] is False and result['execution']['status']=='unavailable'
    assert not list((tmp_path/'records').rglob('*.inspection.json'))


def test_truncated_input_inspection_is_partial_and_never_claims_full_visibility(tmp_path):
    features()
    from workagent_rsi.office_input_inspector import inspect_office_file
    from openpyxl import Workbook
    path=tmp_path/'large.xlsx';book=Workbook();book.active['A1']='x'*4096;book.save(path);book.close()
    value=inspect_office_file(path,'excel')
    assert value['complete'] is False and len(value['sheets'][0]['values'][0][0])==2048


def test_authored_executor_refuses_arbitrary_script(tmp_path):
    _,module=features();script=tmp_path/'bad.py';script.write_text('raise SystemExit(0)')
    with pytest.raises(ValueError,match='arbitrary'):
        module.TrustedLocalExecutor().run(script,tmp_path,5)
