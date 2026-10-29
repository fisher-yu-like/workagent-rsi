from docx import Document
import pytest
from workagent_rsi.contracts import TaskSpec
from workagent_rsi.word_package_checks import inspect_word_rules
from test_benchmark_blind_pack import load_script


def test_literal_edit_plan_preserves_unrelated_text_without_signing_review(tmp_path):
    source=tmp_path/'source.docx';d=Document();d.add_paragraph('TechVentures Inc.');d.add_paragraph('Delaware');d.add_paragraph('Keep this paragraph.');d.save(source)
    task=TaskSpec(task_id='word-edit/3c94244b1a2888d506c5/native-1',domain='word',
        instruction="Change 'TechVentures Inc.' to 'Aurora Systems Inc.' throughout the document. Also change 'Delaware' to 'New York' as the governing law state.",input_files=[str(source)])
    module=load_script('complete_qualification_specs.py');draft=module.complete_task(task)
    assert draft['admission']['status']=='SPEC_INCOMPLETE' and not draft['spec']['requirement_inventory']['reviewed']
    assert draft['capability_preflight']['status']=='PASS'
    output=tmp_path/'output.docx';d=Document();d.add_paragraph('Aurora Systems Inc.');d.add_paragraph('New York');d.add_paragraph('Keep this paragraph.');d.save(output)
    assert inspect_word_rules(output,draft['native_word_rules'],source)[0]=='PASS'
    d.paragraphs[-1].text='Unauthorized change';d.save(output)
    assert inspect_word_rules(output,draft['native_word_rules'],source)[0]=='FAIL'
    changed=task.model_copy(update={'instruction':'Keep the company name unchanged.'})
    with pytest.raises(ValueError,match='different instruction'):
        module.complete_task(changed)


def test_versioned_api_table_rules_allow_requested_added_rows(tmp_path):
    module=load_script('complete_qualification_specs.py')
    d=Document();d.add_paragraph('API 3.2.1');t=d.add_table(rows=0,cols=5)
    for values in [('Tier','Requests/min','Batch Size','Concurrent Connections','Monthly Cost'),('Professional','1000','500','100','$1,999')]:
        for c,v in zip(t.add_row().cells,values):c.text=v
    source=tmp_path/'source.docx';d.save(source)
    task=TaskSpec(task_id='word-edit/f70a53d16d35d89c086b/native-1',domain='word',instruction='Update API to 3.3.0.',input_files=[str(source)])
    rules,limits=module.word_rules(task)
    assert not any(r['kind']=='tables_equal' for r in rules)
    assert any(r.get('table_contains')=='/accounts' for r in rules)
    assert any('metadata' in limit for limit in limits)


def test_new_article_order_uses_actual_source_heading_case(tmp_path):
    module=load_script('complete_qualification_specs.py')
    d=Document();d.add_paragraph('ARTICLE 8: LIMITATION OF LIABILITY');d.add_paragraph('ARTICLE 9: INDEMNIFICATION')
    source=tmp_path/'source.docx';d.save(source)
    task=TaskSpec(task_id='word-edit/b47b04d19bea10fbc814/native-1',domain='word',instruction='Add Article 8A.',input_files=[str(source)])
    rules,limits=module.word_rules(task)
    order=next(r for r in rules if r['kind']=='ordered_text')
    assert order['texts'][0]=='ARTICLE 8: LIMITATION OF LIABILITY'
    output=tmp_path/'output.docx';d=Document()
    for text in order['texts']:d.add_paragraph(text)
    d.save(output)
    assert inspect_word_rules(output,[order],source)[0]=='PASS'
    d.paragraphs[-1].text='Wrong heading';d.save(output)
    assert inspect_word_rules(output,[order],source)[0]=='FAIL'


def test_resolved_comment_short_quotes_select_full_source_comments_without_other_authors(tmp_path):
    from test_word_package_checks import add_part
    module=load_script('complete_qualification_specs.py');d=Document();d.add_paragraph('Source');source=tmp_path/'s.docx';d.save(source)
    ns='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    xml=f'<w:comments xmlns:w="{ns}">'+''.join(f'<w:comment w:id="{i}" w:author="{author}"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:comment>' for i,author,text in [(0,'Sarah Chen','Fixed - updated date.'),(1,'Sarah Chen','No longer relevant - completed.'),(2,'Other Author','Fixed - keep this one.')])+'</w:comments>'
    add_part(source,'word/comments.xml',xml)
    task=TaskSpec(task_id='word-edit/a4ed292e8ca46156df5b/native-1',domain='word',instruction='Delete two resolved comments.',input_files=[str(source)])
    rules,limits=module.word_rules(task)
    assert rules==[{'kind':'comments_preserved','remove_ids':['0','1']}]
