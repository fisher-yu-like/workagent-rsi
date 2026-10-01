from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from workagent_rsi.assessment_contracts import AcceptanceSpec, Requirement, Location
from workagent_rsi.artifact_assessment import assess


def requirement(rid, check, artifact, expected, **position):
    return Requirement(requirement_id=rid, description=rid, check=check, location=Location(artifact=artifact, **position), expected=expected, critical=True, dimension="content", evidence_source="independent test labels")


def test_word_headings_facts_tables_styles_and_references(tmp_path):
    doc = Document()
    doc.add_heading("Results", level=1)
    doc.add_paragraph("Revenue is 35 USD.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "As of", "2026-09-30"
    relation = doc.part.relate_to("https://example.org/source", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    doc.add_paragraph()._p.append(hyperlink)
    path = tmp_path / "report.docx"
    doc.save(path)
    spec = AcceptanceSpec(task_id="word", version="1", artifacts={"report": "word"}, dimension_weights={"content": 1}, requirements=[
        requirement("headings", "word.headings", "report", ["Results"]),
        requirement("style", "word.heading_style", "report", "Heading 1", heading="Results"),
        requirement("fact", "word.paragraph", "report", "Revenue is 35 USD.", paragraph=2),
        requirement("date", "word.table_cell", "report", "2026-09-30", table=1, row=1, column=2),
        requirement("sources", "word.references", "report", ["https://example.org/source"]),
    ])
    assert assess(spec, {"report": path})[1].total_score == 100
    table.cell(0, 1).text = "2026-08-31"
    doc.save(path)
    _, score, issues = assess(spec, {"report": path})
    assert score.acceptance_status == "FAIL"
    assert (issues.issues[0].location.table, issues.issues[0].location.row, issues.issues[0].location.column) == (1, 1, 2)


def test_powerpoint_chart_conclusion_and_legal_overlap(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.slide_width, deck.slide_height)
    title = slide.shapes.add_textbox(Inches(.3), Inches(.3), Inches(5), Inches(.5))
    title.text = "Total sales: 35 USD"
    data = ChartData()
    data.categories = ["East", "West"]
    data.add_series("Sales", [15, 20])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.3), Inches(1), Inches(5), Inches(3), data)
    path = tmp_path / "deck.pptx"
    deck.save(path)
    spec = AcceptanceSpec(task_id="slides", version="1", artifacts={"deck": "powerpoint"}, dimension_weights={"content": 1}, requirements=[
        requirement("slides", "powerpoint.slide_count", "deck", 1),
        requirement("conclusion", "powerpoint.object_text", "deck", "Total sales: 35 USD", slide=1, object_id=title.shape_id),
        requirement("chart", "powerpoint.chart_values", "deck", [[15, 20]], slide=1, object_id=chart.shape_id),
        requirement("bounds", "powerpoint.bounds", "deck", [], slide=1),
        requirement("overlap", "powerpoint.overlap", "deck", [], slide=1),
    ])
    assert assess(spec, {"deck": path})[1].total_score == 100
    title.text = "Total sales: 99 USD"
    deck.save(path)
    _, score, issues = assess(spec, {"deck": path})
    assert score.acceptance_status == "FAIL"
    assert issues.issues[0].location.object_id == title.shape_id


def test_ambiguous_overlap_is_review_not_automatic_failure(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for x in (1, 2):
        slide.shapes.add_textbox(Inches(x), Inches(1), Inches(3), Inches(2)).text = "Content"
    path = tmp_path / "overlap.pptx"
    deck.save(path)
    spec = AcceptanceSpec(task_id="overlap", version="1", artifacts={"deck": "powerpoint"}, dimension_weights={"content": 1}, requirements=[requirement("overlap", "powerpoint.overlap", "deck", [], slide=1)])
    shared, score, issues = assess(spec, {"deck": path})
    assert shared.checks[0].status == "NEEDS_REVIEW"
    assert score.total_score is None
    assert issues.issues[0].kind == "assessment_gap"


def test_cross_file_fact_uses_independent_expected_value(tmp_path):
    book = Workbook()
    book.active["A1"] = "2026-09-30 | USD | 35"
    book_path = tmp_path / "data.xlsx"
    book.save(book_path)
    doc = Document()
    doc.add_paragraph("2026-09-30 | USD | 99")
    doc_path = tmp_path / "report.docx"
    doc.save(doc_path)
    r = requirement("shared_fact", "facts.consistent", "book", "2026-09-30 | USD | 35", sheet="Sheet", cell="A1")
    r = r.model_copy(update={"options": {"locations": [{"artifact": "book", "sheet": "Sheet", "cell": "A1"}, {"artifact": "report", "paragraph": 1}]}})
    spec = AcceptanceSpec(task_id="cross", version="1", artifacts={"book": "excel", "report": "word"}, dimension_weights={"content": 1}, requirements=[r])
    _, score, issues = assess(spec, {"book": book_path, "report": doc_path})
    assert score.acceptance_status == "FAIL"
    assert score.total_score == 50
    assert issues.issues[0].location.artifact == "report"
    assert issues.issues[0].location.paragraph == 1
    # Agreement on the same wrong value must still fail.
    book.active["A1"] = "2026-09-30 | USD | 99"
    book.save(book_path)
    assert assess(spec, {"book": book_path, "report": doc_path})[1].total_score == 0
