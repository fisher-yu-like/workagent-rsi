import importlib

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
import pytest

from workagent_rsi.assessment_contracts import AcceptanceSpec,Requirement,Location


def spec(check='powerpoint.texts',expected=None):
    return AcceptanceSpec(task_id='visible',version='authored-v1',artifacts={'output':'powerpoint'},dimension_weights={'content':1},
        requirements=[Requirement(requirement_id='visible',description='Required content must be visible',check=check,
            expected=['Required'] if expected is None else expected,location=Location(artifact='output',slide=1),
            critical=True,dimension='content',evidence_source='authored geometry')])


def deck(tmp_path,kind):
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    if kind=='background':
        bg=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height);bg.fill.solid();bg.fill.fore_color.rgb=RGBColor(255,255,255)
    shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(.5));shape.text='Required'
    if kind=='outside':shape.left=prs.slide_width+1
    if kind=='rotated':shape.rotation=45
    if kind=='white':shape.text_frame.paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255)
    if kind=='covered':
        cover=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(5),Inches(3));cover.fill.solid();cover.fill.fore_color.rgb=RGBColor(255,255,255)
    path=tmp_path/'deck.pptx';prs.save(path);return prs,path


def test_findings_distinguish_certain_overflow_and_visual_risk(tmp_path):
    module=importlib.import_module('workagent_rsi.visibility_checks');prs,_=deck(tmp_path,'outside')
    rows=module.inspect_visibility(prs,1)
    assert any(r.kind=='off_canvas' and r.certainty=='certain' for r in rows)


def test_group_coordinates_are_transformed_to_slide_space(tmp_path):
    module=importlib.import_module('workagent_rsi.visibility_checks');prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    group=slide.shapes.add_group_shape();shape=group.shapes.add_textbox(0,0,10,10);shape.text='Required'
    group.left=prs.slide_width+100
    assert any(r.object_id==shape.shape_id and r.kind=='off_canvas' for r in module.inspect_visibility(prs,1))


@pytest.mark.parametrize('rendered,checked,complete',[([1,2],[1,2],True),([1],[1],False),([1,2],[1],False)])
def test_unchecked_required_page_blocks_visual_completion(rendered,checked,complete):
    from workagent_rsi.visibility_checks import VisualCoverageLedger
    ledger=VisualCoverageLedger(required_pages=[1,2],rendered_pages=rendered,checked_pages=checked)
    assert ledger.complete is complete


def test_checked_page_without_render_is_invalid():
    from workagent_rsi.visibility_checks import VisualCoverageLedger
    with pytest.raises(ValueError):VisualCoverageLedger(required_pages=[1],rendered_pages=[],checked_pages=[1])


class AuthoredProvider:
    def __init__(self,*,miss_last=False,fail_last=False):self.calls=[];self.miss_last=miss_last;self.fail_last=fail_last
    def identity_config(self):return {'provider':'authored-page-coverage-unit-fixture'}
    def judge(self,request,work):
        from workagent_rsi.judgement_provider import VisualJudgement
        self.calls.append(request)
        images=request['images'];last=any(i['image_id']=='page-17' for i in images)
        if self.fail_last and last:return None,{'status':'unavailable','error':'authored last-batch failure','model_calls':0}
        selected=images[:-1] if self.miss_last and last else images
        return VisualJudgement(status='PASS',completion=1,observed='Authored mapping fixture',observations=[
            {'image_id':i['image_id'],'artifact_location':i['artifact_location'],'region':'whole page','observation':'Authored citation only'} for i in selected]),{'model_calls':1}


def long_visual(tmp_path,provider,*,cache=False):
    from test_visual_assessment import FakeRenderer,artifact
    from workagent_rsi.judgement_provider import VisualAssessmentChannel
    path=artifact(tmp_path)
    channel=VisualAssessmentChannel(FakeRenderer(tmp_path/'render',pages=17),provider,tmp_path/'cache',max_images_per_check=8,cache_enabled=cache)
    req=Requirement(requirement_id='all-pages',description='Inspect all required pages',check='visual',
        location=Location(artifact='output'),dimension='content',evidence_source='authored coverage',options={'pages':list(range(1,18))})
    return channel,req,path
