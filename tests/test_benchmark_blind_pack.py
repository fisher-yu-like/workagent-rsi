"""Partial actual cohorts cannot satisfy the twenty-file human-review gate."""
import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from workagent_rsi.hashing import sha256_file


def load_script(name):
    path = Path(__file__).parents[1] / 'project_artifacts/phase3_experiments/scripts' / name
    spec = importlib.util.spec_from_file_location(path.stem, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_partial_natural_only_pack_keeps_shortage_and_empty_human_choices(tmp_path, monkeypatch):
    from openpyxl import Workbook
    from docx import Document
    from pptx import Presentation
    source = tmp_path / 'standards'
    source.mkdir()
    Workbook().save(source / 'source.xlsx')
    Document().save(source / 'source.docx')
    Presentation().save(source / 'source.pptx')
    cases = []
    for domain, suffix, defect in [('excel', 'xlsx', 'wrong_range'), ('word', 'docx', 'wrong_value'), ('powerpoint', 'pptx', 'obstruction')]:
        for category in ['correct', 'single_defect', 'equivalent', 'gap']:
            cases.append(SimpleNamespace(case_id=domain+'-'+category, domain=domain, category=category,
                defect_type=defect if category == 'single_defect' else None,
                instruction='Authored pack-builder unit fixture.', artifact_path='source.'+suffix,
                input_path='source.'+suffix, source_family='fixture-'+domain, oracle_status='INCOMPLETE'))
    module = load_script('prepare_human_calibration.py')
    monkeypatch.setattr(module, 'load_calibration_cases', lambda _: cases)
    artifact = tmp_path / 'actual.xlsx'
    Workbook().save(artifact)
    natural = tmp_path / 'natural.json'
    natural.write_text(json.dumps({'cases': [{'case_id': 'one', 'domain': 'excel',
        'source_kind': 'natural_model_output', 'source_family': 'authored-new-family',
        'source_refs': ['authored-unit-provenance'], 'artifact_path': artifact.name,
        'artifact_hash': sha256_file(artifact), 'instruction': 'Authored natural-intake unit fixture.', 'input_paths': []}]}), encoding='utf-8')
    with pytest.raises(ValueError, match='at least20'):
        module.prepare(source / 'manifest.json', tmp_path / 'refused', natural_manifest=natural)
    out = tmp_path / 'partial'
    report = module.prepare(source / 'manifest.json', out, natural_manifest=natural,
                            allow_partial_natural=True, include_standards=False)
    assert report['status'] == 'blocked' and report['natural_cases'] == 1
    assert report['standards_cases'] == 0 and report['actual_human_labels'] == 0
    assert any('actual1' in b for b in report['blockers'])
    manifest = json.loads((out / 'blind-review-pack/review-manifest.json').read_text(encoding='utf-8'))
    assert len(manifest['case_ids']) == 1 and 'oracle' not in json.dumps(manifest)
    page = (out / 'blind-review-pack/index.html').read_text(encoding='utf-8')
    assert page.count('<select ') == 1 and '<option value="">请选择</option>' in page


def test_cached_page_attachment_rejects_changed_actual_page(tmp_path):
    module = load_script('attach_benchmark_review_renders.py')
    public = tmp_path / 'review/blind-review-pack'
    public.mkdir(parents=True)
    source = tmp_path / 'source'
    source.mkdir()
    (source / 'review-manifest.json').write_text('{"cases":[]}', encoding='utf-8')
    (source / 'rendered-pages.json').write_text('[]', encoding='utf-8')
    artifact = public / 'opaque.xlsx'
    artifact.write_bytes(b'authored hash-binding fixture; not an Office qualification')
    manifest = {'cases': [{'case_id': 'opaque', 'artifact': artifact.name, 'artifact_hash': sha256_file(artifact)}]}
    (public / 'review-manifest.json').write_text(json.dumps(manifest), encoding='utf-8')
    (public / 'index.html').write_text('<article id="opaque"><div class="pages"></div></article>', encoding='utf-8')
    (public.parent / 'agreement-report.json').write_text('{"status":"blocked"}', encoding='utf-8')
    cache = tmp_path / 'cache'
    cache.mkdir()
    page = cache / 'page.png'
    page.write_bytes(b'authored page identity')
    (cache / 'render-results.json').write_text(json.dumps([{'artifact_sha256': sha256_file(artifact), 'status': 'available',
        'render': {'artifact_sha256': sha256_file(artifact), 'pages': [{'path': str(page), 'page': 1, 'sha256': sha256_file(page)}]}}]), encoding='utf-8')
    page.write_bytes(b'tampered page')
    with pytest.raises(ValueError, match='cached actual page changed'):
        module.attach(public.parent, source, cache)
    assert not (public / 'opaque-page-1.png').exists()
