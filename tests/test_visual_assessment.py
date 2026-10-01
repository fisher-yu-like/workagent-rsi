import json
from pathlib import Path
import subprocess

from PIL import Image

from workagent_rsi.artifact_assessment import assess
from workagent_rsi.artifact_renderer import RenderResult, RenderedPage
from workagent_rsi.assessment_contracts import AcceptanceSpec, Location, Requirement
from workagent_rsi.hashing import canonical_json_hash, sha256_file
from workagent_rsi.judgement_provider import CommandVisualProvider, VisualAssessmentChannel


class FakeRenderer:
    def __init__(self, root, status="available", identity="r1", pages=1):
        self.root, self.status, self.identity, self.pages = Path(root), status, identity, pages
    def identity_config(self):
        return {"renderer": self.identity}
    def render(self, path, domain, artifact_sha256):
        if self.status != "available":
            return RenderResult(status=self.status, artifact_sha256=artifact_sha256, renderer_identity=self.identity, elapsed_seconds=0, error="controlled renderer failure")
        self.root.mkdir(parents=True, exist_ok=True)
        pages = []
        for number in range(1, self.pages + 1):
            image = self.root / f"page-{number}.png"
            Image.new("RGB", (64, 32), "white").save(image)
            pages.append(RenderedPage(image_id=f"page-{number}", path=str(image), sha256=sha256_file(image), page=number, artifact_location=f"slide {number}", width=64, height=32))
        return RenderResult(status="available", artifact_sha256=artifact_sha256, renderer_identity=self.identity, elapsed_seconds=.01,
                            pages=pages)


def provider(root, calls, *, invalid_mapping=False, identity="vlm-1", code=0):
    def runner(command, cwd, timeout):
        request = json.loads(Path(command[1]).read_text())
        calls.append(request)
        observation = {"image_id": "unknown" if invalid_mapping else "page-1", "artifact_location": "slide 1", "region": "title", "observation": "Title is readable"}
        Path(command[2]).write_text(json.dumps({"status": "PASS", "completion": 1, "observed": "Readable", "observations": [observation]}), encoding="utf-8")
        return subprocess.CompletedProcess(command, code, stdout="", stderr="controlled" if code else "")
    return CommandVisualProvider(["fake-vlm", "{request}", "{output}"], model_identity=identity, prompt_version="visual-v1", runner=runner)


def spec():
    return AcceptanceSpec(task_id="visual", version="1", artifacts={"deck": "powerpoint"}, dimension_weights={"presentation": 1}, requirements=[Requirement(requirement_id="legibility", description="Title is readable without clipping", check="visual", location=Location(artifact="deck", slide=1), expected="readable", critical=True, dimension="presentation", evidence_source="real rendered slide", options={"pages": [1]})])


def artifact(tmp_path):
    # Shared parser needs a genuine PPTX. The renderer itself is injected here.
    from pptx import Presentation
    path = tmp_path / "deck.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    return path


def test_visual_channel_passes_real_image_mapping_and_caches(tmp_path):
    calls = []
    channel = VisualAssessmentChannel(FakeRenderer(tmp_path / "render"), provider(tmp_path, calls), tmp_path / "cache")
    path = artifact(tmp_path)
    first = assess(spec(), {"deck": path}, visual=channel)
    second = assess(spec(), {"deck": path}, visual=channel)
    assert first[1].total_score == second[1].total_score == 100
    assert len(calls) == 1
    assert Path(calls[0]["images"][0]["path"]).is_file()
    assert calls[0]["images"][0]["sha256"] == sha256_file(calls[0]["images"][0]["path"])
    assert second[0].telemetry["visual"][0]["model_calls"] == 0


def test_file_or_model_identity_change_misses_cache(tmp_path):
    calls = []
    path = artifact(tmp_path)
    first = VisualAssessmentChannel(FakeRenderer(tmp_path / "render"), provider(tmp_path, calls, identity="v1"), tmp_path / "cache")
    assess(spec(), {"deck": path}, visual=first)
    second = VisualAssessmentChannel(FakeRenderer(tmp_path / "render"), provider(tmp_path, calls, identity="v2"), tmp_path / "cache")
    assess(spec(), {"deck": path}, visual=second)
    from pptx import Presentation
    deck = Presentation(path)
    deck.slides[0].shapes.add_textbox(0, 0, 10, 10).text = "changed"
    deck.save(path)
    assess(spec(), {"deck": path}, visual=second)
    assert len(calls) == 3


def test_renderer_and_provider_failures_are_incomplete(tmp_path):
    path = artifact(tmp_path)
    calls = []
    unavailable = VisualAssessmentChannel(FakeRenderer(tmp_path / "r1", "unavailable"), provider(tmp_path, calls), tmp_path / "c1")
    shared, score, issues = assess(spec(), {"deck": path}, visual=unavailable)
    assert shared.checks[0].status == "UNAVAILABLE" and score.total_score is None and issues.issues[0].kind == "assessment_gap"
    broken = VisualAssessmentChannel(FakeRenderer(tmp_path / "r2"), provider(tmp_path, calls, code=1), tmp_path / "c2")
    shared, score, _ = assess(spec(), {"deck": path}, visual=broken)
    assert shared.checks[0].status == "ERROR" and score.acceptance_status == "INCOMPLETE"


def test_invalid_visual_location_is_checker_error(tmp_path):
    path = artifact(tmp_path)
    calls = []
    channel = VisualAssessmentChannel(FakeRenderer(tmp_path / "r"), provider(tmp_path, calls, invalid_mapping=True), tmp_path / "c")
    shared, score, _ = assess(spec(), {"deck": path}, visual=channel)
    assert shared.checks[0].status == "ERROR"
    assert score.total_score is None


def test_visual_identity_is_part_of_fixed_assessment_identity(tmp_path):
    calls = []
    a = VisualAssessmentChannel(FakeRenderer(tmp_path / "a", identity="a"), provider(tmp_path, calls, identity="v1"), tmp_path / "ca")
    b = VisualAssessmentChannel(FakeRenderer(tmp_path / "b", identity="b"), provider(tmp_path, calls, identity="v1"), tmp_path / "cb")
    from workagent_rsi.artifact_assessment import SharedAssessment
    assert SharedAssessment(visual=a).identity() != SharedAssessment(visual=b).identity()


def test_visual_budget_stops_before_model_call(tmp_path):
    path = artifact(tmp_path)
    calls = []
    limited = VisualAssessmentChannel(FakeRenderer(tmp_path / "r"), provider(tmp_path, calls), tmp_path / "c", max_images_per_check=0)
    shared, score, _ = assess(spec(), {"deck": path}, visual=limited)
    assert shared.checks[0].status == "UNAVAILABLE"
    assert score.total_score is None
    assert calls == []


def test_visual_judgement_must_cite_every_selected_page(tmp_path):
    path = artifact(tmp_path)
    calls = []
    two_page_spec = spec().model_copy(update={"requirements": [spec().requirements[0].model_copy(update={"options": {"pages": [1, 2]}})]})
    channel = VisualAssessmentChannel(FakeRenderer(tmp_path / "r", pages=2), provider(tmp_path, calls), tmp_path / "c")
    shared, score, _ = assess(two_page_spec, {"deck": path}, visual=channel)
    assert shared.checks[0].status == "NEEDS_REVIEW"
    assert score.total_score is None
